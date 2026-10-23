"""Sharing recovery and media tests use temporary stores and never send real messages."""
import asyncio
import copy
import hashlib
import io
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, Mock, patch

from bin import test_learning_sharing as fixture
from bin.test_learning_sharing_materials import local_only
from lan_bitable_template_portal.learning import LearningError
from lan_bitable_template_portal.learning_sharing import evidence_locations
from lan_bitable_template_portal.learning_sharing_routes import SharingRuntime, install
from lan_bitable_template_portal.sharing_images import SharingImages
from lan_bitable_template_portal.sharing_material_images import extract_images
from lan_bitable_template_portal.sharing_sources import SharingMaterials


class SharingResilienceTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.enterContext(patch('socket.socket.connect', local_only))
        self.fixture = fixture.SharingTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.core, self.learning, self.bid = self.fixture.core, self.fixture.learning, self.fixture.bid

    def client(self):
        from fastapi import FastAPI
        from fastapi.responses import JSONResponse
        from fastapi.testclient import TestClient
        app = FastAPI()
        controller = SimpleNamespace(_auth_required_response=lambda: JSONResponse({}, status_code=401),
            _json_ok=lambda request, session, data: JSONResponse({'ok': True, 'data': data}))
        async def body(request): return await request.json()
        install(app, controller, SimpleNamespace(service=object()), lambda: self.learning,
            lambda _: (object(), fixture.ADMIN), body, Mock(),
            lambda error: JSONResponse({'error': str(error)}, status_code=error.status if isinstance(error, LearningError) else 500))
        client = TestClient(app)
        self.addCleanup(client.close)
        return client

    async def test_short_valid_generation_can_retry_without_rewriting_good_questions(self):
        source = {'text': fixture.SOURCE, 'title': 'Fixture source', 'participants': self.core.batch(self.bid)['participants']}
        with self.learning.transaction() as conn:
            batch = self.core.batch(self.bid, conn)
            batch.update(status='queued', questions=[], question_types=['single'])
            self.learning._put('sharing_batch', self.bid, batch, conn)
        values = [{**q, 'type': 'single', 'correct': [0]} for q in fixture.questions(5)]
        model = AsyncMock(return_value=values[:3])
        worker = SharingRuntime(SimpleNamespace(service=object(), generate_sharing_questions=model), lambda: self.learning)
        with patch('lan_bitable_template_portal.learning_sharing_routes.read_source', return_value=source):
            await worker.generate(self.core.batch(self.bid))
        saved = self.core.batch(self.bid)
        self.assertEqual((saved['status'], len(saved['questions'])), ('failed', 3))
        response = self.client().post(f'/api/learning/sharing/{self.bid}/retry', json={'version': saved['version']})
        self.assertEqual(response.status_code, 200, response.text)
        model.return_value = values[3:]
        await worker.generate(self.core.batch(self.bid))
        result = self.core.batch(self.bid)
        self.assertEqual((result['status'], len(result['questions'])), ('draft', 5))
        self.assertEqual(result['questions'][:3], saved['questions'])
        self.assertEqual(model.await_count, 2)
        self.assertEqual(self.learning._all('sharing_delivery'), [])

    def test_admin_gets_own_paper_without_changing_other_persons_paper(self):
        with self.learning.transaction() as conn:
            batch = self.core.batch(self.bid, conn)
            batch['participants'][0]['open_id'] = fixture.ADMIN['id']
            self.learning._put('sharing_batch', self.bid, batch, conn)
        self.fixture.publish()
        result = self.core.detail(self.bid, fixture.ADMIN)
        self.assertEqual(result['paper']['open_id'], fixture.ADMIN['id'])
        self.assertEqual(len(result['my_questions']), 5)
        self.assertTrue(all('correct' not in q and 'evidence' not in q for q in result['my_questions']))
        self.assertTrue(all('correct' in q for q in result['questions']))
        self.assertNotIn('paper', self.core.detail(self.bid, {'id': 'ou_duty', 'shared_account': True, 'scope': 'A'}))

    async def test_reminder_never_overwrites_paper_receipt_and_old_status_is_repaired(self):
        self.fixture.publish()
        attempt = self.core.attempts(self.bid)[0]
        with self.learning.transaction() as conn:
            attempt.update(delivery_status='sent', submitted_at=fixture.TODAY.isoformat())
            self.learning._put('sharing_attempt', attempt['id'], attempt, conn)
            job = self.learning._get('sharing_delivery', 'paper:' + attempt['id'], conn)
            job.update(status='sent', sent_at=fixture.TODAY.isoformat(), message_id='fixture-message')
            self.learning._put('sharing_delivery', job['id'], job, conn)
        worker = SharingRuntime(SimpleNamespace(), lambda: self.learning)
        worker.core = self.core
        await worker._send({'id': 'reminder-fixture', 'batch_id': self.bid, 'attempt_id': attempt['id'],
            'kind': 'reminder', 'date': fixture.TODAY.date().isoformat(), 'failures': 0})
        current = self.learning._get('sharing_attempt', attempt['id'])
        self.assertEqual((current['delivery_status'], current['reminder_delivery_status']), ('sent', 'skipped'))
        with self.learning.transaction() as conn:
            current['delivery_status'] = 'skipped'
            self.learning._put('sharing_attempt', current['id'], current, conn)
        row = next(p for p in self.core.detail(self.bid, fixture.ADMIN)['participants'] if p['open_id'] == attempt['open_id'])
        self.assertEqual((row['delivery_status'], row['reminder_delivery_status']), ('sent', 'skipped'))

    def test_light_listing_does_not_call_detail_and_progress_excludes_content(self):
        self.fixture.publish()
        with patch.object(self.core, 'detail', side_effect=AssertionError('list must not fetch full detail')):
            self.assertEqual(self.core.listing(fixture.ADMIN)['items'][0]['total'], 3)
            self.assertEqual(self.core.listing({'id': 'ou_person0'})['items'][0]['total'], 1)
            self.assertEqual(self.core.listing({'id': 'ou_stranger'})['total'], 0)
        progress = self.core.detail(self.bid, fixture.ADMIN, progress_only=True)
        self.assertTrue({'questions', 'results', 'source_text', 'source_images', 'evidence_locations'}.isdisjoint(progress))
        compact = self.core.detail(self.bid, fixture.ADMIN, compact=True)
        self.assertTrue(all('answers' not in row for row in compact['results']))
        self.assertTrue(all({'answers', 'draft_answers', 'saved_operations', 'ticket'}.isdisjoint(row) for row in self.core.attempts(self.bid, summary=True)))
        paper = self.core.answer_detail(self.bid, fixture.ADMIN, 'ou_person0')
        self.assertEqual(len(paper['question_ids']), 5)
        with self.assertRaises(LearningError): self.core.answer_detail(self.bid, {'id': 'ou_person0'}, 'ou_person1')

    def test_draft_can_add_directory_person_once_and_published_roster_is_immutable(self):
        with self.learning.transaction() as conn:
            self.learning._put('person', 'p_added', {'id': 'p_added', 'name': 'Added', 'active': True, 'login_ids': ['ou_added'], 'scopes': ['A']}, conn, False)
        data = {'version': 1, 'questions': self.core.batch(self.bid)['questions'], 'participant_ids': ['p_added']}
        result = self.core.edit(self.bid, data, fixture.ADMIN)
        self.assertEqual(len(result['participants']), 4)
        data['version'] = result['version']
        result = self.core.edit(self.bid, data, fixture.ADMIN)
        self.assertEqual(len(result['participants']), 4)
        self.core.publish(self.bid, result['version'], fixture.ADMIN, 'cli_bot')
        with self.assertRaises(LearningError): self.core.edit(self.bid, {**data, 'version': result['version'] + 1}, fixture.ADMIN)

    def test_repeated_material_upload_is_private_and_idempotent(self):
        materials = SharingMaterials(self.learning)
        first = materials.upload(fixture.ADMIN, 'fixture.pdf', b'synthetic bytes')
        second = materials.upload(fixture.ADMIN, 'fixture.pdf', b'synthetic bytes')
        other = materials.upload({**fixture.ADMIN, 'id': 'ou_other'}, 'fixture.pdf', b'synthetic bytes')
        self.assertEqual(first['id'], second['id'])
        self.assertNotEqual(first['id'], other['id'])
        self.assertEqual(len(list(materials.root.iterdir())), 2)
        path = Path(self.learning._get('sharing_material', first['id'])['path'])
        path.write_bytes(b'corrupted')
        repaired = materials.upload(fixture.ADMIN, 'fixture.pdf', b'synthetic bytes')
        self.assertEqual(repaired['id'], first['id'])
        self.assertEqual(path.read_bytes(), b'synthetic bytes')

    async def test_material_retry_keeps_successful_snapshot(self):
        first, second = 'https://vnet.feishu.cn/docx/First', 'https://vnet.feishu.cn/docx/Second'
        with self.learning.transaction() as conn:
            batch = self.core.batch(self.bid, conn)
            batch.update(status='queued', url='', material_urls=[first, second], auto_generate=False,
                         requested_participants=batch['participants'])
            self.learning._put('sharing_batch', self.bid, batch, conn)
        worker = SharingRuntime(SimpleNamespace(service=object()), lambda: self.learning)
        reader = Mock(side_effect=[{'title': 'First', 'text': fixture.SOURCE}, LearningError('second unavailable', 503)])
        with patch('lan_bitable_template_portal.learning_sharing_routes.read_material_link', reader):
            await worker.generate(self.core.batch(self.bid))
            saved = self.core.batch(self.bid)
            self.assertEqual(saved['material_state'][first]['status'], 'ready')
            self.assertEqual(saved['material_state'][second]['status'], 'failed')
            reader.reset_mock(side_effect=True)
            reader.return_value = {'title': 'Second', 'text': fixture.SOURCE}
            await worker.generate(saved)
        self.assertEqual(reader.call_count, 1)
        self.assertEqual(reader.call_args.args[1], second)
        self.assertEqual(self.core.batch(self.bid)['status'], 'draft')

    def test_evidence_points_to_original_page_and_is_not_in_unanswered_view(self):
        batch = self.core.batch(self.bid)
        quote = batch['questions'][0]['evidence']
        batch['source'] = {'sources': [{'title': 'Slides', 'text': quote, 'references': [{'kind': 'slide', 'slide': 2, 'start': 0, 'end': len(quote)}]}]}
        self.assertEqual(evidence_locations(batch, batch['questions'][0])[0]['references'][0]['slide'], 2)
        self.fixture.publish()
        self.assertEqual(self.core.detail(self.bid, {'id': 'ou_person0'})['evidence_locations'], {})

    def test_real_office_illustrations_are_protected_local_snapshots(self):
        from PIL import Image
        from docx import Document
        from pptx import Presentation
        from pptx.util import Inches
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as SheetImage
        from reportlab.pdfgen.canvas import Canvas
        folder = Path(self.fixture.folder.name)
        image_path = folder / 'fixture.png'
        Image.new('RGB', (80, 60), 'green').save(image_path)
        document = Document(); document.add_picture(str(image_path)); document.save(folder / 'fixture.docx')
        slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[6]); slide.shapes.add_picture(str(image_path), Inches(1), Inches(1)); slides.save(folder / 'fixture.pptx')
        book = Workbook(); book.active.add_image(SheetImage(io.BytesIO(image_path.read_bytes())), 'A1'); book.save(folder / 'fixture.xlsx'); book.close()
        pdf = Canvas(str(folder / 'fixture.pdf')); pdf.drawImage(str(image_path), 20, 20); pdf.showPage(); pdf.save()
        store = SharingImages(self.learning, None)
        for suffix in ('.docx', '.pptx', '.xlsx', '.pdf'):
            with self.subTest(suffix=suffix):
                path = folder / ('fixture' + suffix)
                refs = extract_images(path, hashlib.sha256(path.read_bytes()).hexdigest(), store)
                self.assertEqual(len(refs), 1)
                prepared = store.prepare(refs, ocr=False)
                self.assertEqual(len(prepared), 1)
                self.assertTrue(store.content(prepared[0]).startswith(b'\xff\xd8'))
                self.assertNotIn('path', prepared[0])

    def test_picture_only_material_retains_images_and_corruption_is_actionable(self):
        source = Path(self.fixture.folder.name) / 'broken.pdf'
        source.write_bytes(b'not a pdf')
        with self.assertRaisesRegex(LearningError, '图片无法完整读取'):
            extract_images(source, 'fixture', SharingImages(self.learning, None))
        images = [{'local_sha256': 'a' * 64, 'label': 'Slide photo'}]
        with patch('openclaw_service.assistant.lighthouse_document_tools.run_local', return_value={
                'text': '', 'empty': True, 'extraction_complete': True}):
            parsed = SharingMaterials.parse(source.with_suffix('.pptx'), 'fixture', images=images)
        self.assertEqual(parsed['images'], images)
        self.assertEqual(parsed['text'], '')


if __name__ == '__main__':
    unittest.main()
