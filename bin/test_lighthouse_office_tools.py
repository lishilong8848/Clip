"""Local synthetic workbook/PDF/diagram workflows; no production/cloud writes."""
import copy
import io
import json
import os
import sys
import tempfile
import unittest
from contextlib import asynccontextmanager
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent))
from openclaw_service.assistant.lighthouse_ai import AssistantError
from openclaw_service.assistant.lighthouse_files import LighthouseFiles
from openclaw_service.assistant.lighthouse_document_tools import run_local
from openclaw_service.assistant.lighthouse_office_tools import prepare_inputs, analyze_tables, organize_pdf
from openclaw_service.assistant.lighthouse_diagrams import flowchart_source

ACTOR = {'id': 'office-fixture', 'scopes': ['D']}


class Store:
    def __init__(self, root):
        self.db_path, self.docs = root / 'state.sqlite3', {}

    def get_document(self, namespace, key):
        return copy.deepcopy(self.docs.get((namespace, key)))

    def put_document(self, namespace, key, value):
        self.docs[namespace, key] = copy.deepcopy(value)


def workbook(rows, second_sheet=False):
    from openpyxl import Workbook
    book = Workbook()
    for row in rows:
        book.active.append(row)
    if second_sheet:
        book.create_sheet('另一张表').append(['编号'])
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def pdf_bytes(labels):
    from reportlab.pdfgen.canvas import Canvas
    output = io.BytesIO()
    canvas = Canvas(output)
    for text in labels:
        canvas.drawString(40, 700, text)
        canvas.showPage()
    canvas.save()
    return output.getvalue()


class OfficeToolsTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.files = LighthouseFiles(Store(Path(temporary.name)))

    def upload(self, name, content):
        return self.files.upload(ACTOR, name, content, extract=False, source_scopes=['D'])['id']

    def inputs(self, ids, kind='spreadsheet', **options):
        return prepare_inputs(self.files, ACTOR, [{'file_id': identity, **options} for identity in ids], kind)

    def test_merge_preserves_rows_and_detects_duplicates_without_formula_injection(self):
        from openpyxl import load_workbook
        ids = [self.upload('第一份.csv', '编号,数量\n001,2\n002,3\n'.encode()),
               self.upload('第二份.csv', '编号,数量\n001,4\n003,=1+1\n'.encode())]
        result = run_local('spreadsheet', {'inputs': self.inputs(ids), 'operation': 'merge', 'keys': ['编号']})
        self.assertEqual(result['summary']['output_rows'], 4)
        self.assertEqual(result['summary']['duplicate_rows'], 2)
        book = load_workbook(io.BytesIO(result['files'][0]['content']))
        try:
            self.assertEqual(book['合并数据']['D2'].value, '001')
            self.assertEqual(book['合并数据']['E5'].value, '=1+1')
            self.assertEqual(book['合并数据']['E5'].data_type, 's')
            self.assertEqual(book['重复关键列'].max_row, 3)
        finally:
            book.close()

    def test_compare_requires_unique_keys_and_reports_real_differences(self):
        before = self.upload('原表.xlsx', workbook([['编号', '状态'], ['001', '未开始'], ['002', '进行中']]))
        after = self.upload('新表.xlsx', workbook([['编号', '状态'], ['001', '已完成'], ['003', '进行中']]))
        result = analyze_tables(self.inputs([before, after]), 'compare', keys=['编号'])
        self.assertEqual({k: result['summary'][k] for k in ('added', 'removed', 'changed', 'unchanged')},
                         {'added': 1, 'removed': 1, 'changed': 1, 'unchanged': 0})
        duplicate = self.upload('重复.xlsx', workbook([['编号', '状态'], ['001', '未开始'], ['001', '进行中']]))
        with self.assertRaisesRegex(AssistantError, '重复'):
            analyze_tables(self.inputs([before, duplicate]), 'compare', keys=['编号'])

    def test_pivot_keeps_decimal_precision_and_missing_groups_unknown(self):
        from openpyxl import load_workbook
        identity = self.upload('汇总.csv', '楼栋,数量\nD,0.1\nD,0.2\nE,\n'.encode())
        result = analyze_tables(self.inputs([identity]), 'pivot', group_by=['楼栋'], value_column='数量')
        book = load_workbook(io.BytesIO(result['files'][0]['content']))
        try:
            self.assertEqual(book.active['B2'].value, '0.3')
            self.assertIsNone(book.active['B3'].value)
            self.assertEqual(book.active['D3'].value, 1)
        finally:
            book.close()

    def test_multisheet_headers_and_missing_formula_caches_require_clarification(self):
        identity = self.upload('多表.xlsx', workbook([['编号'], ['001']], second_sheet=True))
        found = analyze_tables(self.inputs([identity]), 'inspect')
        self.assertTrue(found['summary']['tables'][0]['requires_sheet'])
        with self.assertRaisesRegex(AssistantError, '多个工作表'):
            analyze_tables(self.inputs([identity]), 'merge')
        identity = self.upload('缺缓存.xlsx', workbook([['编号', '合计'], ['001', '=1+1']]))
        with self.assertRaisesRegex(AssistantError, '计算缓存'):
            analyze_tables(self.inputs([identity]), 'merge')
        identity = self.upload('空表头.csv', '编号,\n001,2\n'.encode())
        with self.assertRaisesRegex(AssistantError, '表头'):
            analyze_tables(self.inputs([identity]), 'merge')

    def test_formatted_empty_columns_do_not_count_as_data_columns(self):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill
        book = Workbook()
        book.active.append(['编号', '数量'])
        book.active.append(['001', 2])
        book.active.cell(2, 200).fill = PatternFill('solid', fgColor='FF0000')
        output = io.BytesIO()
        book.save(output)
        book.close()
        identity = self.upload('空白格式.xlsx', output.getvalue())
        result = analyze_tables(self.inputs([identity]), 'inspect')
        self.assertEqual(result['summary']['tables'][0]['headers'], ['编号', '数量'])
        self.assertEqual(result['summary']['tables'][0]['row_count'], 1)

    def test_no_arbitrary_paths_or_cross_account_files(self):
        identity = self.upload('资料.pdf', pdf_bytes(['A']))
        with self.assertRaises(AssistantError):
            prepare_inputs(self.files, ACTOR, [{'file_id': identity, 'path': 'C:/private'}], 'pdf')
        with self.assertRaises(AssistantError):
            prepare_inputs(self.files, {**ACTOR, 'id': 'other', 'is_admin': True}, [{'file_id': identity}], 'pdf')
        with self.assertRaises(AssistantError):
            prepare_inputs(self.files, {**ACTOR, 'scopes': ['A']}, [{'file_id': identity}], 'pdf')

    def test_pdf_merge_extract_rotate_split_preserve_original(self):
        from pypdf import PdfReader
        original = pdf_bytes(['Page A', 'Page B', 'Page C'])
        identity = self.upload('资料.pdf', original)
        inputs = self.inputs([identity], 'pdf', pages=[3, 1])
        result = run_local('pdf', {'inputs': inputs, 'operation': 'extract', 'rotation': 90})
        reader = PdfReader(io.BytesIO(result['files'][0]['content']))
        self.assertEqual(len(reader.pages), 2)
        self.assertIn('Page C', reader.pages[0].extract_text())
        self.assertEqual(reader.pages[0].rotation, 90)
        result = organize_pdf(self.inputs([identity], 'pdf'), 'split', split_every=2)
        self.assertEqual([len(PdfReader(io.BytesIO(file['content'])).pages) for file in result['files']], [2, 1])
        second = self.upload('另一份.pdf', pdf_bytes(['Page D']))
        result = organize_pdf(self.inputs([identity, second], 'pdf'), 'merge')
        self.assertEqual(result['summary']['pages'], 4)
        self.assertEqual(Path(self.files.get(ACTOR, identity)['path']).read_bytes(), original)
        with self.assertRaisesRegex(AssistantError, '页码'):
            organize_pdf(self.inputs([identity], 'pdf', pages=[0]), 'extract')
        with self.assertRaisesRegex(AssistantError, '额度'):
            organize_pdf(self.inputs([identity], 'pdf'), 'split', max_outputs=2)

    def test_pdf_actions_are_removed_and_encrypted_files_rejected(self):
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, TextStringObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=100, height=100)
        page[NameObject('/AA')] = DictionaryObject({NameObject('/O'): DictionaryObject({NameObject('/S'): NameObject('/JavaScript'), NameObject('/JS'): TextStringObject('app.alert("fixture")')})})
        writer.add_js('app.alert("fixture")')
        output = io.BytesIO()
        writer.write(output)
        identity = self.upload('交互.pdf', output.getvalue())
        result = organize_pdf(self.inputs([identity], 'pdf'), 'extract')
        reader = PdfReader(io.BytesIO(result['files'][0]['content']))
        self.assertNotIn('/AA', reader.pages[0])
        self.assertNotIn('/OpenAction', reader.root_object)
        self.assertNotIn('/JavaScript', reader.root_object.get('/Names', {}))
        writer.encrypt('fixture')
        output = io.BytesIO()
        writer.write(output)
        writer.close()
        identity = self.upload('加密.pdf', output.getvalue())
        with self.assertRaisesRegex(AssistantError, '未加密'):
            organize_pdf(self.inputs([identity], 'pdf'), 'extract')

    def test_flowchart_is_structured_and_rendered_offline(self):
        from PIL import Image
        nodes = [{'id': 'start', 'label': '接收现场信息', 'kind': 'terminal'}, {'id': 'check', 'label': '资料是否完整？', 'kind': 'decision'},
                 {'id': 'fill', 'label': '补充缺少的资料'}, {'id': 'review', 'label': '负责人核对'}, {'id': 'done', 'label': '归档', 'kind': 'terminal'}]
        edges = [{'from': 'start', 'to': 'check'}, {'from': 'check', 'to': 'fill', 'label': '否'},
                 {'from': 'fill', 'to': 'check'}, {'from': 'check', 'to': 'review', 'label': '是'}, {'from': 'review', 'to': 'done'}]
        result = run_local('flowchart', {'title': '资料核对示例（测试流程）', 'nodes': nodes, 'edges': edges})
        content = result['files'][0]['content']
        with Image.open(io.BytesIO(content)) as image:
            self.assertGreater(image.width, 400)
            self.assertGreater(image.height, 300)
        if os.environ.get('LIGHTHOUSE_TEST_ARTIFACTS'):
            target = Path(os.environ['LIGHTHOUSE_TEST_ARTIFACTS'])
            target.mkdir(parents=True, exist_ok=True)
            (target / 'flowchart.png').write_bytes(content)
        escaped = flowchart_source('测试', [{'id': 'x', 'label': '"]\nclick x "https://example.com"'}], [])
        self.assertEqual(len(escaped.splitlines()), 2)
        self.assertNotIn('<script', escaped)
        with self.assertRaises(AssistantError):
            flowchart_source('测试', nodes, [{'from': 'missing', 'to': 'done'}])


class OfficeConversationTests(unittest.IsolatedAsyncioTestCase):
    async def test_pdf_operation_uses_private_files_and_returns_an_actual_attachment(self):
        from bin import test_lighthouse_stream as fixtures
        from pydantic_ai.models.function import FunctionModel, DeltaToolCall
        from openclaw_service.assistant.lighthouse_model import LighthouseModel
        from openclaw_service.assistant import lighthouse_document_tools
        await fixtures.StreamTests.asyncSetUp(self)
        file = self.portal.files.upload(self.actor, '资料.pdf', pdf_bytes(['A']), extract=False)
        calls, delivered = [], []

        async def stream(messages, info):
            self.assertTrue({'analyze_spreadsheets', 'process_pdf', 'create_flowchart'} <= {tool.name for tool in info.function_tools})
            calls.append(1)
            if len(calls) == 1:
                yield {0: DeltaToolCall(name='process_pdf', json_args=json.dumps({'operation': 'extract', 'inputs': [{'file_id': file['id'], 'pages': [1]}]}))}
            else:
                yield '提取的PDF已生成，原件未改动。'

        @asynccontextmanager
        async def model(*_):
            yield FunctionModel(stream_function=stream)

        async def authorize():
            return self.actor

        async def emit(event, data):
            if event == 'artifact':
                delivered.append(data['file'])

        def worker(kind, payload):
            self.assertEqual(kind, 'pdf')
            return organize_pdf(**payload)

        engine = LighthouseModel(self.portal, model_factory=model)
        try:
            with patch.object(lighthouse_document_tools, 'run_local', side_effect=worker):
                await engine.answer(self.actor, {'question': '请提取我上传PDF的第一页', 'file_ids': [file['id']], '_profile': self.model.profile()}, [], self.request, emit, authorize, {})
            self.assertEqual(len(delivered), 1)
            self.assertTrue(Path(self.portal.files.get(self.actor, delivered[0]['id'])['path']).read_bytes().startswith(b'%PDF'))
            with self.assertRaises(AssistantError):
                self.portal.files.get({**self.actor, 'id': 'other'}, delivered[0]['id'])
            self.assertEqual(self.reads, [])
        finally:
            await engine.close()
            await self.runtime.close()


if __name__ == '__main__':
    unittest.main()
