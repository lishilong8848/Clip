"""Isolated boundary checks for the lighthouse agent and its supporting modules.

These tests only exercise fake in-process APIs and an in-memory store.  They never
make real cloud/provider requests, never start a server, and never mutate
production data, config, or dist assets.

Coverage requested for verification:
  1. Real POST /api/notice-attachments sends raw image bytes plus MIME/content-type
     and query `file_name` (not JSON, not multipart).
  2. Cabinet `/exports` job polls GET /api/cabinet-power/jobs/{job_id} with
     status=succeeded, not the generic /api/jobs endpoint.
  3. Submitted task resume never repeats the already-completed first write.
  4. Image recognition still in progress (batch status=pending but an image is
     recognizing) must not be claimed as completed.
  5. JSON file with 500 rows preserves all rows while redacting sensitive keys.
  6. Images after five text attachments still reach vision.
  7. Private Qt / credential API routes are excluded from the catalogue.
  8. A pending persisted plan recovered on the first GET after service restart is
     clearly marked failed (no replay, no being stuck).
"""
import copy
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

sys.path.insert(0, str(Path(__file__).resolve().parent))

from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse
from PIL import Image

from lan_bitable_template_portal.lighthouse_agent import PLAN_NAMESPACE, PortalAgent
from lan_bitable_template_portal.lighthouse_ai import LighthouseAssistant
from lan_bitable_template_portal.lighthouse_api import PortalAPICatalog
from lan_bitable_template_portal.lighthouse_files import LighthouseFiles

ACTOR = {"id": "boundary-fixture-a", "scopes": ["A"], "is_admin": False}


class Store:
    def __init__(self, path):
        self.db_path, self.docs = path, {}

    def get_document(self, namespace, key):
        return copy.deepcopy(self.docs.get((namespace, key)))

    def put_document(self, namespace, key, value):
        self.docs[namespace, key] = copy.deepcopy(value)


def _png_bytes():
    output = io.BytesIO()
    Image.new("RGB", (12, 12), "white").save(output, "PNG")
    return output.getvalue()


async def _gather(agent):
    import asyncio
    if agent.tasks:
        await asyncio.gather(*tuple(agent.tasks))


def _make_assistant(store, *, model=None, search=True):
    if model is None:
        model = Mock()
        model.settings.return_value = {
            "configured": True,
            "enabled": True,
            "active_model_id": "default",
            "models": [{"id": "default", "name": "测试模型", "model": "fixture-model", "configured": True}],
        }
        model.profile.return_value = {"id": "default", "name": "测试模型", "model": "fixture-model"}
    search_mock = Mock()
    if search:
        search_mock = Mock(side_effect=AssertionError("boundary test must use backend APIs, not DOM scraping"))
    return LighthouseAssistant(store, search_mock, model=model)


def _request():
    return Request({
        "type": "http",
        "scheme": "http",
        "server": ("testserver", 80),
        "client": ("127.0.0.1", 4567),
        "path": "/api/assistant/agent",
        "root_path": "",
        "query_string": b"",
        "headers": [(b"cookie", b"fixture=a"), (b"origin", b"http://testserver")],
    })


class NoticeAttachmentsBytesBoundary(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.store = Store(Path(self.tmp.name) / "state.sqlite3")
        self.assistant = _make_assistant(self.store)
        self.files = LighthouseFiles(self.store)
        self.captured = None

        app = FastAPI()

        @app.post("/api/notice-attachments")
        async def notice_attachments(request: Request):
            raw = await request.body()
            self.captured = {
                "raw": raw,
                "content_type": request.headers.get("content-type", ""),
                "file_name": request.query_params.get("file_name", ""),
            }
            return {"ok": True, "data": {"upload_id": "up-1", "file_token": "private-file-token"}}

        self.catalog = PortalAPICatalog(app)
        self.agent = PortalAgent(self.assistant, self.catalog, self.files)

    async def test_notice_attachments_raw_image_bytes_and_filename_query(self):
        original = _png_bytes()
        file = self.files.upload(ACTOR, "确认截图.png", original)
        file_provider = lambda fid: self.files.get(ACTOR, fid)

        result = await self.catalog.invoke({
            "api_id": "POST /api/notice-attachments",
            "files": {"file": [file["id"]]},
        }, _request(), file_provider=file_provider)

        self.assertTrue(result["ok"], result)
        self.assertIsNotNone(self.captured)
        # Original image bytes are sent verbatim (not JSON, not multipart-wrapped).
        self.assertEqual(self.captured["raw"], original)
        # Raw binary upload reaches the route with the image MIME in content-type.
        self.assertTrue(self.captured["content_type"].startswith("image/"))
        # The original file name travels as a query parameter, not in the body.
        self.assertEqual(self.captured["file_name"], "确认截图.png")
        self.assertNotIn("multipart/form-data", self.captured["content_type"])


class CabinetExportsJobBoundary(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.store = Store(Path(self.tmp.name) / "state.sqlite3")
        self.assistant = _make_assistant(self.store)
        self.files = LighthouseFiles(self.store)
        self.cabinet_job_calls = []
        self.generic_job_calls = []

        app = FastAPI()

        @app.get("/api/cabinet-power/jobs/{job_id}")
        async def cabinet_job(job_id: str, request: Request):
            self.cabinet_job_calls.append(job_id)
            return {"ok": True, "data": {"status": "succeeded"}}

        @app.get("/api/jobs/{job_id}")
        async def generic_job(job_id: str, request: Request):
            self.generic_job_calls.append(job_id)
            # Deliberately fail so any accidental generic-job polling is visible.
            return {"ok": True, "data": {"phase": "failed"}}

        self.catalog = PortalAPICatalog(app)
        self.agent = PortalAgent(self.assistant, self.catalog, self.files)

    async def test_cabinet_exports_job_polls_cabinet_jobs_succeeded_not_generic(self):
        operation = {"api_id": "POST /api/cabinet-power/export-batches"}
        result = {"_raw": {"job_id": "cab-export-1"}, "data": {"job_id": "cab-export-1"}, "ok": True}
        spec = self.agent._task_spec(operation, result)

        self.assertIsNotNone(spec)
        self.assertEqual(spec["kind"], "cabinet_job")
        self.assertEqual(spec["operation"]["api_id"], "GET /api/cabinet-power/jobs/{job_id}")
        self.assertEqual(spec["operation"]["path_params"], {"job_id": "cab-export-1"})

        poll, done = await self.agent._task_result(ACTOR, spec, _request())
        self.assertTrue(done, poll)
        # The cabinet endpoint was contacted and the generic /api/jobs was not.
        self.assertEqual(self.cabinet_job_calls, ["cab-export-1"])
        self.assertEqual(self.generic_job_calls, [])


class SubmittedResumeBoundary(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.store = Store(Path(self.tmp.name) / "state.sqlite3")
        self.assistant = _make_assistant(self.store)
        self.files = LighthouseFiles(self.store)
        self.export_writes, self.finalize_writes = [], []

        app = FastAPI()

        @app.post("/api/cabinet-power/export-batches")
        async def submit_export(request: Request):
            self.export_writes.append("export")
            return {"ok": True, "data": {"job_id": "cab-resume-job"}}

        @app.get("/api/cabinet-power/jobs/{job_id}")
        async def cabinet_job(job_id: str, request: Request):
            return {"ok": True, "data": {"status": "succeeded"}}

        @app.post("/api/finalize")
        async def finalize(request: Request):
            self.finalize_writes.append("finalize")
            return {"ok": True, "data": {"saved": True}}

        self.catalog = PortalAPICatalog(app)
        self.agent = PortalAgent(self.assistant, self.catalog, self.files)

    async def test_submitted_resume_never_repeats_write(self):
        decision = {
            "operations": [
                {"api_id": "POST /api/cabinet-power/export-batches"},
                {"api_id": "POST /api/finalize"},
            ],
        }
        turn_id = "resume_boundary_00001"
        plan = self.agent.prepare(ACTOR, decision, turn_id, [])
        # Simulate: first write already ran and returned an async job still
        # processing at the time the original agent task exited.
        plan["status"] = "submitted"
        plan["error"] = "后台任务仍在处理，可查询原任务结果；不会重复提交。"
        plan["results"] = [{
            "ok": True,
            "status": 200,
            "data": {"job_id": "cab-resume-job"},
            "_raw": {"job_id": "cab-resume-job"},
            "api_id": "POST /api/cabinet-power/export-batches",
            "truncated": False,
        }]
        self.store.put_document(PLAN_NAMESPACE, plan["id"], plan)

        refreshed = await self.agent.refresh(ACTOR, plan, self.request_or_placeholder())
        self.assertIsNotNone(refreshed)
        await _gather(self.agent)

        # The already-completed submit write is never repeated; only the
        # remaining step is retried after the resumed job is seen as succeeded.
        self.assertEqual(self.export_writes, [])
        self.assertEqual(self.finalize_writes, ["finalize"])
        finished = self.agent.get_plan(ACTOR, plan["id"])
        self.assertEqual(finished["status"], "completed")
        self.assertEqual(len(finished["results"]), 2)

    def request_or_placeholder(self):
        return _request()


class RecognitionInProgressBoundary(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.store = Store(Path(self.tmp.name) / "state.sqlite3")
        self.assistant = _make_assistant(self.store)
        self.files = LighthouseFiles(self.store)

        app = FastAPI()

        @app.get("/api/cabinet-power/batches/{batch_id}/status")
        async def batch_status(batch_id: str, request: Request):
            # Batch-level status is *pending*, but one image is still recognizing.
            return {
                "ok": True,
                "data": {
                    "status": "pending",
                    "batch_id": batch_id,
                    "images": [
                        {"status": "recognizing", "name": "a.png", "deleted_at": None},
                        {"status": "recognized", "name": "b.png", "deleted_at": None},
                    ],
                },
            }

        self.catalog = PortalAPICatalog(app)
        self.agent = PortalAgent(self.assistant, self.catalog, self.files)

    async def test_recognition_still_in_progress_despite_batch_pending_is_not_done(self):
        operation = {"api_id": "POST /api/cabinet-power/batches"}
        submit_result = {
            "_raw": {"batch_id": "batch-pending", "status": "recognizing",
                     "images": [{"status": "recognizing"}]},
            "data": {"batch_id": "batch-pending", "status": "recognizing",
                     "images": [{"status": "recognizing"}]},
            "ok": True,
        }
        spec = self.agent._task_spec(operation, submit_result)
        self.assertEqual(spec["kind"], "recognition")

        poll, done = await self.agent._task_result(ACTOR, spec, _request())
        # The in-progress recognition must NOT be reported as completed.
        self.assertFalse(done, poll)
        self.assertTrue(poll["ok"])


class JsonRowsPreservedBoundary(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.store = Store(Path(self.tmp.name) / "state.sqlite3")
        self.files = LighthouseFiles(self.store)

    def test_json_with_500_rows_preserves_all_rows_and_redacts_sensitive_keys(self):
        rows = [
            {
                "scope": "A",
                "name": f"item-{index}",
                "token": f"secret-token-{index}",
                "phone": f"1{index % 1000000000:09d}",
            }
            for index in range(500)
        ]
        # `document_data` is supposed to keep every row while dropping the
        # sensitive token/phone keys. Extraction re-serialises the doc as a
        # single-line JSON string, and the per-line safe_text cap (6000 chars)
        # truncates it, so asserting all 500 rows survive exposes any row loss.
        content = json.dumps(rows, ensure_ascii=False, indent=2).encode("utf-8")
        file = self.files.upload(ACTOR, "rows.json", content)

        stored = self.files.get(ACTOR, file["id"])
        text = stored.get("text", "")
        self.assertTrue(text, "JSON text extraction returned no content")

        # Every one of the 500 rows survives extraction.
        for index in range(500):
            self.assertIn(f'"name": "item-{index}"', text)

        # Sensitive values are never preserved.
        for index in (0, 123, 499):
            self.assertNotIn(f"secret-token-{index}", text)
            self.assertNotIn(f"1{index % 1000000000:09d}", text)
        # Sensitive key names are dropped entirely.
        self.assertNotIn("token", text)
        self.assertNotIn("phone", text)

    async def test_incomplete_long_text_cannot_be_used_as_business_payload(self):
        for content in ("a" * 6001, "text\n" * 20001):
            with self.subTest(size=len(content)):
                file = self.files.upload(ACTOR, "long.txt", content.encode())
                stored = self.files.get(ACTOR, file["id"])
                self.assertEqual(stored["text"], "")
                self.assertIn("完整读取上限", stored["error"])
                self.assertTrue(Path(stored["path"]).is_file())
                catalog = Mock()
                agent = PortalAgent(_make_assistant(self.store), catalog, self.files)
                from lan_bitable_template_portal.lighthouse_ai import AssistantError
                with self.assertRaises(AssistantError):
                    await agent._invoke(ACTOR, {"api_id": "POST /api/import", "body": {"text": {"$file_text": file["id"]}}}, _request())
                catalog.invoke.assert_not_called()

    def test_excel_nonempty_column_beyond_read_limit_is_not_silently_lost(self):
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.cell(1, 1, "first")
        workbook.active.cell(1, 81, "must-not-be-lost")
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        file = self.files.upload(ACTOR, "wide.xlsx", output.getvalue())
        stored = self.files.get(ACTOR, file["id"])
        self.assertEqual(stored["text"], "")
        self.assertIn("读取上限", stored["error"])
        self.assertTrue(Path(stored["path"]).is_file())


class VisionAfterTextBoundary(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.store = Store(Path(self.tmp.name) / "state.sqlite3")
        self.files = LighthouseFiles(self.store)

    @patch("lan_bitable_template_portal.lighthouse_files.recognize_text", return_value="text attached")
    def test_images_after_five_text_attachments_still_reach_vision(self, _recognize):
        text_ids = []
        for index in range(5):
            text_ids.append(self.files.upload(ACTOR, f"note-{index}.txt", f"内容 {index}\n".encode())["id"])
        image = self.files.upload(ACTOR, "eye.png", _png_bytes())

        parts = self.files.image_parts(ACTOR, [*text_ids, image["id"]])
        # Text attachments consume none of the five-image slot; the PNG still
        # reaches the vision parts.
        self.assertEqual(len(parts), 1)
        self.assertEqual(parts[0]["type"], "image_url")
        self.assertIn("data:image/jpeg;base64,", parts[0]["image_url"]["url"])


class PrivateApiExcludedBoundary(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.store = Store(Path(self.tmp.name) / "state.sqlite3")
        self.files = LighthouseFiles(self.store)

        app = FastAPI()

        @app.get("/api/qt/backend/status")
        async def qt_status(request: Request):
            return {"ok": True}

        @app.get("/api/admin/credential/list")
        async def credential_list(request: Request):
            return {"ok": True}

        @app.post("/api/admin/credential/save")
        async def credential_save(request: Request):
            return {"ok": True}

        @app.get("/api/auth/login")
        async def auth_login(request: Request):
            return {"ok": True}

        @app.get("/api/backend/stats")
        async def backend_stats(request: Request):
            return {"ok": True}

        @app.get("/api/health")
        async def health(request: Request):
            return {"ok": True}

        @app.get("/api/notice/list")
        async def notice_list(request: Request):
            return {"ok": True, "data": []}

        self.catalog = PortalAPICatalog(app)

    def test_qt_and_credential_api_paths_excluded(self):
        result = self.catalog.discover(page_size=200)
        ids = {item["id"] for item in result["items"]}

        self.assertNotIn("GET /api/qt/backend/status", ids)
        self.assertNotIn("GET /api/admin/credential/list", ids)
        self.assertNotIn("POST /api/admin/credential/save", ids)
        self.assertNotIn("GET /api/auth/login", ids)
        self.assertNotIn("GET /api/backend/stats", ids)
        self.assertNotIn("GET /api/health", ids)
        # Normal business route still present, so the exclusion is not hiding everything.
        self.assertIn("GET /api/notice/list", ids)


class PendingPlanRecoveryBoundary(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.store = Store(Path(self.tmp.name) / "state.sqlite3")

    def test_pending_persisted_plan_fails_on_first_get_after_restart(self):
        model = Mock()
        model.settings.return_value = {
            "configured": True,
            "enabled": True,
            "active_model_id": "default",
            "models": [{"id": "default", "name": "测试模型", "model": "fixture-model", "configured": True}],
        }
        model.profile.return_value = {"id": "default", "name": "测试模型", "model": "fixture-model"}
        search = Mock(side_effect=AssertionError("no DOM scraping in boundary tests"))

        # Persist a conversation that contains a pending plan (simulating an
        # interrupted run whose answer never completed).
        state = {
            "id": "conv-boundary-1",
            "turns": [
                {
                    "operation_id": "pending_restart_0001",
                    "question": "当前有哪些未完成工作",
                    "status": "pending",
                    "scopes": ["A"],
                    "at": 1700000000,
                }
            ],
        }
        self.store.put_document("lighthouse_ai", "conversation:" + ACTOR["id"], state)

        # First GET after "service restart": a fresh assistant instance with no
        # active work must clearly fail the stale pending plan instead of leaving
        # it stuck and never replay the write.
        restarted = LighthouseAssistant(self.store, search, model=model)
        result = restarted.conversation(ACTOR)
        turns = result["turns"]
        self.assertEqual(len(turns), 1)
        turn = turns[0]
        self.assertEqual(turn["status"], "failed")
        self.assertIn("中断", turn["error"])
        # No answer generation / replay was triggered.
        model.complete.assert_not_called()

        persisted = self.store.get_document("lighthouse_ai", "conversation:" + ACTOR["id"])
        self.assertEqual(persisted["turns"][0]["status"], "failed")

    async def test_previous_generation_plan_is_invalid_after_conversation_clear(self):
        assistant = _make_assistant(self.store)
        app = FastAPI()
        @app.delete("/api/drills/{drill_id}")
        async def delete(drill_id: str):
            self.fail("a cleared plan must never write")
        agent = PortalAgent(assistant, PortalAPICatalog(app), LighthouseFiles(self.store))
        turn_id = "legacy_plan_000000001"
        plan = agent.prepare(ACTOR, {"operations": [{"api_id": "DELETE /api/drills/{drill_id}", "path_params": {"drill_id": "fixture"}}]}, turn_id, [])
        plan.pop("conversation_id")
        self.store.put_document(PLAN_NAMESPACE, plan["id"], plan)
        state = assistant._state(ACTOR)
        state["turns"].append({"operation_id": turn_id, "question": "删除演练", "status": "completed", "scopes": ["A"], "plan": agent.public_plan(plan)})
        self.store.put_document("lighthouse_ai", assistant._key(ACTOR), state)
        self.assertEqual(agent.get_plan(ACTOR, plan["id"])["id"], plan["id"])
        assistant.clear(ACTOR)
        from lan_bitable_template_portal.lighthouse_ai import AssistantError
        with self.assertRaises(AssistantError) as error:
            await agent.confirm(ACTOR, plan["id"], {"version": plan["version"], "stage": "review"}, _request())
        self.assertEqual(error.exception.status, 409)
        self.assertFalse(agent.tasks)

    def test_gateway_marks_string_and_depth_caps_as_incomplete(self):
        catalog = PortalAPICatalog(FastAPI())
        result = catalog._consume_json(200, {"ok": True, "data": {"content": "a" * 6001}}, "GET /api/read")
        self.assertTrue(result["truncated"])
        self.assertEqual(len(result["data"]["content"]), 6000)
        nested = {"value": "deep"}
        for _ in range(13):
            nested = {"child": nested}
        result = catalog._consume_json(200, {"ok": True, "data": nested}, "GET /api/read")
        self.assertTrue(result["truncated"])

    async def test_completed_plan_recovers_failed_chat_save_without_replaying_write(self):
        assistant = _make_assistant(self.store)
        catalog = Mock()
        catalog.get.return_value = {"read_only": False, "risk": "high", "name": "测试操作", "schema": {"body": {}}}
        catalog.validate_operation.side_effect = lambda operation: (operation, [])
        agent = PortalAgent(assistant, catalog, LighthouseFiles(self.store))
        turn_id = "completed_plan_000001"
        plan = agent.prepare(ACTOR, {"operations": [{"api_id": "POST /api/write"}]}, turn_id, [])
        plan["status"] = "running"
        state = assistant._state(ACTOR)
        state["turns"] = [{"operation_id": turn_id, "question": "测试操作", "status": "completed", "scopes": ["A"], "plan": agent.public_plan(plan)}]
        self.store.put_document("lighthouse_ai", assistant._key(ACTOR), state)
        plan["status"] = "completed"
        original_put = self.store.put_document
        def fail_chat(namespace, key, value):
            if namespace == "lighthouse_ai":
                raise RuntimeError("chat persistence interrupted")
            original_put(namespace, key, value)
        with patch.object(self.store, "put_document", side_effect=fail_chat), self.assertRaises(RuntimeError):
            agent._save_plan(ACTOR, plan)
        self.assertEqual(self.store.get_document(PLAN_NAMESPACE, plan["id"])["status"], "completed")
        self.assertEqual(assistant._state(ACTOR)["turns"][0]["plan"]["status"], "running")
        await agent.refresh(ACTOR, agent.get_plan(ACTOR, plan["id"]), _request())
        recovered = assistant.conversation(ACTOR)["turns"][0]
        self.assertEqual(recovered["plan"]["status"], "completed")
        self.assertEqual(recovered["answer"], "操作已完成。")
        catalog.invoke.assert_not_called()


if __name__ == "__main__":
    unittest.main()
