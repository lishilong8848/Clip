"""Offline ACL, version, contextual evidence and image-routing acceptance checks."""
import copy
import hashlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parent))

from PIL import Image
from bin.test_lighthouse_knowledge import FakeEmbedder, FakeMemoryIndex, actor
from bin.test_lighthouse_knowledge_extract import make_pdf_text
from openclaw_service.assistant.lighthouse_ai import AssistantError
from openclaw_service.assistant.lighthouse_knowledge import KnowledgeBase
from openclaw_service.assistant import lighthouse_knowledge_evidence as evidence

TEXT = {'endpoint': 'https://models.example.com/v1/chat/completions',
        'model': 'wanwu/glm-auto', 'image_input': True}
VISION = {**TEXT, 'id': 'vision', 'model': 'wanwu/MiniMax-M2.5', 'image_input': True}


def png(color='red'):
    buffer = io.BytesIO()
    with Image.new('RGB', (48, 36), color) as image:
        image.save(buffer, 'PNG')
    return buffer.getvalue()


class EvidenceTests(unittest.TestCase):
    def setUp(self):
        self.enterContext(patch('socket.socket.connect', side_effect=AssertionError('No network')))
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.allowed, self.source_version = True, 1
        self.kb = KnowledgeBase(self.temp.name, embedder=FakeEmbedder(),
            memory_index_factory=FakeMemoryIndex, start_worker=False,
            source_acl=lambda user, source: self.allowed and set(source['scopes']) <= set(user.get('scopes', [])),
            source_current=lambda source: source['revision'] == self.source_version)
        self.addCleanup(self.kb.close)
        self.actor = {**actor(admin=True), 'scopes': ['E'], 'modules': ['knowledge']}
        self.model = Mock()
        self.model.settings.return_value = {'enabled': True, 'models': [], 'active_model_id': 'text'}
        self.model.complete.return_value = json.dumps({'text': 'The red box connects to the left line.'})
        self.ocr = self.enterContext(patch.object(evidence, 'recognize_text', return_value='OCR text only'))

    def upload(self, name='policy.txt', content=b'needle source text', *, sections=None, managed=False):
        self.kb.extractor = (lambda *_: sections) if sections is not None else None
        if managed:
            doc = self.kb.upsert_source('fixture/' + name, {'scopes': ['E'], 'revision': 1, 'kind': 'attachment'}, name, content)
        else:
            doc = self.kb.upload(self.actor, name, content)
        self.assertTrue(self.kb.process_one())
        with self.kb.connect(readonly=True) as db:
            current = dict(db.execute('SELECT * FROM documents WHERE id=?', (doc['id'],)).fetchone())
        self.assertEqual(current['status'], 'ready', current['error'])
        return {'document_id': doc['id'], 'version': current['active_version'], 'ordinal': 0, 'name': name}

    def image(self, *, managed=False, name='diagram.png'):
        return self.upload(name, png(), managed=managed,
                           sections=[{'text': 'needle diagram text', 'location': name + ':图片OCR'}])

    def read(self, items, profile=VISION):
        return self.kb.answer_evidence(self.actor, items, profile=profile, model=self.model)

    def test_context_uses_authorized_originals_not_supplied_text_and_is_read_only(self):
        parts = [{'text': 'needle paragraph %d' % i, 'location': 'policy.txt:段落%d' % (i + 1)} for i in range(30)]
        item = self.upload(sections=parts)
        item.update(ordinal=10, text='forged text', location='fake page')
        with self.kb.connect(readonly=True) as db:
            before = list(db.iterdump())
        found = self.read([item], TEXT)
        current = found['items'][0]
        self.assertIn('needle paragraph 2', current['text'])
        self.assertIn('needle paragraph 18', current['text'])
        self.assertNotIn('forged', current['text'])
        self.assertEqual(current['location'], 'policy.txt:段落11（片段1）')
        self.assertEqual(current['coverage']['indexed_document_chunks'], 30)
        self.assertFalse(current['coverage']['complete'])
        self.assertEqual(found['warning'], '')
        self.assertEqual(current['source_hash'], hashlib.sha256(b'needle source text').hexdigest())
        self.model.settings.assert_not_called()
        self.model.complete.assert_not_called()
        with self.kb.connect(readonly=True) as db:
            self.assertEqual(list(db.iterdump()), before)

    def test_table_header_and_later_rows_are_preserved(self):
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.title = 'Equipment'
        workbook.active.append(['Device', 'Limit'])
        for index in range(100):
            workbook.active.append(['needle-%d' % index, index])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        item = self.upload('limits.xlsx', output.getvalue())
        item['ordinal'] = 1
        found = self.read([item], TEXT)['items'][0]
        self.assertIn('Device\tLimit', found['text'])
        self.assertIn('needle-99', found['text'])
        self.assertTrue(all(part['table_header'] == 'Device、Limit' for part in found['context']))
        self.assertIn('工作表「Equipment」', found['text'])

    def test_model_summary_expands_to_saved_original_text(self):
        self.kb.source_summarizer = lambda *_: 'Model summary needle'
        original = 'needle original operating limits and warning'
        item = self.upload('policy.txt', original.encode(), managed=True)
        found = self.read([item], TEXT)['items'][0]
        self.assertIn(original, found['text'])
        self.assertNotIn('Model summary', found['text'])
        self.assertEqual(found['context'][0]['kind'], 'original_text')

    def test_context_budgets_reserve_each_hit_and_keep_coverage(self):
        parts = [{'text': ('needle paragraph %d ' % index) * 240,
                  'location': 'policy.txt:段落%d' % index} for index in range(18)]
        item = self.upload(sections=parts)
        items = [{**item, 'ordinal': index * 5} for index in range(12)]
        result = self.read(items, TEXT)
        self.assertEqual(len(result['items']), 12)
        self.assertLessEqual(sum(len(i['text']) for i in result['items']), evidence.MAX_TEXT)
        self.assertTrue(all(i['text'] for i in result['items']))
        self.assertTrue(any('text_budget' in i['coverage']['missing'] for i in result['items']))
        self.assertIn('expansion_budget', result['items'][-1]['coverage']['missing'])

    def test_scoped_acl_denies_before_open_or_model(self):
        item = self.image(managed=True)
        self.actor['scopes'] = ['A']
        with patch.object(self.kb, 'file', wraps=self.kb.file) as file:
            with self.assertRaises(AssistantError):
                self.read([item])
            file.assert_not_called()
        self.model.complete.assert_not_called()

    def test_revocation_during_picture_model_discards_result(self):
        item = self.image(managed=True)
        def revoke(*args, **kwargs):
            self.allowed = False
            return '{"text":"Do not leak this observation"}'
        self.model.complete.side_effect = revoke
        with self.assertRaises(AssistantError):
            self.read([item])
        self.model.complete.assert_called_once()

    def test_version_change_during_model_discards_even_for_owner(self):
        item = self.image()
        def replace(*args, **kwargs):
            with self.kb.connect() as db:
                db.execute('UPDATE documents SET active_version=2 WHERE id=?', (item['document_id'],))
            return '{"text":"Old picture"}'
        self.model.complete.side_effect = replace
        with self.assertRaises(AssistantError):
            self.read([item])

    def test_source_current_fails_closed_before_and_after_model(self):
        item = self.image(managed=True)
        self.source_version = 2
        with self.assertRaises(AssistantError):
            self.read([item])
        self.model.complete.assert_not_called()
        self.source_version = 1
        def replace(*args, **kwargs):
            self.source_version = 2
            return '{"text":"Old source"}'
        self.model.complete.side_effect = replace
        with self.assertRaises(AssistantError):
            self.read([item])

    def test_file_hash_change_during_model_is_rejected(self):
        item = self.image()
        path, _ = self.kb.file(self.actor, item['document_id'], item['version'])
        def replace(*args, **kwargs):
            path.write_bytes(png('blue'))
            return '{"text":"Old content"}'
        self.model.complete.side_effect = replace
        with self.assertRaises(AssistantError):
            self.read([item])

    def test_model_allowlist_revocation_discards_picture_result(self):
        item = self.image()
        def revoke(*args, **kwargs):
            self.kb.save_settings(self.actor, {'approved_origins': ['https://other.example.com']})
            return '{"text":"Do not retain"}'
        self.model.complete.side_effect = revoke
        with self.assertRaises(AssistantError):
            self.read([item])

    def test_only_minimax_receives_actual_images_with_source_anchor(self):
        item = self.image()
        result = self.read([item])
        visual = result['items'][0]['visual_evidence'][0]
        messages = self.model.complete.call_args.args[0]
        self.assertEqual(self.model.complete.call_args.kwargs['profile'], VISION)
        self.assertIn(item['document_id'], messages[1]['content'][0]['text'])
        self.assertIn('data:image/jpeg;base64,', messages[1]['content'][1]['image_url']['url'])
        self.assertTrue(visual['visual_understanding'])
        self.assertEqual(visual['kind'], 'model_visual_interpretation')
        self.assertEqual(visual['source_sha256'], result['items'][0]['source_hash'])
        self.assertNotIn(visual['text'], result['items'][0]['text'])
        self.assertNotIn('base64', json.dumps(result))
        self.assertNotIn(self.temp.name, json.dumps(result))
        self.ocr.assert_not_called()

    def test_text_models_get_no_images_even_if_flag_is_true(self):
        item = self.image()
        for name in ('wanwu/glm-auto', 'wanwu/qwen-auto', 'wanwu/deepseek-auto', 'other-vision', 'fake-minimax'):
            with self.subTest(model=name):
                result = self.read([item], {**TEXT, 'model': name})
                visual = result['items'][0]['visual_evidence'][0]
                self.assertFalse(visual['visual_understanding'])
                self.assertEqual(visual['kind'], 'ocr_transcription')
                self.assertIn('OCR', visual['note'])
        self.model.complete.assert_not_called()

    def test_configured_minimax_can_assist_text_model_but_must_be_allowed(self):
        item = self.image()
        self.model.settings.return_value['models'] = [{**VISION, 'configured': True}]
        self.model.profile.return_value = VISION
        self.read([item], TEXT)
        self.assertEqual(self.model.complete.call_args.kwargs['profile'], VISION)
        self.model.complete.reset_mock()
        self.kb.save_settings(self.actor, {'approved_origins': ['https://models.example.com']})
        self.model.profile.return_value = {**VISION, 'endpoint': 'https://unapproved.example.com/v1'}
        self.read([item], TEXT)
        self.model.complete.assert_not_called()

    def test_disabled_minimax_and_failed_vision_use_explicit_ocr_fallback(self):
        item = self.image()
        result = self.read([item], {**VISION, 'image_input': False})
        self.model.complete.assert_not_called()
        self.assertFalse(result['items'][0]['visual_evidence'][0]['visual_understanding'])
        self.model.complete.side_effect = AssistantError('vision unsupported', 502)
        result = self.read([item])
        self.assertTrue(result['warning'])
        self.assertEqual(result['items'][0]['visual_evidence'][0]['kind'], 'ocr_transcription')

    def test_receipts_reuse_unchanged_evidence_and_reject_tampering(self):
        item = self.image()
        first = self.read([item])
        second = self.read(first['items'])
        self.assertEqual(second['items'], first['items'])
        self.model.complete.assert_called_once()
        forged = copy.deepcopy(first['items'])
        forged[0]['visual_evidence'][0]['text'] = 'forged result'
        third = self.read(forged)
        self.assertNotIn('forged result', json.dumps(third))
        self.assertEqual(self.model.complete.call_count, 2)

    def test_cached_evidence_rechecks_acl_and_model_configuration(self):
        item = self.image(managed=True)
        first = self.read([item])
        self.allowed = False
        with self.assertRaises(AssistantError):
            self.read(first['items'])
        self.model.complete.assert_called_once()
        self.allowed = True
        second = self.read(first['items'], TEXT)
        self.assertEqual(second['items'][0]['visual_evidence'][0]['kind'], 'ocr_transcription')
        self.model.complete.assert_called_once()

    def test_docx_reads_only_matched_nearby_embedded_images(self):
        from docx import Document
        document = Document()
        document.add_paragraph('needle chapter')
        document.add_picture(io.BytesIO(png()))
        for i in range(8):
            document.add_paragraph('unrelated paragraph %d' % i)
        document.add_picture(io.BytesIO(png('blue')))
        buffer = io.BytesIO()
        document.save(buffer)
        item = self.upload('chapter.docx', buffer.getvalue())
        result = self.read([item])
        self.model.complete.assert_called_once()
        visual = result['items'][0]['visual_evidence'][0]
        self.assertIn('段落2', visual['location'])
        self.assertIn('word/media/', visual['location'])

    def test_pptx_reads_matched_slide_not_other_slides(self):
        from pptx import Presentation
        from pptx.util import Inches
        document = Presentation()
        for index in range(2):
            slide = document.slides.add_slide(document.slide_layouts[6])
            box = slide.shapes.add_textbox(0, 0, Inches(2), Inches(1))
            box.text = 'needle slide %d' % index
            slide.shapes.add_picture(io.BytesIO(png('red' if index == 0 else 'blue')), 0, Inches(1))
        buffer = io.BytesIO()
        document.save(buffer)
        item = self.upload('slides.pptx', buffer.getvalue())
        result = self.read([item])
        self.model.complete.assert_called_once()
        self.assertIn('幻灯片1', result['items'][0]['visual_evidence'][0]['location'])

    def test_xlsx_embedded_picture_keeps_sheet_and_row_anchor(self):
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as ExcelImage
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Needle'
        sheet.append(['Device', 'Limit'])
        sheet.append(['needle', 'one'])
        sheet.add_image(ExcelImage(io.BytesIO(png())), 'C2')
        other = workbook.create_sheet('Other')
        other.append(['other'])
        other.add_image(ExcelImage(io.BytesIO(png('blue'))), 'A1')
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        item = self.upload('equipment.xlsx', buffer.getvalue())
        result = self.read([item])
        self.model.complete.assert_called_once()
        self.assertIn('工作表「Needle」图片行2', result['items'][0]['visual_evidence'][0]['location'])

    def test_pdf_renders_matched_page_with_pixel_bound(self):
        item = self.upload('policy.pdf', make_pdf_text('needle page'))
        result = self.read([item])
        self.assertIn('第1页', result['items'][0]['visual_evidence'][0]['location'])
        data_url = self.model.complete.call_args.args[0][1]['content'][1]['image_url']['url']
        import base64
        with Image.open(io.BytesIO(base64.b64decode(data_url.split(',')[1]))) as image:
            self.assertLessEqual(max(image.size), 1600)

    def test_missing_renderer_and_bounded_image_count_are_reported(self):
        item = self.upload('policy.pdf', make_pdf_text('needle page'))
        with patch.dict(sys.modules, {'pypdfium2': None}):
            result = self.read([item])
        self.assertIn('visual_read_failed', result['items'][0]['coverage']['missing'])
        self.assertTrue(result['warning'])
        self.model.complete.assert_not_called()
        from docx import Document
        document = Document()
        paragraph = document.add_paragraph('needle')
        for color in ('red', 'blue', 'green'):
            paragraph.add_run().add_picture(io.BytesIO(png(color)))
        output = io.BytesIO()
        document.save(output)
        item = self.upload('many.docx', output.getvalue())
        result = self.read([item])
        self.assertEqual(result['coverage']['images_read'], evidence.MAX_IMAGES)
        self.assertEqual(self.model.complete.call_count, evidence.MAX_IMAGES)
        self.assertFalse(result['items'][0]['coverage']['visual_complete'])

    def test_oversized_source_and_xml_entities_fail_without_inference(self):
        item = self.image()
        with patch.object(evidence, 'MAX_FILE_BYTES', 8):
            result = self.read([item])
        self.assertIn('visual_read_failed', result['items'][0]['coverage']['missing'])
        self.model.complete.assert_not_called()
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            archive.writestr('word/document.xml', '<!DOCTYPE a [<!ENTITY x "x">]><a/>')
        with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as archive:
            with self.assertRaises(ValueError):
                evidence._Package(archive).xml('word/document.xml')


if __name__ == '__main__':
    unittest.main()
