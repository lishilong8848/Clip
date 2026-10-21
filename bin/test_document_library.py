"""Temporary-directory document library acceptance; no provider or model calls."""
import copy
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent))

from lan_bitable_template_portal.document_library import DocumentLibraryService, validate_source
from openclaw_service.assistant.lighthouse_ai import AssistantError
from openclaw_service.assistant.lighthouse_knowledge import KnowledgeBase
from bin.test_lighthouse_knowledge import FakeEmbedder, FakeMemoryIndex


def actor(scopes=('A',), *, admin=False):
    return {'id': 'ouFixtureAdmin' if admin else 'ouFixturePerson', 'name': 'Fixture',
        'modules': ['knowledge'], 'scopes': list(scopes), 'is_admin': admin}


def configuration(**changes):
    return {'id': 'manuals', 'name': 'Manuals', 'base_token': 'baseFixture123', 'table_id': 'tblFixture',
        'display_fields': ['fldTitle', 'fldScope', 'fldNumber'], 'attachment_fields': ['fldFiles'],
        'scope_field': 'fldScope', 'read_scopes': ['A', 'B'], 'write_scopes': ['A', 'B'], **changes}


class Remote:
    def __init__(self):
        self.rows = {
            'recOne': {'record_id': 'recOne', 'fields': {'Title': 'Alpha policy', 'Building': 'A', 'Number': 4,
                'Files': [{'file_token': 'tokenFixtureOne', 'name': 'alpha.txt', 'size': 24}]}},
            'recTwo': {'record_id': 'recTwo', 'fields': {'Title': 'Beta policy', 'Building': 'B', 'Number': 7,
                'Files': [{'file_token': 'tokenFixtureTwo', 'name': 'beta.txt', 'size': 24}]}},
        }
        self.files = {'tokenFixtureOne': b'Alpha policy service steps', 'tokenFixtureTwo': b'Beta policy service steps'}
        self.metadata = {}
        self.creates = self.updates = self.uploads = self.downloads = 0
        self.fail_write_after_commit = self.fail_write_before_commit = False
        self.fail_scan = False

    def fields(self):
        return [{'field_id': fid, 'field_name': name, 'type': typ, 'property': {}}
            for fid, name, typ in [('fldTitle', 'Title', 1), ('fldScope', 'Building', 3),
                ('fldNumber', 'Number', 2), ('fldFiles', 'Files', 17), ('fldOpaque', 'Opaque', 999), ('fldFormula', 'Formula', 20)]]

    def records(self, cursor='', page_size=100):
        if self.fail_scan:
            raise OSError('fixture scan failure')
        return {'items': copy.deepcopy(list(self.rows.values())), 'has_more': False, 'page_token': ''}

    def get(self, identity):
        return copy.deepcopy(self.rows.get(identity))

    def create(self, fields, operation_id):
        self.creates += 1
        identity = 'recCreated' + str(self.creates)
        self.rows[identity] = {'record_id': identity, 'fields': copy.deepcopy(fields)}
        return self.get(identity)

    def update(self, identity, fields):
        self.updates += 1
        if self.fail_write_before_commit:
            raise OSError('fixture timeout')
        fields = copy.deepcopy(fields)
        if 'Files' in fields:
            previous = {v['file_token']: v for v in self.rows[identity]['fields'].get('Files', [])}
            fields['Files'] = [{**previous.get(v['file_token'], {}), **self.metadata.get(v['file_token'], {}), **v} for v in fields['Files']]
        self.rows[identity]['fields'].update(fields)
        if self.fail_write_after_commit:
            raise OSError('fixture response loss')
        return self.get(identity)

    def upload(self, path, name):
        self.uploads += 1
        token = 'tokenFixtureUpload' + str(self.uploads)
        self.files[token] = Path(path).read_bytes()
        self.metadata[token] = {'name': name, 'size': len(self.files[token])}
        return {'file_token': token, **self.metadata[token]}

    def download(self, token):
        self.downloads += 1
        return self.files[token]

    def close(self):
        pass


class LibraryFixture(unittest.TestCase):
    def setUp(self):
        self.temp = self.enterContext(tempfile.TemporaryDirectory())
        self.enterContext(patch.dict(os.environ, {'CLIPFLOW_DATA_DIR': self.temp}))
        from upload_event_module import logger
        trace = getattr(logger, '_crash_trace_file', None)
        if trace and Path(trace.name).resolve().is_relative_to(Path(self.temp).resolve()):
            def close_logging():
                import faulthandler
                import logging
                logger._stop_logging_listener()
                logging.shutdown()
                faulthandler.disable()
                trace.close()
                logger._crash_trace_file = None
            self.addCleanup(close_logging)
        self.remote = Remote()
        self.now = 1000.0
        self.kb = KnowledgeBase(Path(self.temp) / 'knowledge', embedder=FakeEmbedder(),
            memory_index_factory=FakeMemoryIndex, start_worker=False)
        self.addCleanup(self.kb.close)
        self.service = DocumentLibraryService(self.temp, remote_factory=lambda _: self.remote,
            knowledge=self.kb, clock=lambda: self.now)
        self.admin = actor(('A', 'B'), admin=True)
        self.service.save_table(self.admin, configuration())
        self.drain()

    def drain(self):
        for _ in range(80):
            changed = self.service.process_one()
            indexed = self.kb.process_one()
            if not changed and not indexed:
                return
        self.fail('worker failed to settle')

    def record(self):
        return self.service.record(actor(), 'manuals', 'recOne')

    def mutation(self, fields, **changes):
        return {'fields': fields, 'expected_revision': self.record()['revision'],
            'operation_id': 'fixture-operation-1', 'reason': 'verified fixture', **changes}


class LibraryTests(LibraryFixture):
    def test_native_fields_schema_and_scope_filtered_pagination(self):
        schema = self.service.schema(actor(), 'manuals')
        self.assertEqual(schema['fields'][2]['type'], 2)
        self.assertTrue(schema['fields'][-1]['read_only'])
        self.assertTrue(schema['fields'][-2]['read_only'])
        listed = self.service.list_records(actor(), 'manuals', {'q': 'policy', 'page_size': 1})
        self.assertEqual(listed['total'], 1)
        self.assertFalse(listed['has_more'])
        self.assertIsInstance(listed['items'][0]['fields']['fldNumber'], int)
        self.assertTrue(listed['items'][0]['capabilities']['edit'])
        self.assertNotIn('file_token', json.dumps(listed))
        with self.assertRaises(AssistantError):
            self.service.record(actor(), 'manuals', 'recTwo')

    def test_knowledge_module_and_buildings_required_even_for_admin(self):
        for denied in ({**actor(), 'modules': []}, {**actor(), 'is_guest': True}, actor(('B',), admin=True)):
            with self.assertRaises(AssistantError):
                self.service.record(denied, 'manuals', 'recOne')
        with self.assertRaises(AssistantError):
            self.service.save_table(actor(), configuration(id='other', table_id='tblOther'))

    def test_typed_create_patch_and_operation_idempotency(self):
        payload = self.mutation({'fldNumber': 12, 'fldTitle': 'Updated'})
        saved = self.service.save_record(actor(), 'manuals', payload, identity='recOne')
        repeated = self.service.save_record(actor(), 'manuals', payload, identity='recOne')
        self.assertEqual(saved['revision'], repeated['revision'])
        self.assertEqual(self.remote.updates, 1)
        self.assertEqual(self.remote.rows['recOne']['fields']['Number'], 12)
        created = self.service.save_record(actor(), 'manuals', {'fields': {'fldScope': 'A', 'fldNumber': 23},
            'reason': 'create fixture', 'operation_id': 'fixture-create-1'})
        self.assertEqual(created['fields']['fldNumber'], 23)
        with self.assertRaises(AssistantError):
            self.service.save_record(actor(), 'manuals', {**payload, 'fields': {'fldNumber': 99}}, identity='recOne')
        self.assertEqual(self.remote.updates, 1)

    def test_stale_revision_readonly_fields_and_scope_moves_do_not_write(self):
        for fields in ({'fldFormula': 'bad'}, {'fldOpaque': 'bad'}, {'fldFiles': []}, {'fldUnknown': 1}, {'fldScope': 'B'}):
            with self.assertRaises(AssistantError):
                self.service.save_record(actor(), 'manuals', self.mutation(fields), identity='recOne')
        with self.assertRaises(AssistantError):
            self.service.save_record(actor(), 'manuals', self.mutation({'fldNumber': 9}, expected_revision='stale'), identity='recOne')
        self.assertEqual(self.remote.updates, 0)

    def test_lost_update_response_is_confirmed_without_resending(self):
        self.remote.fail_write_after_commit = True
        payload = self.mutation({'fldNumber': 18})
        saved = self.service.save_record(actor(), 'manuals', payload, identity='recOne')
        self.assertEqual(saved['fields']['fldNumber'], 18)
        self.service.save_record(actor(), 'manuals', payload, identity='recOne')
        self.assertEqual(self.remote.updates, 1)

    def test_uncertain_update_remains_blocked_and_retry_never_mutates(self):
        self.remote.fail_write_before_commit = True
        payload = self.mutation({'fldNumber': 18})
        for _ in range(2):
            with self.assertRaises(AssistantError):
                self.service.save_record(actor(), 'manuals', payload, identity='recOne')
        self.assertEqual(self.remote.updates, 1)
        self.assertEqual(self.service.list_records(actor(), 'manuals', {})['items'], [])
        self.assertEqual(self.kb.list(actor(), {})['items'], [])

    def test_write_guard_denial_does_not_poison_record_or_replay(self):
        payload = self.mutation({'fldNumber': 18})
        with patch.object(self.remote, 'update', side_effect=AssistantError('Blocked fixture', 403, category='write_blocked')) as blocked:
            for _ in range(2):
                with self.assertRaises(AssistantError) as error:
                    self.service.save_record(actor(), 'manuals', payload, identity='recOne')
                self.assertEqual(error.exception.status, 403)
            self.assertEqual(blocked.call_count, 1)
        self.assertEqual(self.record()['fields']['fldNumber'], 4)
        self.assertEqual(len(self.kb.list(actor(), {})['items']), 1)

    def test_attachment_replace_preserves_identity_versions_audit_and_old_bytes(self):
        old = self.record()
        attachment = old['attachments'][0]['attachment_id']
        payload = {'field_id': 'fldFiles', 'replace_attachment_id': attachment,
            'expected_revision': old['revision'], 'operation_id': 'fixture-upload-1', 'reason': 'replace fixture'}
        saved = self.service.upload_attachment(actor(), 'manuals', 'recOne', payload, 'new.txt', io.BytesIO(b'New alpha procedure'))
        self.assertEqual(saved['attachments'][0]['attachment_id'], attachment)
        self.service.upload_attachment(actor(), 'manuals', 'recOne', payload, 'new.txt', io.BytesIO(b'New alpha procedure'))
        self.assertEqual((self.remote.uploads, self.remote.updates), (1, 1))
        versions = self.service.versions(actor(), 'manuals', 'recOne')['items']
        self.assertEqual(len(versions), 2)
        self.assertEqual(versions[0]['reason'], 'replace fixture')
        self.assertIn('?version=1', versions[-1]['attachments'][0]['download_url'])
        del self.remote.files['tokenFixtureOne']
        path, _ = self.service.download(actor(), 'manuals', 'recOne', attachment, version=1)
        self.assertEqual(path.read_bytes(), b'Alpha policy service steps')

    def test_initial_incremental_source_deletion_and_partial_scan_recovery(self):
        self.assertEqual(len(self.kb.list(actor(), {})['items']), 1)
        del self.remote.rows['recOne']
        self.remote.fail_scan = True
        self.service.request_sync(actor(), 'manuals')
        self.service.process_one()
        self.assertEqual(len(self.kb.list(actor(), {})['items']), 1)
        self.remote.fail_scan = False
        self.now += 10
        self.drain()
        self.assertEqual(self.kb.list(actor(), {})['items'], [])
        with self.assertRaises(AssistantError):
            self.service.record(actor(), 'manuals', 'recOne')

    def test_background_recovery_and_dedup_survive_reopen(self):
        with self.kb.connect() as db:
            before = db.execute('SELECT count(*) FROM versions').fetchone()[0]
        reopened = DocumentLibraryService(self.temp, remote_factory=lambda _: self.remote, clock=lambda: self.now)
        reopened.knowledge = self.kb
        reopened.request_sync(actor(), 'manuals')
        for _ in range(10):
            if not reopened.process_one():
                break
        with self.kb.connect() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM versions').fetchone()[0], before)

    def test_acl_change_disables_index_without_waiting_for_worker(self):
        source = self.service.tables(self.admin)['items'][0]
        self.service.save_table(self.admin, configuration(read_scopes=['B'], write_scopes=['B'], expected_revision=source['revision']), key='manuals')
        self.assertEqual(self.kb.list(actor(), {})['items'], [])
        self.assertFalse(self.kb.search(actor(), 'Alpha policy')['items'])

    def test_remote_scope_change_denies_old_cache_and_historical_download(self):
        attachment = self.record()['attachments'][0]['attachment_id']
        self.remote.rows['recOne']['fields']['Building'] = 'B'
        with self.assertRaises(AssistantError):
            self.record()
        self.assertEqual(self.kb.list(actor(), {})['items'], [])
        with self.assertRaises(AssistantError):
            self.service.download(actor(('B',)), 'manuals', 'recOne', attachment, version=1)

    def test_source_table_configuration_validation(self):
        for changes in ({'module': 'assistant'}, {'url': 'https://evil.example/base/baseFixture123'},
                {'write_scopes': ['B'], 'read_scopes': ['A']}, {'scope_field': '', 'fixed_scopes': []},
                {'url': 'https://team.feishu.cn/base/differentBase?table=tblFixture'}):
            with self.assertRaises(AssistantError):
                validate_source(configuration(**changes))

    def test_missing_and_empty_scope_grants_never_default_open(self):
        payload = configuration()
        del payload['read_scopes'], payload['write_scopes']
        for source in (validate_source(payload), validate_source(configuration(read_scopes=[], write_scopes=[]))):
            self.assertEqual(source['read_scopes'], [])
            self.assertEqual(source['write_scopes'], [])
            self.assertFalse(self.service._allowed(self.admin, source, ['A']))
            self.assertFalse(self.service._allowed(self.admin, source, ['A'], write=True))
        payload = configuration()
        del payload['write_scopes']
        self.assertEqual(validate_source(payload)['write_scopes'], [])
        table = self.service.tables(self.admin)['items'][0]
        self.service.save_table(self.admin, configuration(read_scopes=[], write_scopes=[],
            expected_revision=table['revision']), key='manuals')
        self.assertEqual(self.service.tables(actor())['items'], [])
        with self.assertRaises(AssistantError):
            self.record()
        self.assertEqual(self.kb.list(actor(), {})['items'], [])
        self.assertFalse(self.service.tables(self.admin)['items'][0]['capabilities']['create'])

    def test_record_scope_requires_every_component_to_be_known(self):
        source = self.service._table('manuals')
        for value in (None, '', [], ['A', 'Z'], ['A', {}], ['A', None], 'A/Z',
                'A/110unknown', 'A (unknown)', {'text': 'A', 'name': 'unknown'}):
            with self.subTest(value=value):
                self.assertEqual(self.service._record_scopes(source, {'fldScope': value}), [])
        for value in ('A', 'A\u697c', [{'name': 'A'}], [{'type': 'text', 'text': 'A'},
                {'type': 'text', 'text': '\u697c'}]):
            with self.subTest(value=value):
                self.assertEqual(self.service._record_scopes(source, {'fldScope': value}), ['A'])
        for value in (['A', 'B'], 'A/B', 'A-B', 'AB\u697c'):
            with self.subTest(value=value):
                self.assertEqual(self.service._record_scopes(source, {'fldScope': value}), ['A', 'B'])
        self.assertEqual(self.service._record_scopes(source, {'fldScope': 'A/110'}), ['110', 'A'])
        source['fixed_scopes'] = ['A']
        self.assertEqual(self.service._record_scopes(source, {}), [])
        self.assertEqual(self.service._record_scopes(source, {'fldScope': ['A', 'B']}), [])

    def test_mixed_unknown_scope_retires_visibility_without_any_write(self):
        old = self.record()
        attachment = old['attachments'][0]['attachment_id']
        payload = self.mutation({'fldNumber': 8})
        self.remote.rows['recOne']['fields']['Building'] = ['A', 'Unknown']
        with self.assertRaises(AssistantError):
            self.record()
        self.assertEqual(self.service.list_records(actor(), 'manuals', {})['items'], [])
        self.assertEqual(self.kb.list(actor(), {})['items'], [])
        self.assertFalse(self.kb.search(actor(), 'Alpha policy')['items'])
        with self.assertRaises(AssistantError):
            self.service.versions(actor(), 'manuals', 'recOne')
        with self.assertRaises(AssistantError):
            self.service.download(actor(), 'manuals', 'recOne', attachment, version=1)
        with self.assertRaises(AssistantError):
            self.service.save_record(actor(), 'manuals', payload, identity='recOne')
        self.assertEqual((self.remote.updates, self.remote.uploads), (0, 0))

    def test_write_scopes_only_narrow_existing_read_scope(self):
        table = self.service.tables(self.admin)['items'][0]
        self.service.save_table(self.admin, configuration(write_scopes=['B'], expected_revision=table['revision']), key='manuals')
        current = self.record()
        self.assertFalse(current['capabilities']['edit'])
        with self.assertRaises(AssistantError):
            self.service.save_record(actor(), 'manuals', self.mutation({'fldNumber': 42}), identity='recOne')
        self.assertEqual(self.remote.updates, 0)

    def test_fixed_scope_create_capability_requires_all_record_scopes(self):
        previous = self.service.tables(self.admin)['items'][0]
        self.service.save_table(self.admin, configuration(scope_field='', fixed_scopes=['A', 'B'],
            expected_revision=previous['revision']), key='manuals')
        self.assertFalse(self.service.tables(actor())['items'][0]['capabilities']['create'])
        self.assertTrue(self.service.tables(self.admin)['items'][0]['capabilities']['create'])

    def test_same_record_operation_claim_is_cross_instance_cas(self):
        payload = self.mutation({'fldNumber': 9})
        source = self.service._table('manuals')
        self.service._claim_write(actor(), source, 'recOne', payload, fields=payload['fields'])
        peer = DocumentLibraryService(self.temp, remote_factory=lambda _: self.remote)
        with self.assertRaises(AssistantError):
            peer._claim_write(actor(), source, 'recOne', {**payload, 'operation_id': 'fixture-peer-operation'}, fields=payload['fields'])
        self.assertEqual(self.remote.updates, 0)

    def test_scan_does_not_retire_a_record_saved_after_scan_started(self):
        self.service.request_sync(actor(), 'manuals')
        task = self.service._claim('sync_jobs')
        self.now += 1
        raw = {'record_id': 'recNew', 'fields': {'Title': 'New', 'Building': 'A', 'Number': 3}}
        source = self.service._table('manuals')
        self.service._remember(source, self.service._normalize(source, raw, self.remote.fields()))
        self.service._scan_page(task)
        self.assertFalse(self.service._cached('manuals', 'recNew')['deleted'])

    def test_expired_scan_cannot_overwrite_a_new_record_snapshot(self):
        self.service.request_sync(actor(), 'manuals')
        task = self.service._claim('sync_jobs')
        self.now += 181
        source = self.service._table('manuals')
        raw = self.remote.get('recOne')
        record = self.service._normalize(source, raw, self.remote.fields())
        with self.assertRaises(AssistantError):
            self.service._remember(source, record, generation=task['id'], job=(task['id'], task['lease']))


class KnowledgeSourceACLTests(LibraryFixture):
    def test_pptx_source_parser_retains_slide_and_unredacted_security_check_text(self):
        from pptx import Presentation
        from openclaw_service.assistant.lighthouse_knowledge_extract import extract_sections
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = 'Fixture procedure'
        slide.shapes.add_textbox(0, 0, 1000000, 1000000).text_frame.text = 'password: fixture-secret-for-test'
        content = io.BytesIO()
        deck.save(content)
        sections = extract_sections(content.getvalue(), 'fixture.pptx')
        self.assertTrue(any('fixture-secret-for-test' in section['text'] for section in sections))
        self.assertTrue(all('1' in section['location'] for section in sections))

    def test_source_list_search_file_versions_groups_and_citation_revocation(self):
        a, b = actor(), actor(('B',))
        listed = self.kb.list(a, {})
        identity = listed['items'][0]['id']
        found = self.kb.search(a, 'Alpha policy')['items']
        self.assertTrue(found)
        profile = {'endpoint': 'https://company.example/v1'}
        self.assertTrue(self.kb.evidence_current(found, profile, actor=a))
        self.assertFalse(self.kb.evidence_current(found, profile))
        self.assertFalse(self.kb.evidence_current(found, profile, actor=b))
        for call in (lambda: self.kb.file(b, identity), lambda: self.kb.document(b, identity, {}),
                lambda: self.kb.file(b, identity, 1)):
            with self.assertRaises(AssistantError):
                call()
        group = self.kb.create_group(a, 'Private alpha')
        with self.kb.connect() as db:
            db.execute('INSERT INTO doc_group_members VALUES(?,?)', (identity, group['id']))
        self.assertEqual(self.kb.groups(a)['groups'][0]['count'], 1)
        self.assertEqual(self.kb.groups(b)['groups'], [])
        self.service.refresh_actor = lambda prior: {**prior, 'modules': []}
        self.assertFalse(self.kb.evidence_current(found, profile, actor=a))
        self.assertFalse(self.kb.result_current(a, listed))

    def test_shared_documents_keep_visibility_and_do_not_merge_with_source_hash(self):
        uploaded = self.kb.upload(actor(), 'shared.txt', b'Alpha policy service steps')
        self.kb.process_one()
        self.assertFalse(uploaded.get('duplicate', False))
        self.assertEqual(len(self.kb.list(actor(), {})['items']), 2)
        self.assertEqual(len(self.kb.list(actor(('B',)), {})['items']), 2)
        self.service.refresh_actor = lambda prior: {**prior, 'modules': []}
        self.assertEqual([v['id'] for v in self.kb.list(actor(), {})['items']], [uploaded['id']])

    def test_source_cannot_be_edited_restored_or_regrouped_as_shared_even_by_admin(self):
        item = self.kb.list(actor(), {})['items'][0]
        self.assertFalse(item['can_edit'])
        for action in ('delete', 'retry', 'restore'):
            with self.assertRaises(AssistantError):
                self.kb.change(self.admin, item['id'], action, item['version'])
        with self.assertRaises(AssistantError):
            self.kb.upload(self.admin, 'replace.txt', b'changed text', document_id=item['id'], revision=item['version'])
        with self.assertRaises(AssistantError):
            self.kb.set_documents_group(self.admin, None, [item['id']])

    def test_missing_source_authority_or_provenance_fails_closed(self):
        self.kb.source_acl = None
        self.assertEqual(self.kb.list(actor(), {})['items'], [])
        self.kb.source_acl = self.service.can_read_source
        with self.kb.connect() as db:
            db.execute('DELETE FROM document_sources')
        self.assertEqual(self.kb.list(actor(), {})['items'], [])

    def test_unchanged_attachment_content_reuses_knowledge_version_on_record_edit(self):
        identity = self.kb.list(actor(), {})['items'][0]['id']
        with self.kb.connect() as db:
            before = db.execute('SELECT count(*) FROM versions WHERE document_id=?', (identity,)).fetchone()[0]
        self.service.save_record(actor(), 'manuals', self.mutation({'fldNumber': 9}), identity='recOne')
        self.drain()
        with self.kb.connect() as db:
            after = db.execute('SELECT count(*) FROM versions WHERE document_id=?', (identity,)).fetchone()[0]
        self.assertEqual(before, after)
        self.assertEqual(self.kb.file(actor(), identity)[1], 'alpha.txt')

    def test_old_ingestion_completion_is_rejected_after_new_canonical_revision(self):
        with self.kb.connect() as db:
            old = dict(db.execute('SELECT * FROM document_sources WHERE document_id=?', (self.kb.list(actor(), {})['items'][0]['id'],)).fetchone())
        source = json.loads(old['payload'])
        self.service.save_record(actor(), 'manuals', self.mutation({'fldNumber': 9}), identity='recOne')
        self.drain()
        with self.assertRaises(AssistantError):
            self.kb.upsert_source(old['source_key'], source, 'alpha.txt', b'Stale replacement')
        self.assertTrue(self.kb.search(actor(), 'Alpha policy')['items'])

    def test_late_invalidation_does_not_disable_already_synced_current_version(self):
        self.service.save_record(actor(), 'manuals', self.mutation({'fldNumber': 9}), identity='recOne')
        self.drain()
        self.kb.invalidate_sources('manuals/recOne/')
        self.assertTrue(self.kb.search(actor(), 'Alpha policy')['items'])

    def test_source_retry_recovers_on_existing_index_worker_without_new_version(self):
        current = self.record()
        self.kb.extractor = lambda *args: (_ for _ in ()).throw(OSError('temporary fixture parser failure'))
        self.service.upload_attachment(actor(), 'manuals', 'recOne', {'field_id': 'fldFiles',
            'expected_revision': current['revision'], 'operation_id': 'fixture-index-retry', 'reason': 'retry fixture',
            'replace_attachment_id': current['attachments'][0]['attachment_id']}, 'alpha.txt', io.BytesIO(b'Alpha policy new edition'))
        self.drain()
        self.assertEqual(self.record()['knowledge_status'], 'failed')
        self.kb.extractor = None
        with self.kb.connect() as db:
            db.execute('UPDATE source_retry SET next_attempt=0')
            before = db.execute('SELECT count(*) FROM versions').fetchone()[0]
        self.assertTrue(self.kb.process_one())
        self.assertEqual(self.record()['knowledge_status'], 'ready')
        with self.kb.connect() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM versions').fetchone()[0], before)

    def test_cached_answer_artifact_and_stream_hook_do_not_survive_revocation(self):
        from types import SimpleNamespace
        from openclaw_service.assistant.lighthouse_stream import LighthouseStream
        item = self.kb.search(actor(), 'Alpha policy')['items'][0]
        turn = {'question': 'Fixture question', 'answer': 'Protected alpha body', 'sources': [{
            'kind': 'company_knowledge', 'title': 'Protected source', 'data': item}],
            'output_files': [{'id': 'fixture-artifact'}]}
        self.kb.bind_source_artifacts(turn)
        self.assertTrue(self.kb.cached_turn_current(actor(), turn))
        self.assertTrue(self.kb.artifact_current(actor(), 'fixture-artifact'))
        stream = LighthouseStream.__new__(LighthouseStream)
        stream.portal = SimpleNamespace(get_knowledge=lambda: self.kb)
        self.service.refresh_actor = lambda prior: {**prior, 'scopes': ['B']}
        hidden = stream._knowledge_turn(actor(), turn)
        self.assertTrue(hidden['knowledge_access_changed'])
        self.assertNotIn('Protected alpha body', json.dumps(hidden))
        self.assertEqual(hidden['sources'], [])
        self.assertFalse(self.kb.artifact_current(actor(), 'fixture-artifact'))
        self.assertEqual(turn['answer'], 'Protected alpha body')

    def test_same_content_from_different_source_acls_never_merges(self):
        self.remote.files['tokenFixtureNewBeta'] = b'Alpha policy service steps'
        self.remote.rows['recTwo']['fields']['Files'] = [{'file_token': 'tokenFixtureNewBeta', 'name': 'beta.txt', 'size': 24}]
        self.service.request_sync(actor(), 'manuals')
        self.drain()
        a = self.kb.list(actor(), {})['items']
        b = self.kb.list(actor(('B',)), {})['items']
        self.assertEqual((len(a), len(b)), (1, 1))
        self.assertNotEqual(a[0]['id'], b[0]['id'])

    def test_revoked_completed_sse_never_replays_protected_text(self):
        import asyncio
        from types import SimpleNamespace
        from openclaw_service.assistant.lighthouse_stream import LighthouseStream
        item = self.kb.search(actor(), 'Alpha policy')['items'][0]
        sources = [{'kind': 'company_knowledge', 'data': item}]
        stream = LighthouseStream.__new__(LighthouseStream)
        stream.portal = SimpleNamespace(get_knowledge=lambda: self.kb)
        run = {'status': 'completed', 'scopes': ['A'], 'events': [
            {'type': 'text-delta', 'delta': 'Protected alpha body'},
            {'type': 'data-turn', 'data': {'sources': sources, 'answer': 'Protected alpha body'}}]}
        stream.get_run = lambda *args: run
        self.service.refresh_actor = lambda prior: {**prior, 'scopes': ['B']}
        async def authorize():
            return actor()
        async def read():
            return ''.join([chunk async for chunk in stream.stream(actor(), 'fixture', authorize)])
        data = asyncio.run(read())
        self.assertNotIn('Protected alpha body', data)


class LibraryRoutesTests(LibraryFixture):
    def setUp(self):
        super().setUp()
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from lan_bitable_template_portal.document_library_routes import install_document_library_routes
        self.current_actor = self.admin
        self.authorization_count = 0
        self.revoke_on_second = False
        async def authorize(request):
            self.authorization_count += 1
            if self.revoke_on_second and self.authorization_count % 2 == 0:
                return {**self.current_actor, 'modules': []}
            return self.current_actor
        app = FastAPI()
        install_document_library_routes(app, authorize, lambda: self.service)
        self.client = self.enterContext(TestClient(app))
        self.prefix = '/api/document-library/tables/manuals'

    def test_rest_schema_list_read_write_and_no_delete_routes(self):
        self.assertEqual(self.client.get(self.prefix + '/schema').status_code, 200)
        listed = self.client.get(self.prefix + '/records', params={'scope': 'A', 'page_size': 1}).json()['data']
        self.assertEqual(listed['items'][0]['record_id'], 'recOne')
        self.assertIsInstance(listed['has_more'], bool)
        saved = self.client.patch(self.prefix + '/records/recOne', json=self.mutation({'fldNumber': 10}))
        self.assertEqual(saved.status_code, 200, saved.text)
        for target in ('/records/recOne', '/records/recOne/attachments/any'):
            self.assertIn(self.client.delete(self.prefix + target).status_code, (404, 405))

    def test_multipart_upload_and_protected_version_download(self):
        current = self.record()
        upload = self.client.post(self.prefix + '/records/recOne/attachments', data={
            'field_id': 'fldFiles', 'expected_revision': current['revision'],
            'operation_id': 'fixture-route-upload', 'reason': 'route fixture',
            'replace_attachment_id': current['attachments'][0]['attachment_id']}, files={'file': ('new.txt', b'New procedure', 'text/plain')})
        self.assertEqual(upload.status_code, 200, upload.text)
        versions = self.client.get(self.prefix + '/records/recOne/versions').json()['data']['items']
        old_url = versions[-1]['attachments'][0]['download_url']
        downloaded = self.client.get(old_url)
        self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(downloaded.content, b'Alpha policy service steps')
        self.assertEqual(downloaded.headers['cache-control'], 'private, no-store')
        self.current_actor = actor(('B',))
        self.assertEqual(self.client.get(old_url).status_code, 404)

    def test_revocation_before_response_is_not_served(self):
        self.revoke_on_second = True
        response = self.client.get(self.prefix + '/records')
        self.assertEqual(response.status_code, 403)
        self.assertNotIn('Alpha policy', response.text)

    def test_config_write_requires_admin_and_optimistic_revision(self):
        self.current_actor = actor()
        self.assertEqual(self.client.post('/api/document-library/tables', json=configuration(id='new', table_id='tblNew')).status_code, 403)
        self.current_actor = self.admin
        self.assertEqual(self.client.put(self.prefix, json=configuration(expected_revision='stale')).status_code, 409)


if __name__ == '__main__':
    unittest.main()
