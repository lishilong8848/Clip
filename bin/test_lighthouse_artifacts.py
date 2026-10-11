"""Generated deliverables, previews and permissions; no cloud writes."""
import asyncio
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock, patch

sys.path.insert(0, str(Path(__file__).parent))
from test_lighthouse_stream import Store
from openclaw_service.assistant.lighthouse_files import LighthouseFiles
from openclaw_service.assistant.lighthouse_artifacts import create_document, document_bytes, package_long_reply, preview_file
from openclaw_service.assistant.lighthouse_ai import AssistantError
from lan_bitable_template_portal.feishu_assistant import FeishuMessenger, FeishuAssistant

ACTOR = {'id': 'test-author', 'scopes': ['A']}


class ArtifactTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.files = LighthouseFiles(Store(Path(temp.name) / 'test.sqlite3'))

    def test_word_contains_real_chinese_document_and_preview(self):
        from docx import Document
        file = create_document(self.files, ACTOR, '交接说明.docx', {'title': '交接说明', 'paragraphs': ['设备运行正常。', '次日检查执行情况。']})
        raw = self.files.get(ACTOR, file['id'])
        self.assertEqual(Document(raw['path']).paragraphs[1].text, '设备运行正常。')
        self.assertIn('次日检查', preview_file(self.files, ACTOR, file['id'])['text'])

    def test_excel_preserves_numbers_and_literal_formula_like_text(self):
        from openpyxl import load_workbook
        file = create_document(self.files, ACTOR, '数据.xlsx', {'sheets': [{'name': 'A楼', 'headers': ['名称', '数量'], 'rows': [['=HYPERLINK("https://example.com")', 3], ['设备', 0]]}]})
        raw = self.files.get(ACTOR, file['id'])
        book = load_workbook(raw['path'])
        self.assertEqual(book.active['A2'].data_type, 's')
        self.assertEqual(book.active['B2'].value, 3)
        book.close()
        preview = preview_file(self.files, ACTOR, file['id'])
        self.assertEqual(preview['sheets'][0]['rows'][1][1], '3')

    def test_ppt_and_pdf_are_openable_not_renamed_text(self):
        from pptx import Presentation
        from pypdf import PdfReader
        ppt = document_bytes('介绍.pptx', {'slides': [{'title': '设备检查', 'bullets': ['检查电源', '记录结果']}]})
        deck = Presentation(io.BytesIO(ppt))
        self.assertEqual(deck.slides[0].shapes.title.text, '设备检查')
        pdf = document_bytes('介绍.pdf', {'title': '设备检查', 'paragraphs': ['记录设备运行情况。']})
        self.assertIn('记录设备运行情况', PdfReader(io.BytesIO(pdf)).pages[0].extract_text())

    def test_invalid_structure_not_silently_omitted(self):
        bad = [('bad.xls', {}), ('bad.docx', {'sheets': []}), ('bad.xlsx', {'sheets': [{'name': 'x', 'headers': ['a'], 'rows': [[1, 2]]}]}),
               ('bad.pptx', {'slides': [{'title': 'x', 'bullets': ['x'] * 7}]}), ('bad.pdf', {'paragraphs': ['\x00']})]
        for name, data in bad:
            with self.subTest(name=name), self.assertRaises(AssistantError):
                document_bytes(name, data)

    def test_duplicate_sheet_names_and_nonfinite_values_rejected(self):
        sheet = {'name': 'Data', 'headers': ['x'], 'rows': [[1]]}
        with self.assertRaises(AssistantError): document_bytes('bad.xlsx', {'sheets': [sheet, {**sheet, 'name': 'data'}]})
        with self.assertRaises(AssistantError): document_bytes('bad.xlsx', {'sheets': [{**sheet, 'rows': [[float('inf')]]}]})

    def test_long_reply_retains_complete_answer_references_and_existing_files(self):
        text = '这是需要保留的完整回答。\n' * 400
        source = {'number': 1, 'title': '原业务记录', 'url': '/workbench-lite?scope=A'}
        result = package_long_reply(self.files, ACTOR, {'answer': text, 'sources': [source], 'file_ids': ['existing'], 'plan': {'status': 'needs_input'}})
        self.assertLess(len(result['answer']), 200)
        file = result['output_files'][0]
        original = Path(self.files.get(ACTOR, file['id'])['path']).read_text(encoding='utf-8')
        self.assertTrue(original.startswith(text))
        self.assertIn('[1] 原业务记录', original)
        self.assertEqual(result['plan']['status'], 'needs_input')
        self.assertEqual(result['file_ids'], ['existing', file['id']])

    def test_small_reply_has_no_file_and_sensitive_file_is_rejected(self):
        self.assertEqual(package_long_reply(self.files, ACTOR, {'answer': '你好'}), {'answer': '你好'})
        with self.assertRaises(AssistantError): create_document(self.files, ACTOR, '不允许.docx', {'paragraphs': ['身份证号：11010519491231002X']})
        self.assertFalse(self.files.store.docs)

    def test_long_reply_file_inherits_referenced_file_scope(self):
        result = package_long_reply(self.files, {**ACTOR, 'scopes': ['D'], 'allowed_scopes': ['A', 'D']},
                                    {'answer': '设备记录说明。' * 700, '_artifact_source_scopes': ['A', 'D']})
        identity = result['output_files'][0]['id']
        self.assertEqual(self.files.get({**ACTOR, 'scopes': ['A', 'D']}, identity)['source_scopes'], ['A', 'D'])
        with self.assertRaises(AssistantError): self.files.get({**ACTOR, 'scopes': ['D']}, identity)

    def test_preview_and_download_are_owner_and_scope_isolated(self):
        file = create_document(self.files, ACTOR, 'test.docx', {'paragraphs': ['原始内容']})
        for actor in ({'id': 'other', 'scopes': ['A'], 'is_admin': True}, {'id': ACTOR['id'], 'scopes': ['B']}):
            with self.subTest(actor=actor), self.assertRaises(AssistantError): preview_file(self.files, actor, file['id'])

    def test_code_preview_and_read_preserve_original_code_without_execution(self):
        text = 'const message = "中文，保留标点";\nconsole.log(message);'
        file = self.files.upload(ACTOR, 'example.js', text.encode(), extract=False, source_scopes=['A'])
        self.assertEqual(preview_file(self.files, ACTOR, file['id'])['text'], text)
        self.assertIn('console.log', self.files.text(ACTOR, file['id'])['text'])

    def test_generated_ppt_can_be_read_in_followup(self):
        file = create_document(self.files, ACTOR, '继续修改.pptx', {'slides': [{'title': '交接', 'bullets': ['核对设备情况']}]})
        self.assertIn('核对设备情况', self.files.text(ACTOR, file['id'])['text'])

    def test_spreadsheet_preview_is_bounded(self):
        file = create_document(self.files, ACTOR, 'long.xlsx', {'sheets': [{'name': '工作表', 'headers': ['内容'], 'rows': [[str(i)] for i in range(350)]}]})
        result = preview_file(self.files, ACTOR, file['id'])
        self.assertTrue(result['truncated'])
        self.assertEqual(len(result['sheets'][0]['rows']), 200)


class FeishuArtifactTests(unittest.IsolatedAsyncioTestCase):
    async def test_document_workers_bound_concurrent_cpu_work(self):
        import threading
        import time
        from openclaw_service.assistant.lighthouse_artifacts import document_task
        running = maximum = 0
        lock = threading.Lock()
        def work(index):
            nonlocal running, maximum
            with lock:
                running += 1
                maximum = max(maximum, running)
            time.sleep(.02)
            with lock: running -= 1
            return index
        self.assertEqual(await asyncio.gather(*(document_task(work, i) for i in range(20))), list(range(20)))
        self.assertEqual(maximum, 2)

    async def test_file_upload_failure_does_not_claim_delivery(self):
        messenger = FeishuMessenger({})
        client = SimpleNamespace(post=AsyncMock(return_value=SimpleNamespace(json=lambda: {'code': 99991663})))
        messenger.client, messenger.headers = AsyncMock(return_value=client), AsyncMock(return_value={})
        reply = {'text': '这是完整报告。' * 900}
        with patch('lan_bitable_template_portal.portal_service.external_real_write_guard', return_value={'real_write_allowed': True}):
            with self.assertRaises(AssistantError): await messenger.send({'message_id': 'om_test'}, reply)
        self.assertEqual(client.post.await_count, 1)
        self.assertNotIn('sent_parts', reply)
        self.assertNotIn('long_text_file_key', reply)

    async def test_file_upload_permission_error_is_not_a_transient_retry(self):
        messenger = FeishuMessenger({})
        client = SimpleNamespace(post=AsyncMock(return_value=SimpleNamespace(status_code=403, json=lambda: {'code': 99991672})))
        messenger.client, messenger.headers = AsyncMock(return_value=client), AsyncMock(return_value={})
        with patch('lan_bitable_template_portal.portal_service.external_real_write_guard', return_value={'real_write_allowed': True}):
            with self.assertRaises(AssistantError) as error:
                await messenger.upload_file('example.py', b'print("example")')
        self.assertEqual(error.exception.status, 403)
        self.assertEqual(error.exception.category, 'feishu_file_upload_permission')
        self.assertIn('文件上传权限', str(error.exception))

    async def test_long_answer_becomes_file_and_retry_does_not_resend_finished_parts(self):
        messenger = FeishuMessenger({})
        client = SimpleNamespace(post=AsyncMock(return_value=SimpleNamespace(json=lambda: {'code': 0, 'data': {'file_key': 'file_test'}})))
        messenger.client = AsyncMock(return_value=client)
        messenger.headers = AsyncMock(return_value={'Authorization': 'test'})
        message, reply = {'message_id': 'om_fixture'}, {'text': '完整说明。' * 900}
        checkpoint = AsyncMock()
        with patch('lan_bitable_template_portal.portal_service.external_real_write_guard', return_value={'real_write_allowed': True}):
            await messenger.send(message, reply, checkpoint=checkpoint)
            await messenger.send(message, reply, checkpoint=checkpoint)
        self.assertEqual(client.post.await_count, 3)
        self.assertEqual(client.post.call_args_list[0].kwargs['files']['file'][1].decode(), '完整说明。' * 900)
        parts = [call.kwargs['json'] for call in client.post.call_args_list[1:]]
        self.assertEqual([p['msg_type'] for p in parts], ['interactive', 'file'])
        self.assertLess(len(parts[0]['content']), 500)
        self.assertEqual(json.loads(parts[1]['content']), {'file_key': 'file_test'})

    async def test_generated_file_reply_uses_original_message(self):
        messenger = FeishuMessenger({})
        client = SimpleNamespace(post=AsyncMock(return_value=SimpleNamespace(json=lambda: {'code': 0})))
        messenger.client, messenger.headers = AsyncMock(return_value=client), AsyncMock(return_value={})
        with patch('lan_bitable_template_portal.portal_service.external_real_write_guard', return_value={'real_write_allowed': True}):
            await messenger.send({'message_id': 'om_original'}, {'text': '已生成文件', 'files': [{'id': 'abc', 'file_key': 'file_private'}]})
        self.assertTrue(all(call.args[0].endswith('/om_original/reply') for call in client.post.call_args_list))
        self.assertEqual(client.post.call_args_list[1].kwargs['json']['msg_type'], 'file')

    async def test_partial_file_reply_retries_only_unsent_file_with_the_same_uuid(self):
        messenger = FeishuMessenger({})
        success = SimpleNamespace(json=lambda: {'code': 0})
        client = SimpleNamespace(post=AsyncMock(side_effect=[success, success, TimeoutError(), success]))
        messenger.client, messenger.headers = AsyncMock(return_value=client), AsyncMock(return_value={})
        reply = {'text': '已生成两个文件', 'files': [{'id': 'a', 'file_key': 'file_a'}, {'id': 'b', 'file_key': 'file_b'}]}
        with patch('lan_bitable_template_portal.portal_service.external_real_write_guard', return_value={'real_write_allowed': True}):
            with self.assertRaises(TimeoutError): await messenger.send({'message_id': 'om_files'}, reply)
            self.assertEqual(reply['sent_parts'], ['answer:0', 'file:a'])
            await messenger.send({'message_id': 'om_files'}, reply)
        parts = [call.kwargs['json'] for call in client.post.call_args_list]
        self.assertEqual(parts[2], parts[3])
        self.assertEqual(reply['sent_parts'], ['answer:0', 'file:a', 'file:b'])

    async def test_attachment_stop_notice_is_not_deduplicated_as_the_original_answer(self):
        messenger = FeishuMessenger({})
        client = SimpleNamespace(post=AsyncMock(return_value=SimpleNamespace(json=lambda: {'code': 0})))
        messenger.client, messenger.headers = AsyncMock(return_value=client), AsyncMock(return_value={})
        message = {'message_id': 'om_stopped'}
        with patch('lan_bitable_template_portal.portal_service.external_real_write_guard', return_value={'real_write_allowed': True}):
            await messenger.send(message, {'text': '文件已经生成'})
            stopped = {'text': '本次文件发送已停止', 'attachment_stopped': True}
            await messenger.send(message, stopped)
            await messenger.send(message, stopped)
        self.assertEqual(client.post.await_count, 2)
        self.assertNotEqual(client.post.call_args_list[0].kwargs['json']['uuid'], client.post.call_args_list[1].kwargs['json']['uuid'])

    async def test_attachment_permission_rechecked_even_when_uploaded(self):
        bot = object.__new__(FeishuAssistant)
        bot._request_context = AsyncMock(return_value=(None, ACTOR))
        bot._call = AsyncMock(side_effect=[{'conversation_id': 'same'}, AssistantError('文件权限失效', 403)])
        bot.messenger = SimpleNamespace(upload_file=AsyncMock())
        with self.assertRaises(AssistantError):
            await bot._prepare_reply_files({'message_id': 'msg'}, {'actor_id': ACTOR['id'], 'conversation_id': 'same', 'files': [{'id': 'a' * 32, 'file_key': 'old'}]})
        bot.messenger.upload_file.assert_not_awaited()

    async def test_file_upload_wait_cannot_outlive_conversation_or_permission(self):
        for change in ('conversation', 'permission', 'identity'):
            with self.subTest(change=change):
                bot = object.__new__(FeishuAssistant)
                contexts = [(None, ACTOR), (None, {**ACTOR, 'id': 'other'} if change == 'identity' else ACTOR)]
                bot._request_context = AsyncMock(side_effect=contexts)
                responses = [{'conversation_id': 'same'}, {'name': 'test.docx'},
                             {'conversation_id': 'new' if change == 'conversation' else 'same'},
                             AssistantError('文件权限失效', 403) if change == 'permission' else {'name': 'test.docx'}]
                bot._call = AsyncMock(side_effect=responses)
                bot.messenger = SimpleNamespace(upload_file=AsyncMock(return_value='uploaded_private_key'))
                bot.inbox = SimpleNamespace(set=Mock())
                reply = {'actor_id': ACTOR['id'], 'conversation_id': 'same', 'files': [{'id': 'a' * 32}]}
                with patch('lan_bitable_template_portal.feishu_assistant_files._assistant_download', AsyncMock(return_value=b'file')), self.assertRaises(AssistantError):
                    await bot._prepare_reply_files({'message_id': 'msg'}, reply)
                bot.messenger.upload_file.assert_awaited_once()

    async def test_successful_file_upload_revalidates_without_uploading_again_on_retry(self):
        bot = object.__new__(FeishuAssistant)
        bot._request_context = AsyncMock(return_value=(None, ACTOR))
        async def call(request, actor, path, **kwargs):
            return {'conversation_id': 'same'} if path == 'conversation' else {'name': 'test.docx'}
        bot._call = AsyncMock(side_effect=call)
        bot.messenger = SimpleNamespace(upload_file=AsyncMock(return_value='uploaded_private_key'))
        bot.inbox = SimpleNamespace(set=Mock())
        reply = {'actor_id': ACTOR['id'], 'conversation_id': 'same', 'files': [{'id': 'a' * 32}]}
        download = AsyncMock(return_value=b'file')
        with patch('lan_bitable_template_portal.feishu_assistant_files._assistant_download', download):
            await bot._prepare_reply_files({'message_id': 'msg'}, reply)
            await bot._prepare_reply_files({'message_id': 'msg'}, reply)
        download.assert_awaited_once()
        bot.messenger.upload_file.assert_awaited_once()
        self.assertEqual(bot._request_context.await_count, 3)
        self.assertEqual(reply['files'][0]['file_key'], 'uploaded_private_key')


if __name__ == '__main__':
    unittest.main()
