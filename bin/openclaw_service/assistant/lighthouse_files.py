"""Private assistant attachments; originals never become arbitrary filesystem tools."""
import asyncio
import base64
import hashlib
import io
import json
import mimetypes
import multiprocessing
import os
import re
import time
import uuid
import zipfile
from itertools import zip_longest
from pathlib import Path
from xml.etree import ElementTree

from .lighthouse_ai import AssistantError, safe_data, safe_text
from .lighthouse_artifacts import CODE_EXTENSIONS

FILE_NAMESPACE = "lighthouse_files"
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_BATCH_BYTES = 100 * 1024 * 1024
DRILL_TEMPLATE_BYTES = 64 * 1024 * 1024
WATER_PHOTO_BYTES = 8 * 1024 * 1024
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".log", ".json", ".xml"}
AUDIO_EXTENSIONS = {'.wav', '.mp3', '.m4a', '.ogg', '.opus', '.webm', '.flac'}
ALLOWED_EXTENSIONS = IMAGE_EXTENSIONS | TEXT_EXTENSIONS | CODE_EXTENSIONS | AUDIO_EXTENSIONS | {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx"}
GENERATED_EXTENSIONS = {".zip", ".doc", ".xls", ".bin", ".pptx"} | CODE_EXTENSIONS
EXTRACTION_VERSION = 2


def upload_limit(actor, purpose=None):
    """Central per-purpose upload policy shared by the file service and HTTP route."""
    if purpose is not None and not isinstance(purpose, str):
        raise AssistantError("不支持的上传用途。", 400)
    purpose = (purpose or "").strip()
    if not purpose:
        return {"purpose": "default", "max_items": 10, "max_bytes": MAX_FILE_BYTES,
                "batch_bytes": MAX_BATCH_BYTES, "extensions": ALLOWED_EXTENSIONS, "extract": True}
    if purpose == "drill_template":
        if not actor.get("is_admin"):
            raise AssistantError("仅管理员可上传演练模板。", 403)
        return {"purpose": "drill_template", "max_items": 1, "max_bytes": DRILL_TEMPLATE_BYTES,
                "batch_bytes": MAX_BATCH_BYTES, "extensions": {".xlsx"}, "extract": False}
    if purpose in {"water_photo", "notice_photo"}:
        return {"purpose": purpose, "max_items": 20 if purpose == "water_photo" else 10, "max_bytes": WATER_PHOTO_BYTES,
                "batch_bytes": MAX_BATCH_BYTES, "extensions": IMAGE_EXTENSIONS, "extract": False}
    raise AssistantError("不支持的上传用途。", 400)


def _ocr_worker(content, sender):
    try:
        if os.name == "nt":
            import ctypes
            kernel = ctypes.WinDLL("kernel32", use_last_error=True)
            kernel.SetPriorityClass(kernel.GetCurrentProcess(), 0x00004000)
        from PIL import Image
        from winocr import recognize_pil
        with Image.open(io.BytesIO(content)) as image:
            image = image.convert("RGB")
            async def recognize():
                return await recognize_pil(image, "zh-Hans-CN")
            result = asyncio.run(recognize())
        sender.send((True, "\n".join(str(line.text or "") for line in result.lines)))
    except Exception:
        sender.send((False, "当前系统图片文字识别不可用，可继续使用图片模型或手动补充文字。"))
    finally:
        sender.close()


def recognize_text(content, timeout=30):
    context = multiprocessing.get_context("spawn")
    receiver, sender = context.Pipe(duplex=False)
    process = context.Process(target=_ocr_worker, args=(content, sender), daemon=True)
    try:
        process.start()
        if os.name == 'nt':
            from upload_event_module.services.process_lifetime import register_child_process
            if not register_child_process(process.pid) and process.is_alive():
                raise AssistantError("图片识别进程生命周期绑定失败，原图已保留。", 503)
        sender.close()
        if not receiver.poll(timeout):
            raise AssistantError("图片文字识别超时，原图已保留，可继续使用图片模型。", 408)
        ok, result = receiver.recv()
        if not ok:
            raise AssistantError(result, 422)
        return result
    finally:
        if process.is_alive():
            process.terminate()
        if process.pid:
            process.join(timeout=2)
        sender.close()
        receiver.close()


def _check_archive(content):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        info = archive.infolist()
        if len(info) > 5000 or sum(item.file_size for item in info) > MAX_BATCH_BYTES:
            raise AssistantError("文档解压后过大，请拆分文件。", 413)


def safe_document(text):
    text = str(text)
    if len(text) > 100000 or any(len(line) > 6000 for line in text.splitlines()):
        raise AssistantError("文档文字超过完整读取上限，请拆分后上传；原文件仍可下载。", 413)
    text = "\n".join(safe_text(line, limit=None) for line in text.splitlines())
    if len(text) > 100000 or any(len(line) > 6000 for line in text.splitlines()):
        raise AssistantError("文档文字超过完整读取上限，请拆分后上传；原文件仍可下载。", 413)
    return text


def document_data(value, depth=0):
    if depth > 12:
        raise AssistantError("文档数据层级过深，请简化后上传。")
    if isinstance(value, dict):
        return {key: document_data(item, depth + 1) for key, item in value.items() if key in safe_data({key: None})}
    if isinstance(value, list):
        return [document_data(item, depth + 1) for item in value]
    return safe_data(value)


def text_excerpt(document, offset, length):
    text = document['text']
    end = min(len(text), offset + length)
    if end < len(text):
        # safe_text drops a trailing newline; leave it for the next nonempty chunk.
        trimmed = end
        while trimmed > offset and text[trimmed - 1] == '\n':
            trimmed -= 1
        if trimmed > offset:
            end = trimmed
    references = [ref for ref in document.get('references', [])
                  if ref['start'] < end and ref['end'] > offset]
    complete = document.get('extraction_complete', False)
    return {'text': text[offset:end], 'offset': offset, 'end_offset': max(offset, end),
            'total_characters': len(text), 'next_offset': end if end < len(text) else None,
            'truncated': end < len(text), 'partial': offset > 0 or end < len(text) or not complete,
            'extraction_complete': complete, 'empty': document.get('empty', not text.strip()),
            'references': [{**ref, 'partial': ref['start'] < offset or ref['end'] > end}
                           for ref in references[:40]],
            'reference_count': len(references), 'references_truncated': len(references) > 40,
            'next_reference_offset': references[40]['start'] if len(references) > 40 else None,
            **{key: document[key] for key in ('total_pages', 'pages_without_text',
                'pages_without_text_count', 'pages_without_text_truncated') if key in document},
            'coverage': 'extractable_text', 'untrusted_document': True}


def extract_text(content, suffix):
    return extract_document(content, suffix)['text']


def extract_document(content, suffix, *, markdown=False):
    lines, references, characters = [], [], 0
    metadata = {'extraction_complete': True,
                'note': '仅提取可读文字；不代表已读取图片、批注或其他嵌入内容。来源位置对应原件，字符位置对应提取文本。'}

    def add(text, source, label=''):
        nonlocal characters
        text = safe_document(text)
        if not text.strip():
            return
        block = safe_document(label + text)
        start = characters + bool(lines)
        characters = start + len(block)
        if characters > 100000:
            raise AssistantError('文档文字超过完整读取上限，请拆分后上传；原文件仍可下载。', 413)
        lines.append(block)
        references.append({**source, 'start': start, 'end': characters})

    if suffix in TEXT_EXTENSIONS | CODE_EXTENSIONS:
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = content.decode(encoding)
            except UnicodeError:
                continue
            if suffix == ".json":
                text = json.dumps(document_data(json.loads(text)), ensure_ascii=False, indent=2)
            safe_document(text)
            for index, line in enumerate(text.splitlines(), 1):
                block = safe_text(line)
                start = characters + bool(lines)
                characters = start + len(block)
                lines.append(block)
                if block.strip():
                    references.append({'kind': 'normalized_json' if suffix == '.json' else 'line',
                                       'line': index, 'start': start, 'end': characters})
            break
        else:
            raise AssistantError("无法识别文本编码，请另存为UTF-8。")
    elif suffix in IMAGE_EXTENSIONS:
        add(recognize_text(content), {'kind': 'image_ocr', 'page': 1})
        metadata['note'] = '图片OCR可能识别错误，须核对原图；不代表已识别图形含义或签字。'
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise AssistantError("加密PDF无法读取，请上传未加密文件。")
        if len(reader.pages) > 100:
            raise AssistantError("PDF最多100页，请拆分文件。", 413)
        missing = []
        for index, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            if not text.strip():
                missing.append(index)
            add(text, {'kind': 'page', 'page': index}, f'[page {index}]\n')
        metadata.update(total_pages=len(reader.pages), pages_without_text=missing[:40],
                        pages_without_text_count=len(missing), pages_without_text_truncated=len(missing) > 40,
                        extraction_complete=not missing)
    elif suffix == '.pptx':
        _check_archive(content)
        from pptx import Presentation
        slides = Presentation(io.BytesIO(content)).slides
        if len(slides) > 100:
            raise AssistantError('PPT最多读取100页，请拆分文件。', 413)
        for index, slide in enumerate(slides, 1):
            for shape_index, shape in enumerate(slide.shapes, 1):
                text = shape.text if shape.has_text_frame else ''
                if shape.has_table:
                    text = '\n'.join(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
                add(text, {'kind': 'slide', 'page': index, 'shape': shape_index}, f'[slide {index}, shape {shape_index}]\n')
        metadata['total_pages'] = len(slides)
    elif suffix == ".docx":
        _check_archive(content)
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            xml = archive.read("word/document.xml")
        root = ElementTree.fromstring(xml)
        word = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        section, paragraph, table = '', 0, 0

        def paragraph_text(p):
            return ''.join(node.text or '' if node.tag == word + 't' else
                           '\t' if node.tag == word + 'tab' else '\n' if node.tag in {word + 'br', word + 'cr'} else ''
                           for node in p.iter())

        def body_blocks(parent):
            for child in parent:
                if child.tag in {word + 'p', word + 'tbl'}:
                    yield child
                else:
                    yield from body_blocks(child)

        for block in body_blocks(root.find(word + 'body')):
            if block.tag == word + 'p':
                paragraph += 1
                text = paragraph_text(block)
                style = block.find('./' + word + 'pPr/' + word + 'pStyle')
                if style is not None and re.fullmatch(r'heading\d+|title|subtitle', style.get(word + 'val', ''), re.I):
                    section = safe_text(text)
                add(text, {'kind': 'paragraph', 'paragraph': paragraph, 'section': section}, f'[paragraph {paragraph}] ')
            elif block.tag == word + 'tbl':
                table += 1
                for row_index, row in enumerate(block.findall(word + 'tr'), 1):
                    cells = [' '.join(paragraph_text(p) for p in cell.iter(word + 'p'))
                             for cell in row.findall(word + 'tc')]
                    paragraph += sum(1 for _ in row.iter(word + 'p'))
                    add('| ' + ' | '.join(cells) + ' |',
                        {'kind': 'table_row', 'table': table, 'row': row_index, 'section': section},
                        f'[table {table}, row {row_index}] ')
        metadata['note'] += ' Word按段落、标题及表格行定位，不推测排版页码。'
    elif suffix in {'.xlsx', '.xlsm'}:
        _check_archive(content)
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        books = []
        try:
            for data_only in (False, True):
                books.append(load_workbook(io.BytesIO(content), read_only=True, data_only=data_only, keep_links=False))
            for sheet, cached in zip(books[0].worksheets, books[1].worksheets):
                sheet.reset_dimensions()
                cached.reset_dimensions()
                title = safe_text(sheet.title)
                add('工作表：' + title, {'kind': 'sheet', 'sheet': title})
                populated_rows = 0
                for index, pair in enumerate(zip_longest(sheet.iter_rows(), cached.iter_rows()), 1):
                    if index > 100000:
                        raise AssistantError('表格行范围过大，请删除多余空行后重试；原文件仍可下载。', 413)
                    row, values = pair
                    if row is None or values is None or len(row) != len(values):
                        raise AssistantError('工作表结构读取不完整，未使用部分数据。')
                    populated = [i for i, cell in enumerate(row) if cell.value is not None and str(cell.value).strip()]
                    if not populated:
                        continue
                    populated_rows += 1
                    if populated_rows > 5000 or populated[-1] >= 80:
                        raise AssistantError('表格文字超过读取上限，请拆分后上传；原文件仍可下载。', 413)
                    cells = []
                    for cell, value in zip(row[:populated[-1] + 1], values):
                        text = '' if cell.value is None else str(cell.value)
                        if cell.data_type == 'f':
                            if not isinstance(cell.value, str):
                                text = getattr(cell.value, 'text', None) or f'[formula at {cell.coordinate}; expression unavailable]'
                            result = 'missing' if value.value is None else str(value.value)
                            text += f' [cached={result}; not recalculated]'
                        cells.append(text)
                    cell_range = f'A{index}:{get_column_letter(populated[-1] + 1)}{index}'
                    line = '| ' + ' | '.join(cells) + ' |' if markdown else '\t'.join(cells)
                    add(line, {'kind': 'row', 'sheet': title, 'row': index, 'cells': cell_range},
                        f'[{title}!{cell_range}] ')
            metadata['note'] += ' Excel保留空列和原始行号；公式显示原式及保存的缓存，missing表示没有结果，不当作0，不重新计算。'
        finally:
            for book in books:
                book.close()
    else:
        raise AssistantError('不支持此文件格式的文字提取。')
    return {'text': '\n'.join(lines), 'references': references,
            'empty': not any(ref['kind'] != 'sheet' for ref in references), **metadata}


class LighthouseFiles:
    def __init__(self, store, root=None):
        self.store = store
        self.root = (Path(root) if root else Path(store.db_path).parent / "lighthouse_assistant" / "files").resolve()

    def upload(self, actor, name, content, *, extract=True, source_scopes=None, source_modules=None, purpose=None):
        limits = upload_limit(actor, purpose)
        purpose = limits["purpose"] if limits["purpose"] != "default" else None
        if purpose:
            extract = False
        name = re.sub(r"[\x00-\x1f]", "", str(name or "文件").replace("\\", "/").split("/")[-1])[:180]
        suffix = Path(name).suffix.lower()
        if suffix in AUDIO_EXTENSIONS:
            extract = False
        if suffix not in limits["extensions"] and not (not purpose and not extract and suffix in GENERATED_EXTENSIONS):
            if purpose == "drill_template":
                raise AssistantError("演练模板仅支持 .xlsx 文件。")
            if purpose == "water_photo":
                raise AssistantError("水表照片仅支持图片。")
            raise AssistantError("支持图片、文本、代码、PDF、Word、Excel和PPT文件。")
        if not content or len(content) > limits["max_bytes"]:
            if purpose == "drill_template":
                raise AssistantError("演练模板须为1字节至64MiB。", 413)
            if purpose == "water_photo":
                raise AssistantError("水表照片须非空，单张不超过8MiB。", 413)
            raise AssistantError("单文件须为1字节至20MiB。", 413)
        mime = mimetypes.guess_type(name)[0] or "application/octet-stream"
        if suffix in IMAGE_EXTENSIONS:
            from PIL import Image
            with Image.open(io.BytesIO(content)) as image:
                mime = Image.MIME.get(image.format, mime)
                if image.width * image.height > 40000000:
                    raise AssistantError("图片像素过大，请缩小后上传。", 413)
                image.verify()
        identity = uuid.uuid4().hex
        directory = self.root / hashlib.sha256(actor["id"].encode()).hexdigest()[:24]
        directory.mkdir(parents=True, exist_ok=True)
        target = directory / (identity + suffix)
        target.write_bytes(content)
        item = {"id": identity, "owner": actor["id"], "name": name, "mime": mime, "size": len(content),
                "path": str(target), "relative_path": target.relative_to(self.root).as_posix(),
                "sha256": hashlib.sha256(content).hexdigest(), "created_at": time.time(), "text": "", "error": "", "extracted": extract}
        if suffix in AUDIO_EXTENSIONS:
            item.update(extracted=True, extraction_version=EXTRACTION_VERSION,
                        error='语音附件未转写，请使用已配置的语音识别服务。')
        topic = actor.get('_feishu_topic') if actor.get('channel', '').startswith('feishu:') else None
        if isinstance(topic, dict) and topic.get('id'):
            item['topic_id'] = topic['id']
        if source_scopes:
            item["source_scopes"] = sorted(set(source_scopes))
        if source_modules is not None or source_scopes:
            item['source_modules'] = sorted(set(source_modules if source_modules is not None else actor.get('modules', [])))
        if purpose:
            item["purpose"] = str(purpose)
        if extract:
            item['extraction_version'] = EXTRACTION_VERSION
            try:
                item.update(extract_document(content, suffix))
                if item['empty']:
                    item["error"] = "未提取到可读文字；图片可继续识别，扫描PDF请补充截图。"
            except Exception as exc:
                item["error"] = str(exc) if isinstance(exc, AssistantError) else "文件文字提取未完成，原文件已保留，可补充文字。"
        try:
            self.store.put_document(FILE_NAMESPACE, identity, item)
        except Exception:
            target.unlink(missing_ok=True)
            raise AssistantError("附件保存未完成，请重新上传。", 503) from None
        return self.public(item)

    @staticmethod
    def public(item):
        return {k: item[k] for k in ("id", "name", "mime", "size", "error")} | {"url": "/api/assistant/files/" + item["id"], "is_image": Path(item["name"]).suffix.lower() in IMAGE_EXTENSIONS}

    def get(self, actor, identity):
        if not isinstance(identity, str) or not re.fullmatch(r"[a-f0-9]{32}", identity):
            raise AssistantError("附件编号无效。", 404)
        item = self.store.get_document(FILE_NAMESPACE, identity)
        grant = actor.get('_shared_plan') or {}
        shared = identity in grant.get('file_ids', []) and actor.get('channel') == 'feishu:' + str(grant.get('chat_id'))
        if not item or item.get("owner") != actor["id"] and not shared or item.get("deleted_at"):
            raise AssistantError("附件不存在或没有访问权限。", 404)
        topic = actor.get('_feishu_topic') if actor.get('channel', '').startswith('feishu:') else None
        topic_id = topic.get('id') if isinstance(topic, dict) else None
        if item.get('topic_id') != topic_id and not (shared and topic_id and grant.get('topic_id') == topic_id):
            raise AssistantError('附件不属于当前话题或个人会话。', 404)
        if set(item.get("source_scopes", [])) - set(actor.get("allowed_scopes", actor["scopes"])):
            raise AssistantError("当前账号已无权读取该业务文件。", 403)
        if set(item.get('source_modules', [])) - set(actor.get('modules', [])):
            raise AssistantError('当前账号已无权读取该模块的文件。', 403)
        relative = item.get('relative_path')
        if relative is not None:
            if not isinstance(relative, str) or Path(relative).is_absolute() or '..' in Path(relative).parts:
                raise AssistantError("附件原文件位置无效。", 404)
            target = (self.root / relative).resolve()
        else:
            target = Path(item["path"]).resolve()
        if not target.is_relative_to(self.root) or not target.is_file():
            raise AssistantError("附件原文件不可用，请重新上传。", 404)
        return {**item, 'path': str(target)}

    def text(self, actor, identity, offset=0, length=4000):
        if type(offset) is not int or offset < 0 or type(length) is not int or length < 1:
            raise AssistantError('读取位置须为非负整数，读取长度须为正整数。')
        item = self.get(actor, identity)
        if not item.get("extracted", True) or item.get('extraction_version') != EXTRACTION_VERSION:
            if Path(item["name"]).suffix.lower() in GENERATED_EXTENSIONS - ALLOWED_EXTENSIONS:
                raise AssistantError("该文件可下载，但不支持直接提取文字，请另存为PDF或Excel新格式。")
            try:
                from .lighthouse_document_tools import checked_file
                suffix = Path(item["name"]).suffix.lower()
                item.update(extract_document(checked_file(item), suffix))
                item['error'] = '' if not item['empty'] else '未提取到可读文字，请核对原件。'
            except Exception as exc:
                item.update(text='', references=[], empty=True, extraction_complete=False)
                item["error"] = str(exc) if isinstance(exc, AssistantError) else "文件文字提取未完成，可下载原文件核对。"
            item["extracted"] = True
            item['extraction_version'] = EXTRACTION_VERSION
            self.store.put_document(FILE_NAMESPACE, identity, item)
        # safe_data also limits strings to 6000; advancing farther would silently skip text.
        return {"file_id": identity, "name": safe_text(item["name"]),
                **text_excerpt(item, offset, min(6000, length)),
                'note': item.get('note', ''), "error": item.get("error", "")}

    def context(self, actor, identities):
        if not isinstance(identities, list) or len(identities) > 10 or len(set(identities)) != len(identities):
            raise AssistantError("每次最多附带10个不同文件。")
        items = [self.get(actor, identity) for identity in identities]
        if sum(item["size"] for item in items) > MAX_BATCH_BYTES:
            raise AssistantError("附件合计不得超过100MiB。", 413)
        return [self.text(actor, item["id"], length=max(300, 5000 // max(1, len(items)))) for item in items]

    def image_parts(self, actor, identities):
        from PIL import Image
        parts = []
        for identity in identities:
            if len(parts) >= 5:
                break
            item = self.get(actor, identity)
            if Path(item["name"]).suffix.lower() not in IMAGE_EXTENSIONS:
                continue
            with Image.open(item["path"]) as image:
                image = image.convert("RGB")
                image.thumbnail((1600, 1600))
                output = io.BytesIO()
                image.save(output, format="JPEG", quality=85)
            parts.append({"type": "image_url", "image_url": {"url": "data:image/jpeg;base64," + base64.b64encode(output.getvalue()).decode()}})
        return parts
