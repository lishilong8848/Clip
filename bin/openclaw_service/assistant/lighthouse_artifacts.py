"""Bounded document creation and content previews. Generated code is never executed."""
import asyncio
import io
import json
import math
import re
from concurrent.futures import ThreadPoolExecutor
from functools import partial
from pathlib import Path
from xml.sax.saxutils import escape

from .lighthouse_ai import API_KEY, AssistantError, private_identifier

CODE_EXTENSIONS = {'.py', '.js', '.ts', '.jsx', '.tsx', '.html', '.css', '.sql', '.sh', '.ps1',
                   '.bat', '.yaml', '.yml', '.toml', '.ini', '.c', '.cpp', '.h', '.java', '.go', '.rs', '.rb'}
OFFICE_EXTENSIONS = {'.docx', '.xlsx', '.pptx', '.pdf'}
MAX_CONTENT = 512 * 1024
LONG_REPLY = 3500
PREVIEW_CHARS = 80000
_WORKERS = ThreadPoolExecutor(max_workers=2, thread_name_prefix='assistant-documents')


def attachment_analysis_requested(question):
    return bool(re.search(r'附件|文件|文档|表格|Excel|PDF|Word|PPT|CSV', question, re.I)
        and re.search(r'分析|总结|对比|比较|核对|统计|汇总|提取|整理|读取|阅读|内容|多少|哪些|生成|导出|转换|转成|转为|改写', question)
        and not re.search(r'写入|回填|导入|提交|绑定|(?:上传|更新).{0,20}(?:多维表|台账|系统|平台)|(?:发送|删除).{0,20}通告', question))


def file_generation_requested(question):
    return bool(re.search(r'生成|创建|制作|写(?:个|一个|一份)|做(?:个|一个|一份)|整理成|导出(?:为|成)?|另存为|保存为', question)
        and re.search(r'文件|Word|Excel|PDF|PPT|PowerPoint|\.\s*(?:py|txt|md|csv|docx|xlsx|pptx|pdf)\b', question, re.I)
        and not re.search(r'如何|怎么|怎样|为什么|教程|不要生成|不用生成|不需要文件|只(?:需|要)?(?:解释|说明)', question))


async def document_task(function, *args, **kwargs):
    return await asyncio.get_running_loop().run_in_executor(_WORKERS, partial(function, *args, **kwargs))


def check_content(value):
    try:
        text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, allow_nan=False)
    except (ValueError, TypeError, RecursionError):
        raise AssistantError('文档数据含无效数值或类型。') from None
    if not text.strip() or len(text.encode('utf-8')) > MAX_CONTENT:
        raise AssistantError('生成内容须非空且不超过512KiB，请拆分内容。', 413)
    if private_identifier(text) or API_KEY.search(text):
        raise AssistantError('文件包含敏感个人信息或凭证，已阻止生成。', 403)
    pending = [value]
    while pending:
        part = pending.pop()
        if isinstance(part, dict): pending.extend([*part, *part.values()])
        elif isinstance(part, list): pending.extend(part)
        elif isinstance(part, str) and any(ord(char) < 32 and char not in '\n\r\t' for char in part):
            raise AssistantError('内容包含文档不支持的控制字符。')
    return text


def document_bytes(name, data):
    """A small explicit schema prevents formulas, macros and arbitrary file access."""
    suffix = Path(name).suffix.lower()
    if suffix not in OFFICE_EXTENSIONS or not isinstance(data, dict):
        raise AssistantError('办公文件支持 .docx、.xlsx、.pptx、.pdf。')
    check_content(data)
    allowed = {'sheets'} if suffix == '.xlsx' else {'slides'} if suffix == '.pptx' else {'title', 'paragraphs'}
    if set(data) - allowed:
        raise AssistantError('该格式接受的字段为：' + '、'.join(sorted(allowed)))
    title = data.get('title', '')
    paragraphs = data.get('paragraphs', [])
    if not isinstance(title, str) or len(title) > 200 or not isinstance(paragraphs, list) or len(paragraphs) > 1000 or any(not isinstance(p, str) or len(p) > 10000 for p in paragraphs):
        raise AssistantError('标题最多200字，正文每段最多10000字、合计1000段。')
    output = io.BytesIO()
    if suffix == '.xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        sheets = data.get('sheets')
        if not isinstance(sheets, list) or not 1 <= len(sheets) <= 10:
            raise AssistantError('Excel须提供1至10个 sheets，每个含 name、headers 和 rows。')
        book = Workbook()
        book.remove(book.active)
        names = set()
        for sheet in sheets:
            if not isinstance(sheet, dict) or set(sheet) - {'name', 'headers', 'rows'}:
                raise AssistantError('工作表结构无效。')
            label, headers, rows = sheet.get('name', ''), sheet.get('headers'), sheet.get('rows')
            if not isinstance(label, str) or not label.strip() or len(label) > 31 or re.search(r'[\\/*?:\[\]]', label) or label.casefold() in names:
                raise AssistantError('工作表名称须唯一，最多31字且不含Excel保留符号。')
            if not isinstance(headers, list) or not 1 <= len(headers) <= 50 or any(not isinstance(h, str) or not h for h in headers):
                raise AssistantError('每张表须提供1至50个列标题。')
            if not isinstance(rows, list) or len(rows) > 5000 or any(not isinstance(row, list) or len(row) != len(headers) for row in rows):
                raise AssistantError('每张表最多5000行，每行单元格数须与列标题一致。')
            names.add(label.casefold())
            ws = book.create_sheet(label)
            for row_index, row in enumerate([headers, *rows], 1):
                for column, value in enumerate(row, 1):
                    if value is not None and (type(value) not in {str, int, float, bool} or isinstance(value, str) and len(value) > 32767 or isinstance(value, float) and not math.isfinite(value)):
                        raise AssistantError('单元格仅支持有限数字、文字、布尔值或空值。')
                    cell = ws.cell(row_index, column, value)
                    if isinstance(cell.value, str):
                        cell.data_type = 's'  # Model/user text must never turn into an Excel formula.
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='2261A8')
                ws.column_dimensions[cell.column_letter].width = min(42, max(16, len(cell.value) * 2 + 4))
        book.save(output)
    elif suffix == '.pptx':
        from pptx import Presentation
        from pptx.util import Inches, Pt
        slides = data.get('slides')
        if not isinstance(slides, list) or not 1 <= len(slides) <= 40:
            raise AssistantError('PPT须提供1至40个 slides，每页含 title、bullets。')
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        for item in slides:
            if not isinstance(item, dict) or set(item) - {'title', 'bullets'} or not isinstance(item.get('title'), str) or not 1 <= len(item['title']) <= 70:
                raise AssistantError('每页须提供不超过70字的标题。')
            bullets = item.get('bullets', [])
            if not isinstance(bullets, list) or len(bullets) > 6 or any(not isinstance(b, str) or len(b) > 160 for b in bullets):
                raise AssistantError('每页最多6条要点，每条最多160字，请拆页。')
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = item['title']
            slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
            frame = slide.placeholders[1].text_frame
            frame.clear()
            for i, value in enumerate(bullets):
                para = frame.paragraphs[0] if not i else frame.add_paragraph()
                para.text, para.font.size = value, Pt(18)
                para.font.name = 'Microsoft YaHei'
                para.space_after = Pt(14)
        deck.save(output)
    else:
        if not title.strip() and not any(p.strip() for p in paragraphs):
            raise AssistantError('请提供文档标题或正文。')
        if suffix == '.docx':
            from docx import Document
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            from docx.shared import Pt
            doc = Document()
            normal = doc.styles['Normal']
            normal.font.name, normal.font.size = 'Microsoft YaHei', Pt(11)
            fonts = OxmlElement('w:rFonts')
            fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
            normal.element.get_or_add_rPr().append(fonts)
            normal.paragraph_format.space_after = Pt(8)
            if title: doc.add_heading(title, 0)
            for para in paragraphs: doc.add_paragraph(para)
            doc.save(output)
        else:
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
            style = ParagraphStyle('body', fontName='STSong-Light', fontSize=11, leading=18, wordWrap='CJK')
            heading = ParagraphStyle('title', parent=style, fontSize=18, leading=26, spaceAfter=16)
            elements = [Paragraph(escape(title), heading)] if title else []
            for para in paragraphs:
                elements.extend([Paragraph(escape(para).replace('\n', '<br/>'), style), Spacer(1, 8)])
            SimpleDocTemplate(output, title=title).build(elements)
    return output.getvalue()


def create_document(files, actor, name, data, *, source_scopes=None):
    return files.upload(actor, name, document_bytes(name, data), extract=False, source_scopes=source_scopes or actor['scopes'])


def package_long_reply(files, actor, result):
    text = str(result.get('answer') or '')
    if len(text) <= LONG_REPLY:
        return result
    check_content(text)
    if str(actor.get('channel') or '').startswith('feishu:'):
        # Feishu delivers long answers in resumable message parts, not hidden attachments.
        return result
    sources = result.get('sources') or []
    source_lines = ['\n\n## 参考资料'] if sources else []
    for source in sources:
        source_lines.append(f"[{source.get('number', '')}] {source.get('title', '')} {source.get('url', '')}")
    full = text + '\n'.join(source_lines)
    check_content(full)
    scopes = sorted(set(actor['scopes']) | set(result.get('_artifact_source_scopes') or []))
    file = files.upload(actor, '完整回答.md', full.encode('utf-8'), extract=False, source_scopes=scopes)
    return {**result, 'answer': '内容较长，已整理为「完整回答.md」，可在文件预览中查看或下载。',
            'output_files': [*(result.get('output_files') or []), file],
            'file_ids': list(dict.fromkeys([*(result.get('file_ids') or []), file['id']]))}


def preview_file(files, actor, identity):
    item = files.get(actor, identity)
    path = Path(item['path'])
    suffix = path.suffix.lower()
    result = {'id': identity, 'name': item['name'], 'kind': 'download', 'truncated': False}
    if suffix in CODE_EXTENSIONS | {'.txt', '.md', '.csv', '.log', '.json', '.xml'}:
        with path.open('rb') as source:
            content = source.read(PREVIEW_CHARS + 1)
        result.update(kind='html' if suffix == '.html' else 'markdown' if suffix == '.md' else 'text',
                      text=content[:PREVIEW_CHARS].decode('utf-8-sig', errors='replace'), truncated=len(content) > PREVIEW_CHARS)
    elif suffix in {'.png', '.jpg', '.jpeg', '.webp', '.gif', '.bmp', '.pdf'}:
        result['kind'] = 'pdf' if suffix == '.pdf' else 'image'
    elif suffix in {'.docx', '.xlsx', '.xlsm', '.pptx'}:
        from .lighthouse_files import _check_archive
        _check_archive(path.read_bytes())
        if suffix == '.docx':
            from docx import Document
            doc = Document(path)
            lines = [p.text for p in doc.paragraphs]
            lines.extend('\t'.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
            text = '\n\n'.join(lines)
            result.update(kind='text', text=text[:PREVIEW_CHARS], truncated=len(text) > PREVIEW_CHARS)
        elif suffix == '.pptx':
            from pptx import Presentation
            deck = Presentation(path)
            result.update(kind='slides', slides=[{'title': f'第{i + 1}页', 'text': '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)[:12000]}
                                                for i, slide in enumerate(list(deck.slides)[:40])], truncated=len(deck.slides) > 40)
        else:
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                result.update(kind='spreadsheet', sheets=[])
                remaining = PREVIEW_CHARS
                for ws in list(book)[:10]:
                    rows = []
                    for row in ws.iter_rows(max_row=min(ws.max_row or 1, 200), max_col=min(ws.max_column or 1, 50), values_only=True):
                        values = [str(value) if value is not None else '' for value in row]
                        cost = sum(len(v) for v in values) + len(values)
                        if cost > remaining:
                            result['truncated'] = True
                            break
                        rows.append(values)
                        remaining -= cost
                    result['sheets'].append({'name': ws.title, 'rows': rows})
                    result['truncated'] |= (ws.max_row or 0) > 200 or (ws.max_column or 0) > 50
                result['truncated'] |= len(book.sheetnames) > 10
            finally:
                book.close()
    return result
