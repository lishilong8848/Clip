"""Attachment-only table analysis and static PDF page operations."""
import csv
import datetime as dt
import io
import json
from collections import defaultdict
from decimal import localcontext
from itertools import zip_longest
from pathlib import Path

from .lighthouse_ai import AssistantError, safe_data
from .lighthouse_artifacts import document_bytes
from .lighthouse_mcp_office import number


def prepare_inputs(files, actor, inputs, kind):
    if not isinstance(inputs, list) or not 1 <= len(inputs) <= 6:
        raise AssistantError('每次请选择1至6份本人可访问的附件。')
    allowed = {'file_id', 'sheet', 'header_row'} if kind == 'spreadsheet' else {'file_id', 'pages'}
    result, total = [], 0
    for value in inputs:
        if not isinstance(value, dict) or set(value) - allowed or 'file_id' not in value:
            raise AssistantError('请使用附件编号及允许的工作表、表头行或页码，不接受本机路径。')
        item = files.get(actor, value['file_id'])
        suffix = Path(item['name']).suffix.lower()
        if suffix not in ({'.xlsx', '.xlsm', '.csv'} if kind == 'spreadsheet' else {'.pdf'}):
            raise AssistantError('表格支持xlsx、xlsm和CSV；PDF整理仅接受PDF附件。')
        total += item['size']
        if total > 100 * 1024 * 1024:
            raise AssistantError('所选附件合计最多100MiB。', 413)
        result.append({**value, 'name': item['name'], 'path': item['path'], 'sha256': item['sha256'],
                       'source_scopes': item.get('source_scopes', [])})
    return result


def _cell(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, str) and len(value) > 6000:
        raise AssistantError('单元格文字过长，请精简后处理。', 413)
    return value


def _table(item, inspect=False):
    from .lighthouse_document_tools import checked_file
    from .lighthouse_files import _check_archive
    content = checked_file(item)
    header_row = item.get('header_row', 1)
    if type(header_row) is not int or not 1 <= header_row <= 100:
        raise AssistantError('表头行须为1至100之间的整数。')
    suffix = Path(item['name']).suffix.lower()
    books = []
    try:
        if suffix == '.csv':
            text = None
            for encoding in ('utf-8-sig', 'gb18030'):
                try:
                    text = content.decode(encoding)
                    break
                except UnicodeError:
                    continue
            if text is None:
                raise AssistantError('CSV编码无法识别，请另存为UTF-8。')
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            names = ['CSV']
            rows = ((i, [v if v != '' else None for v in row]) for i, row in enumerate(csv.reader(io.StringIO(text), dialect), 1))
        else:
            from openpyxl import load_workbook
            _check_archive(content)
            for data_only in (True, False):
                books.append(load_workbook(io.BytesIO(content), read_only=True, data_only=data_only, keep_links=False))
            names = books[0].sheetnames
        sheet = item.get('sheet', '')
        if not isinstance(sheet, str):
            raise AssistantError('工作表名称格式无效。')
        if not sheet and len(names) > 1:
            if inspect:
                return {'file_id': item['file_id'], 'name': item['name'], 'sheets': names[:40], 'total_sheets': len(names),
                        'truncated': len(names) > 40, 'partial': True, 'requires_sheet': True,
                        'coverage': 'sheet_catalog_only'}
            raise AssistantError('附件有多个工作表，请先选择工作表：' + '、'.join(names[:20]))
        sheet = sheet or names[0]
        if sheet not in names:
            raise AssistantError('工作表不存在，请先检查附件目录。')
        if books:
            sheets = [book[sheet] for book in books]
            for ws in sheets:
                ws.reset_dimensions()

            def excel_rows():
                for index, pair in enumerate(zip_longest(sheets[0].iter_rows(), sheets[1].iter_rows()), 1):
                    cached, formula = pair
                    if cached is None or formula is None or len(cached) != len(formula):
                        raise AssistantError('工作表结构读取不完整，未使用部分数据。')
                    values = []
                    for old, source in zip(cached, formula):
                        if index >= header_row and source.data_type == 'f' and old.value is None:
                            raise AssistantError('存在没有计算缓存的公式，请先用Excel计算并保存；不会把缺失结果当成0。')
                        if source.data_type == 'e' or old.data_type == 'e':
                            raise AssistantError('工作表包含公式错误，请修正后处理。')
                        values.append(_cell(old.value))
                    yield index, values
            rows = excel_rows()
        headers, indexes, data = None, [], []
        for line, row in rows:
            if line > 100000:
                raise AssistantError('工作表行范围过大，请删除多余空行或拆分文件。', 413)
            if line < header_row:
                continue
            populated = {i for i, value in enumerate(row) if value is not None and value != ''}
            if line == header_row:
                indexes = sorted(populated)
                headers = [str(row[i]).strip() for i in indexes]
                if not 1 <= len(headers) <= 40 or any(not h for h in headers) or len(set(headers)) != len(headers):
                    raise AssistantError('表头须为1至40个非空且不重复的列名，请选择正确表头行。')
                continue
            if not populated:
                continue
            if populated - set(indexes):
                raise AssistantError('数据列没有对应表头，请选择正确表头行，不会丢弃无表头数据。')
            values = [_cell(row[i]) if i < len(row) else None for i in indexes]
            data.append({'line': line, 'values': dict(zip(headers, values))})
            if len(data) > 5000:
                raise AssistantError('每次分析最多5000条非空数据，请拆分文件。', 413)
        if headers is None:
            raise AssistantError('没有找到指定表头行。')
        from openpyxl.utils import get_column_letter
        return {'file_id': item['file_id'], 'name': item['name'], 'sheet': sheet, 'headers': headers, 'rows': data,
                'header_row': header_row, 'column_letters': [get_column_letter(i + 1) for i in indexes],
                'first_data_row': data[0]['line'] if data else None, 'last_data_row': data[-1]['line'] if data else None,
                'row_count': len(data), 'partial': False, 'coverage': 'selected_sheet',
                'note': '使用完整所选工作表中保存的值；公式读取缓存，不重新计算。不代表读取了其他工作表。'}
    finally:
        for book in books:
            book.close()


def _key(row, names):
    values = [row[name] for name in names]
    if any(value is None or value == '' for value in values):
        raise AssistantError('关键列存在空值，请先补全，不按空值关联记录。')
    return tuple((type(value).__name__, json.dumps(value, ensure_ascii=False, allow_nan=False)) for value in values)


def analyze_tables(inputs, operation, keys=None, group_by=None, value_column='', aggregate='sum'):
    if operation not in {'inspect', 'merge', 'compare', 'pivot'}:
        raise AssistantError('表格操作支持inspect、merge、compare、pivot。')
    tables = [_table(item, inspect=operation == 'inspect') for item in inputs]
    if operation == 'inspect':
        return {'files': [], 'summary': safe_data({'tables': [{k: v for k, v in table.items() if k != 'rows'} for table in tables]})}
    if sum(table['row_count'] for table in tables) > 5000:
        raise AssistantError('所选工作表合计最多5000条数据，请拆分处理。', 413)
    columns = list(dict.fromkeys(column for table in tables for column in table['headers']))
    if len(columns) > 40:
        raise AssistantError('合并后超过40个数据列，请缩小范围。', 413)
    for names in (keys or [], group_by or []):
        if not isinstance(names, list) or len(names) > 3 or any(not isinstance(k, str) or any(k not in t['headers'] for t in tables) for k in names) or len(names) != len(set(names)):
            raise AssistantError('关键列或分组列须为各表共有的1至3个不同列名。')
    rows = [(table, row) for table in tables for row in table['rows']]
    sheets, summary = [], {'input_rows': len(rows), 'operation': operation,
        'sources': [{k: v for k, v in table.items() if k != 'rows'} for table in tables],
        'note': '使用完整所选数据，关键列按原值及类型匹配，不删除前导零；缺失值不当作0。'}
    if operation == 'merge':
        if {'来源文件', '来源工作表', '原始行号'} & set(columns):
            raise AssistantError('原表已有来源文件、来源工作表或原始行号列，请更名后合并。')
        headers = ['来源文件', '来源工作表', '原始行号', *columns]
        export = [[t['name'], t['sheet'], r['line'], *[r['values'].get(c) for c in columns]] for t, r in rows]
        sheets.append({'name': '合并数据', 'headers': headers, 'rows': export})
        if keys:
            groups = defaultdict(list)
            for index, (_, row) in enumerate(rows):
                groups[_key(row['values'], keys)].append(index)
            repeated = [index for indexes in groups.values() if len(indexes) > 1 for index in indexes]
            sheets.append({'name': '重复关键列', 'headers': headers, 'rows': [export[index] for index in repeated]})
            summary['duplicate_rows'] = len(repeated)
        summary['output_rows'] = len(export)
    elif operation == 'compare':
        if len(tables) != 2 or not keys:
            raise AssistantError('核对差异须选择前后两张表，并提供关键列。')
        indexed = []
        for table in tables:
            values = {}
            for row in table['rows']:
                key = _key(row['values'], keys)
                if key in values:
                    raise AssistantError('关键列存在重复，无法唯一关联；请先用合并工具列出重复项。')
                values[key] = row
            indexed.append(values)
        before, after = indexed
        added = [r for k, r in after.items() if k not in before]
        removed = [r for k, r in before.items() if k not in after]
        changes, changed_rows = [], 0
        for key, old in before.items():
            if key not in after:
                continue
            new = after[key]
            changed = False
            for column in columns:
                left, right = old['values'].get(column), new['values'].get(column)
                if type(left) != type(right) or left != right or (column in old['values']) != (column in new['values']):
                    changed = True
                    changes.append([*[new['values'][k] for k in keys], column, left, right,
                                    old['line'], new['line']])
            changed_rows += changed
        for name, data in (('新增', added), ('移除', removed)):
            sheets.append({'name': name, 'headers': ['原始行号', *columns],
                           'rows': [[r['line'], *[r['values'].get(c) for c in columns]] for r in data]})
        sheets.append({'name': '字段变动', 'headers': [*keys, '变动字段', '原值', '新值', '原表行号', '新表行号'], 'rows': changes})
        summary.update(added=len(added), removed=len(removed), changed=changed_rows, changed_cells=len(changes),
                       unchanged=len(before.keys() & after.keys()) - changed_rows)
    else:
        if not group_by or aggregate not in {'sum', 'mean', 'count', 'min', 'max'}:
            raise AssistantError('分类汇总须提供分组列，聚合支持sum、mean、count、min、max。')
        if aggregate != 'count' and (not value_column or any(value_column not in t['headers'] for t in tables)):
            raise AssistantError('请提供各表共有的数值列。')
        groups = {}
        for _, row in rows:
            record = row['values']
            group = tuple(record[c] for c in group_by)
            encoded = tuple((type(v).__name__, str(v)) for v in group)
            entry = groups.setdefault(encoded, {'group': list(group), 'rows': 0, 'values': [], 'missing': 0})
            entry['rows'] += 1
            if aggregate != 'count':
                value = record[value_column]
                if value is None or value == '':
                    entry['missing'] += 1
                else:
                    try:
                        entry['values'].append(number(value))
                    except ValueError:
                        raise AssistantError('数值列包含无法可靠计算的内容，请先核对原数据。') from None
        export = []
        for entry in groups.values():
            values = entry['values']
            result = entry['rows'] if aggregate == 'count' else None
            if values:
                with localcontext() as context:
                    context.prec = 50
                    result = str({'sum': lambda: sum(values), 'mean': lambda: sum(values) / len(values),
                                  'min': lambda: min(values), 'max': lambda: max(values)}[aggregate]())
            export.append([*entry['group'], result, entry['rows'], entry['missing']])
        sheets.append({'name': '分类汇总', 'headers': [*group_by, aggregate + (':' + value_column if value_column else ''), '原始条数', '数值缺失条数'], 'rows': export})
        summary['groups'] = len(groups)
    content = document_bytes('表格分析.xlsx', {'sheets': sheets})
    return {'summary': safe_data(summary), 'files': [{'name': '表格分析.xlsx', 'content': content}]}


def organize_pdf(inputs, operation, split_every=1, rotation=0, max_outputs=10):
    from pypdf import PdfReader, PdfWriter
    from .lighthouse_document_tools import checked_file
    if operation not in {'merge', 'extract', 'split', 'rotate'} or type(rotation) is not int or rotation not in {0, 90, 180, 270}:
        raise AssistantError('PDF支持merge、extract、split、rotate；旋转角度为0/90/180/270。')
    if operation != 'merge' and len(inputs) != 1:
        raise AssistantError('提取、拆分或旋转每次请选择一份PDF。')
    if type(split_every) is not int or not 1 <= split_every <= 300:
        raise AssistantError('每份页数须为1至300之间的整数。')
    selected = []
    for item in inputs:
        reader = PdfReader(io.BytesIO(checked_file(item)))
        if reader.is_encrypted or not 1 <= len(reader.pages) <= 300:
            raise AssistantError('仅处理未加密且不超过300页的PDF。')
        pages = item.get('pages', list(range(1, len(reader.pages) + 1)))
        if not isinstance(pages, list) or not 1 <= len(pages) <= 300 or any(type(page) is not int or not 1 <= page <= len(reader.pages) for page in pages):
            raise AssistantError('页码从1开始，不能超出文档范围；省略pages表示全部页。')
        selected.extend(reader.pages[page - 1] for page in pages)
    if len(selected) > 300:
        raise AssistantError('输出合计最多300页，请分批处理。', 413)
    groups = [selected[i:i + split_every] for i in range(0, len(selected), split_every)] if operation == 'split' else [selected]
    if len(groups) > max_outputs:
        raise AssistantError('拆分文件数量超出本轮剩余额度，请增加每份页数或缩小页码范围。')
    outputs = []
    try:
        for index, pages in enumerate(groups, 1):
            writer = PdfWriter()
            try:
                for page in pages:
                    copied = writer.add_page(page, excluded_keys=('/Annots', '/AA'))
                    if rotation:
                        copied.rotate(rotation)
                output = io.BytesIO()
                writer.write(output)
                content = output.getvalue()
                if len(content) > 20 * 1024 * 1024:
                    raise AssistantError('生成文件超过20MiB，请拆分页码范围。', 413)
                outputs.append({'name': f'整理结果_{index:02}.pdf', 'content': content})
            finally:
                writer.close()
        return {'summary': {'operation': operation, 'pages': len(selected), 'output_files': len(outputs),
            'note': '原件未改动；新文件保留静态页面，不保留交互表单、链接与数字签名效力。'}, 'files': outputs}
    finally:
        selected.clear()
