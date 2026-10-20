"""Synthetic provenance regressions; private temp files and fake models/Feishu only."""
import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree

sys.path.insert(0, str(Path(__file__).resolve().parent))
from bin.test_lighthouse_document_tools import Store, ACTOR
from openclaw_service.assistant.lighthouse_ai import AssistantError, safe_data
from openclaw_service.assistant.lighthouse_files import FILE_NAMESPACE, LighthouseFiles, extract_document
from openclaw_service.assistant.lighthouse_document_tools import read_attachment
from openclaw_service.assistant.lighthouse_office_tools import prepare_inputs, analyze_tables


def sparse_workbook(cached=False):
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = 'Inventory'
    sheet['A3'], sheet['C3'] = 'Device', 'Count'
    sheet['A7'], sheet['C7'] = '001', 0
    sheet['A11'], sheet['C11'] = '002', '=1+2'
    book.create_sheet('Other')['B9'] = 'separate source'
    output = io.BytesIO()
    book.save(output)
    book.close()
    if not cached:
        return output.getvalue()
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as original, zipfile.ZipFile(result, 'w') as target:
        for entry in original.infolist():
            content = original.read(entry.filename)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                root = ElementTree.fromstring(content)
                ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                root.find(".//s:c[@r='C11']/s:v", ns).text = '3'
                content = ElementTree.tostring(root)
            target.writestr(entry, content)
    return result.getvalue()


class FileProvenanceTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.files = LighthouseFiles(Store(Path(temporary.name) / 'state.sqlite3'))

    def upload(self, name, content, **options):
        return self.files.upload(ACTOR, name, content, **options)['id']

    def test_excel_sparse_rows_columns_formula_and_original_are_preserved(self):
        from openpyxl import load_workbook
        for suffix in ('.xlsx', '.xlsm'):
            original = sparse_workbook()
            identity = self.upload('source' + suffix, original)
            result = self.files.text(ACTOR, identity)
            rows = [ref for ref in result['references'] if ref['kind'] == 'row']
            self.assertEqual([(ref['sheet'], ref['row'], ref['cells']) for ref in rows],
                [('Inventory', 3, 'A3:C3'), ('Inventory', 7, 'A7:C7'),
                 ('Inventory', 11, 'A11:C11'), ('Other', 9, 'A9:B9')])
            self.assertIn('001\t\t0', result['text'])
            self.assertIn('=1+2 [cached=missing; not recalculated]', result['text'])
            self.assertFalse(result['partial'])
            self.assertEqual(result['coverage'], 'extractable_text')
            path = Path(self.files.get(ACTOR, identity)['path'])
            self.assertEqual(path.read_bytes(), original)
            book = load_workbook(io.BytesIO(path.read_bytes()), data_only=False)
            self.assertEqual(book['Inventory']['C11'].value, '=1+2')
            book.close()

    def test_enhanced_excel_retains_formula_cache_and_cell_references(self):
        identity = self.upload('source.xlsx', sparse_workbook(cached=True), extract=False)
        result = read_attachment(self.files, ACTOR, identity)
        self.assertIn('=1+2 [cached=3; not recalculated]', result['text'])
        self.assertIn('| 001 |  | 0 |', result['text'])
        row = next(ref for ref in result['references'] if ref.get('row') == 11)
        self.assertEqual(row['cells'], 'A11:C11')
        self.assertIn('=1+2', result['text'][row['start']:row['end']])

    def test_word_sections_paragraphs_and_tables_have_real_locations(self):
        from docx import Document
        document = Document()
        document.add_heading('Checks', 1)
        paragraph = document.add_paragraph('Device')
        paragraph.add_run().add_tab()
        paragraph.add_run('D-101')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = 'D-101', 'OK'
        document.add_paragraph('After table')
        output = io.BytesIO()
        document.save(output)
        identity = self.upload('checks.docx', output.getvalue())
        for result in (self.files.text(ACTOR, identity), read_attachment(self.files, ACTOR, identity)):
            table_ref = next(ref for ref in result['references'] if ref['kind'] == 'table_row')
            self.assertEqual((table_ref['table'], table_ref['row'], table_ref['section']), (1, 1, 'Checks'))
            self.assertTrue(all('page' not in ref for ref in result['references']))
            self.assertIn('Device\tD-101', result['text'])
            self.assertEqual(result['references'][-1]['paragraph'], 5)

    def test_pdf_missing_text_page_is_partial_even_without_more_characters(self):
        from bin.test_lighthouse_office_tools import pdf_bytes
        identity = self.upload('mixed.pdf', pdf_bytes(['Page one', '', 'Page three']))
        for result in (self.files.text(ACTOR, identity), read_attachment(self.files, ACTOR, identity)):
            self.assertEqual([ref['page'] for ref in result['references']], [1, 3])
            self.assertEqual(result['pages_without_text'], [2])
            self.assertEqual(result['total_pages'], 3)
            self.assertFalse(result['extraction_complete'])
            self.assertTrue(result['partial'])
            self.assertIsNone(result['next_offset'])
            self.assertFalse(result['truncated'])

    def test_continuation_does_not_skip_characters_in_safe_data(self):
        content = '\n'.join('line %03d %s' % (index, 'x' * 90) for index in range(90))
        identity = self.upload('long.txt', content.encode())
        first = safe_data(self.files.text(ACTOR, identity, length=8000))
        self.assertLessEqual(len(first['text']), 6000)
        self.assertEqual(first['next_offset'], len(first['text']))
        tail = self.files.text(ACTOR, identity, offset=first['next_offset'], length=8000)
        self.assertEqual(first['text'] + tail['text'], content)
        self.assertTrue(tail['partial'])
        self.assertFalse(tail['truncated'])
        self.assertTrue(first['references_truncated'])
        self.assertEqual(len(first['references']), 40)
        self.assertGreater(first['reference_count'], 40)
        self.assertGreater(first['next_reference_offset'], first['references'][-1]['start'])

    def test_excerpt_mid_row_retains_source_and_marks_cut_reference(self):
        identity = self.upload('source.xlsx', sparse_workbook())
        first = self.files.text(ACTOR, identity)
        ref = next(ref for ref in first['references'] if ref.get('row') == 11)
        part = self.files.text(ACTOR, identity, offset=ref['start'] + 3, length=12)
        self.assertEqual(part['references'][0]['cells'], 'A11:C11')
        self.assertTrue(part['references'][0]['partial'])

    def test_old_cached_text_is_reextracted_from_authorized_original(self):
        identity = self.upload('source.xlsx', sparse_workbook())
        item = self.files.get(ACTOR, identity)
        item.pop('extraction_version')
        item.update(text='old flat text', references=[])
        self.files.store.put_document(FILE_NAMESPACE, identity, item)
        self.assertIn('Inventory!A11:C11', self.files.text(ACTOR, identity)['text'])
        identity = self.upload('changed.txt', b'original', extract=False)
        Path(self.files.get(ACTOR, identity)['path']).write_bytes(b'changed')
        result = self.files.text(ACTOR, identity)
        self.assertTrue(result['partial'])
        self.assertFalse(result['extraction_complete'])
        self.assertEqual(result['text'], '')
        self.assertIn('重新上传', result['error'])

    def test_cross_file_context_is_bounded_and_each_file_has_coverage(self):
        identities = [self.upload('source%d.txt' % i, ('x' * 100 + '\n').encode() * 70) for i in range(10)]
        results = self.files.context(ACTOR, identities)
        self.assertLessEqual(sum(len(result['text']) for result in results), 5000)
        self.assertEqual([result['file_id'] for result in results], identities)
        self.assertTrue(all(result['partial'] and result['next_offset'] for result in results))
        with self.assertRaises(AssistantError):
            self.files.context(ACTOR, identities + [identities[0]])

    def test_acl_is_rechecked_for_both_readers_and_channels_after_cache(self):
        identity = self.upload('source.xlsx', sparse_workbook(), source_scopes=['D'])
        read_attachment(self.files, ACTOR, identity)
        for channel in ('', 'feishu:fixture'):
            for actor in ({**ACTOR, 'id': 'other', 'is_admin': True}, {**ACTOR, 'scopes': ['A']},
                          {**ACTOR, '_feishu_topic': {'id': 'other-topic'}, 'channel': 'feishu:fixture'}):
                actor = {**actor, 'channel': actor.get('channel', channel)}
                for reader in (self.files.text, lambda a, f: read_attachment(self.files, a, f)):
                    with self.assertRaises(AssistantError):
                        reader(actor, identity)

    def test_ocr_last_page_is_still_only_one_page(self):
        identity = self.upload('scan.pdf', b'fixture', extract=False)
        result = {'text': 'recognized page', 'total_pages': 3, 'page': 3, 'next_page': None,
                  'extraction_complete': False, 'references': [{'kind': 'page_ocr', 'page': 3, 'start': 0, 'end': 15}]}
        with patch('openclaw_service.assistant.lighthouse_document_tools.run_local', return_value=result):
            excerpt = read_attachment(self.files, ACTOR, identity, mode='ocr', page=3)
        self.assertTrue(excerpt['partial'])
        self.assertEqual(excerpt['references'][0]['page'], 3)
        self.assertIsNone(excerpt['next_page'])

    def test_table_inspection_and_comparison_summary_include_original_columns(self):
        identity = self.upload('source.xlsx', sparse_workbook(cached=True), extract=False)
        inputs = prepare_inputs(self.files, ACTOR, [{'file_id': identity, 'sheet': 'Inventory', 'header_row': 3}], 'spreadsheet')
        result = analyze_tables(inputs, 'inspect')['summary']['tables'][0]
        self.assertEqual(result['file_id'], identity)
        self.assertEqual(result['column_letters'], ['A', 'C'])
        self.assertEqual((result['header_row'], result['first_data_row'], result['last_data_row']), (3, 7, 11))
        result = analyze_tables(inputs, 'merge')['summary']
        self.assertEqual(result['sources'][0]['row_count'], 2)
        self.assertEqual(result['sources'][0]['coverage'], 'selected_sheet')

    def test_failed_extraction_and_invalid_ranges_never_claim_complete(self):
        identity = self.upload('broken.pdf', b'not a pdf')
        result = self.files.text(ACTOR, identity)
        self.assertTrue(result['error'])
        self.assertTrue(result['partial'])
        self.assertEqual(result['references'], [])
        for offset, length in ((True, 1), (-1, 10), (0, 0), (0, '100')):
            with self.assertRaises(AssistantError):
                self.files.text(ACTOR, identity, offset, length)

    def test_feishu_document_tail_has_partial_and_text_line_reference(self):
        from bin.test_lighthouse_documents import service, URL, ACTOR as DOCUMENT_ACTOR
        from openclaw_service.assistant.lighthouse_documents import read_document
        content = 'A楼\n' + ('line\n' * 1500)
        first = read_document(service(content=content), DOCUMENT_ACTOR, URL)
        tail = read_document(service(content=content), DOCUMENT_ACTOR, URL, first['next_offset'])
        self.assertTrue(tail['partial'])
        self.assertFalse(tail['truncated'])
        self.assertEqual(tail['references'][0]['basis'], 'sanitized_text')
        self.assertGreater(tail['references'][0]['line_start'], 1)

    def test_word_content_controls_are_read_in_document_order(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Before')
        paragraph = document.add_paragraph('Inside content control')._p
        from docx.oxml import OxmlElement
        wrapper, content = OxmlElement('w:sdt'), OxmlElement('w:sdtContent')
        paragraph.addprevious(wrapper)
        wrapper.append(content)
        content.append(paragraph)
        document.add_paragraph('After')
        output = io.BytesIO()
        document.save(output)
        result = extract_document(output.getvalue(), '.docx')
        self.assertEqual([ref['paragraph'] for ref in result['references']], [1, 2, 3])
        self.assertIn('Inside content control', result['text'])

    def test_array_formulas_use_the_expression_instead_of_object_repr(self):
        from openpyxl import Workbook
        from openpyxl.worksheet.formula import ArrayFormula
        book = Workbook()
        book.active['A1'] = ArrayFormula(ref='A1:A2', text='=ROW(A1:A2)')
        output = io.BytesIO()
        book.save(output)
        book.close()
        result = extract_document(output.getvalue(), '.xlsx')
        self.assertIn('=ROW(A1:A2) [cached=missing;', result['text'])
        self.assertNotIn('object at', result['text'])

    def test_filtered_text_keeps_safe_source_line_numbers(self):
        result = extract_document(b'first\npassword: fixture-only\nthird\n', '.txt')
        self.assertEqual(result['text'], 'first\n\nthird')
        self.assertEqual([ref['line'] for ref in result['references']], [1, 3])
        self.assertNotIn('fixture-only', str(result))

    def test_unicode_normalization_cannot_silently_cut_extracted_text(self):
        from openclaw_service.assistant.lighthouse_files import safe_document
        with self.assertRaises(AssistantError):
            safe_document('\ufb03' * 2500)
        from bin.test_lighthouse_documents import service, URL, ACTOR as DOCUMENT_ACTOR
        from openclaw_service.assistant.lighthouse_documents import read_document
        with self.assertRaises(AssistantError):
            read_document(service(content='A楼' + '\ufb03' * 400000), DOCUMENT_ACTOR, URL)


class SharedToolProvenanceTests(unittest.IsolatedAsyncioTestCase):
    async def test_web_and_feishu_receive_the_same_excel_provenance(self):
        from bin import test_lighthouse_stream as fixtures
        from bin.test_lighthouse_general_tools import GeneralToolTests
        await fixtures.StreamTests.asyncSetUp(self)
        self.addAsyncCleanup(self.runtime.close)
        file = self.portal.files.upload(self.actor, 'source.xlsx', sparse_workbook())
        for channel in ('', 'feishu:fixture'):
            _, calls = await GeneralToolTests.run_tools(self, '分析这个附件中的表格',
                [('read_file', {'file_id': file['id']})], file_ids=[file['id']], channel=channel)
            result = next(part.content for message in calls[-1] for part in message.parts
                          if getattr(part, 'part_kind', '') == 'tool-return')
            self.assertTrue(any(ref.get('cells') == 'A11:C11' for ref in result['references']))
            self.assertIn('cached=missing', result['text'])
        self.assertEqual(self.reads, [])


if __name__ == '__main__':
    unittest.main()
