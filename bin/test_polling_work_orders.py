# -*- coding: utf-8 -*-
from __future__ import annotations

import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
import zipfile
import copy
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import MagicMock, patch

BIN = Path(__file__).resolve().parent
if str(BIN) not in sys.path:
    sys.path.insert(0, str(BIN))

from lan_bitable_template_portal.local_notice_images import LocalNoticeImageStore
from lan_bitable_template_portal.polling_work_orders import (
    PollingSopCloudStore,
    PollingWorkOrderService,
    PollingWorkOrderTokenError,
)
from lan_bitable_template_portal.state_store import LanPortalStateStore
import lan_bitable_template_portal.server as portal_server
from lan_bitable_template_portal.server import PortalRuntime
from lan_bitable_template_portal.portal_service import MaintenancePortalService
from clipflow_backend.main import FastAPIPortalController
from fastapi.testclient import TestClient
from lan_bitable_template_portal.workbench_lite import (
    render_polling_work_order_page,
    render_polling_work_order_steps_page,
    render_workbench_lite,
)
from upload_event_module.services.handlers.base import NoticePayload
from upload_event_module.services.handlers.polling_notice import PollingNoticeHandler
from upload_event_module.services.handlers.maintenance_notice import MaintenanceNoticeHandler


def _png_bytes(color: str = "#1678ff", size: tuple[int, int] = (160, 100)) -> bytes:
    from PIL import Image

    output = io.BytesIO()
    Image.new("RGB", size, color).save(output, format="PNG")
    return output.getvalue()


class _FakePollingSopCloud:
    def __init__(self, sops=None, content: bytes = b"") -> None:
        self.sops = {item["sop_id"]: copy.deepcopy(item) for item in (sops or [])}
        self.content = content
        self.saved = []
        self.uploaded = []
        self.list_calls = 0

    def list_sops(self, *, force: bool = False):
        self.list_calls += 1
        return copy.deepcopy(list(self.sops.values()))

    def get_sop(self, sop_id: str, *, force: bool = False):
        return copy.deepcopy(self.sops.get(sop_id))

    def save_sop(self, sop: dict, *, expected_version: int, allow_create: bool = False):
        self.saved.append((copy.deepcopy(sop), expected_version, allow_create))
        self.sops[sop["sop_id"]] = copy.deepcopy(sop)
        return copy.deepcopy(sop)

    def download_attachment(self, attachment: dict) -> bytes:
        return self.content

    def upload_attachment(self, path: Path, file_name: str) -> str:
        self.uploaded.append((str(path), file_name))
        return "cloud-file-token"

    def delete_sop(self, sop: dict) -> None:
        self.sops.pop(sop["sop_id"], None)


class PollingWorkOrderTests(unittest.TestCase):
    def _refresh_fixture(self, root):
        content = b"old-guide"
        sop = {"sop_id": "refresh_sop_1234", "scope": "A", "work_type": "polling", "name": "同步测试", "version": 1,
               "steps": [{"step_id": "refresh_step_1234", "content": "检查设备", "operator_required": True}],
               "attachments": [{"attachment_id": "refresh_file_1234", "file_token": "refresh_token_1234", "name": "guide.txt", "size": len(content), "sha256": hashlib.sha256(content).hexdigest()}]}
        cloud = _FakePollingSopCloud([sop], content)
        service = PollingWorkOrderService(LanPortalStateStore(root / "state.sqlite3"), cloud)
        service.sop_root = root / "sops"
        service.list_sops("A")
        return service, cloud, sop

    def test_manual_refresh_updates_initialized_cache_without_cloud_writes(self):
        with tempfile.TemporaryDirectory() as temp:
            service, cloud, sop = self._refresh_fixture(Path(temp))
            store = service.state_store
            store.put_document("polling_sop", "local_only_1234", {**sop, "sop_id": "local_only_1234", "name": "本地独有"})
            store.put_document("polling_sop", "other_scope_1234", {**sop, "sop_id": "other_scope_1234", "scope": "B"})
            store.put_document("polling_sop", "other_type_1234", {**sop, "sop_id": "other_type_1234", "work_type": "maintenance"})
            cloud.sops[sop["sop_id"]]["steps"].append({"step_id": "new_step_1234", "content": "新增步骤", "operator_required": True})
            cloud.sops[sop["sop_id"]]["version"] = 2
            cloud.sops["new_sop_123456"] = {**copy.deepcopy(sop), "sop_id": "new_sop_123456", "name": "新增 SOP"}
            self.assertEqual(len(service.get_sop(sop["sop_id"])["steps"]), 1)
            result = service.refresh_sops("A")
            self.assertEqual(result["synced_count"], 2)
            self.assertEqual(len(result["items"]), 3)
            self.assertEqual(len(service.get_sop(sop["sop_id"])["steps"]), 2)
            count = cloud.list_calls
            service.list_sops("A")
            self.assertEqual(cloud.list_calls, count)
            self.assertEqual((cloud.saved, cloud.uploaded), ([], []))
            self.assertEqual(store.get_document("polling_sop", "other_scope_1234")["scope"], "B")
            self.assertEqual(store.get_document("polling_sop", "other_type_1234")["work_type"], "maintenance")

    def test_manual_refresh_failure_keeps_local_data_and_allows_retry(self):
        with tempfile.TemporaryDirectory() as temp:
            service, cloud, sop = self._refresh_fixture(Path(temp))
            before = service.get_sop(sop["sop_id"])
            cloud.sops[sop["sop_id"]]["name"] = "尚未提交的更新"
            cloud.sops["invalid_sop_1234"] = {**sop, "sop_id": "invalid_sop_1234", "name": ""}
            with self.assertRaisesRegex(Exception, "名称"):
                service.refresh_sops("A")
            self.assertEqual(service.get_sop(sop["sop_id"]), before)
            del cloud.sops["invalid_sop_1234"]
            with patch.object(cloud, "list_sops", side_effect=TimeoutError("cloud timeout")), self.assertRaises(TimeoutError):
                service.refresh_sops("A")
            self.assertEqual(service.get_sop(sop["sop_id"]), before)
            self.assertEqual(service.refresh_sops("A")["items"][0]["name"], "尚未提交的更新")

    def test_manual_refresh_changed_attachment_does_not_reuse_old_content(self):
        with tempfile.TemporaryDirectory() as temp:
            service, cloud, sop = self._refresh_fixture(Path(temp))
            service.get_sop_attachment(sop["sop_id"], "refresh_file_1234")
            old = Path(service.get_sop(sop["sop_id"], public=False)["attachments"][0]["path"])
            cloud.content = b"new-guide"
            attachment = cloud.sops[sop["sop_id"]]["attachments"][0]
            attachment.update(file_token="new_token_12345", sha256=hashlib.sha256(cloud.content).hexdigest(), size=len(cloud.content))
            service.refresh_sops("A")
            self.assertEqual(service.get_sop_attachment(sop["sop_id"], "refresh_file_1234")[0], b"new-guide")
            self.assertEqual(old.read_bytes(), b"old-guide")

    def test_manual_refresh_scope_marker_keeps_future_reads_local(self):
        with tempfile.TemporaryDirectory() as temp:
            service, cloud, sop = self._refresh_fixture(Path(temp))
            service.state_store.delete_document("polling_sop_meta", "cloud_initialized_v1")
            service.refresh_sops("A")
            count = cloud.list_calls
            service.list_sops("A")
            service._wait_local_cache_bootstrap()
            self.assertEqual(cloud.list_calls, count)

    def test_manual_refresh_rejects_duplicate_ids_or_older_cloud_version(self):
        with tempfile.TemporaryDirectory() as temp:
            service, cloud, sop = self._refresh_fixture(Path(temp))
            before = service.get_sop(sop["sop_id"])
            with patch.object(cloud, "list_sops", return_value=[sop, sop]), self.assertRaisesRegex(Exception, "重复"):
                service.refresh_sops("A")
            cloud.sops[sop["sop_id"]]["version"] = 0
            with self.assertRaisesRegex(Exception, "版本冲突"):
                service.refresh_sops("A")
            self.assertEqual(service.get_sop(sop["sop_id"]), before)

    def test_manual_refresh_batch_rolls_back_on_database_failure(self):
        with tempfile.TemporaryDirectory() as temp:
            service, cloud, sop = self._refresh_fixture(Path(temp))
            before = service.get_sop(sop["sop_id"])
            cloud.sops[sop["sop_id"]]["name"] = "更新"
            cloud.sops["next_sop_12345"] = {**sop, "sop_id": "next_sop_12345", "name": "序列化失败"}
            original = service.state_store._json
            def encode(payload):
                if payload.get("name") == "序列化失败":
                    raise OSError("disk failure")
                return original(payload)
            with patch.object(service.state_store, "_json", side_effect=encode), self.assertRaises(OSError):
                service.refresh_sops("A")
            self.assertEqual(service.get_sop(sop["sop_id"]), before)
            self.assertIsNone(service.state_store.get_document("polling_sop", "next_sop_12345"))

    def test_manual_refresh_requires_login_and_scope_permission(self):
        controller = FastAPIPortalController(host="127.0.0.1", port=18766)
        client = TestClient(controller._build_app())
        manager = MagicMock()
        manager.refresh_sops.return_value = {"items": [], "synced_count": 0}
        with patch.object(PortalRuntime, "polling_work_orders", return_value=manager):
            with patch.object(controller, "_current_session", return_value=None):
                self.assertEqual(client.post("/api/polling-sops/refresh", json={"scope": "A"}).status_code, 401)
            with patch.object(controller, "_current_session", return_value={"user": {"open_id": "ou_a"}}):
                with patch.object(controller, "_authorized_scope_or_error", side_effect=portal_server.PortalError("无权访问 B 楼")):
                    self.assertEqual(client.post("/api/polling-sops/refresh", json={"scope": "B"}).status_code, 403)
                manager.refresh_sops.assert_not_called()
                with patch.object(controller, "_authorized_scope_or_error", return_value="A"):
                    response = client.post("/api/polling-sops/refresh", json={"scope": "A", "work_type": "maintenance"})
                self.assertEqual(response.status_code, 200, response.text)
                manager.refresh_sops.assert_called_once_with("A", "maintenance")

    def test_legacy_local_sop_is_uploaded_to_cloud_once(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            store = LanPortalStateStore(root / "state.sqlite3")
            cloud = _FakePollingSopCloud()
            service = PollingWorkOrderService(store, cloud)
            service.sop_root = root / "sops"
            sop_id = "legacy_sop_1234"
            attachment_id = "legacy_attachment_1234"
            directory = service.sop_root / sop_id
            directory.mkdir(parents=True)
            path = directory / f"{attachment_id}_guide.txt"
            path.write_bytes(b"legacy-guide")
            store.put_document("polling_sop", sop_id, {
                "sop_id": sop_id,
                "scope": "A",
                "work_type": "maintenance",
                "name": "旧版本地SOP",
                "version": 1,
                "steps": [{
                    "step_id": "legacy_step_1234",
                    "content": "检查设备",
                    "operator_required": True,
                    "reviewer_required": False,
                    "photo_required": False,
                    "time_limit_seconds": 0,
                }],
                "attachments": [{
                    "attachment_id": attachment_id,
                    "name": "guide.txt",
                    "size": path.stat().st_size,
                    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                    "path": str(path),
                }],
            })

            self.assertEqual([item["name"] for item in service.list_sops("A", "maintenance")], ["旧版本地SOP"])
            service._wait_local_cache_bootstrap()
            self.assertEqual((len(cloud.uploaded), len(cloud.saved)), (1, 1))
            self.assertEqual(
                store.get_document("polling_sop", sop_id)["attachments"][0]["file_token"],
                "cloud-file-token",
            )
            cloud_reads = cloud.list_calls
            service.list_sops("A", "maintenance")
            self.assertEqual((len(cloud.uploaded), len(cloud.saved)), (1, 1))
            self.assertEqual(cloud.list_calls, cloud_reads)

    def test_legacy_local_sop_does_not_duplicate_cloud_name(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            store = LanPortalStateStore(Path(temp) / "state.sqlite3")
            cloud_sop = {
                "sop_id": "cloud_sop_12345",
                "scope": "A",
                "work_type": "maintenance",
                "name": "同名SOP",
                "version": 1,
                "steps": [],
                "attachments": [],
            }
            cloud = _FakePollingSopCloud([cloud_sop])
            store.put_document("polling_sop", "local_sop_12345", {
                **cloud_sop,
                "sop_id": "local_sop_12345",
                "version": 2,
            })
            service = PollingWorkOrderService(store, cloud)

            listed = service.list_sops("A", "maintenance")
            self.assertEqual([item["sop_id"] for item in listed], ["local_sop_12345"])
            service._wait_local_cache_bootstrap()
            listed = service.list_sops("A", "maintenance")
            self.assertEqual([item["sop_id"] for item in listed], ["cloud_sop_12345"])
            self.assertEqual(cloud.saved, [])

    def test_cloud_attachment_rejects_untrusted_download_url(self) -> None:
        cloud = PollingSopCloudStore.__new__(PollingSopCloudStore)
        cloud.client = MagicMock()
        with self.assertRaisesRegex(Exception, "下载地址无效"):
            cloud.download_attachment({
                "file_token": "file_token_1234",
                "_cloud_download_url": "https://example.invalid/private",
            })
        cloud.client.request_bytes.assert_not_called()

    def test_cloud_sop_create_uses_stable_client_token(self) -> None:
        cloud = PollingSopCloudStore.__new__(PollingSopCloudStore)
        cloud.ensure_schema = MagicMock()
        cloud.get_sop = MagicMock(return_value=None)
        cloud._request = MagicMock(return_value={"record": {"record_id": "recCloudSop"}})
        cloud.invalidate = MagicMock()
        sop = {
            "sop_id": "stable_sop_1234",
            "scope": "A",
            "work_type": "polling",
            "name": "幂等SOP",
            "steps": [],
            "attachments": [],
            "version": 1,
        }
        cloud.save_sop(sop, expected_version=0, allow_create=True)
        first_token = cloud._request.call_args.kwargs["params"]["client_token"]
        cloud.save_sop(sop, expected_version=0, allow_create=True)
        second_token = cloud._request.call_args.kwargs["params"]["client_token"]
        self.assertEqual(first_token, second_token)
        self.assertEqual(cloud.ensure_schema.call_count, 2)

    def test_cloud_sop_read_does_not_check_schema(self) -> None:
        cloud = PollingSopCloudStore.__new__(PollingSopCloudStore)
        cloud.ensure_schema = MagicMock()
        cloud._request = MagicMock(return_value={"items": [], "has_more": False})

        self.assertEqual(cloud.list_records(), [])
        cloud.ensure_schema.assert_not_called()

    def test_sop_cloud_is_authoritative_and_attachment_is_cached_locally(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            content = b"shared-sop-file"
            sop = {
                "sop_id": "shared_sop_1234",
                "scope": "E",
                "work_type": "polling",
                "name": "共享SOP",
                "version": 1,
                "steps": [{"step_id": "step_12345678", "content": "检查设备", "operator_required": True, "reviewer_required": False, "photo_required": False, "time_limit_seconds": 0}],
                "attachments": [{"attachment_id": "attachment_1234", "name": "说明.txt", "size": len(content), "sha256": hashlib.sha256(content).hexdigest(), "file_token": "secret-token"}],
            }
            root = Path(temp)
            cloud = _FakePollingSopCloud([sop], content)
            service = PollingWorkOrderService(LanPortalStateStore(root / "state.sqlite3"), cloud)
            service.sop_root = root / "sops"

            listed = service.list_sops("E")
            self.assertEqual([item["name"] for item in listed], ["共享SOP"])
            self.assertNotIn("file_token", listed[0]["attachments"][0])
            downloaded, name = service.get_sop_attachment("shared_sop_1234", "attachment_1234")
            self.assertEqual((downloaded, name), (content, "说明.txt"))
            saved = service.save_sop({**listed[0], "name": "共享SOP更新", "expected_version": 1})
            self.assertEqual((saved["name"], saved["version"]), ("共享SOP更新", 2))
            self.assertEqual(cloud.saved[-1][1], 1)

    def test_step_photo_requirement_is_configurable_and_legacy_default_is_safe(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            service = PollingWorkOrderService(LanPortalStateStore(root / "state.sqlite3"))
            service.sop_root = root / "sops"
            service.work_order_root = root / "orders"
            sop = service.save_sop({
                "scope": "A", "name": "拍照规则测试",
                "steps": [
                    {"content": "无需拍照步骤", "operator_required": True, "photo_required": False},
                    {"content": "旧步骤默认拍照", "operator_required": True},
                ],
            })
            self.assertEqual([step["photo_required"] for step in sop["steps"]], [False, True])
            sop = service.add_sop_attachment(
                sop["sop_id"], file_name="SOP.txt", content=b"sop",
                expected_version=sop["version"],
            )
            prepared = service.prepare_start({
                "work_type": "polling", "scope": "A", "action": "start",
                "_web_action_request": True, "polling_sop_id": sop["sop_id"],
                "polling_sop_version": sop["version"], "polling_run_count": 1,
                "polling_runs": [{"from_unit": "1#", "to_unit": "2#"}],
                "polling_operator_record_id": "operator",
                "polling_reviewer_record_id": "reviewer",
            }, job_id="photo-rule", people=[
                {"record_id": "operator", "name": "操作员"},
                {"record_id": "reviewer", "name": "审核员"},
            ])
            group = service.create_group(prepared, target_record_id="recPhotoRule", title="测试", public_base_url="")
            token = service.role_token("recPhotoRule", "operator")
            session = service.activate(token, run_index=1, expected_version=group["version"])
            stored = service.get_group("recPhotoRule")
            stored["relay"]["mode"] = "public_relay"
            service.state_store.put_document("polling_work_order", "recPhotoRule", stored)
            session = service.confirm(token, step_key="1:1", expected_version=session["version"])
            self.assertEqual(session["steps"][1]["photo_required"], True)
            with self.assertRaisesRegex(Exception, "请先拍摄并上传"):
                service.confirm(token, step_key="1:2", expected_version=session["version"])

    def test_maintenance_and_polling_isolate_sops_and_build_generic_work_order(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            service = PollingWorkOrderService(
                LanPortalStateStore(root / "state.sqlite3")
            )
            service.sop_root = root / "sops"
            service.work_order_root = root / "orders"
            polling = service.save_sop(
                {
                    "work_type": "polling",
                    "scope": "A",
                    "name": "同名SOP",
                    "steps": [
                        {
                            "content": "切换{{from}}至{{to}}",
                            "operator_required": True,
                        }
                    ],
                }
            )
            polling = service.add_sop_attachment(
                polling["sop_id"],
                file_name="轮巡说明.txt",
                content=b"polling",
                expected_version=polling["version"],
            )
            maintenance = service.save_sop(
                {
                    "work_type": "maintenance",
                    "scope": "A",
                    "name": "维保通用SOP",
                    "steps": [
                        {
                            "content": "检查设备运行状态",
                            "operator_required": True,
                            "reviewer_required": True,
                            "time_limit_seconds": 0,
                        }
                    ],
                }
            )
            self.assertEqual(
                {item["sop_id"] for item in service.list_sops("A", "polling")},
                {polling["sop_id"]},
            )
            self.assertEqual(
                {item["sop_id"] for item in service.list_sops("A", "maintenance")},
                {maintenance["sop_id"]},
            )
            with self.assertRaisesRegex(Exception, "不适用于当前工单类型"):
                service.prepare_start(
                    {
                        "work_type": "maintenance",
                        "scope": "A",
                        "action": "start",
                        "_web_action_request": True,
                        "polling_sop_id": polling["sop_id"],
                        "polling_sop_version": polling["version"],
                    },
                    job_id="maintenance-incompatible",
                    people=[],
                )
            maintenance = service.save_sop(
                {
                    "sop_id": maintenance["sop_id"],
                    "work_type": "maintenance",
                    "scope": "A",
                    "name": "维保通用SOP",
                    "expected_version": maintenance["version"],
                    "steps": [
                        {
                            "content": "检查并记录设备运行状态",
                            "operator_required": True,
                            "reviewer_required": True,
                            "time_limit_seconds": 0,
                        }
                    ],
                }
            )
            self.assertEqual(
                maintenance["steps"][0]["content"],
                "检查并记录设备运行状态",
            )
            with self.assertRaisesRegex(Exception, "不能使用设备指向占位符"):
                service.save_sop(
                    {
                        "work_type": "maintenance",
                        "scope": "A",
                        "name": "错误维保SOP",
                        "steps": [
                            {
                                "content": "检查{{from}}",
                                "operator_required": True,
                            }
                        ],
                    }
                )
            maintenance = service.add_sop_attachment(
                maintenance["sop_id"],
                file_name="维保说明.txt",
                content=b"maintenance",
                expected_version=maintenance["version"],
            )
            prepared = service.prepare_start(
                {
                    "work_type": "maintenance",
                    "scope": "A",
                    "action": "start",
                    "_web_action_request": True,
                    "polling_sop_id": maintenance["sop_id"],
                    "polling_sop_version": maintenance["version"],
                    "polling_operator_record_id": "operator",
                    "polling_reviewer_record_id": "reviewer",
                },
                job_id="maintenance-job",
                people=[
                    {"record_id": "operator", "name": "操作员"},
                    {"record_id": "reviewer", "name": "审核员"},
                ],
            )
            group = service.create_group(
                prepared,
                target_record_id="recMaintenance",
                title="维保测试",
                public_base_url="http://127.0.0.1:18766",
            )
            self.assertEqual(group["work_type"], "maintenance")
            self.assertEqual(group["notice_type"], "维保通告")
            self.assertEqual(group["runs"], [{"run_index": 1, "label": "维保作业"}])
            self.assertEqual(group["steps"][0]["run_label"], "维保作业")
            self.assertEqual(
                group["steps"][0]["content"],
                "检查并记录设备运行状态",
            )
            output_name = service._workbook_output_name(group)
            self.assertIn("维保通用SOP-操作人-操作员-审核人-审核员-操作记录", output_name)
            self.assertNotIn("轮巡至", output_name)
            operator_token = service.role_token("recMaintenance", "operator")
            reviewer_token = service.role_token("recMaintenance", "reviewer")
            session = service.activate(
                operator_token,
                run_index=1,
                expected_version=group["version"],
            )
            session = service.add_step_photo(
                operator_token,
                step_key="1:1",
                expected_version=session["version"],
                file_name="维保步骤.png",
                mime_type="image/png",
                content=_png_bytes(),
            )
            session = service.confirm(
                operator_token,
                step_key="1:1",
                expected_version=session["version"],
            )
            session = service.confirm(
                reviewer_token,
                step_key="1:1",
                expected_version=session["version"],
            )
            self.assertEqual(session["state"], "upload_pending")
            workbook = service.build_execution_workbook("recMaintenance")
            self.assertEqual(workbook["name"], output_name)
            from openpyxl import load_workbook

            generated = load_workbook(workbook["path"], read_only=True)
            try:
                self.assertEqual(generated.sheetnames, ["工单1 维保作业"])
                self.assertEqual(
                    generated.active["B7"].value,
                    "维保通用SOP · 维保作业",
                )
            finally:
                generated.close()
            deleted = service.delete_sop(
                maintenance["sop_id"],
                expected_version=maintenance["version"],
            )
            self.assertTrue(deleted["deleted"])
            self.assertEqual(
                [item["sop_id"] for item in service.list_sops("A", "maintenance")],
                [],
            )
            self.assertEqual(
                [item["sop_id"] for item in service.list_sops("A", "polling")],
                [polling["sop_id"]],
            )

    def test_maintenance_workbench_uses_generic_work_order_selector(self) -> None:
        html = render_workbench_lite(
            payload={"records": [], "ongoing": []},
            session={"role": "admin"},
            scope="A",
            work_type="maintenance",
        )
        polling_html = render_workbench_lite(
            payload={"records": [], "ongoing": []},
            session={"role": "admin"},
            scope="A",
            work_type="polling",
        )
        self.assertIn("维保工单", html)
        self.assertIn("本次维保不使用工单", html)
        self.assertIn("work_type=${encodeURIComponent(workType)}", html)
        self.assertIn("countLabel.hidden=maintenance", html)
        self.assertIn("directionTitle.hidden=maintenance", html)
        self.assertIn("if(pollingSopWorkType()==='polling')for", html)
        self.assertIn("['maintenance', 'polling', 'adjust'].includes(patch.work_type)", html)
        self.assertIn("不可用于维保或设备调整工单", html)
        self.assertIn("function pollingSopHasDevicePlaceholders", html)
        self.assertIn("button.disabled=blocked", html)
        self.assertIn("含设备指向，通用工单不可选", html)
        button_pattern = re.compile(
            r'<h2 class="inbox-title"><span>通告处理</span>'
            r'(<button class="btn ghost" id="lite-polling-sop-open".*?</button>)'
            r'</h2>'
        )
        maintenance_button = button_pattern.search(html)
        polling_button = button_pattern.search(polling_html)
        self.assertIsNotNone(maintenance_button)
        self.assertIsNotNone(polling_button)
        self.assertEqual(
            maintenance_button.group(1),
            polling_button.group(1),
        )
        self.assertEqual(html.count('id="lite-polling-sop-modal"'), 1)
        self.assertEqual(polling_html.count('id="lite-polling-sop-modal"'), 1)

    def test_maintenance_handler_writes_work_order_fields(self) -> None:
        fields = MaintenanceNoticeHandler().build_create_fields(
            NoticePayload(
                text=(
                    "【维保通告】状态：开始\n"
                    "【名称】维保测试\n"
                    "【时间】2026-08-31 09:00~2026-08-31 18:00"
                ),
                polling_work_order_required=True,
                polling_operator_name="操作员",
                polling_reviewer_name="审核员",
            )
        )
        self.assertTrue(fields["是否涉及重要操作"])
        self.assertEqual(fields["操作人"], "操作员")
        self.assertEqual(fields["现场复核人"], "审核员")

    def test_maintenance_end_uses_same_exact_attachment_guard(self) -> None:
        prepared = {
            "action": "end",
            "work_type": "maintenance",
            "notice_type": "维保通告",
            "record_id": "recMaintenanceEnd",
            "target_record_id": "recMaintenanceEnd",
            "site_photo_count": 1,
            "text": "【维保通告】状态：结束\n【名称】维保测试",
        }
        manager = MagicMock()
        manager.get_group.return_value = {
            "work_type": "maintenance",
            "state": "completed",
            "uploaded_file_tokens": ["expected-token"],
        }
        with patch.object(
            portal_server,
            "external_real_write_guard",
            return_value={"mock_external": False, "real_write_allowed": True, "reason": ""},
        ), patch.object(
            portal_server,
            "query_record_by_id",
            return_value=(
                True,
                {
                    "fields": {
                        "是否涉及重要操作": True,
                        "操作人": "操作员",
                        "现场复核人": "审核员",
                        "工单附件": [{"file_token": "other-token"}],
                    }
                },
            ),
        ), patch.object(PortalRuntime, "polling_work_orders", return_value=manager):
            ok, message, _record_id = PortalRuntime._execute_backend_prepared_upload(
                prepared
            )
        self.assertFalse(ok)
        self.assertIn("维保工单附件与本次完成记录不一致", message)

    def test_qt_maintenance_end_is_blocked_by_shared_work_order_guard(self) -> None:
        manager = MagicMock()
        manager.get_group.return_value = {
            "work_type": "maintenance",
            "state": "active",
            "uploaded_file_tokens": [],
        }
        with patch.object(
            portal_server,
            "query_record_by_id",
            return_value=(
                True,
                {
                    "record_version": "1",
                    "fields": {
                        "是否涉及重要操作": True,
                        "操作人": "操作员",
                        "现场复核人": "审核员",
                        "工单附件": [],
                    },
                },
            ),
        ), patch.object(PortalRuntime, "polling_work_orders", return_value=manager):
            result = PortalRuntime.execute_local_notice_upload(
                {
                    "action_type": "end",
                    "data_dict": {
                        "active_item_id": "active-maintenance-end",
                        "record_id": "recMaintenanceQtEnd",
                        "target_record_id": "recMaintenanceQtEnd",
                        "notice_type": "维保通告",
                        "work_type": "maintenance",
                        "text": "【维保通告】状态：结束\n【名称】维保测试",
                    },
                }
            )
        self.assertFalse(result["ok"])
        self.assertIn("维保工单尚未全部完成", result["message"])

    def test_maintenance_work_order_schema_creates_attachment_once(self) -> None:
        service = object.__new__(MaintenancePortalService)
        service._maintenance_work_order_schema_lock = __import__("threading").RLock()
        service._maintenance_work_order_schema_ready = False
        service._write_http_client = MagicMock()
        before = [
            {"field_name": "是否涉及重要操作", "type": 7},
            {"field_name": "操作人", "type": 1},
            {"field_name": "现场复核人", "type": 1},
        ]
        after = [*before, {"field_name": "工单附件", "type": 17}]
        with patch.object(
            service, "_load_raw_table_fields", side_effect=[before, after]
        ), patch.object(
            service, "_request_payload", return_value={"code": 0, "msg": "ok"}
        ) as create, patch.object(service, "_auth_headers", return_value={}):
            first = service.ensure_maintenance_work_order_schema()
            second = service.ensure_maintenance_work_order_schema()
        self.assertEqual(first["created_fields"], ["工单附件"])
        self.assertTrue(second["cached"])
        self.assertEqual(create.call_count, 1)
        self.assertEqual(create.call_args.kwargs["json_payload"]["type"], 17)

    def test_polling_step_photo_rejects_oversized_dimensions_before_storage(self) -> None:
        service = object.__new__(PollingWorkOrderService)
        with self.assertRaisesRegex(Exception, "像素或尺寸超过限制"):
            service.add_step_photo(
                "unused-token",
                step_key="1:1",
                expected_version=1,
                file_name="too-wide.png",
                mime_type="image/png",
                content=_png_bytes(size=(12001, 1)),
            )

    def test_polling_sop_rejects_more_than_thirty_steps(self) -> None:
        with self.assertRaisesRegex(Exception, "SOP 步骤不能超过 30 条"):
            PollingWorkOrderService._normalized_steps(
                [
                    {
                        "content": f"步骤{index + 1}",
                        "operator_required": True,
                    }
                    for index in range(31)
                ]
            )

    def test_attachment_result_keeps_unsaved_sop_name_and_steps(self) -> None:
        node = shutil.which("node")
        if not node:
            self.skipTest("node is unavailable")
        html = render_workbench_lite(
            payload={"records": [], "ongoing": []},
            session={"role": "admin"},
            scope="A",
            work_type="polling",
        )
        match = re.search(
            r"function pollingMergeSopAttachmentResult\(.*?\n    \}",
            html,
            re.S,
        )
        self.assertIsNotNone(match)
        self.assertEqual(
            html.count(
                "litePollingEditingSop=pollingMergeSopAttachmentResult(draft,remote)"
            ),
            2,
        )
        script = "\n".join(
            [
                match.group(0),
                "const draft={name:'修改后的名称',steps:[{content:'修改后的步骤'}]};",
                "const remote={name:'旧名称',steps:[{content:'旧步骤'}],version:3,attachments:[{name:'附件.pdf'}]};",
                "console.log(JSON.stringify(pollingMergeSopAttachmentResult(draft,remote)));",
            ]
        )
        result = subprocess.run(
            [node, "-e", script], check=True, capture_output=True, text=True
        )

        merged = json.loads(result.stdout)
        self.assertEqual(merged["name"], "修改后的名称")
        self.assertEqual(merged["steps"][0]["content"], "修改后的步骤")
        self.assertEqual(merged["version"], 3)
        self.assertEqual(merged["attachments"][0]["name"], "附件.pdf")

    def test_work_order_overview_and_steps_are_separate_pages(self) -> None:
        overview = render_polling_work_order_page()
        steps = render_polling_work_order_steps_page()

        self.assertIn('id="overview"', overview)
        self.assertNotIn('id="steps"', overview)
        self.assertIn("function workOrderCard(order)", overview)
        self.assertIn("另一工单执行中", overview)
        self.assertIn("button.disabled=!order.selectable||busy", overview)
        self.assertIn("/api/polling-work-orders/activate", overview)
        self.assertIn("/polling-work-order/steps?token=", overview)
        self.assertIn("取消当前选择", overview)
        self.assertIn("/api/polling-work-orders/release", overview)
        self.assertIn("全部工单已完成，工单表格已上传。", overview)

        self.assertIn('id="steps"', steps)
        self.assertNotIn('id="overview"', steps)
        self.assertNotIn("function workOrderCard(order)", steps)
        self.assertIn("第 ${Number(step.step_index||0)} 步", steps)
        self.assertNotIn("工单 ${step.run_index}/${step.run_count}", steps)
        self.assertNotIn("步骤 ${step.step_index}/${step.step_count}", steps)
        self.assertIn("倒计时 ${wait} 秒", steps)
        self.assertIn("Number(step.remaining_seconds||0)", steps)
        self.assertIn("function refreshCountdown()", steps)
        self.assertIn("setInterval(refreshCountdown,1000)", steps)
        self.assertIn("const next=body.data,sameView=", steps)
        self.assertNotIn("setInterval(()=>{if(current&&!busy&&!navigating)render(current)},1000)", steps)
        self.assertIn("function leaveWorkOrder()", steps)
        self.assertIn("/api/polling-work-orders/release", steps)
        self.assertIn("退出当前工单并重新选择", steps)
        self.assertIn("浏览器返回不会取消当前选择", steps)
        self.assertIn("拍照/上传操作照片", steps)
        self.assertIn("继续添加照片", steps)
        self.assertIn("操作照片 ${photoCount} 张", steps)
        self.assertIn("操作照片待拍", steps)
        self.assertIn("无需拍照", steps)
        self.assertIn("photoRequired=step.photo_required!==false", steps)
        self.assertIn("回退上一步", steps)
        self.assertIn("/api/polling-work-orders/rollback", steps)
        self.assertIn("请按步骤要求重新执行并确认", steps)
        self.assertIn("busy=false;if(current&&!navigating)render(current)", steps)
        self.assertIn("capture','environment", steps)
        self.assertIn("image.loading='eager'", steps)
        self.assertIn("缩略图加载失败", steps)

        controller = FastAPIPortalController(host="127.0.0.1", port=18766)
        client = TestClient(controller._build_app())
        overview_response = client.get("/polling-work-order?token=test")
        steps_response = client.get(
            "/polling-work-order/steps?token=test&run_index=1"
        )
        self.assertEqual(overview_response.status_code, 200)
        self.assertEqual(steps_response.status_code, 200)
        self.assertNotEqual(overview_response.text, steps_response.text)

    def test_polling_step_photo_api_accepts_image_and_serves_preview(self) -> None:
        manager = MagicMock()
        manager.add_step_photo.return_value = {"version": 2, "steps": []}
        manager.step_photo_content.return_value = (
            b"photo-bytes",
            "image/png",
            "step.png",
        )
        controller = FastAPIPortalController(host="127.0.0.1", port=18766)
        client = TestClient(controller._build_app())
        with patch.object(PortalRuntime, "polling_work_orders", return_value=manager):
            response = client.post(
                "/api/polling-work-orders/photo?token=role-token"
                "&step_key=1%3A1&expected_version=1&file_name=step.png",
                content=b"photo-bytes",
                headers={"Content-Type": "image/png"},
            )
            preview = client.get(
                "/api/polling-work-orders/photos/photo-id?token=role-token"
            )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(preview.status_code, 200, preview.text)
        self.assertEqual(preview.content, b"photo-bytes")
        self.assertEqual(
            manager.add_step_photo.call_args.kwargs["expected_version"], 1
        )
        self.assertEqual(manager.add_step_photo.call_args.kwargs["content"], b"photo-bytes")

    def test_polling_work_order_activate_api_starts_selected_run(self) -> None:
        manager = MagicMock()
        manager.activate.return_value = {
            "version": 2,
            "current_run_index": 1,
            "steps": [{"timer_started": True, "remaining_seconds": 30}],
        }
        manager.release_selection.return_value = {
            "version": 3,
            "current_run_index": 0,
            "steps": [],
        }
        controller = FastAPIPortalController(host="127.0.0.1", port=18766)
        client = TestClient(controller._build_app())
        with patch.object(PortalRuntime, "polling_work_orders", return_value=manager):
            response = client.post(
                "/api/polling-work-orders/activate",
                json={"token": "r" * 32, "run_index": 1, "expected_version": 1},
            )
            release_response = client.post(
                "/api/polling-work-orders/release",
                json={"token": "r" * 32, "run_index": 1, "expected_version": 2},
            )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.json()["data"]["steps"][0]["timer_started"])
        self.assertEqual(manager.activate.call_args.kwargs["run_index"], 1)
        self.assertEqual(manager.activate.call_args.kwargs["expected_version"], 1)
        self.assertEqual(release_response.status_code, 200, release_response.text)
        self.assertEqual(manager.release_selection.call_args.kwargs["run_index"], 1)

    def test_last_confirmation_queues_attachment_upload(self) -> None:
        manager = MagicMock()
        manager.confirm.return_value = {
            "state": "upload_pending",
            "group_id": "recQueued",
        }
        controller = FastAPIPortalController(host="127.0.0.1", port=18766)
        client = TestClient(controller._build_app())
        with patch.object(
            PortalRuntime, "polling_work_orders", return_value=manager
        ), patch.object(
            controller, "_submit_background", return_value=True
        ) as submit, patch.object(
            PortalRuntime, "finalize_polling_work_order_group"
        ) as finalize:
            response = client.post(
                "/api/polling-work-orders/confirm",
                json={"token": "r" * 32, "step_key": "1:1", "expected_version": 1},
            )

        self.assertEqual(response.status_code, 200, response.text)
        completion = response.json()["data"]["completion"]
        self.assertEqual(completion["state"], "upload_pending")
        self.assertTrue(completion["queued"])
        finalize.assert_not_called()
        self.assertEqual(submit.call_args.args[0], "PollingWorkOrderFinalize")
        self.assertEqual(submit.call_args.args[2], "recQueued")

    def test_polling_work_order_rollback_api_uses_reviewer_token(self) -> None:
        manager = MagicMock()
        manager.rollback_previous.return_value = {
            "version": 3,
            "current_index": 0,
            "can_rollback_previous": False,
        }
        controller = FastAPIPortalController(host="127.0.0.1", port=18766)
        client = TestClient(controller._build_app())
        with patch.object(PortalRuntime, "polling_work_orders", return_value=manager):
            response = client.post(
                "/api/polling-work-orders/rollback",
                json={"token": "r" * 32, "step_key": "1:2", "expected_version": 2},
            )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["data"]["current_index"], 0)
        self.assertEqual(manager.rollback_previous.call_args.kwargs["step_key"], "1:2")

    def test_reviewer_can_rollback_repeatedly_and_operator_session_tracks_cursor(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            store = LanPortalStateStore(Path(temp) / "state.sqlite3")
            service = PollingWorkOrderService(store)
            service.work_order_root = Path(temp) / "orders"
            target_record_id = "recRepeatedRollback"
            photo_root = service._group_directory(target_record_id) / "photos"
            photo_root.mkdir(parents=True)
            steps = []
            for index in range(3):
                photo_path = photo_root / f"step-{index + 1}.png"
                photo_path.write_bytes(_png_bytes())
                done = index < 2
                steps.append(
                    {
                        "step_key": f"1:{index + 1}",
                        "global_index": index,
                        "run_index": 1,
                        "step_index": index + 1,
                        "content": f"步骤{index + 1}",
                        "operator_required": True,
                        "reviewer_required": True,
                        "activated_at_ts": time.time(),
                        "photos": [{"photo_id": str(index), "path": str(photo_path)}],
                        "operator_confirmation": {"confirmed_at": "2026-08-27 09:00:00"} if done else {},
                        "reviewer_confirmation": {"confirmed_at": "2026-08-27 09:01:00"} if done else {},
                    }
                )
            operator_token = service.role_token(target_record_id, "operator")
            reviewer_token = service.role_token(target_record_id, "reviewer")
            store.put_document(
                "polling_work_order",
                target_record_id,
                {
                    "group_id": target_record_id,
                    "target_record_id": target_record_id,
                    "state": "active",
                    "version": 1,
                    "current_index": 2,
                    "selected_run_index": 1,
                    "selected_by_role": "reviewer",
                    "runs": [{"from_unit": "1#", "to_unit": "2#"}],
                    "steps": steps,
                    "operator": {"name": "操作员"},
                    "reviewer": {"name": "审核员"},
                    "token_hashes": {
                        "operator": hashlib.sha256(operator_token.encode()).hexdigest(),
                        "reviewer": hashlib.sha256(reviewer_token.encode()).hexdigest(),
                    },
                },
            )

            reviewer_session = service.rollback_previous(
                reviewer_token,
                step_key="1:3",
                expected_version=1,
            )
            self.assertEqual(reviewer_session["current_index"], 1)
            self.assertTrue(reviewer_session["can_rollback_previous"])
            operator_session = service.session(operator_token)
            self.assertEqual(operator_session["current_index"], 1)
            self.assertEqual(operator_session["version"], reviewer_session["version"])

            reviewer_session = service.rollback_previous(
                reviewer_token,
                step_key="1:2",
                expected_version=reviewer_session["version"],
            )
            self.assertEqual(reviewer_session["current_index"], 0)
            self.assertFalse(reviewer_session["can_rollback_previous"])
            operator_session = service.session(operator_token)
            self.assertEqual(operator_session["current_index"], 0)
            self.assertEqual(operator_session["version"], reviewer_session["version"])

    def test_notice_content_prefills_polling_runs(self) -> None:
        node = shutil.which("node")
        if not node:
            self.skipTest("node is unavailable")
        html = render_workbench_lite(
            payload={"records": [], "ongoing": []},
            session={"role": "admin"},
            scope="A",
            work_type="polling",
        )
        functions = []
        for name in (
            "pollingShiftFromHour",
            "pollingUnitGroup",
            "pollingUnitNumbersFromContent",
            "pollingRunsFromContent",
        ):
            match = re.search(rf"function {name}\(.*?\n    \}}", html, re.S)
            self.assertIsNotNone(match, name)
            functions.append(match.group(0))
        script = "\n".join(
            [
                "const litePollingUnits=['1#','2#','3#','4#','5#','6#'];",
                *functions,
                "console.log(JSON.stringify([pollingRunsFromContent('起点1终点3'),pollingRunsFromContent('1 3 4 6'),pollingRunsFromContent('1# 4号 3 6'),pollingRunsFromContent('1 4 6 3'),[8,9,17,18,23,0].map(pollingShiftFromHour)]));",
            ]
        )
        result = subprocess.run(
            [node, "-e", script], check=True, capture_output=True, text=True
        )
        self.assertEqual(
            json.loads(result.stdout),
            [
                [{"from_unit": "1#", "to_unit": "3#"}],
                [
                    {"from_unit": "1#", "to_unit": "3#"},
                    {"from_unit": "4#", "to_unit": "6#"},
                ],
                [
                    {"from_unit": "1#", "to_unit": "3#"},
                    {"from_unit": "4#", "to_unit": "6#"},
                ],
                [
                    {"from_unit": "1#", "to_unit": "3#"},
                    {"from_unit": "4#", "to_unit": "6#"},
                ],
                ["夜", "白", "白", "夜", "夜", "夜"],
            ],
        )

    def test_polling_sop_button_is_in_notice_panel_and_uses_floating_editor(self) -> None:
        html = render_workbench_lite(
            payload={"records": [], "ongoing": []},
            session={"role": "admin"},
            scope="A",
            work_type="polling",
        )
        self.assertIn(
            '<h2 class="inbox-title"><span>通告处理</span><button class="btn ghost" id="lite-polling-sop-open"',
            html,
        )
        self.assertIn(
            'class="end-check-backdrop" id="lite-polling-sop-modal" hidden',
            html,
        )
        self.assertIn("/api/polling-sops?scope=", html)
        self.assertIn("插入其他点", html)
        self.assertIn("['插入起点','{{from}}']", html)
        self.assertIn("['插入终点','{{to}}']", html)
        self.assertIn("['插入其他点','{{other}}']", html)
        self.assertIn("可先配置设备指向", html)
        self.assertIn("设备指向（1#–3# / 4#–6# 组内选择）", html)
        self.assertIn("第${index+1}次起点", html)
        self.assertIn("SOP 必须至少包含一个附件", html)
        self.assertIn("非制冷单元/二次泵轮巡", html)
        self.assertIn(".polling-cooling-mode[hidden] { display:none; }", html)
        self.assertIn("workTypeText.textContent='适用类型'", html)
        self.assertIn("[['adjust','调整'],['polling','轮巡'],['maintenance','维保']]", html)
        self.assertIn("workType.disabled=Boolean(sop.sop_id)", html)
        self.assertIn('name="polling_work_order_exempt"', html)
        self.assertIn("function pollingWorkOrderExempt(form)", html)
        self.assertIn("patch.polling_work_order_exempt = pollingWorkOrderExempt(form)", html)
        self.assertIn("&& !pollingWorkOrderExempt(form)", html)
        self.assertIn('id="lite-polling-sop-feedback"', html)
        self.assertIn('id="lite-polling-sop-delete-confirm"', html)
        self.assertIn("确认删除这份 SOP", html)
        self.assertIn("已经生成的工单不受影响", html)
        self.assertIn("openPollingSopDeleteConfirm(sop,del)", html)
        self.assertIn("setPollingSopFeedback(`已保存：${saved.name}`)", html)
        self.assertNotIn("confirm('确认删除该 SOP", html)
        self.assertIn("本地附件（必填）", html)
        self.assertIn("点击选择或将附件拖到这里", html)
        self.assertIn("drop.ondrop", html)
        self.assertIn("function pollingBindPersonSearch", html)
        self.assertIn("()=>[reviewer.dataset.recordId,'h_duty_account']", html)
        self.assertIn("()=>[operator.dataset.recordId]", html)
        self.assertIn("搜索姓名、工号、岗位或楼栋", html)
        self.assertIn("快捷选择H楼值班账号", html)
        self.assertNotIn("快捷筛选 H楼人员", html)
        self.assertIn("h_duty_account", html)
        self.assertIn("后续工单不能再选择前面已使用的设备编号", html)
        self.assertIn("时间限制（秒）", html)
        self.assertIn("time_limit_seconds", html)
        self.assertIn("是否需要拍照", html)
        self.assertIn("photo_required", html)
        self.assertIn("正在读取 SOP 和人员", html)
        self.assertIn("人员加载失败", html)
        self.assertIn("SOP 加载失败", html)
        self.assertIn("renderPollingSopList();", html)
        self.assertIn("retry.onclick=()=>openPollingSopModal(mode)", html)
        self.assertNotIn("showLiteError(error.message);closePollingSopModal()", html)
        self.assertIn("pollingUnitGroup(run.from_unit).includes(run.to_unit)", html)
        self.assertIn("run.from_unit}→${run.to_unit", html)
        open_source = html.split("async function openPollingSopModal", 1)[1].split(
            "function currentUrlScope", 1
        )[0]
        self.assertLess(open_source.index("modal.hidden=false"), open_source.index("await loadPollingSops()"))
        self.assertLess(open_source.index("const peoplePromise="), open_source.index("await loadPollingSops()"))
        self.assertIn("const peopleError=await peoplePromise", open_source)
        self.assertIn("Date.now()-cached.loadedAt<60000", html)

    def test_polling_end_is_blocked_before_write_until_work_order_attachment_exists(self) -> None:
        prepared = {
            "action": "end",
            "work_type": "polling",
            "notice_type": "设备轮巡",
            "record_id": "recTarget1",
            "target_record_id": "recTarget1",
            "text": "【设备轮巡】状态：结束\n【标题】轮巡测试",
        }
        with patch.object(
            portal_server,
            "external_real_write_guard",
            return_value={"mock_external": False, "real_write_allowed": True, "reason": ""},
        ), patch.object(
            portal_server,
            "query_record_by_id",
            return_value=(
                True,
                {
                    "fields": {
                        "是否涉及重要操作": True,
                        "操作人": "操作员",
                        "现场复核人": "审核员",
                        "工单附件": [],
                    }
                },
            ),
        ):
            ok, message, _record_id = PortalRuntime._execute_backend_prepared_upload(
                prepared
            )
        self.assertFalse(ok)
        self.assertIn("尚未全部完成", message)

    def test_polling_end_requires_completed_local_group_and_exact_attachment(self) -> None:
        prepared = {
            "action": "end",
            "work_type": "polling",
            "notice_type": "设备轮巡",
            "record_id": "recTargetExact",
            "target_record_id": "recTargetExact",
            "text": "【设备轮巡】状态：结束\n【标题】轮巡测试",
        }
        manager = MagicMock()
        manager.get_group.return_value = {
            "state": "completed",
            "uploaded_file_tokens": ["expected-workbook-token"],
        }
        with patch.object(
            portal_server,
            "external_real_write_guard",
            return_value={"mock_external": False, "real_write_allowed": True, "reason": ""},
        ), patch.object(
            portal_server,
            "query_record_by_id",
            return_value=(
                True,
                {
                    "fields": {
                        "是否涉及重要操作": True,
                        "操作人": "操作员",
                        "现场复核人": "审核员",
                        "工单附件": [{"file_token": "manual-file-token"}],
                    }
                },
            ),
        ), patch.object(
            PortalRuntime,
            "polling_work_orders",
            return_value=manager,
        ):
            ok, message, _record_id = PortalRuntime._execute_backend_prepared_upload(
                prepared
            )
        self.assertFalse(ok)
        self.assertIn("与本次完成记录不一致", message)

    def test_exempt_polling_start_creates_directly_without_stale_target_lookup(self) -> None:
        prepared = {
            "action": "start",
            "work_type": "polling",
            "notice_type": "设备轮巡",
            "record_id": "recTarget1",
            "target_record_id": "recTarget1",
            "polling_work_order_exempt": True,
            "text": "【设备轮巡】状态：开始\n【标题】非工单轮巡",
        }
        manager = MagicMock()

        with patch.object(
            portal_server,
            "external_real_write_guard",
            return_value={"mock_external": False, "real_write_allowed": True, "reason": ""},
        ), patch.object(
            PortalRuntime,
            "_existing_target_for_prepared_start",
            return_value="recTarget1",
        ) as existing_target, patch.object(
            PortalRuntime,
            "_upload_change_confirmation_images",
            return_value=(True, "", [], []),
        ), patch.object(
            portal_server,
            "create_bitable_record_by_payload",
            return_value=(True, "recNewTarget"),
        ) as create_record, patch.object(
            portal_server,
            "query_record_by_id",
            return_value=(True, {"fields": {}}),
        ) as query_record, patch.object(
            portal_server,
            "update_bitable_record_fields",
        ), patch.object(
            PortalRuntime,
            "_create_backend_undo_checkpoint",
            return_value="",
        ), patch.object(
            PortalRuntime,
            "_mark_local_notice_images_target_written",
        ), patch.object(
            PortalRuntime,
            "polling_work_orders",
            return_value=manager,
        ):
            ok, _message, record_id = PortalRuntime._execute_backend_prepared_upload(
                prepared
            )

        self.assertTrue(ok)
        self.assertEqual(record_id, "recNewTarget")
        existing_target.assert_not_called()
        query_record.assert_not_called()
        create_record.assert_called_once()
        manager.cancel_group.assert_not_called()

    def test_polling_start_writes_work_order_fields_but_not_h_confirmation(self) -> None:
        fields = PollingNoticeHandler().build_create_fields(
            NoticePayload(
                text=(
                    "【设备轮巡】状态：开始\n"
                    "【标题】轮巡测试\n【时间】2026-08-24 09:00~2026-08-24 18:00"
                ),
                polling_work_order_required=True,
                polling_operator_name="操作员",
                polling_reviewer_name="审核员",
            )
        )
        self.assertTrue(fields["是否涉及重要操作"])
        self.assertEqual(fields["操作人"], "操作员")
        self.assertEqual(fields["现场复核人"], "审核员")
        self.assertNotIn("重要操作H楼确认", fields)

    def test_work_order_message_falls_back_to_current_account(self) -> None:
        manager = MagicMock()
        manager.group_with_links.return_value = {
            "target_record_id": "recTarget1",
            "title": "轮巡测试",
            "sop_name": "切换 SOP",
            "runs": [{"from_unit": "1#", "to_unit": "2#"}],
            "operator": {"name": "操作员", "open_id": "ou_unreachable"},
            "reviewer": {"name": "审核员", "open_id": "ou_reviewer"},
            "initiator_open_id": "ou_current",
            "operator_link": "http://127.0.0.1/operator",
            "reviewer_link": "http://127.0.0.1/reviewer",
        }
        manager.update_notifications.side_effect = lambda _record_id, data: data
        recipients = []

        def send(_text, open_ids):
            recipients.append(open_ids)
            return open_ids != ["ou_unreachable"], "ok", []

        with patch.object(
            PortalRuntime, "polling_work_orders", return_value=manager
        ), patch.object(portal_server, "_send_text_to_open_ids_guarded", side_effect=send):
            notifications = PortalRuntime._send_polling_work_order_links({})

        self.assertEqual(recipients[:2], [["ou_unreachable"], ["ou_current"]])
        self.assertTrue(notifications["operator"]["sent"])
        self.assertTrue(notifications["operator"]["fallback"])
        self.assertEqual(
            notifications["operator"]["recipient_open_id"], "ou_current"
        )

    def test_failed_work_order_link_notifications_are_throttled(self) -> None:
        manager = MagicMock()
        manager.group_with_links.side_effect = lambda group, _base: group
        manager.update_notifications.side_effect = lambda _record_id, data: data
        group = {
            "target_record_id": "recNotifyThrottle",
            "title": "轮巡测试",
            "sop_name": "SOP",
            "runs": [{}],
            "operator": {"name": "操作员", "open_id": "ou_op"},
            "reviewer": {"name": "审核员", "open_id": "ou_re"},
            "initiator_open_id": "ou_sender",
            "operator_link": "https://relay.example/op",
            "reviewer_link": "https://relay.example/re",
        }
        with patch.object(
            PortalRuntime, "polling_work_orders", return_value=manager
        ), patch.object(
            portal_server,
            "_send_text_to_open_ids_guarded",
            return_value=(False, "network down", []),
        ) as send:
            notifications = PortalRuntime._send_polling_work_order_links(group)
            PortalRuntime._send_polling_work_order_links(
                {**group, "notifications": notifications}
            )
        self.assertEqual(send.call_count, 4)
        self.assertGreater(notifications["operator"]["next_retry_at"], time.time())

    def test_sending_existing_public_work_order_keeps_settings_port(self) -> None:
        group = {
            'target_record_id':'recPortTest','title':'Port test','sop_name':'SOP','runs':[{}],
            'operator':{'open_id':'ou_fixture_operator'},'reviewer':{'open_id':'ou_fixture_reviewer'},
            'relay':{'mode':'public_relay','registration_state':'registered',
                     'relay_url':'https://public.example:8787',
                     'operator_link':'https://public.example/work#link_id=op&secret=fixture',
                     'reviewer_link':'https://public.example/work#link_id=re&secret=fixture'},
        }
        manager=MagicMock()
        manager.get_group.return_value=group
        manager.group_with_links.side_effect=lambda value,base:PollingWorkOrderService.group_with_links(None,value,base)
        with patch.object(PortalRuntime,'polling_work_orders',return_value=manager), patch.object(PortalRuntime,'polling_work_order_relay',return_value=MagicMock(enabled=True)), patch.object(portal_server,'_send_text_to_open_ids_guarded',return_value=(True,'ok',[])) as send:
            PortalRuntime._send_polling_work_order_links(group)
        self.assertEqual(send.call_count,2)
        self.assertIn('https://public.example:8787/work#link_id=op&secret=fixture',send.call_args_list[0].args[0])
        self.assertIn('https://public.example:8787/work#link_id=re&secret=fixture',send.call_args_list[1].args[0])

    def test_concurrent_automatic_link_send_only_notifies_each_role_once(self) -> None:
        manager = MagicMock()
        state = {
            "target_record_id": "recNotifyConcurrent",
            "title": "轮巡测试",
            "sop_name": "SOP",
            "runs": [{}],
            "operator": {"name": "操作员", "open_id": "ou_op"},
            "reviewer": {"name": "审核员", "open_id": "ou_re"},
            "initiator_open_id": "ou_sender",
            "operator_link": "https://relay.example/op",
            "reviewer_link": "https://relay.example/re",
            "notifications": {},
        }
        manager.group_with_links.side_effect = lambda group, _base: copy.deepcopy(group)
        manager.get_group.side_effect = lambda _record_id: copy.deepcopy(state)

        def save_notifications(_record_id, notifications):
            state["notifications"] = copy.deepcopy(notifications)
            return copy.deepcopy(state)

        manager.update_notifications.side_effect = save_notifications

        def send(_text, _open_ids):
            time.sleep(0.02)
            return True, "ok", []

        with patch.object(
            PortalRuntime, "polling_work_orders", return_value=manager
        ), patch.object(
            PortalRuntime, "_polling_work_order_public_base_url", return_value=""
        ), patch.object(
            portal_server, "_send_text_to_open_ids_guarded", side_effect=send
        ) as send_mock, ThreadPoolExecutor(max_workers=2) as executor:
            list(executor.map(
                lambda _index: PortalRuntime._send_polling_work_order_links(copy.deepcopy(state)),
                range(2),
            ))

        self.assertEqual(send_mock.call_count, 2)
        self.assertEqual(manager.update_notifications.call_count, 1)

    def test_polling_start_persists_lan_group_without_waiting_for_messages(self) -> None:
        manager = MagicMock()
        group = {"target_record_id": "recLanStart", "state": "active"}
        manager.create_group.return_value = group
        prepared = {
            "work_type": "polling",
            "action": "start",
            "polling_work_order_required": True,
            "title": "局域网轮巡测试",
        }
        with patch.object(
            PortalRuntime, "polling_work_orders", return_value=manager
        ), patch.object(
            PortalRuntime,
            "_polling_work_order_public_base_url",
            return_value="http://192.168.1.10:18766",
        ), patch.object(
            PortalRuntime,
            "polling_work_order_public_relay_enabled",
            return_value=False,
        ), patch.object(
            PortalRuntime,
            "polling_work_order_relay",
        ) as relay_connector, patch.object(
            PortalRuntime, "_send_polling_work_order_links"
        ) as send_links, patch.dict(
            os.environ,
            {"CLIPFLOW_POLLING_RELAY_URL": "https://relay.example"},
        ):
            result = PortalRuntime._create_polling_work_order_group(
                prepared, "recLanStart"
            )
        self.assertEqual(result, group)
        self.assertEqual(
            manager.create_group.call_args.kwargs["public_base_url"],
            "http://192.168.1.10:18766",
        )
        self.assertNotIn("public_relay", manager.create_group.call_args.kwargs)
        relay_connector.assert_not_called()
        send_links.assert_not_called()

    def test_unavailable_public_relay_fixes_start_to_local_fallback(self) -> None:
        prepared = {
            "work_type": "polling",
            "notice_type": "设备轮巡",
            "action": "start",
            "status": "开始",
            "title": "公网降级测试",
            "polling_work_order_required": True,
        }
        with patch.object(
            PortalRuntime,
            "polling_work_order_public_relay_enabled",
            return_value=True,
        ), patch.object(
            PortalRuntime,
            "polling_work_order_public_relay_health",
            return_value={
                "ready": False,
                "error": "连接超时",
                "checked_at": 123.0,
                "url": "https://relay.example",
            },
        ):
            resolved = PortalRuntime._resolve_polling_work_order_mode(prepared)
        self.assertEqual(resolved["polling_work_order_mode"], "local_fallback")
        self.assertEqual(resolved["polling_work_order_fallback_reason"], "连接超时")
        text = PortalRuntime.service._synchronize_prepared_notice_text(resolved)["text"]
        self.assertIn("【工单模式】局域网（公网工单不可用，已自动切换）", text)

        manager = MagicMock()
        group = {"target_record_id": "recFallbackStart", "state": "active"}
        manager.create_group.return_value = group
        with patch.object(
            PortalRuntime, "polling_work_orders", return_value=manager
        ), patch.object(
            PortalRuntime,
            "_polling_work_order_public_base_url",
            return_value="http://192.168.1.10:18766",
        ), patch.object(
            PortalRuntime, "polling_work_order_relay"
        ) as relay_connector, patch.object(
            PortalRuntime, "_send_polling_work_order_links"
        ) as send_links:
            PortalRuntime._create_polling_work_order_group(
                resolved, "recFallbackStart"
            )
        self.assertNotIn("public_relay", manager.create_group.call_args.kwargs)
        relay_connector.assert_not_called()
        send_links.assert_not_called()

    def test_ready_public_relay_fixes_start_to_public_mode(self) -> None:
        prepared = {
            "work_type": "polling",
            "action": "start",
            "polling_work_order_required": True,
        }
        with patch.object(
            PortalRuntime,
            "polling_work_order_public_relay_enabled",
            return_value=True,
        ), patch.object(
            PortalRuntime,
            "polling_work_order_public_relay_health",
            return_value={
                "ready": True,
                "checked_at": 123.0,
                "url": "https://relay.example",
            },
        ):
            resolved = PortalRuntime._resolve_polling_work_order_mode(dict(prepared))
        self.assertEqual(resolved["polling_work_order_mode"], "public_relay")

    def test_work_order_mode_is_not_rechecked_after_it_is_fixed(self) -> None:
        prepared = {
            "work_type": "polling",
            "action": "start",
            "polling_work_order_required": True,
        }
        previous = {
            "polling_work_order_mode": "local_fallback",
            "polling_work_order_fallback_reason": "此前检测失败",
            "polling_work_order_relay_checked_at": 123.0,
            "polling_work_order_relay_url": "https://relay.example",
        }
        with patch.object(
            PortalRuntime, "polling_work_order_public_relay_health"
        ) as health:
            resolved = PortalRuntime._resolve_polling_work_order_mode(
                prepared, previous_prepared=previous
            )
        health.assert_not_called()
        self.assertEqual(resolved["polling_work_order_mode"], "local_fallback")
        self.assertEqual(
            resolved["polling_work_order_fallback_reason"], "此前检测失败"
        )

    def test_public_work_order_setting_selects_local_simulation_relay(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            store = LanPortalStateStore(Path(temp) / "state.sqlite3")
            store.put_settings(
                {
                    "polling_work_order_public_relay_enabled": True,
                    "polling_work_order_public_relay_url": "http://192.168.224.122:18767",
                }
            )
            store.put_document(
                "polling_work_order",
                "recOldDirectClient",
                {
                    "target_record_id": "recOldDirectClient",
                    "state": "active",
                    "relay": {
                        "mode": "public_service",
                        "registration_state": "registration_pending",
                        "public_group_id": "old-direct-id",
                    },
                },
            )
            previous_store = PortalRuntime.state_store
            previous_connector = PortalRuntime._polling_relay_connector
            previous_signature = PortalRuntime._polling_relay_connector_signature
            PortalRuntime.state_store = store
            PortalRuntime._polling_relay_connector = None
            PortalRuntime._polling_relay_connector_signature = ""
            try:
                self.assertTrue(
                    PortalRuntime.polling_work_order_public_relay_enabled()
                )
                relay = PortalRuntime.polling_work_order_relay()
                self.assertTrue(relay.enabled)
                self.assertEqual(relay.config.base_url, "http://192.168.224.122:18767")
                self.assertTrue(relay.config.allow_insecure_http)
                migrated = store.get_document(
                    "polling_work_order", "recOldDirectClient"
                )
                self.assertEqual(migrated["relay"]["mode"], "public_relay")
                self.assertEqual(
                    migrated["relay"]["registration_state"],
                    "registration_pending",
                )
                self.assertFalse(migrated["relay"].get("public_group_id"))
            finally:
                PortalRuntime.state_store = previous_store
                PortalRuntime._polling_relay_connector = previous_connector
                PortalRuntime._polling_relay_connector_signature = previous_signature

    def test_existing_public_group_keeps_connector_running_after_toggle_off(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            store = LanPortalStateStore(Path(temp) / "state.sqlite3")
            store.put_settings(
                {
                    "polling_work_order_public_relay_enabled": False,
                    "polling_work_order_public_relay_url": "http://127.0.0.1:18767",
                }
            )
            store.put_document(
                "polling_work_order",
                "recPublicStillRunning",
                {
                    "target_record_id": "recPublicStillRunning",
                    "state": "active",
                    "relay": {"mode": "public_relay", "registration_state": "registered"},
                },
            )
            previous_store = PortalRuntime.state_store
            PortalRuntime.state_store = store
            try:
                self.assertFalse(
                    PortalRuntime.polling_work_order_public_relay_enabled()
                )
                self.assertTrue(
                    PortalRuntime.polling_work_order_public_relay_should_run()
                )
            finally:
                PortalRuntime.state_store = previous_store

    def test_legacy_public_relay_environment_is_used_before_first_setting_save(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            store = LanPortalStateStore(Path(temp) / "state.sqlite3")
            previous_store = PortalRuntime.state_store
            PortalRuntime.state_store = store
            try:
                with patch.dict(
                    os.environ,
                    {
                        "CLIPFLOW_POLLING_RELAY_ENABLED": "1",
                        "CLIPFLOW_POLLING_RELAY_URL": "https://legacy.example.com",
                    },
                ):
                    self.assertTrue(
                        PortalRuntime.polling_work_order_public_relay_enabled()
                    )
                    self.assertEqual(
                        PortalRuntime._polling_work_order_public_relay_url(),
                        "https://legacy.example.com",
                    )
            finally:
                PortalRuntime.state_store = previous_store

    def test_polling_start_can_create_public_relay_order(self) -> None:
        manager = MagicMock()
        relay = MagicMock(enabled=True)
        group = {"target_record_id": "recPublicStart", "state": "active"}
        projected = {
            **group,
            "operator_link": "https://relay.example/operator",
            "reviewer_link": "https://relay.example/reviewer",
        }
        manager.create_group.return_value = group
        manager.get_group.return_value = group
        manager.group_with_links.return_value = projected
        prepared = {
            "work_type": "polling",
            "action": "start",
            "polling_work_order_required": True,
            "title": "公网轮巡测试",
        }
        with patch.object(
            PortalRuntime, "polling_work_orders", return_value=manager
        ), patch.object(
            PortalRuntime,
            "polling_work_order_public_relay_enabled",
            return_value=True,
        ), patch.object(
            PortalRuntime,
            "polling_work_order_relay",
            return_value=relay,
        ), patch.object(
            PortalRuntime, "_send_polling_work_order_links"
        ) as send_links:
            result = PortalRuntime._create_polling_work_order_group(
                prepared, "recPublicStart"
            )
        self.assertEqual(result, group)
        self.assertEqual(manager.create_group.call_args.kwargs["public_base_url"], "")
        self.assertTrue(manager.create_group.call_args.kwargs["public_relay"])
        relay.register_group.assert_not_called()
        send_links.assert_not_called()

    def test_existing_public_group_keeps_its_creation_mode_when_setting_is_off(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            store = LanPortalStateStore(Path(temp) / "state.sqlite3")
            previous_store = PortalRuntime.state_store
            PortalRuntime.state_store = store
            manager = MagicMock()
            group = {
                "target_record_id": "recLegacyRelay",
                "state": "active",
                "relay": {"mode": "public_relay"},
            }
            store.put_document("polling_work_order", "recLegacyRelay", group)
            manager.open_groups.return_value = [group]
            try:
                with patch.object(
                    PortalRuntime, "polling_work_orders", return_value=manager
                ), patch.object(
                    PortalRuntime, "_send_polling_work_order_links"
                ) as send_links, patch.object(
                    portal_server,
                    "query_record_by_id",
                    return_value=(True, {"fields": {}}),
                ), patch.object(
                    PortalRuntime.service,
                    "_target_record_lifecycle",
                    return_value={"active": True, "finished": False},
                ):
                    PortalRuntime.process_polling_work_order_uploads()
                migrated = store.get_document(
                    "polling_work_order", "recLegacyRelay"
                )
                self.assertEqual(migrated["relay"]["mode"], "public_relay")
                send_links.assert_not_called()
            finally:
                PortalRuntime.state_store = previous_store

    def test_resend_links_checks_scope_and_never_returns_capabilities(self) -> None:
        controller = FastAPIPortalController(host="127.0.0.1", port=18766)
        client = TestClient(controller._build_app())
        manager = MagicMock()
        manager.get_group.return_value = {
            "target_record_id": "recRelaySafe",
            "scope": "E",
        }
        sensitive_group = {
            "notifications": {
                "operator": {"sent": True, "fallback": False},
                "reviewer": {"sent": True, "fallback": True},
            },
            "operator_link": "https://relay.example/#secret",
            "reviewer_link": "https://relay.example/#secret2",
            "attachments": [{"path": "D:/private/file.xlsx"}],
        }
        with patch.object(
            controller, "_current_session", return_value={"user": {"open_id": "ou_a"}}
        ), patch.object(
            PortalRuntime, "polling_work_orders", return_value=manager
        ), patch.object(
            controller,
            "_authorized_scope_or_error",
            side_effect=portal_server.PortalError("无权访问 E 楼"),
        ), patch.object(
            PortalRuntime, "_send_polling_work_order_links"
        ) as denied_send:
            denied = client.post(
                "/api/polling-work-orders/recRelaySafe/resend-links"
            )
        self.assertEqual(denied.status_code, 403, denied.text)
        denied_send.assert_not_called()
        with patch.object(
            controller, "_current_session", return_value={"user": {"open_id": "ou_a"}}
        ), patch.object(
            PortalRuntime, "polling_work_orders", return_value=manager
        ), patch.object(
            controller, "_authorized_scope_or_error", return_value="E"
        ), patch.object(
            PortalRuntime,
            "_send_polling_work_order_links",
            return_value=sensitive_group,
        ):
            response = client.post(
                "/api/polling-work-orders/recRelaySafe/resend-links"
            )
        self.assertEqual(response.status_code, 200, response.text)
        response_text = response.text
        self.assertNotIn("operator_link", response_text)
        self.assertNotIn("reviewer_link", response_text)
        self.assertNotIn("D:/private", response_text)

    def test_sop_snapshot_and_serial_role_confirmations(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            store = LanPortalStateStore(root / "state.sqlite3")
            service = PollingWorkOrderService(store)
            service.sop_root = root / "sops"
            service.work_order_root = root / "orders"
            sop = service.save_sop(
                {
                    "scope": "A",
                    "name": "制冷单元切换",
                    "steps": [
                        {"content": "将{{from}}切换至{{to}}，检查{{other}}", "operator_required": True, "reviewer_required": True},
                        {"content": "复查{{to}}", "operator_required": False, "reviewer_required": True},
                    ],
                }
            )
            sop = service.add_sop_attachment(
                sop["sop_id"],
                file_name="原始SOP.txt",
                content=b"sop",
                expected_version=sop["version"],
            )
            prior_version = sop["version"]
            sop = service.save_sop(
                {
                    "sop_id": sop["sop_id"],
                    "scope": "A",
                    "name": "制冷单元切换",
                    "expected_version": prior_version,
                    "steps": [
                        {
                            "content": "将{{from}}切换至{{to}}，检查{{other}}",
                            "operator_required": True,
                            "reviewer_required": True,
                        },
                        {
                            "content": "修改后复查{{to}}",
                            "operator_required": False,
                            "reviewer_required": True,
                        },
                    ],
                }
            )
            self.assertEqual(sop["version"], prior_version + 1)
            self.assertEqual(sop["steps"][1]["content"], "修改后复查{{to}}")
            exempt = service.prepare_start(
                {
                    "work_type": "polling",
                    "scope": "A",
                    "action": "start",
                    "_web_action_request": True,
                    "polling_work_order_exempt": True,
                },
                job_id="job-exempt",
                people=[],
            )
            self.assertTrue(exempt["polling_work_order_exempt"])
            self.assertNotIn("polling_work_order_required", exempt)
            self.assertNotIn("polling_work_order_spec", exempt)
            prepared = service.prepare_start(
                {
                    "work_type": "polling",
                    "scope": "A",
                    "action": "start",
                    "_web_action_request": True,
                    "polling_sop_id": sop["sop_id"],
                    "polling_sop_version": sop["version"],
                    "polling_run_count": 1,
                    "polling_runs": [{"from_unit": "2#", "to_unit": "3#"}],
                    "polling_operator_record_id": "operator",
                    "polling_reviewer_record_id": "reviewer",
                },
                job_id="job-1",
                people=[
                    {"record_id": "operator", "name": "操作员", "open_id": "ou_operator"},
                    {"record_id": "reviewer", "name": "审核员", "open_id": "ou_reviewer"},
                ],
            )
            with self.assertRaisesRegex(Exception, "不属于当前楼栋"):
                service.prepare_start(
                    {
                        "work_type": "polling",
                        "scope": "B",
                        "action": "start",
                        "_web_action_request": True,
                        "polling_sop_id": sop["sop_id"],
                        "polling_sop_version": sop["version"],
                    },
                    job_id="job-wrong-building",
                    people=[],
                )
            with self.assertRaisesRegex(Exception, "不能跨越"):
                service.prepare_start(
                    {
                        "work_type": "polling",
                        "scope": "A",
                        "action": "start",
                        "_web_action_request": True,
                        "polling_sop_id": sop["sop_id"],
                        "polling_sop_version": sop["version"],
                        "polling_run_count": 1,
                        "polling_runs": [{"from_unit": "1#", "to_unit": "5#"}],
                    },
                    job_id="job-cross-group",
                    people=[],
                )
            with self.assertRaisesRegex(Exception, "已选择的设备编号"):
                service.prepare_start(
                    {
                        "work_type": "polling",
                        "scope": "A",
                        "action": "start",
                        "_web_action_request": True,
                        "polling_sop_id": sop["sop_id"],
                        "polling_sop_version": sop["version"],
                        "polling_run_count": 2,
                        "polling_runs": [
                            {"from_unit": "1#", "to_unit": "2#"},
                            {"from_unit": "2#", "to_unit": "3#"},
                        ],
                    },
                    job_id="job-reused-unit",
                    people=[],
                )
            h_duty = service._person_by_id(
                [], "h_duty_account", allow_h_duty=True
            )
            self.assertEqual(h_duty["name"], "H楼值班账号")
            self.assertEqual(h_duty["building"], "H楼")
            with self.assertRaisesRegex(Exception, "不能是同一人"):
                service.prepare_start(
                    {
                        "work_type": "polling",
                        "scope": "A",
                        "action": "start",
                        "_web_action_request": True,
                        "polling_sop_id": sop["sop_id"],
                        "polling_sop_version": sop["version"],
                        "polling_run_count": 1,
                        "polling_runs": [{"from_unit": "1#", "to_unit": "2#"}],
                        "polling_operator_record_id": "same-person",
                        "polling_reviewer_record_id": "same-person",
                    },
                    job_id="job-same-person",
                    people=[
                        {
                            "record_id": "same-person",
                            "name": "同一人员",
                            "open_id": "ou_same",
                        }
                    ],
                )
            group = service.create_group(
                prepared,
                target_record_id="recTarget1",
                title="轮巡测试",
                public_base_url="http://127.0.0.1:18766",
            )
            self.assertEqual(group["steps"][0]["content"], "将2#切换至3#，检查1#")
            self.assertEqual(group["steps"][1]["content"], "修改后复查3#")
            operator_token = service.role_token("recTarget1", "operator")
            reviewer_token = service.role_token("recTarget1", "reviewer")
            with self.assertRaisesRegex(Exception, "工单总览"):
                service.confirm(reviewer_token, step_key="1:1", expected_version=1)
            session = service.activate(
                operator_token, run_index=1, expected_version=1
            )
            with self.assertRaises(Exception):
                service.confirm(
                    reviewer_token,
                    step_key="1:1",
                    expected_version=session["version"],
                )
            session = service.add_step_photo(
                operator_token,
                step_key="1:1",
                expected_version=session["version"],
                file_name="第一步.png",
                mime_type="image/png",
                content=_png_bytes(),
            )
            session = service.confirm(
                operator_token,
                step_key="1:1",
                expected_version=session["version"],
            )
            session = service.confirm(reviewer_token, step_key="1:1", expected_version=session["version"])
            self.assertEqual(session["current_index"], 1)
            first_photo_path = Path(
                service.get_group("recTarget1")["steps"][0]["photos"][0]["path"]
            )
            with self.assertRaisesRegex(Exception, "只有现场审核人"):
                service.rollback_previous(
                    operator_token,
                    step_key="1:2",
                    expected_version=session["version"],
                )
            session = service.rollback_previous(
                reviewer_token,
                step_key="1:2",
                expected_version=session["version"],
            )
            self.assertEqual(session["current_index"], 0)
            self.assertFalse(first_photo_path.exists())
            self.assertFalse(service.get_group("recTarget1")["steps"][0]["photos"])
            session = service.add_step_photo(
                operator_token,
                step_key="1:1",
                expected_version=session["version"],
                file_name="第一步重做.png",
                mime_type="image/png",
                content=_png_bytes("#ff8a00"),
            )
            session = service.confirm(
                operator_token, step_key="1:1", expected_version=session["version"]
            )
            session = service.confirm(
                reviewer_token, step_key="1:1", expected_version=session["version"]
            )
            with self.assertRaisesRegex(Exception, "拍摄并上传"):
                service.confirm(
                    reviewer_token,
                    step_key="1:2",
                    expected_version=session["version"],
                )
            session = service.add_step_photo(
                reviewer_token,
                step_key="1:2",
                expected_version=session["version"],
                file_name="第二步.png",
                mime_type="image/png",
                content=_png_bytes("#12a150"),
            )
            session = service.confirm(reviewer_token, step_key="1:2", expected_version=session["version"])
            self.assertEqual(session["state"], "upload_pending")
            service.mark_upload_result("recTarget1", success=True, file_tokens=["file-token"])
            self.assertEqual(service.session(operator_token)["state"], "completed")
            self.assertEqual(service.open_groups()[0]["state"], "completed")
            service.cancel_group("recTarget1", reason="target_terminal")
            with self.assertRaises(PollingWorkOrderTokenError):
                service.session(operator_token)

            prepared_two_runs = service.prepare_start(
                {
                    "work_type": "polling",
                    "scope": "A",
                    "action": "start",
                    "_web_action_request": True,
                    "polling_sop_id": sop["sop_id"],
                    "polling_sop_version": sop["version"],
                    "polling_run_count": 2,
                    "polling_runs": [
                        {"from_unit": "1#", "to_unit": "2#"},
                        {"from_unit": "4#", "to_unit": "5#"},
                    ],
                    "polling_operator_record_id": "operator",
                    "polling_reviewer_record_id": "reviewer",
                },
                job_id="job-two-runs",
                people=[
                    {"record_id": "operator", "name": "操作员"},
                    {"record_id": "reviewer", "name": "审核员"},
                ],
            )
            service.create_group(
                prepared_two_runs,
                target_record_id="recTargetTwoRuns",
                title="两次轮巡测试",
                public_base_url="",
            )
            operator_two = service.role_token("recTargetTwoRuns", "operator")
            reviewer_two = service.role_token("recTargetTwoRuns", "reviewer")
            overview = service.session(operator_two)
            self.assertEqual(
                [item["state"] for item in overview["work_orders"]],
                ["available", "available"],
            )
            second_run_started = service.activate(
                operator_two,
                run_index=2,
                expected_version=overview["version"],
            )
            with self.assertRaisesRegex(Exception, "已选择其他工单"):
                service.activate(
                    reviewer_two,
                    run_index=1,
                    expected_version=overview["version"],
                )
            self.assertTrue(service.session(operator_two)["can_release_selection"])
            self.assertFalse(service.session(reviewer_two)["can_release_selection"])
            with self.assertRaisesRegex(Exception, "不能代为退出"):
                service.release_selection(
                    reviewer_two,
                    run_index=2,
                    expected_version=second_run_started["version"],
                )
            released = service.release_selection(
                operator_two,
                run_index=2,
                expected_version=second_run_started["version"],
            )
            self.assertEqual(
                [item["state"] for item in released["work_orders"]],
                ["available", "available"],
            )
            self.assertEqual(
                service.get_group("recTargetTwoRuns")["steps"][2]["activated_at_ts"],
                0.0,
            )
            second_run_started = service.activate(
                reviewer_two,
                run_index=2,
                expected_version=released["version"],
            )
            sop = service.save_sop(
                {
                    "sop_id": sop["sop_id"],
                    "scope": "A",
                    "name": "制冷单元切换",
                    "expected_version": sop["version"],
                    "steps": [
                        {**step, "time_limit_seconds": seconds}
                        for step, seconds in zip(sop["steps"], (4, 6))
                    ],
                }
            )
            two_run_group = service.get_group("recTargetTwoRuns")
            self.assertEqual(
                [step["time_limit_seconds"] for step in two_run_group["steps"][:2]],
                [4, 6],
            )
            self.assertEqual(
                [step["time_limit_seconds"] for step in two_run_group["steps"][2:]],
                [0, 0],
            )
            self.assertEqual(two_run_group["version"], second_run_started["version"])
            first_run = service.session(operator_two)
            self.assertEqual(
                [item["state"] for item in first_run["work_orders"]],
                ["locked", "active"],
            )
            self.assertTrue(first_run["steps"])
            self.assertTrue(
                all(step["run_index"] == 2 for step in first_run["steps"])
            )
            first_run = service.add_step_photo(
                operator_two,
                step_key="2:1",
                expected_version=first_run["version"],
                file_name="工单2第一步.png",
                mime_type="image/png",
                content=_png_bytes(),
            )
            first_run = service.confirm(
                operator_two, step_key="2:1", expected_version=first_run["version"]
            )
            first_run = service.confirm(
                reviewer_two, step_key="2:1", expected_version=first_run["version"]
            )
            first_run = service.add_step_photo(
                reviewer_two,
                step_key="2:2",
                expected_version=first_run["version"],
                file_name="工单2第二步.png",
                mime_type="image/png",
                content=_png_bytes("#12a150"),
            )
            second_run = service.confirm(
                reviewer_two, step_key="2:2", expected_version=first_run["version"]
            )
            self.assertEqual(
                [item["state"] for item in second_run["work_orders"]],
                ["available", "completed"],
            )
            self.assertEqual(second_run["current_run_index"], 0)
            self.assertFalse(second_run["steps"])
            activated_second_run = service.activate(
                operator_two,
                run_index=1,
                expected_version=second_run["version"],
            )
            self.assertEqual(activated_second_run["steps"][0]["step_index"], 1)
            self.assertEqual(
                activated_second_run["steps"][0]["content"],
                "将1#切换至2#，检查3#",
            )
            self.assertTrue(activated_second_run["steps"][0]["timer_started"])
            self.assertGreater(activated_second_run["steps"][0]["remaining_seconds"], 0)
            idempotent_activation = service.activate(
                reviewer_two,
                run_index=1,
                expected_version=second_run["version"],
            )
            self.assertEqual(
                idempotent_activation["version"], activated_second_run["version"]
            )

            with self.assertRaisesRegex(Exception, "必须使用"):
                service.save_sop(
                    {
                        "scope": "A",
                        "name": "旧占位符",
                        "steps": [
                            {
                                "content": "将{from}切换至{to}",
                                "operator_required": True,
                            }
                        ],
                    }
                )

            other_building = service.save_sop(
                {
                    "scope": "B",
                    "name": "制冷单元切换",
                    "steps": [
                        {"content": "B楼步骤", "operator_required": True, "reviewer_required": False}
                    ],
                }
            )
            self.assertEqual([item["sop_id"] for item in service.list_sops("A")], [sop["sop_id"]])
            self.assertEqual([item["sop_id"] for item in service.list_sops("B")], [other_building["sop_id"]])

    def test_step_timer_photo_and_execution_workbook(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            store = LanPortalStateStore(root / "state.sqlite3")
            service = PollingWorkOrderService(store)
            service.sop_root = root / "sops"
            service.work_order_root = root / "orders"
            sop = service.save_sop(
                {
                    "scope": "A",
                    "name": "倒计时拍照 SOP",
                    "steps": [
                        {
                            "content": "将{{from}}切换到{{to}}",
                            "operator_required": True,
                            "reviewer_required": True,
                            "time_limit_seconds": 2,
                        }
                    ],
                }
            )
            sop = service.add_sop_attachment(
                sop["sop_id"],
                file_name="原SOP.pdf",
                content=b"sop",
                expected_version=sop["version"],
            )
            prepared = service.prepare_start(
                {
                    "work_type": "polling",
                    "scope": "A",
                    "action": "start",
                    "_web_action_request": True,
                    "polling_sop_id": sop["sop_id"],
                    "polling_sop_version": sop["version"],
                    "polling_run_count": 1,
                    "polling_runs": [{"from_unit": "1#", "to_unit": "2#"}],
                    "polling_operator_record_id": "operator",
                    "polling_reviewer_record_id": "reviewer",
                },
                job_id="timer-photo-job",
                people=[
                    {"record_id": "operator", "name": "操作员", "open_id": "ou_operator"},
                    {"record_id": "reviewer", "name": "审核员", "open_id": "ou_reviewer"},
                ],
            )
            service.create_group(
                prepared,
                target_record_id="recTimerPhoto",
                title="轮巡倒计时拍照测试",
                public_base_url="http://127.0.0.1:18766",
            )
            operator_token = service.role_token("recTimerPhoto", "operator")
            reviewer_token = service.role_token("recTimerPhoto", "reviewer")
            first_session = service.session(operator_token)
            self.assertFalse(first_session["steps"])
            self.assertEqual(first_session["work_orders"][0]["state"], "available")
            with self.assertRaisesRegex(Exception, "工单总览"):
                service.confirm(
                    operator_token,
                    step_key="1:1",
                    expected_version=first_session["version"],
                )
            activated_session = service.activate(
                operator_token,
                run_index=1,
                expected_version=first_session["version"],
            )
            self.assertTrue(activated_session["steps"][0]["timer_started"])
            self.assertGreater(activated_session["steps"][0]["remaining_seconds"], 0)
            with self.assertRaisesRegex(Exception, "还需等待"):
                service.confirm(
                    operator_token,
                    step_key="1:1",
                    expected_version=activated_session["version"],
                )
            with self.assertRaisesRegex(Exception, "内容损坏"):
                service.add_step_photo(
                    operator_token,
                    step_key="1:1",
                    expected_version=activated_session["version"],
                    file_name="损坏.png",
                    mime_type="image/png",
                    content=b"not-an-image",
                )
            photo_bytes = _png_bytes()
            photo_session = service.add_step_photo(
                operator_token,
                step_key="1:1",
                expected_version=activated_session["version"],
                file_name="现场.png",
                mime_type="image/png",
                content=photo_bytes,
            )
            photo = photo_session["steps"][0]["photos"][0]
            self.assertIn("/api/polling-work-orders/photos/", photo["preview_url"])
            self.assertEqual(
                service.step_photo_content(
                    operator_token, photo_id=photo["photo_id"]
                )[0],
                photo_bytes,
            )
            with self.assertRaisesRegex(Exception, "已有操作记录"):
                service.release_selection(
                    operator_token,
                    run_index=1,
                    expected_version=photo_session["version"],
                )
            stored_before_replace = service.get_group("recTimerPhoto")["steps"][0][
                "photos"
            ][0]
            photo_directory = Path(stored_before_replace["path"]).parent
            files_before_failed_replace = set(photo_directory.iterdir())
            with patch.object(
                store, "put_document", side_effect=RuntimeError("forced write failure")
            ), self.assertRaisesRegex(RuntimeError, "forced write failure"):
                service.add_step_photo(
                    operator_token,
                    step_key="1:1",
                    expected_version=photo_session["version"],
                    file_name="替换照片.png",
                    mime_type="image/png",
                    content=_png_bytes("#ff0000"),
                )
            stored_photo = service.get_group("recTimerPhoto")["steps"][0]["photos"][0]
            self.assertEqual(stored_photo["photo_id"], photo["photo_id"])
            self.assertEqual(set(photo_directory.iterdir()), files_before_failed_replace)
            multi_photo_session = service.add_step_photo(
                operator_token,
                step_key="1:1",
                expected_version=photo_session["version"],
                file_name="补充照片.png",
                mime_type="image/png",
                content=_png_bytes("#ff0000"),
            )
            self.assertEqual(len(multi_photo_session["steps"][0]["photos"]), 2)
            self.assertTrue(Path(stored_photo["path"]).is_file())
            duplicate_session = service.add_step_photo(
                operator_token,
                step_key="1:1",
                expected_version=multi_photo_session["version"],
                file_name="重复照片.png",
                mime_type="image/png",
                content=photo_bytes,
            )
            self.assertEqual(len(duplicate_session["steps"][0]["photos"]), 2)
            with patch(
                "lan_bitable_template_portal.polling_work_orders.POLLING_STEP_MAX_PHOTOS",
                2,
            ), self.assertRaisesRegex(Exception, "最多上传 2 张"):
                service.add_step_photo(
                    operator_token,
                    step_key="1:1",
                    expected_version=multi_photo_session["version"],
                    file_name="超限照片.png",
                    mime_type="image/png",
                    content=_png_bytes("#00ff00"),
                )
            with patch(
                "lan_bitable_template_portal.polling_work_orders.POLLING_WORK_ORDER_MAX_PHOTOS",
                2,
            ), self.assertRaisesRegex(Exception, "整个工单组最多上传 2 张"):
                service.add_step_photo(
                    operator_token,
                    step_key="1:1",
                    expected_version=multi_photo_session["version"],
                    file_name="工单总量超限照片.png",
                    mime_type="image/png",
                    content=_png_bytes("#00ff00"),
                )
            group = service.get_group("recTimerPhoto")
            group["steps"][0]["activated_at_ts"] = time.time() - 3
            store.put_document("polling_work_order", "recTimerPhoto", group)
            operator_confirmed_session = service.confirm(
                operator_token,
                step_key="1:1",
                expected_version=multi_photo_session["version"],
            )
            with self.assertRaisesRegex(Exception, "已有确认"):
                service.add_step_photo(
                    operator_token,
                    step_key="1:1",
                    expected_version=operator_confirmed_session["version"],
                    file_name="确认后替换.png",
                    mime_type="image/png",
                    content=_png_bytes("#00ff00"),
                )
            completed_session = service.confirm(
                reviewer_token,
                step_key="1:1",
                expected_version=operator_confirmed_session["version"],
            )
            self.assertEqual(completed_session["state"], "upload_pending")
            workbook_info = service.build_execution_workbook("recTimerPhoto")
            workbook = load_workbook(workbook_info["path"], data_only=False)
            sheet = workbook["工单1 1#→2#"]
            values = [cell.value for row in sheet.iter_rows() for cell in row]
            self.assertEqual(sheet.max_row, 9)
            self.assertEqual(sheet["C3"].value, "操作员")
            self.assertEqual(sheet["D3"].value, "审核员")
            self.assertIsNotNone(sheet["E3"].value)
            self.assertEqual(sheet["C9"].value, 1)
            self.assertEqual(sheet["D9"].value, "将1#切换到2#")
            self.assertEqual(len(sheet._images), 3)
            self.assertIn(5, [image.anchor._from.col for image in sheet._images])
            self.assertIn("$A$1:$F$9", str(sheet.print_area))
            workbook.close()
            self.assertIn("倒计时拍照 SOP · 1#→2#", values)
            self.assertIn("将1#切换到2#", values)
            self.assertRegex(
                workbook_info["name"],
                r"^A楼\d{4}年\d{1,2}月\d{1,2}日-倒计时拍照 SOP-1#轮巡至2#-操作人-操作员-审核人-审核员-操作记录\.xlsx$",
            )
            with patch(
                "lan_bitable_template_portal.polling_work_orders.POLLING_WORK_ORDER_MAX_BYTES",
                1,
            ), self.assertRaisesRegex(Exception, "超过允许大小"):
                service.build_execution_workbook("recTimerPhoto")

    def test_template_workbook_expands_steps_and_uses_one_sheet_per_run(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            store = LanPortalStateStore(root / "state.sqlite3")
            service = PollingWorkOrderService(store)
            service.work_order_root = root / "orders"
            target_record_id = "recTemplateRuns"
            photo_root = service._group_directory(target_record_id) / "photos"
            photo_root.mkdir(parents=True)
            first_photo = photo_root / "first.png"
            second_photo = photo_root / "second.png"
            first_photo.write_bytes(_png_bytes("#1678ff", (240, 120)))
            second_photo.write_bytes(_png_bytes("#12a150", (120, 240)))
            steps = []
            for run_index, label in ((1, "1#→2#"), (2, "4#→5#")):
                step_count = 2 if run_index == 1 else 3
                for step_index in range(1, step_count + 1):
                    content = (
                        '=WEBSERVICE("https://example.invalid")'
                        if run_index == 1 and step_index == 1
                        else f"{label}操作步骤{step_index}"
                    )
                    photo_path = first_photo if run_index == 1 else second_photo
                    steps.append(
                        {
                            "step_key": f"{run_index}:{step_index}",
                            "run_index": run_index,
                            "run_label": label,
                            "step_index": step_index,
                            "content": content,
                            "operator_required": True,
                            "reviewer_required": False,
                            "operator_confirmation": {
                                "confirmed_at": (
                                    f"2026-08-26 {9 + run_index:02d}:0{step_index}:00"
                                )
                            },
                            "reviewer_confirmation": {},
                            "photos": (
                                [
                                    {
                                        "photo_id": f"photo-{run_index}",
                                        "name": photo_path.name,
                                        "path": str(photo_path),
                                    }
                                ]
                                if photo_path
                                else []
                            ),
                        }
                    )
            store.put_document(
                "polling_work_order",
                target_record_id,
                {
                    "group_id": target_record_id,
                    "target_record_id": target_record_id,
                    "state": "upload_pending",
                    "scope": "A",
                    "title": "双工单模板测试",
                    "sop_name": "冷机切换SOP",
                    "operator": {"name": "操作员甲"},
                    "reviewer": {"name": "审核员乙"},
                    "runs": [
                        {"from_unit": "1#", "to_unit": "2#"},
                        {"from_unit": "4#", "to_unit": "5#"},
                    ],
                    "steps": steps,
                    "updated_at": "2026-08-26 11:03:00",
                },
            )

            workbook_info = service.build_execution_workbook(target_record_id)
            workbook = load_workbook(workbook_info["path"], data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["工单1 1#→2#", "工单2 4#→5#"],
            )
            self.assertEqual(
                workbook_info["name"],
                "A楼2026年8月26日-冷机切换SOP-1#4#轮巡至2#5#-操作人-操作员甲-审核人-审核员乙-操作记录.xlsx",
            )
            first_sheet, second_sheet = workbook.worksheets
            self.assertEqual(first_sheet.max_row, 10)
            self.assertEqual(second_sheet.max_row, 11)
            self.assertEqual(
                [first_sheet.cell(row=row, column=3).value for row in range(9, 11)],
                [1, 2],
            )
            self.assertEqual(
                [second_sheet.cell(row=row, column=3).value for row in range(9, 12)],
                [1, 2, 3],
            )
            self.assertIn(
                "A11:B11", {str(item) for item in second_sheet.merged_cells.ranges}
            )
            self.assertIn(
                "D11:E11", {str(item) for item in second_sheet.merged_cells.ranges}
            )
            self.assertEqual(second_sheet["C10"].style_id, second_sheet["C11"].style_id)
            self.assertEqual(
                second_sheet.row_dimensions[10].height,
                second_sheet.row_dimensions[11].height,
            )
            for sheet, last_row, image_count in (
                (first_sheet, 10, 3),
                (second_sheet, 11, 4),
            ):
                self.assertEqual(len(sheet._images), image_count)
                self.assertIn(f"$A$1:$E${last_row}", str(sheet.print_area))
                self.assertEqual(sheet.print_title_rows, "$8:$8")
            self.assertEqual(first_sheet["C3"].value, "操作员甲")
            self.assertEqual(first_sheet["D3"].value, "审核员乙")
            self.assertEqual(first_sheet["D9"].value, "'=WEBSERVICE(\"https://example.invalid\")")
            self.assertEqual(second_sheet["D9"].value, "4#→5#操作步骤1")
            self.assertEqual(first_sheet["E3"].number_format, "yyyy-mm-dd hh:mm:ss")
            self.assertLess(first_sheet["E3"].value, second_sheet["E3"].value)
            workbook.close()
            from PIL import Image

            with zipfile.ZipFile(workbook_info["path"]) as archive:
                embedded_sizes = []
                for name in archive.namelist():
                    if not name.startswith("xl/media/"):
                        continue
                    with Image.open(io.BytesIO(archive.read(name))) as image:
                        embedded_sizes.append(image.size)
            self.assertTrue(embedded_sizes)
            self.assertTrue(
                all(width <= 230 and height <= 190 for width, height in embedded_sizes)
            )
            self.assertNotIn((240, 120), embedded_sizes)
            self.assertNotIn((120, 240), embedded_sizes)

    def test_template_workbook_fails_when_photo_metadata_file_is_missing(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            store = LanPortalStateStore(root / "state.sqlite3")
            service = PollingWorkOrderService(store)
            service.work_order_root = root / "orders"
            target_record_id = "recMissingPhoto"
            missing_path = service._group_directory(target_record_id) / "photos" / "missing.png"
            store.put_document(
                "polling_work_order",
                target_record_id,
                {
                    "group_id": target_record_id,
                    "target_record_id": target_record_id,
                    "state": "upload_pending",
                    "sop_name": "照片缺失SOP",
                    "operator": {"name": "操作员"},
                    "reviewer": {"name": "审核员"},
                    "runs": [{"from_unit": "1#", "to_unit": "2#"}],
                    "steps": [
                        {
                            "run_index": 1,
                            "step_index": 1,
                            "content": "检查照片",
                            "photos": [{"name": "missing.png", "path": str(missing_path)}],
                        }
                    ],
                    "updated_at": "2026-08-26 12:00:00",
                },
            )

            with self.assertRaisesRegex(Exception, "本地照片不存在"):
                service.build_execution_workbook(target_record_id)
            self.assertFalse(
                list(service._group_directory(target_record_id).glob("*.xlsx"))
            )

    def test_finalize_uploads_only_generated_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            store = LanPortalStateStore(root / "state.sqlite3")
            service = PollingWorkOrderService(store)
            service.sop_root = root / "sops"
            service.work_order_root = root / "orders"
            sop = service.save_sop(
                {
                    "scope": "A",
                    "name": "最终上传 SOP",
                    "steps": [
                        {
                            "content": "执行{{from}}到{{to}}",
                            "operator_required": True,
                            "reviewer_required": False,
                        }
                    ],
                }
            )
            sop = service.add_sop_attachment(
                sop["sop_id"],
                file_name="不会直接上传.pdf",
                content=b"source-sop",
                expected_version=sop["version"],
            )
            prepared = service.prepare_start(
                {
                    "work_type": "polling",
                    "scope": "A",
                    "action": "start",
                    "_web_action_request": True,
                    "polling_sop_id": sop["sop_id"],
                    "polling_sop_version": sop["version"],
                    "polling_run_count": 1,
                    "polling_runs": [{"from_unit": "1#", "to_unit": "2#"}],
                    "polling_operator_record_id": "operator",
                    "polling_reviewer_record_id": "reviewer",
                },
                job_id="finalize-job",
                people=[
                    {"record_id": "operator", "name": "操作员"},
                    {"record_id": "reviewer", "name": "审核员"},
                ],
            )
            service.create_group(
                prepared,
                target_record_id="recFinalize",
                title="最终上传测试",
                public_base_url="",
            )
            token = service.role_token("recFinalize", "operator")
            active_session = service.activate(
                token,
                run_index=1,
                expected_version=service.session(token)["version"],
            )
            photo_session = service.add_step_photo(
                token,
                step_key="1:1",
                expected_version=active_session["version"],
                file_name="步骤照片.png",
                mime_type="image/png",
                content=_png_bytes(),
            )
            service.confirm(
                token,
                step_key="1:1",
                expected_version=photo_session["version"],
            )
            service.mark_upload_progress(
                "recFinalize",
                token_by_sha256={"old-workbook-digest": "old-workbook-token"},
            )
            expected_workbook_name = service.build_execution_workbook("recFinalize")[
                "name"
            ]
            uploaded_names = []
            patches = []

            def upload_file(*, file_path, file_name, app_token):
                del file_path, app_token
                uploaded_names.append(file_name)
                return "workbook-token" if file_name.endswith(".xlsx") else "photo-token"

            active_fields = {
                "轮巡状态": "开始",
                "工单附件": [
                    {"file_token": "manual-attachment-token"},
                    {"file_token": "old-workbook-token"},
                ],
                "操作照片": [{"file_token": "historical-photo-token"}],
            }
            verified_fields = {
                **active_fields,
                "工单附件": [
                    {"file_token": "manual-attachment-token"},
                    {"file_token": "workbook-token"},
                ],
            }
            with patch.object(
                PortalRuntime, "polling_work_orders", return_value=service
            ), patch.object(
                PortalRuntime.service,
                "_upload_bitable_file",
                side_effect=upload_file,
            ), patch.object(
                portal_server,
                "query_record_by_id",
                side_effect=[
                    (True, {"fields": active_fields}),
                    (True, {"fields": verified_fields}),
                ],
            ), patch.object(
                portal_server,
                "update_bitable_record_fields",
                side_effect=lambda _record_id, _notice_type, fields: (
                    patches.append(fields) or (True, "recFinalize")
                ),
            ):
                result = PortalRuntime.finalize_polling_work_order_group(
                    "recFinalize"
                )
            self.assertTrue(result["ok"], result)
            self.assertEqual(uploaded_names, [expected_workbook_name])
            self.assertEqual(
                patches[0]["工单附件"],
                [
                    {"file_token": "manual-attachment-token"},
                    {"file_token": "workbook-token"},
                ],
            )
            self.assertNotIn("操作照片", patches[0])

    def test_local_notice_image_survives_store_reload(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            store = LanPortalStateStore(root / "state.sqlite3")
            images = LocalNoticeImageStore(store)
            images.root = root / "images"
            saved = images.save(
                identity="change:rec1",
                kind="ali",
                content=b"png-bytes",
                file_name="confirm.png",
                mime_type="image/png",
                scope="A",
            )
            images.mark_feishu_uploaded(
                saved["local_image_id"], file_token="file-token", target_written=True
            )
            reloaded = LocalNoticeImageStore(store)
            reloaded.root = root / "images"
            item = reloaded.list("change:rec1", scope="A")[0]
            self.assertEqual(item["preview_url"], f"/api/notice-images/{saved['local_image_id']}")
            self.assertTrue(item["target_written"])
            self.assertNotIn("file_token", item)
            self.assertNotIn("feishu_file_token", item)
            self.assertEqual(reloaded.content(saved["local_image_id"])[0], b"png-bytes")


if __name__ == "__main__":
    unittest.main()
