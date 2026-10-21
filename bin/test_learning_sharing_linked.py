"""Offline contracts for Feishu linked materials; all HTTP uses MockTransport."""
import hashlib
import io
import json
from pathlib import Path
import socket
import unittest
from types import SimpleNamespace
from unittest.mock import Mock, patch

import httpx

from bin.test_learning_sharing_materials import local_only
from lan_bitable_template_portal.learning import LearningError
from lan_bitable_template_portal.sharing_sources import read_material_link, resolve_document, source_url


class LinkedMaterialTests(unittest.TestCase):
    def setUp(self):
        self.enterContext(patch('socket.socket.connect', local_only))
        from upload_event_module.services import feishu_token_manager as auth
        self.credentials = ('fixture-app', 'fixture-secret')
        self.manager = SimpleNamespace(_credential_identity=self.credentials,
            _credentials=Mock(return_value=self.credentials), get_tenant_token=Mock(return_value='fixture-renewed'))
        self.enterContext(patch.object(auth, 'token_manager', self.manager))
        self.enterContext(patch.object(auth, 'config', SimpleNamespace(user_token='fixture-token')))
        self.calls, self.paths = [], []
        self.name, self.revision, self.node_type = 'material.pdf', 1, 'file'
        self.blocks = [{'items': [], 'has_more': False}]
        self.meta_override = None
        self.download = lambda request: httpx.Response(200, content=b'fixture-content')
        self.client = self.enterContext(httpx.Client(transport=httpx.MockTransport(self.transport), follow_redirects=True))
        def payload(method, url, **kwargs):
            response = self.client.request(method, url, headers=kwargs['headers'], params=kwargs.get('params'),
                json=kwargs.get('json_payload'), follow_redirects=False)
            response.raise_for_status()
            return response.json()
        self.portal = SimpleNamespace(_auth_headers=lambda: {'Authorization': 'Bearer fixture-token'},
            _request_payload=payload, _http_client=SimpleNamespace(_client_for_request=lambda: self.client))
        self.parser = self.enterContext(patch('openclaw_service.assistant.lighthouse_document_tools.run_local', side_effect=self.parse))

    def transport(self, request):
        self.assertEqual(request.url.host, 'open.feishu.cn')
        self.calls.append(request)
        path = request.url.path
        if path.endswith('/metas/batch_query'):
            self.assertEqual(request.method, 'POST')
            self.assertEqual(json.loads(request.content), {'request_docs': [{'doc_token': 'FileToken', 'doc_type': 'file'}]})
            data = self.meta_override if self.meta_override is not None else {'metas': [
                {'doc_token': 'FileToken', 'doc_type': 'file', 'title': self.name, 'latest_modify_time': self.revision,
                 'url': 'https://untrusted.invalid/ignored'}], 'failed_list': []}
        elif path.endswith('/spaces/get_node'):
            data = {'node': {'obj_type': self.node_type, 'obj_token': 'FileToken' if self.node_type == 'file' else 'DocToken'}}
        elif path.endswith('/download'):
            return self.download(request)
        elif path.endswith('/raw_content'):
            data = {'content': 'Document introduction'}
        elif path.endswith('/blocks'):
            self.assertEqual(request.url.params['document_revision_id'], '1')
            data = self.blocks.pop(0)
        elif path.endswith('/documents/DocToken'):
            data = {'document': {'title': 'Doc material', 'revision_id': 1}}
        else:
            raise AssertionError('Unexpected fixture API path: ' + path)
        return httpx.Response(200, json={'code': 0, 'data': data})

    def parse(self, kind, payload):
        self.assertEqual(kind, 'markdown')
        path = Path(payload['path']); self.paths.append(path)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), payload['sha256'])
        return {'text': 'parsed complete text', 'extraction_complete': True,
                'references': [{'kind': 'slide' if path.suffix == '.pptx' else 'page', 'page': 2, 'start': 0, 'end': 20}]}

    def test_direct_and_wiki_files_all_supported_formats_share_parser_and_cleanup(self):
        for suffix in ('pdf', 'docx', 'xlsx', 'xlsm', 'pptx'):
            self.name = '材料.' + suffix
            for kind in ('file', 'wiki'):
                with self.subTest(suffix=suffix, kind=kind):
                    source = read_material_link(self.portal, f'https://vnet.feishu.cn/{kind}/FileToken')
                    self.assertEqual(source['title'], self.name)
                    self.assertEqual(source['text'], 'parsed complete text')
                    self.assertEqual(source['references'][0]['page'], 2)
                    self.assertEqual(self.paths[-1].suffix, '.' + suffix)
                    self.assertFalse(self.paths[-1].exists(), 'linked bytes are temporary, snapshot remains')
        self.assertEqual(self.parser.call_count, 10)
        self.assertTrue(all(request.headers['authorization'] == 'Bearer fixture-token' for request in self.calls))

    def test_single_docx_attachment_pagination_preserves_body_and_slide_offsets(self):
        self.blocks = [{'items': [], 'has_more': True, 'page_token': 'next'},
                       {'items': [{'file': {'token': 'Embedded', 'name': 'deck.pptx'}}], 'has_more': False}]
        source = read_material_link(self.portal, 'https://vnet.feishu.cn/docx/DocToken')
        ref = source['references'][0]
        self.assertEqual(source['text'][ref['start']:ref['end']], 'parsed complete text')
        self.assertEqual(ref['kind'], 'slide')
        self.assertTrue(source['text'].startswith('Document introduction\n\n'))
        self.assertEqual(source['attachment']['name'], 'deck.pptx')
        self.assertTrue(any(r.url.path.endswith('/medias/Embedded/download') for r in self.calls))

    def test_multiple_invalid_or_incomplete_attachments_never_download(self):
        for blocks in (
            [{'items': [{'file': {'token': 'One', 'name': 'a.pdf'}}, {'file': {'token': 'Two', 'name': 'b.pptx'}}], 'has_more': False}],
            [{'items': [{'file': {'token': '../escape', 'name': 'a.pdf'}}], 'has_more': False}],
            [{'items': [{'file': {'token': 'One', 'name': 'a.exe'}}], 'has_more': False}],
            [{'items': [], 'has_more': True}],
            [{'items': [], 'has_more': True, 'page_token': 'repeat'}, {'items': [], 'has_more': True, 'page_token': 'repeat'}],
        ):
            with self.subTest(blocks=blocks):
                self.blocks = blocks
                with self.assertRaises(LearningError):
                    read_material_link(self.portal, 'https://vnet.feishu.cn/docx/DocToken')
        self.assertFalse(any(r.url.path.endswith('/download') for r in self.calls))

    def test_missing_metadata_unrecognized_format_and_revision_change_fail_explicitly(self):
        for metadata in ({}, {'metas': [], 'failed_list': [{'code': 1}]},
                         {'metas': [{'doc_token': 'Other', 'doc_type': 'file', 'latest_modify_time': 1}]},
                         {'metas': [{'doc_token': 'FileToken', 'doc_type': 'file'}]}):
            self.meta_override = metadata
            with self.assertRaises(LearningError): read_material_link(self.portal, 'https://vnet.feishu.cn/file/FileToken')
        self.meta_override = None; self.name = 'no-extension'
        with self.assertRaises(LearningError): read_material_link(self.portal, 'https://vnet.feishu.cn/file/FileToken')
        self.assertFalse(any(r.url.path.endswith('/download') for r in self.calls))
        self.name = 'valid.pdf'
        def changed(request):
            self.revision = 2
            return httpx.Response(200, content=b'fixture-content')
        self.download = changed
        with self.assertRaises(LearningError) as error: read_material_link(self.portal, 'https://vnet.feishu.cn/file/FileToken')
        self.assertEqual(error.exception.status, 409)

    def test_redirect_permission_and_unknown_responses_do_not_parse_or_forward_credentials(self):
        for response in (httpx.Response(302, headers={'location': 'https://untrusted.invalid/stolen'}),
                         httpx.Response(403, json={'code': 99991672}),
                         httpx.Response(200, json={'code': 0, 'data': {}}), httpx.Response(200, content=b'')):
            self.download = lambda request, response=response: response
            with self.assertRaises(LearningError): read_material_link(self.portal, 'https://vnet.feishu.cn/file/FileToken')
        self.parser.assert_not_called()
        self.assertTrue(all(r.url.host == 'open.feishu.cn' for r in self.calls))

    def test_stream_size_limit_checks_declared_and_actual_bytes_and_closes(self):
        class Chunks(httpx.SyncByteStream):
            count, closed = 0, False
            def __iter__(self):
                for _ in range(4):
                    self.count += 1
                    yield b'x' * 65536
            def close(self): self.closed = True
        with patch('lan_bitable_template_portal.sharing_sources.MAX_FILE', 65537):
            for headers in ({'content-length': '1000000'}, {}):
                stream = Chunks()
                self.download = lambda request: httpx.Response(200, headers=headers, stream=stream)
                with self.assertRaises(LearningError) as error: read_material_link(self.portal, 'https://vnet.feishu.cn/file/FileToken')
                self.assertEqual(error.exception.status, 413)
                self.assertTrue(stream.closed)
                self.assertEqual(stream.count, 0 if headers else 2)
        self.parser.assert_not_called()

    def test_download_transport_errors_are_retryable_and_redacted(self):
        class BrokenStream(httpx.SyncByteStream):
            def __init__(self, error):
                self.error, self.closed = error, False
            def __iter__(self):
                yield b'partial' * 10000
                raise self.error
            def close(self): self.closed = True
        for error_type in (httpx.ConnectTimeout, httpx.ConnectError, httpx.ReadTimeout,
                           httpx.ReadError, httpx.RemoteProtocolError):
            for streaming in (False, True):
                with self.subTest(error=error_type.__name__, streaming=streaming):
                    streams = []
                    def failed(request):
                        error = error_type('https://private.invalid/file/PrivateToken?token=fixture-secret '
                                           'Authorization: Bearer fixture-token', request=request)
                        if not streaming:
                            raise error
                        streams.append(BrokenStream(error))
                        return httpx.Response(200, stream=streams[-1])
                    self.download = failed
                    with self.assertRaises(LearningError) as caught:
                        read_material_link(self.portal, 'https://vnet.feishu.cn/file/FileToken')
                    self.assertEqual(caught.exception.status, 503)
                    self.assertEqual(str(caught.exception), '分享材料读取时网络超时或连接中断，请稍后重试。')
                    for private in ('https://', 'PrivateToken', 'fixture-secret', 'fixture-token', 'Authorization'):
                        self.assertNotIn(private, str(caught.exception))
                    self.assertIsNone(caught.exception.__cause__)
                    self.assertTrue(caught.exception.__suppress_context__)
                    self.assertTrue(all(stream.closed for stream in streams))
        self.parser.assert_not_called()
        self.manager.get_tenant_token.assert_not_called()

    def test_non_transport_error_is_not_reclassified_as_network(self):
        def failed(request):
            raise RuntimeError('internal failure fixture-secret')
        self.download = failed
        with self.assertRaises(LearningError) as caught:
            read_material_link(self.portal, 'https://vnet.feishu.cn/file/FileToken')
        self.assertEqual(caught.exception.status, 502)
        self.assertNotIn('网络', str(caught.exception))
        self.assertNotIn('fixture-secret', str(caught.exception))
        self.parser.assert_not_called()

    def test_qt_refresh_once_and_application_change_discards_download(self):
        downloads = []
        def expired(request):
            downloads.append(request.headers['authorization'])
            return httpx.Response(401, json={'code': 99991663}) if len(downloads) == 1 else httpx.Response(200, content=b'fixture-content')
        self.download = expired
        read_material_link(self.portal, 'https://vnet.feishu.cn/file/FileToken')
        self.assertEqual(downloads, ['Bearer fixture-token', 'Bearer fixture-renewed'])
        self.manager.get_tenant_token.assert_called_once_with(force_refresh=True)
        self.parser.reset_mock()
        self.download = lambda request: httpx.Response(401, json={'code': 99991663})
        prior = len([r for r in self.calls if r.url.path.endswith('/download')])
        with self.assertRaises(LearningError): read_material_link(self.portal, 'https://vnet.feishu.cn/file/FileToken')
        self.assertEqual(len([r for r in self.calls if r.url.path.endswith('/download')]) - prior, 2)
        self.parser.assert_not_called()
        def changed(request):
            self.manager._credentials.return_value = ('other-app', 'other-secret')
            return httpx.Response(200, content=b'fixture-content')
        self.download = changed
        with self.assertRaises(LearningError): read_material_link(self.portal, 'https://vnet.feishu.cn/file/FileToken')
        self.parser.assert_not_called()

    def test_parser_failure_cleans_temporary_linked_file(self):
        def fail(kind, payload):
            self.paths.append(Path(payload['path']))
            raise LearningError('Unsupported file contents')
        self.parser.side_effect = fail
        with self.assertRaises(LearningError): read_material_link(self.portal, 'https://vnet.feishu.cn/file/FileToken')
        self.assertFalse(self.paths[-1].exists())

    def test_source_allowlist_and_minutes_do_not_accept_file_nodes(self):
        for url in ('https://evil.invalid/file/X', 'https://vnet.feishu.cn@evil.invalid/file/X',
                    'https://vnet.feishu.cn:443/file/X', 'https://vnet.feishu.cn/file/../X',
                    'https://open.feishu.cn/open-apis/drive/v1/files/X/download'):
            with self.assertRaises(LearningError): source_url(url)
        self.assertEqual(source_url('https://vnet.feishu.cn/file/X?ignored=1'), 'https://vnet.feishu.cn/file/X')
        with self.assertRaises(LearningError): source_url('https://vnet.feishu.cn/file/X', minutes=True)
        with self.assertRaises(LearningError): resolve_document(self.portal, 'https://vnet.feishu.cn/wiki/FileToken')
        self.assertFalse(any(r.url.path.endswith('/download') for r in self.calls))

    def test_linked_pdf_ocr_reuses_missing_page_handling(self):
        def ocr(kind, payload):
            if kind == 'markdown':
                return {'text': 'first', 'references': [{'kind': 'page', 'page': 1, 'start': 0, 'end': 5}],
                        'total_pages': 2, 'extraction_complete': False}
            self.assertEqual(payload['page'], 2)
            return {'text': 'scanned last page', 'low_confidence_count': 1}
        self.parser.side_effect = ocr
        source = read_material_link(self.portal, 'https://vnet.feishu.cn/file/FileToken')
        self.assertEqual(source['low_confidence_count'], 1)
        self.assertEqual([ref['page'] for ref in source['references']], [1, 2])
        self.assertIn('scanned last page', source['text'])

    def test_real_linked_pptx_uses_existing_extractor(self):
        try:
            from pptx import Presentation
        except ImportError:
            self.skipTest('Use the project runtime for Office extraction')
        from openclaw_service.assistant.lighthouse_document_tools import document_markdown
        slides = Presentation()
        for text in ('first slide', 'last slide'):
            slides.slides.add_slide(slides.slide_layouts[1]).shapes.title.text = text
        buffer = io.BytesIO(); slides.save(buffer)
        self.name = 'linked.pptx'
        self.download = lambda request: httpx.Response(200, content=buffer.getvalue())
        self.parser.side_effect = lambda kind, payload: document_markdown(Path(payload['path']).read_bytes(), '.pptx')
        source = read_material_link(self.portal, 'https://vnet.feishu.cn/file/FileToken')
        self.assertIn('last slide', source['text'])
        self.assertEqual({ref['page'] for ref in source['references']}, {1, 2})


if __name__ == '__main__':
    unittest.main()
