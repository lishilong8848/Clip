import copy
import datetime as dt
import io
import hashlib
import json
import os
import re
import sys
import tempfile
import threading
import time
import unittest
import uuid
from collections import Counter
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from xml.etree import ElementTree as ET
from .lan_bitable_template_portal.cabinet_power_data import source_rows, from_feishu, to_fields,source_evidence,complete_source_record
from .lan_bitable_template_portal.cabinet_power_excel import CabinetError, Workbook, T, baseline_correction_operations, calculate, completed_state_event, derive_records, export_workbook, dates, digest, map_state_baseline, room_code, system_name
from .lan_bitable_template_portal.cabinet_power import CabinetFeishu, CabinetPowerService, EXPORT_ARCHIVE_APP_TOKEN, EXPORT_ARCHIVE_FIELDS, EXPORT_ARCHIVE_TABLE_ID, EXPORT_FORMAT_VERSION, equivalent
from .lan_bitable_template_portal.cabinet_power_batches import CabinetBatchService, CabinetBatchStore, POWER_ACTIONS_BY_STATE
from .lan_bitable_template_portal.cabinet_power_evidence import _rows_from_ocr
TEMPLATES=Path(__file__).parent/"lan_bitable_template_portal/templates/cabinet_power"

class CabinetBatchRecognitionTests(unittest.TestCase):
    def test_export_timestamp_text_uses_iso_minutes_and_keeps_missing_values(self):
        from .lan_bitable_template_portal.cabinet_power_excel import export_time_text
        self.assertEqual(export_time_text('1、2026/9/22 9:01:45\n2、\n3、2026年9月23日 08：02：59'),
                         '1、2026-09-22 09:01\n2、\n3、2026-09-23 08:02')
        self.assertEqual(export_time_text('时间未知 / 2026-99-22 09:01'),'时间未知 / 2026-99-22 09:01')

    def test_pasted_confirmation_text_accepts_minimal_full_markdown_and_wrapped_rows(self):
        from .lan_bitable_template_portal.cabinet_power_text import parse_confirmation_text
        minimal=['EA118-E2-2','A11','测试电转正式电','2026-09-16 16:02:59','2026-09-14 16:03:16']
        full=['EA118','南通综保区基地A','E2-2.EA118',*minimal,'A11','Success']
        header='机房\t机房系统名称\t包间\t包间系统名称\t机架\t操作类型\t期望完成时间\t实际完成时间\t运营商机柜编号\t结果\n'
        for text in ('\t'.join(minimal),' '.join(minimal),'\n'.join(minimal),' | '.join(minimal),header+'\t'.join(full),'| '+' | '.join(full)+' |\n| --- | --- |'):
            with self.subTest(text=text):
                rows=parse_confirmation_text(text)
                self.assertEqual(len(rows),1)
                row=rows[0]
                self.assertEqual([row[key] for key in ('scope','room','rack','action','expected','actual')],['E','202',*minimal[1:]])
                self.assertEqual(row['result'],'成功' if 'Success' in text else '')
                self.assertEqual(row['supplier_rack'],'A11' if 'Success' in text else '')

    def test_pasted_text_preserves_missing_cells_and_normalizes_without_guessing_success(self):
        from .lan_bitable_template_portal.cabinet_power_text import parse_confirmation_text
        rows=parse_confirmation_text('EA118-B-216运营商机房\tB06\t上正式电\t2026/9/16 10:00\t2026/9/14 16：03：16\tB06\t功\n'
                                     'EA118-B-247运营商机房 B07 下测试电 2026/9/16 10:00 2026/9/14 09:00 B07 Failed')
        self.assertEqual([(row['scope'],row['room'],row['rack']) for row in rows],[('B','216','B06'),('B','247','B07')])
        self.assertEqual((rows[0]['expected'],rows[0]['actual'],rows[0]['result']),('2026-09-16 10:00:00','2026-09-14 16:03:16',''))
        self.assertEqual(rows[1]['result'],'失败')
        self.assertEqual(rows[1]['actual'],'2026-09-14 09:00:00')
        with self.assertRaises(CabinetError): parse_confirmation_text('只有邮件标题，没有机柜明细')
        with self.assertRaises(CabinetError): parse_confirmation_text('x'*200001)
        with self.assertRaisesRegex(CabinetError,'期望完成时间'):
            parse_confirmation_text('EA118-E2-2\tA11\t上正式电\t\t2026-09-14 16:03:16')
        with self.assertRaisesRegex(CabinetError,'包间系统名称'):
            parse_confirmation_text('A11 上正式电 2026-09-16 16:02:59 2026-09-14 16:03:16')
        with self.assertRaisesRegex(CabinetError,'机柜编号'):
            parse_confirmation_text('EA118-E2-2 测试电转正式电 2026-09-16 16:02:59 2026-09-14 16:03:16')
        with self.assertRaisesRegex(CabinetError,'第2条缺少机柜编号、操作类型'):
            parse_confirmation_text('EA118-E2-2 A11 上正式电 2026-09-16 16:02:59 2026-09-14 16:03:16\n'
                                    'EA118-E2-1 2026-09-16 16:02:59 2026-09-14 16:03:16')

    def test_ten_thousand_batch_summaries_paginate_without_reading_details(self):
        with tempfile.TemporaryDirectory() as root:
            store=CabinetBatchStore(Path(root)/'batches.sqlite3')
            records=[];views=[]
            for index in range(10000):
                bid=f'batch-{index:05d}'
                records.append((bid,'owner','pending',None,'["A"]','{"source":"image"}',1,'2026-09-20','2026-09-20'))
                summary={'batch_id':bid,'scopes':['A'],'rooms':['A楼 203'],'stats':{'total':1},'is_todo':True}
                views.append((bid,'*',1,0,json.dumps(summary)))
            with store._connect() as conn,conn:
                conn.executemany('INSERT INTO batches VALUES(?,?,?,?,?,?,?,?,?)',records)
                conn.executemany('INSERT INTO batch_list_views VALUES(?,?,?,?,?)',views)
            with patch.object(store,'_decode',side_effect=AssertionError('details must not be decoded')):
                page=store.list_page('owner',[],False,'A','todo','','',9999,20)
                self.assertEqual((page['total'],page['pending_count'],page['page'],len(page['items'])),(10000,10000,500,20))
                self.assertEqual(store.list_page("owner' OR 1=1",[],False,'A','todo','','',1,20)['total'],0)

    def test_xlsm_package_parts_use_default_namespace(self):
        from .lan_bitable_template_portal.cabinet_power_excel import xml_bytes
        for namespace, name in (("http://schemas.openxmlformats.org/package/2006/content-types", "Types"),
                                ("http://schemas.openxmlformats.org/package/2006/relationships", "Relationships")):
            root=ET.Element("{"+namespace+"}"+name)
            raw=xml_bytes(root)
            self.assertIn(("<"+name+" xmlns=\"").encode(),raw)
            self.assertEqual(ET.fromstring(raw).tag,root.tag)

    def test_empty_batch_delete_route_is_registered(self):
        source=(Path(__file__).parent/"lan_bitable_template_portal/cabinet_power_routes.py").read_text(encoding="utf-8")
        self.assertIn('"batches/{batch_id}":["GET","PATCH","DELETE"]',source)

    def test_split_ocr_cells_produce_a_complete_batch_row(self):
        lines=[
            ("包间",[("包",245,45),("间",262,45)]),("包间系统名",[("包",350,34)]),
            ("操作类型",[("操作类型",521,45)]),("期望完成时间",[("期望完成时间",649,45)]),
            ("实际完成时间",[("实际完成时间",808,45)]),("营商机柜编",[("营商机柜编",983,34)]),
            ("E2-1EAI1",[("E2",246,112),("一",265,120),("1",272,112),("EAI",285,112),("1",316,112)]),
            ("8",[("8",280,131)]),("EAI18-E2-",[("EAI",354,112),("1",386,112),("8",394,112),("一",403,120),("E2",409,112),("一",428,120)]),
            ("AI6",[("AI",464,122),("6",485,122)]),("正式电转丬则试",[("正式电转丬则试",522,108)]),("电",[("电",565,131)]),
            ("2026-04-0420℃",[("2026",656,112),("04",697,112),("04",721,112),("20",743,112),("℃",762,112)]),
            ("0.22",[("0",700,131),("22",713,131)]),("2026-04-0320℃",[("2026",815,112),("04",856,112),("03",880,112),("20",902,112),("℃",921,112)]),
            ("0.41",[("0",859,131),("41",872,131)]),("AI6",[("AI",1002,122),("6",1023,122)]),("SUCCES",[("SUCCES",1097,112)]),
        ]
        rows=_rows_from_ocr(lines,1176)
        self.assertEqual(rows,[{"scope":"E","room":"201","rack":"A16","supplier_rack":"A16",
            "action":"正式电转测试电","expected":"2026-04-04 20:00:22","actual":"2026-04-03 20:00:41",
            "result":"成功","raw":rows[0]["raw"]}])

    def test_ocr_accepts_misread_second_room_header_and_wrapped_minutes(self):
        lines=[
            ("包间",[("包",311,51),("间",334,51)]),("包问系统名",[("包",447,35),("问",470,35)]),
            ("操作类",[("操",671,35),("作",694,35),("类",716,35)]),("期望完成时间",[("期",784,51)]),
            ("实际完成时间",[("实",989,51)]),("运营商机柜",[("运",1195,35)]),("编号",[("编",1195,66)]),
            ("E2.2.EAI",[("E2",314,139),(".",340,149),("2",348,139),(".",361,153),("EAI",367,139)]),("18",[("18",349,166)]),
            ("EAI18.E2.",[("EAI",449,139),("18",491,139),(".",514,149),("E2",522,139),(".",548,149)]),("2",[("2",496,166)]),
            ("B17",[("B17",595,153)]),("上测试",[("上",676,136),("测",698,137),("试",720,136)]),("电",[("电",700,166)]),
            ("2026.07.1700:0",[("2026",787,139),("07",844,139),("17",877,139),("00",906,139),("0",937,139)]),("0:00",[("0",847,166),("00",865,166)]),
            ("2026.07.1815:5",[("2026",992,139),("07",1048,139),("18",1082,139),("15",1113,139),("5",1142,139)]),("1:55",[("1",1053,166),("55",1070,166)]),
            ("B17",[("B17",1239,153)]),("功",[("功",1369,167)]),
        ]
        rows=_rows_from_ocr(lines,1420)
        self.assertEqual([(row["scope"],row["room"],row["rack"],row["supplier_rack"],row["action"],row["expected"],row["actual"],row["result"]) for row in rows],
                         [("E","202","B17","B17","上测试电","2026-07-17 00:00:00","2026-07-18 15:51:55","")])

    def test_b_carrier_legacy_aliases_display_as_separate_named_rooms(self):
        for room, old_name in (("216", "EA118-B2-16"), ("247", "EA118-B2-47")):
            with self.subTest(room=room):
                expected = f"EA118-B-{room}运营商机房"
                self.assertEqual(system_name("B", room), expected)
                self.assertEqual(room_code(old_name, "B"), (room, False))
                self.assertEqual(room_code(expected, "B"), (room, False))
                op = from_feishu({"record_id": "recCarrier", "fields": {
                    "楼栋": "B楼", "包间系统名称": old_name, "机架": "B02",
                    "操作类型": "上测试电", "实际完成时间": 1741859439000,
                    "结果": "成功",
                }})
                self.assertEqual((op["room"], op["system_name"]), (room, expected))
                self.assertEqual(op["raw_fields"]["包间系统名称"], old_name)
        self.assertEqual(system_name("B", "201"), "EA118-B2-1")

    def test_pdf_rows_keep_missing_values_and_split_glued_result(self):
        text=(
            "机柜下测试电确认单\n申请时间： 2026-08-31 17:49:03\n操作类型： 下测试电\n"
            "EA118 B4-2.EA118 B16 B16 服务器机柜 WholeRack 2026-08-31 15:03:55\n"
            "EA118 A4-2.EA118 G02 G02 服务器机柜 WholeRack 2026-09-03 10:12:20成功 2026-08-31 14:48:26"
        )
        batch=object.__new__(CabinetBatchService)
        batch.import_root=Path("unused")
        batch._pdf_reader=lambda _path:SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda **_kwargs:text)])
        batch._set_file_progress=lambda *_args,**_kwargs:None
        rows=batch._parse_pdf("batch",{"file_id":"file","sha256":"hash","name":"机柜下测试电确认单.pdf"})
        self.assertEqual(len(rows),2)
        self.assertEqual((rows[0]["scope"],rows[0]["room"],rows[0]["actual"],rows[0]["result"]),("B","402","",""))
        self.assertEqual((rows[1]["actual"],rows[1]["result"]),("2026-09-03 10:12:20","成功"))
        self.assertTrue(all("order_time" not in row and "application_time" not in row for row in rows))

    def test_pdf_action_validation_uses_latest_successful_state(self):
        inventory=[{"room":"201","rack":rack,"rack_type":"服务器机柜"} for rack in ("A01","A02","A03")]
        events={"A02":[{"action":"上测试电","actual":"2026-09-14 10:00:00","result":"成功"}],
                "A03":[{"action":"上正式电","actual":"2026-09-14 10:00:00","result":"成功"}]}
        snapshot={"config":{"inventory":inventory},"operations":[{"room":"201","rack":rack,"events":items} for rack,items in events.items()]}
        batch=object.__new__(CabinetBatchService)
        batch.cabinet=SimpleNamespace(_snapshot=lambda _scope:snapshot)
        self.assertEqual(POWER_ACTIONS_BY_STATE["off"],{"上测试电","上正式电"})
        self.assertEqual(POWER_ACTIONS_BY_STATE["test"],{"测试电转正式电","下测试电"})
        self.assertEqual(POWER_ACTIONS_BY_STATE["formal"],{"正式电转测试电","下正式电"})
        rows=[]
        for index,(rack,action) in enumerate((("A01","上正式电"),("A02","下测试电"),("A03","正式电转测试电"),("A01","下正式电"),("A02","上正式电"),("A03","上测试电"))):
            actual="2026-09-15 11:00:00" if index<3 else "2026-09-15 11:01:00"
            rows.append({"row_id":f"row_{index}","operation_id":f"operation_{index}","scope":"A","room":"201","rack":rack,"rack_type":"服务器机柜","action":action,
                         "expected":actual,"actual":actual,"result":"成功","status":"ready"})
        payload={"source":"pdf","rows":rows}
        batch._validate_rows(payload)
        self.assertEqual([row["current_power_state"] for row in rows],["off","test","formal","off","test","formal"])
        self.assertEqual([row["status"] for row in rows],["ready"]*3+["invalid"]*3)
        with self.assertRaisesRegex(CabinetError,"当前为测试电"):
            batch._row_payload({"source":"pdf","batch_id":"test"},rows[4])

class MemoryStore:
    db_path=Path("unused.sqlite")
    def __init__(self): self.documents={}
    def get_document(self,ns,key): return copy.deepcopy(self.documents.get((ns,key)))
    def put_document(self,ns,key,value): self.documents[(ns,key)]=copy.deepcopy(value)

class FakeFeishu:
    def __init__(self,records=()):
        self.records={r["record_id"]:copy.deepcopy(r) for r in records}; self.creates=0; self.fail_after_create=False; self.fail_after_delete=False; self.list_calls=0
        self.batch_create_calls=0; self.batch_update_calls=0; self.fail_after_batch_create=False; self.fail_after_batch_update=False
        self.attachments={}
    def list_all(self,path="records",filters=None): self.list_calls+=1; return copy.deepcopy(list(self.records.values()))
    def ensure_fields(self): return True
    def get(self,rid): return copy.deepcopy(self.records[rid])
    def create(self,fields,operation_id):
        self.creates+=1; rid="recNew"+str(self.creates)
        self.records[rid]={"record_id":rid,"fields":copy.deepcopy(fields)}
        if self.fail_after_create: raise TimeoutError("response lost")
        return self.get(rid)
    def update(self,rid,fields):
        self.records[rid]["fields"].update(copy.deepcopy(fields)); return self.get(rid)
    def batch_create(self,rows):
        self.batch_create_calls+=1
        created=[self.create(fields,"batch") for fields in rows]
        if self.fail_after_batch_create: raise TimeoutError("batch response lost")
        return created
    def batch_update(self,rows):
        self.batch_update_calls+=1
        updated=[self.update(row["record_id"],row["fields"]) for row in rows]
        if self.fail_after_batch_update: raise TimeoutError("batch response lost")
        return updated
    def batch_get(self,record_ids):
        return [self.get(rid) for rid in record_ids if rid in self.records]
    def delete(self,rid):
        self.records.pop(rid)
        if self.fail_after_delete: raise TimeoutError("delete response lost")
    def upload_attachment(self,path,file_name):
        token="fileEvidence"+str(len(self.attachments)+1).zfill(12)
        self.attachments[token]=Path(path).read_bytes()
        return token
    def download_attachment(self,token): return self.attachments[token]

class FakeExportFeishu:
    def __init__(self):
        self.fields={"自动编号":{"field_name":"自动编号","type":1005},
                     "导出文件":{"field_name":"导出文件","field_id":"fldLegacyAttachment","type":17}}
        self.records={}; self.upload_calls=0; self.creates=0; self.fail_after_create=False; self.fail_share=False; self.lock=threading.RLock()
    def list_all(self,path="records",filters=None):
        with self.lock:
            items=list(self.fields.values()) if path=="fields" else list(self.records.values())
            if filters and path=="records":
                value=json.loads(filters.split("=",1)[1])
                items=[item for item in items if item.get("fields",{}).get("导出标识")==value]
            return copy.deepcopy(items)
    def request(self,method,path,body=None,params=None):
        if method=="POST" and path=="fields":
            with self.lock:
                self.fields[body["field_name"]]={"field_name":body["field_name"],"type":body["type"],"field_id":"fld"+str(len(self.fields)),"property":copy.deepcopy(body.get("property") or {})}
                return {"field":copy.deepcopy(self.fields[body["field_name"]])}
        if method=="PUT" and path.startswith("fields/"):
            with self.lock:
                old=next(name for name,field in self.fields.items() if field.get("field_id")==path.split("/")[1])
                field=self.fields.pop(old)
                field.update(field_name=body["field_name"],type=body["type"],property=copy.deepcopy(body.get("property") or {}))
                self.fields[field["field_name"]]=field
                if old!=field["field_name"]:
                    for record in self.records.values():
                        if old in record["fields"]: record["fields"][field["field_name"]]=record["fields"].pop(old)
                return {"field":copy.deepcopy(field)}
        raise AssertionError((method,path,body,params))
    def upload_attachment(self,path,file_name):
        with self.lock: self.upload_calls+=1
        self.uploaded_path=Path(path); self.uploaded_name=file_name
        return "file-export-token-"+str(self.upload_calls)
    def create(self,fields,operation_id):
        with self.lock:
            existing=next((record for record in self.records.values() if record.get("operation_id")==operation_id),None)
            if existing: return copy.deepcopy(existing)
            self.creates+=1; rid="recExport"+str(self.creates)
            self.records[rid]={"record_id":rid,"fields":copy.deepcopy(fields),"operation_id":operation_id}
        if self.fail_after_create: raise TimeoutError("response lost")
        return self.get(rid)
    def update(self,rid,fields):
        with self.lock:
            self.records[rid]["fields"].update(copy.deepcopy(fields)); return self.get(rid)
    def get(self,rid):
        with self.lock: return copy.deepcopy(self.records[rid])
    def record_share_link(self,rid):
        if self.fail_share: raise TimeoutError("share link unavailable")
        return "https://vnet.feishu.cn/record/"+rid

def fixtures(with_power_baseline=False):
    models={}; records=[]; directory=[]; configs={}
    for scope in "ABCDE":
        content=(TEMPLATES/(scope+".xlsm")).read_bytes(); model,rows=source_rows(content,scope); models[scope]=model
        records.extend({"record_id":"rec"+scope+str(i),"fields":r["fields"]} for i,r in enumerate(rows))
        if with_power_baseline and scope=='C':
            evidence=source_evidence(content,model)
            records=[(complete_source_record(r,evidence) or r) if r['fields']['楼栋']=='C楼' else r for r in records]
        book=Workbook(content)
        data={"hash":model["template_hash"],"formats":model["formats"],"summary_cells":{n:{ref:book.value(c) for ref,c in book.cells(n).items()} for n in book.sheets if "汇总" in n}}
        configs[scope]={**model,"template_data":data}
        map_values={n:{ref:book.value(c) for ref,c in book.cells(n).items()} for n in book.sheets if "平面图" in n}
        configs[scope]["map_values"]=map_values
        directory.append({"record_id":"mapvals"+scope,"fields":{"类别":"平面图数值","楼栋":scope+"楼","布局资料":json.dumps(map_values,ensure_ascii=False)}})
        directory.append({"record_id":"template"+scope,"fields":{"类别":"模板资料","楼栋":scope+"楼","布局资料":json.dumps(data,ensure_ascii=False)}})
        for room in model["rooms"]:
            geometry={k:v for k,v in room.items() if k!="layout"}
            if with_power_baseline:
                if 'power_baseline' not in configs[scope]: configs[scope]['power_baseline']=map_state_baseline(content,configs[scope],[from_feishu(r) for r in records if r['fields']['楼栋']==scope+'楼'])[0]
                geometry['power_baseline']={'version':1,'template_hash':model['template_hash'],'racks':{k.split('/')[1]:v for k,v in configs[scope]['power_baseline'].items() if k.startswith(room['id']+'/')}}
            directory.append({"record_id":"room"+scope+room["id"],"fields":{"类别":"房间","楼栋":scope+"楼","包间":room["id"],"名称":room["name"],"数量":room["total"],"布局资料":json.dumps(geometry,ensure_ascii=False)}})
        for rack in model["inventory"]:
            directory.append({"record_id":"rack"+scope+rack["room"]+rack["rack"],"fields":{"类别":"机柜","楼栋":scope+"楼","包间":rack["room"],"名称":rack["rack"],"机柜类型":rack["rack_type"],"布局资料":json.dumps({"positions":rack["positions"],"template_color":rack["template_color"]},ensure_ascii=False)}})
    return models,records,directory,configs

class CabinetPowerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.models,cls.source_records,cls.directory_records,cls.configs=fixtures()

    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.store=MemoryStore()
        self.store.db_path=Path(self.tmp.name)/"state.sqlite"
        self.remote=FakeFeishu(self.source_records); self.service=CabinetPowerService(self.store,self.remote,self.tmp.name)
        self.service._directory=FakeFeishu(self.directory_records)
        self.service.do_refresh("",{}, {})

    def tearDown(self):
        self.service.shutdown(); self.tmp.cleanup()

    def test_rack_power_corrects_only_selected_cabinet_and_exports_all_its_records(self):
        def profile(): return next(r for r in self.service.racks('B')['items'] if (r['room'],r['rack'])==('302','B04'))
        rows=[op for op in self.service._snapshot('B')['operations'] if (op['room'],op['rack'])==('302','B04') and not op.get('meta',{}).get('baseline_correction')]
        duplicate=copy.deepcopy(self.remote.records[rows[0]['record_id']])
        duplicate['record_id']='recSecondPower'; duplicate['fields'].update({'来源行号':None,'数据标识':'power_duplicate_test'})
        self.remote.records[duplicate['record_id']]=duplicate
        self.service.do_refresh('B',{}, {})
        before=copy.deepcopy(self.remote.records)
        rows=[op for op in self.service._snapshot('B')['operations'] if (op['room'],op['rack'])==('302','B04') and not op.get('meta',{}).get('baseline_correction')]
        self.assertGreater(len(rows),1)
        expected_ids={op['record_id'] for op in rows}
        payload={'room':'302','rack':'B04','power':4000,'expected_version':profile()['power_version'],'operation_id':uuid.uuid4().hex}
        baseline=copy.deepcopy(self.service._snapshot('B')['config']['power_baseline'])
        self.service.save_rack_power('B',payload,'owner')
        self.service.save_rack_power('B',payload,'owner')
        self.assertEqual(profile()['power'],4000)
        for rid,record in self.remote.records.items():
            expected=copy.deepcopy(before[rid])
            if rid in expected_ids: expected['fields']['机柜功率（W）']=4000
            self.assertEqual(record,expected,rid)
        snap=self.service._snapshot('B')
        self.assertEqual(snap['config']['power_baseline'],baseline)
        self.assertTrue(all(op['power']==4000 for op in snap['operations'] if op['record_id'] in expected_ids))
        book=Workbook(export_workbook((TEMPLATES/'B.xlsm').read_bytes(),snap['config'],snap['operations']))
        self.assertIn('机柜上电汇总表',book.sheets)
        self.assertFalse(any('（邮件）' in name for name in book.sheets))
        for fmt in snap['config']['template_data']['formats']:
            for _,row in book.rows(fmt['sheet']):
                if row.get(fmt['room'])=='EA118-B3-2' and row.get(fmt['rack'])=='B04':
                    self.assertEqual(row.get(fmt['power']),4000)
        styles=ET.fromstring(book.archive.read('xl/styles.xml'))
        xfs=styles.find(T('cellXfs'))
        from openpyxl.styles.numbers import BUILTIN_FORMATS,is_date_format
        from .lan_bitable_template_portal.cabinet_power_excel import coord
        codes={**BUILTIN_FORMATS,**{int(node.get('numFmtId')):node.get('formatCode') for node in styles.find(T('numFmts'))}}
        for fmt in snap['config']['template_data']['formats']:
            for ref,cell in book.cells(fmt['sheet']).items():
                col,rn=coord(ref)
                if col==fmt['power'] and rn>fmt['header'] and book.value(cell) not in (None,''):
                    self.assertFalse(is_date_format(codes.get(int(xfs[int(cell.get('s','0'))].get('numFmtId','0')),'')),ref)
        self.service.do_refresh('B',{}, {})
        self.assertEqual(profile()['power'],4000)
        with self.assertRaisesRegex(CabinetError,'已变化'):
            self.service.save_rack_power('B',{**payload,'operation_id':uuid.uuid4().hex},'owner')

    def test_rack_power_without_history_zero_clear_and_new_operation_inheritance(self):
        profile=lambda: next(r for r in self.service.racks('A')['items'] if (r['room'],r['rack'])==('203','A02'))
        self.assertIsNone(profile()['power'])
        count=len(self.remote.records)
        for power in (4000,0,''):
            self.service.save_rack_power('A',{'room':'203','rack':'A02','power':power,'expected_version':profile()['power_version'],'operation_id':uuid.uuid4().hex},'owner')
            self.assertEqual(profile()['power'],None if power=='' else power)
            self.assertEqual(len(self.remote.records),count)
            self.service.do_refresh('A',{}, {})
            self.assertEqual(profile()['power'],None if power=='' else power)
        for value in (-1,True,{},'nan','inf'):
            with self.assertRaises(CabinetError): self.service.save_rack_power('A',{'power':value,'operation_id':uuid.uuid4().hex},'owner')
        self.service.save_rack_power('A',{'room':'203','rack':'A02','power':5000,'expected_version':profile()['power_version'],'operation_id':uuid.uuid4().hex},'owner')
        saved=self.service.save_operation('A',{'room':'203','rack':'A02','rack_type':'网络机柜','groups':[{'action':'上正式电','actual':'2026-09-01 10:00:00','expected':'','result':'成功'}],'operation_id':uuid.uuid4().hex},'owner')
        self.assertEqual(saved['power'],5000)

    def test_rack_power_lost_response_local_failure_and_conflict_resume(self):
        profile=lambda: next(r for r in self.service.racks('B')['items'] if (r['room'],r['rack'])==('302','B04'))
        original=profile()['power']; oid=uuid.uuid4().hex
        payload={'room':'302','rack':'B04','power':4000,'expected_version':profile()['power_version'],'operation_id':oid}
        update=self.remote.update
        def lost(rid,fields): update(rid,fields); raise TimeoutError('response lost')
        with patch.object(self.remote,'update',side_effect=lost), self.assertRaises(TimeoutError):
            self.service.save_rack_power('B',payload,'owner')
        self.assertEqual(profile()['power'],original)
        changed=self.service.local.document('B','write:'+oid)['stages'][0]['record_id']
        self.remote.update(changed,{'机架':'B05'})
        with self.assertRaisesRegex(CabinetError,'位置发生冲突'): self.service.resume_write('B',oid,'owner')
        self.remote.update(changed,{'机架':'B04'})
        with self.assertRaises(CabinetError): self.service.resume_write('B',oid,'another')
        with patch.object(self.service.local,'commit_operation',side_effect=OSError('disk unavailable')), self.assertRaises(OSError):
            self.service.resume_write('B',oid,'owner')
        self.assertEqual(profile()['power'],original)
        self.service._cache.clear()
        self.service.resume_write('B',oid,'owner')
        self.assertEqual(profile()['power'],4000)
        payload.update(power=4500,expected_version=profile()['power_version'],operation_id=uuid.uuid4().hex)
        with patch.object(self.remote,'update',side_effect=lost), self.assertRaises(TimeoutError):
            self.service.save_rack_power('B',payload,'owner')
        journal=self.service.local.document('B','write:'+payload['operation_id'])
        changed=journal['stages'][0]['record_id']
        self.remote.update(changed,{'机柜功率（W）':4700})
        with self.assertRaises(CabinetError): self.service.resume_write('B',payload['operation_id'],'owner')
        self.service.reconcile_write('B',payload['operation_id'],'owner')
        self.assertFalse(self.service.pending_writes('B'))
        self.assertEqual(next(op for op in self.service._snapshot('B')['operations'] if op['record_id']==changed)['power'],4700)

    def test_record_power_editor_also_updates_other_records_for_that_rack(self):
        rows=[op for op in self.service._snapshot('B')['operations'] if (op['room'],op['rack'])==('302','B04')]
        old=next(op for op in rows if not op.get('meta',{}).get('baseline_correction'))
        self.service.save_operation('B',{'power':4000,'operation_id':uuid.uuid4().hex,'expected_version':old['version']},'owner',old['record_id'])
        for op in self.service._snapshot('B')['operations']:
            if op['record_id'] in {row['record_id'] for row in rows if not row.get('meta',{}).get('baseline_correction')}:
                self.assertEqual(op['power'],4000)

    def test_create_uses_stable_uuid4_client_token_helper(self):
        self.assertIn("_stable_uuid4_client_token",CabinetFeishu.create.__code__.co_names)
        remote=CabinetFeishu(); captured={}
        def request(method,path,body=None,params=None): captured.update(method=method,path=path,body=body,params=params); return {"record":{"record_id":"recToken","fields":body["fields"]}}
        remote.request=request
        remote.create({"机架":"A01"},"manual_operation_123456")
        token=captured["params"]["client_token"]
        self.assertEqual((captured["method"],captured["path"],uuid.UUID(token).version),("POST","records",4))

    def test_attachment_readback_checks_file_tokens(self):
        expected={"上下电确认截图":[{"file_token":"fileEvidence12345"}]}
        self.assertFalse(equivalent(expected,{"上下电确认截图":[]}))
        self.assertTrue(equivalent(expected,{"上下电确认截图":[{"file_token":"fileEvidence12345","name":"proof.png"}]}))

    def test_b_carrier_rooms_stay_separate_from_standard_rooms(self):
        rooms = {room["id"]: room for room in self.service.overview("B")["rooms"]}
        for room in ("216", "247"):
            self.assertEqual(rooms[room]["name"], f"EA118-B-{room}运营商机房")
            self.assertTrue(rooms[room]["carrier"])
            self.assertEqual(rooms[room]["total"], 10)
        self.assertEqual(rooms["201"]["name"], "EA118-B2-1")
        self.assertFalse(rooms["201"]["carrier"])

    def test_text_preview_is_read_only_and_creation_is_idempotent_with_audit(self):
        service=self.service.batches
        text='EA118-A2-3\tA02\t上正式电\t2026-09-16 16:02:59\t2026-09-14 16:03:16\tA02\tSuccess'
        sources=[{'id':'text_sample_01','text':text}]
        before_records=copy.deepcopy(self.remote.records)
        preview=service.text_preview(sources,['A'])
        self.assertEqual(len(preview['rows']),1)
        self.assertEqual(service.store.list(None),[])
        row=preview['rows'][0]
        request={'request_id':'text_create_request_0001','sources':sources,'rows':[{'text_id':row['text_id'],'text_row':row['text_row'],'expected':'2026-09-17 16:02:59'}]}
        created=service.create_text(request,'owner',['A'])
        self.assertEqual(service.create_text(request,'owner',['A'])['batch_id'],created['batch_id'])
        self.assertEqual(len(service.store.list(None)),1)
        self.assertEqual(created['source'],'text')
        self.assertEqual(created['rows'][0]['expected'],'2026-09-17 16:02:59')
        self.assertEqual(created['rows'][0]['original']['expected'],'2026-09-16 16:02:59')
        self.assertTrue(created['rows'][0]['edits'])
        self.assertEqual(self.remote.records,before_records)
        payload,_=service._row_payload(created,created['rows'][0])
        self.assertEqual(payload[0]['batch_meta']['text_source']['raw'],text)
        with self.assertRaisesRegex(CabinetError,'请求内容已改变'):
            service.create_text({**request,'rows':[{'text_id':row['text_id'],'text_row':row['text_row'],'result':'失败'}]},'owner',['A'])
        with self.assertRaises(CabinetError): service.text_preview(sources,['B'])
        with self.assertRaises(CabinetError): service.create_text({**request,'request_id':'text_bad_ref_00000001','rows':[{'text_id':'missing1','text_row':1}]},'owner',['A'])

    def test_text_combines_pastes_and_keeps_missing_results_for_review(self):
        service=self.service.batches
        text='EA118-E2-2 A11 测试电转正式电 2026-09-16 16:02:59 2026-09-14 16:03:16'
        sources=[{'id':'text_fragment_01','text':text},{'id':'text_fragment_02','text':text+' A11 Success'}]
        preview=service.text_preview(sources,['E'])
        self.assertEqual(len(preview['rows']),2)
        self.assertEqual([row['result'] for row in preview['rows']],['','成功'])
        self.assertTrue(any(issue['code']=='result' for issue in preview['rows'][0]['issues']))
        self.assertEqual(preview['rows'][1]['status'],'duplicate')
        self.assertEqual(service.store.list(None),[])

    def test_batch_text_fill_only_changes_selected_existing_rows_with_audit(self):
        service=self.service.batches
        original={'scope':'E','room':'202','rack':'A11','action':'上正式电','expected':'','actual':'',
                  'result':'失败','failure_reason':'保留原因','rack_type':'服务器机柜'}
        sources=[{'id':'fill_text_sample','text':'EA118-E2-2 A11 测试电转正式电 2026-09-16 16:02:59 2026-09-14 16:03:16 A11 Success\nEA118-E2-2 A12 上正式电 2026-09-16 16:02:59 2026-09-14 16:03:16'}]
        remote=copy.deepcopy(self.remote.records)
        for source in ('manual','notice','image','pdf','text'):
            with self.subTest(source=source):
                batch=service.create_manual([original,original],'owner')
                batch=service._change(batch['batch_id'],lambda batch:batch.update(source=source))
                before=copy.deepcopy(service.get(batch['batch_id']))
                preview=service.preview_text_fill(batch['batch_id'],{'sources':sources},'owner',['E'])
                self.assertEqual(len(preview['rows'][0]['targets']),2)
                self.assertEqual(preview['rows'][1]['issue'],'本批次无此机柜')
                self.assertEqual(service.store.get(batch['batch_id']),before)
                row_id=batch['rows'][1]['row_id']
                payload={'version':preview['version'],'sources':sources,'rows':[{'text_id':sources[0]['id'],'text_row':1,'row_id':row_id}]}
                saved=service.apply_text_fill(batch['batch_id'],payload,'owner',['E'])
                row=saved['rows'][1]
                self.assertEqual(len(saved['rows']),2)
                self.assertEqual((row['action'],row['expected'],row['actual']),('测试电转正式电','2026-09-16 16:02:59','2026-09-14 16:03:16'))
                for key in ('scope','room','rack','result','failure_reason','original','operation_id'):
                    self.assertEqual(row[key],before['rows'][1][key],key)
                self.assertEqual(saved['rows'][0]['expected'],'')
                self.assertEqual([edit['field'] for edit in row['edits']],['action','expected','actual'])
                self.assertTrue(all(edit['source']=='text_fill' and edit['raw_text'] for edit in row['edits']))
                with self.assertRaises(CabinetError): service.apply_text_fill(batch['batch_id'],payload,'owner',['E'])
                payload['version']=saved['version']
                retried=service.apply_text_fill(batch['batch_id'],payload,'owner',['E'])
                self.assertEqual(retried['rows'][1]['edits'],row['edits'])
        self.assertEqual(self.remote.records,remote)

    def test_batch_text_fill_rejects_mismatches_conflicts_permissions_and_locked_rows(self):
        service=self.service.batches
        batch=service.create_manual([{'scope':'E','room':'202','rack':'A11','result':'成功'},
                                     {'scope':'E','room':'202','rack':'A12','result':'成功'}],'owner')
        batch_id=batch['batch_id']; row_id=batch['rows'][0]['row_id']
        sources=[{'id':'fill_guard_source','text':'EA118-E2-2 A11 上测试电 2026-09-16 16:02:59 2026-09-14 16:03:16'}]
        payload={'version':batch['version'],'sources':sources,'rows':[{'text_id':sources[0]['id'],'text_row':1,'row_id':row_id}]}
        for selections in ([{**payload['rows'][0],'row_id':batch['rows'][1]['row_id']}],payload['rows']*2,[{**payload['rows'][0],'row_id':'missing'}]):
            with self.assertRaises(CabinetError): service.apply_text_fill(batch_id,{**payload,'rows':selections},'owner',['E'])
            self.assertEqual(service.store.get(batch_id),batch)
        with self.assertRaises(CabinetError): service.apply_text_fill(batch_id,payload,'owner',['A'])
        with self.assertRaises(CabinetError): service.preview_text_fill(batch_id,{'sources':sources},'stranger',['A'])
        with self.assertRaisesRegex(CabinetError,'缺少'): service.preview_text_fill(batch_id,{'sources':[{'id':'fill_incomplete','text':'EA118-E2-2 A11 上测试电'}]},'owner',['E'])
        for status in ('queued','writing','completed','excluded_manual','excluded_cancelled','rollback_failed'):
            batch=service._change(batch_id,lambda batch:batch['rows'][0].update(status=status))
            payload['version']=batch['version']
            preview=service.preview_text_fill(batch_id,{'sources':sources},'owner',['E'])
            self.assertFalse(preview['rows'][0]['targets'][0]['editable'])
            with self.assertRaises(CabinetError): service.apply_text_fill(batch_id,payload,'owner',['E'])
        batch=service._change(batch_id,lambda batch:batch['rows'][0].update(status='rolled_back',operation_started=True))
        payload['version']=batch['version']
        saved=service.apply_text_fill(batch_id,payload,'owner',['E'])
        self.assertEqual(saved['rows'][0]['actual'],'2026-09-14 16:03:16')
        batch=service._change(batch_id,lambda batch:batch.update(source_notice={'deleted_at':'2026-09-23'}))
        payload['version']=batch['version']
        with self.assertRaises(CabinetError): service.apply_text_fill(batch_id,payload,'owner',['E'])

    def test_text_hundred_pastes_create_five_hundred_rows(self):
        text='EA118-E2-2 A11 测试电转正式电 2026-09-16 16:02:59 2026-09-14 16:03:16 A11 Success'
        sources=[{'id':f'text_large_{index:04d}','text':'\n'.join([text]*5)} for index in range(100)]
        service=self.service.batches
        preview=service.text_preview(sources,['E'])
        self.assertEqual(len(preview['rows']),500)
        self.assertEqual(service.store.list(None),[])
        rows=[{key:row[key] for key in ('text_id','text_row')} for row in preview['rows']]
        batch=service.create_text({'request_id':'text_large_request_0001','sources':sources,'rows':rows},'owner',['E'])
        self.assertEqual((len(batch['text_sources']),len(batch['rows'])),(100,500))

    def test_b_carrier_drawings_preserve_opposite_columns_and_exclude_copied_helpers(self):
        book=Workbook((TEMPLATES/'B.xlsm').read_bytes())
        for room,left,right in (('216','B05','B10'),('247','B10','B05')):
            name=f'B-{room}-机柜平面图'
            cells=book.cells(name)
            self.assertEqual((book.value(cells['G4']),book.value(cells['K4'])),(left,right))
            self.assertNotIn('V1',cells)
            self.assertEqual([book.value(cells[f'E{row}']) for row in range(22,27)],[10,5,5,1,4])
            model=self.service.layout('B',room)
            racks=[cell for cell in model['layout']['cells'] if cell.get('rack')]
            self.assertEqual({cell['rack'] for cell in racks},{f'B{index:02}' for index in range(1,11)})
            self.assertEqual(len(racks),10)
            self.assertEqual(model['room']['region'],'A1:S26')

    def test_b_carrier_geometry_upgrade_keeps_baseline_history_and_live_summary(self):
        from .lan_bitable_template_portal.cabinet_power import layout_identity
        from .lan_bitable_template_portal.cabinet_power_excel import inventory_state_baseline
        extension=self.service._layout_data('B')['extension']
        legacy=copy.deepcopy(self.configs['B'])
        replacements={room['id']:room for room in extension['previous_rooms']}
        for room in legacy['rooms']:
            if room['id'] in replacements: room.update(replacements[room['id']])
        new_keys={(rack['room'],rack['rack']) for rack in extension['inventory']}
        legacy['inventory']=[rack for rack in legacy['inventory'] if (rack['room'],rack['rack']) not in new_keys]
        for rack in legacy['inventory']:
            if rack['room'] in replacements: rack['positions']=[]
        legacy['template_data']['hash']='legacy-carrier-layout'
        self.assertEqual(layout_identity(legacy),extension['from_identity'])
        records=[record for record in self.source_records if record['fields']['楼栋']=='B楼']
        baseline=inventory_state_baseline(legacy,[from_feishu(record) for record in records])[0]
        self.service.local.replace('B',legacy,records,[])
        self.service.local.document('B','power_baseline:frozen_v1',{'version':1,'layout_identity':layout_identity(legacy),'racks':baseline})
        snapshot=self.service._snapshot('B')
        self.assertEqual({key:value for key,value in snapshot['config']['power_baseline'].items() if key in baseline},baseline)
        self.assertEqual(len(snapshot['config']['inventory']),len(legacy['inventory'])+10)
        self.assertEqual({r['record_id'] for r in self.service.local.load('B')['records']},{r['record_id'] for r in records})
        self.service._directory.records.pop('rackB216B06',None)
        saved=self.service.save_operation('B',{'operation_id':'carrier_room_new_rack_power','room':'216','rack':'B06',
            'rack_type':'服务器机柜','result':'成功','groups':[{'id':'carrier_new_event','action':'上正式电',
            'expected':'2026-09-22 09:00:00','actual':'2026-09-22 09:01:00','result':'成功'}]},'owner')
        fresh=CabinetPowerService(self.store,self.remote,self.tmp.name)
        try:
            current=fresh._snapshot('B')
            config=current['config']
            self.assertEqual({key:value for key,value in config['power_baseline'].items() if key in baseline},baseline)
            self.assertEqual(next(r for r in fresh.overview('B')['racks'] if (r['room'],r['rack'])==('216','B06'))['state'],'formal')
            upgraded=next(r for r in config['inventory'] if (r['room'],r['rack'])==('216','B06'))
            self.assertTrue(upgraded['record_id'])
            self.assertEqual(upgraded['rack_type'],'服务器机柜')
            exported=Workbook(export_workbook((TEMPLATES/'B.xlsm').read_bytes(),config,current['operations']))
            for room,expected in (('216',[10,6,4,1,5]),('247',[10,5,5,1,4])):
                name=f'B-{room}-机柜平面图'; cells=exported.cells(name)
                self.assertEqual([exported.value(cells[f'E{row}']) for row in range(22,27)],expected)
                projected={cell['ref']:cell for cell in fresh.layout('B',room)['layout']['cells']}
                self.assertEqual([int(projected[f'E{row}']['text']) for row in range(22,27)],expected)
                self.assertEqual((exported.value(cells['G19']),exported.value(cells['G20'])),(expected[2],expected[1]))
                self.assertEqual((int(projected['G19']['text']),int(projected['G20']['text'])),(expected[2],expected[1]))
            self.assertEqual(exported.archive.read('xl/vbaProject.bin'),Workbook((TEMPLATES/'B.xlsm').read_bytes()).archive.read('xl/vbaProject.bin'))
            self.assertIn(saved['record_id'],{r['record_id'] for r in fresh.local.load('B')['records']})
        finally: fresh.shutdown()

    def test_current_state_uses_each_floorplan_baseline_then_new_operations(self):
        for scope in "ABCDE":
            content=(TEMPLATES/(scope+".xlsm")).read_bytes()
            operations=[from_feishu(record) for record in self.source_records if record["fields"]["楼栋"]==scope+"楼"]
            config=copy.deepcopy(self.configs[scope])
            config["power_baseline"]=map_state_baseline(content,config,operations)[0]
            expected=Counter(value["state"] for value in config["power_baseline"].values())
            actual=Counter(rack["state"] for rack in derive_records(config,operations)["racks"])
            self.assertEqual((actual["formal"],actual["test"],actual["off"]),
                             (expected["formal"],expected["test"],expected["off"]),scope)
        content=(TEMPLATES/"C.xlsm").read_bytes(); model,rows=source_rows(content,"C")
        evidence=source_evidence(content,model); records=[]
        for index,row in enumerate(rows):
            record={**row,"record_id":f"recCBaseline{index}"}
            records.append(complete_source_record(record,evidence) or record)
        operations=[from_feishu(record) for record in records]
        config={**model,"template_data":{"hash":model["template_hash"],"formats":model["formats"]}}
        config["power_baseline"]=map_state_baseline(content,config,operations)[0]
        corrections=baseline_correction_operations(config,operations)
        self.assertEqual((len(corrections),Counter(item["action"] for item in corrections)),
                         (25,Counter({"测试电转正式电":1,"上测试电":12,"下测试电":12})))
        self.service.local.replace("C",config,records,[]); self.service._cache.pop("C",None)
        snapshot=self.service._snapshot("C")
        self.assertEqual(sum(bool(item.get("meta",{}).get("baseline_correction")) for item in snapshot["operations"]),25)
        first=corrections[0]; visible=self.service.operations("C",{"room":first["room"],"rack":first["rack"],"page_size":100})
        self.assertTrue(any(item.get("meta",{}).get("baseline_correction") for item in visible["items"]))
        self.assertEqual((visible["rack_state"]["latest_success"]["action"],visible["rack_state"]["latest_success"]["actual"],
                          visible["rack_state"]["latest_success"]["baseline_correction"]),(first["action"],"",True))
        self.assertTrue(self.service.operations("C",{"sheet":first["source"],"page_size":100})["items"])
        exported=Workbook(export_workbook(content,snapshot["config"],snapshot["operations"]))
        fmt=next(item for item in model["formats"] if item["sheet"]==first["source"])
        action_column=fmt["groups"][-1 if first["category"]=="down" else 0]["action"]
        self.assertTrue(any(row.get(fmt["rack"])==first["rack"] and row.get(action_column)==first["action"]
                            for _number,row in exported.rows(first["source"])))
        self.assertEqual(derive_records(config,operations)["counts"],{
            "total":998,"formal":939,"test":31,"off":28,"unknown":0,"powered":970})
        no_time=next(op for op in operations if any(event["result"]=="成功" and not event["actual"] for event in op["events"]))
        self.assertFalse(any("未配对" in issue for issue in no_time["issues"]))
        room,rack=next(key.split("/") for key,value in config["power_baseline"].items() if value["state"]=="formal")
        operations.append({"record_id":"recCNewConversion","scope":"C","room":room,"rack":rack,
            "issues":[],"groups":[{}],"events":[{"id":"new_conversion:0","action":"正式电转测试电",
            "actual":"2026-09-21 10:00:00","expected":"","result":"成功","group":0}]})
        changed=derive_records(config,operations)["counts"]
        self.assertEqual((changed["formal"],changed["test"]),(938,32))

    def test_inconsistent_cloud_baseline_is_rebuilt_from_template(self):
        scope="A"; content=(TEMPLATES/"A.xlsm").read_bytes(); config=copy.deepcopy(self.configs[scope])
        records=[copy.deepcopy(record) for record in self.source_records if record["fields"]["楼栋"]=="A楼"]
        operations=[from_feishu(record) for record in records]
        config["power_baseline"]=map_state_baseline(content,config,operations)[0]
        key=next(key for key,value in config["power_baseline"].items() if value["state"]=="formal")
        config["power_baseline"][key].update(state="off",color="#00B050")
        self.service.local.replace(scope,config,records,[]); self.service._cache.pop(scope,None)
        room,rack=key.split("/"); state=self.service.operations(scope,{"room":room,"rack":rack,"page_size":100})["rack_state"]
        current=self.service.config(scope)["power_baseline"][key]
        self.assertEqual((current["color"],current["state"],state["state"],state["latest_success"]["action"]),
                         ("#FF0000","formal","formal","上正式电"))

    def test_floorplan_baseline_is_frozen_across_later_writes_and_restart(self):
        scope="A"; before=self.service._snapshot(scope)
        frozen=copy.deepcopy(self.service.local.document(scope,"power_baseline:frozen_v1"))
        rack=next(item for item in derive_records(before["config"],before["operations"])["racks"] if item["state"]=="off")
        saved=self.service.save_operation(scope,{"operation_id":"baseline_restart_write_01",
            "room":rack["room"],"rack":rack["rack"],"rack_type":rack["rack_type"],"result":"成功",
            "groups":[{"id":"later_event","action":"上正式电","expected":"2026-09-20 09:00:00",
                       "actual":"2026-09-20 09:01:00","result":"成功"}]},"owner")
        fresh=CabinetPowerService(self.store,self.remote,self.tmp.name); fresh._directory=self.service._directory
        try:
            current=next(item for item in fresh.overview(scope)["racks"]
                         if (item["room"],item["rack"])==(rack["room"],rack["rack"]))
            self.assertEqual(current["state"],"formal")
            self.assertEqual(fresh.local.document(scope,"power_baseline:frozen_v1"),frozen)
            self.assertNotIn(saved["events"][0]["id"],str(frozen))
        finally:
            fresh.shutdown()

    def test_export_attachment_uses_archive_base_parent(self):
        path=Path(self.tmp.name)/"sample.xlsm"; path.write_bytes(b"export")
        remote=CabinetFeishu(EXPORT_ARCHIVE_TABLE_ID,EXPORT_ARCHIVE_APP_TOKEN); captured={}
        remote.require_write=lambda:None; remote.token=lambda:"tenant-token"
        remote._http=SimpleNamespace(request_file_json=lambda *args,**kwargs:(captured.update(kwargs) or {"code":0,"data":{"file_token":"file-token"}}))
        self.assertEqual(remote.upload_attachment(path,path.name),"file-token")
        self.assertEqual(captured["data"]["parent_node"],EXPORT_ARCHIVE_APP_TOKEN)
        self.assertEqual(captured["data"]["parent_type"],"bitable_file")

    def test_editor_hides_blank_template_groups_and_names_real_transition(self):
        source=(Path(__file__).parent/"lan_bitable_template_portal/frontend/src/components/CabinetPowerPage.vue").read_text(encoding="utf-8")
        self.assertIn("map(() => newGroup(false))",source)
        self.assertIn("editorGroupLabel(group, i)",source)
        self.assertIn("stateAction(target)",source)
        self.assertIn("!!form.target_state && group._editing",source)

    def test_export_ui_supports_cloud_retry_and_parallel_all_buildings(self):
        source=(Path(__file__).parent/"lan_bitable_template_portal/frontend/src/components/CabinetPowerPage.vue").read_text(encoding="utf-8")
        self.assertIn("一键导出/上传所有楼栋",source)
        self.assertIn("write('export-batches'",source)
        self.assertIn("resumeAllExports",source)
        self.assertNotIn("retryExportUpload",source)

    def test_all_source_rows_and_idle_cabinets_are_retained(self):
        counts={s:sum(r["fields"]["楼栋"]==s+"楼" for r in self.source_records) for s in "ABCDE"}
        self.assertEqual(counts,{"A":1031,"B":1087,"C":1220,"D":988,"E":1272})
        for s,n in (("D",221),("E",870)):
            self.assertEqual(sum(from_feishu(r)["empty"] for r in self.source_records if r["fields"]["楼栋"]==s+"楼"),n)
        self.assertEqual(len({r["fields"]["数据标识"] for r in self.source_records}),5598)

    def test_incomplete_history_remains_lossless(self):
        row=next(r for r in self.source_records if r["fields"]["楼栋"]=="C楼" and r["fields"]["来源工作表"]=="机柜上下电时间统计" and r["fields"]["来源行号"]==690)
        op=from_feishu(row)
        self.assertTrue(op["issues"])
        self.assertEqual(op["meta"]["cells"]["10"],44451.7603240741)
        self.assertTrue(any(g["actual"] and not g["action"] for g in op["groups"]))

    def test_template_history_and_visible_feishu_edits(self):
        row=next(r for r in self.source_records if r["fields"]["楼栋"]=="E楼" and r["fields"]["来源行号"]==2)
        op=from_feishu(row); self.assertEqual(len(op["events"]),3)
        row=copy.deepcopy(row); row["fields"]["操作类型"]="下正式电"; row["fields"]["实际完成时间"]=1788220800000
        changed=from_feishu(row); self.assertEqual(changed["groups"][0]["action"],"下正式电")
        self.assertEqual(changed["events"][0]["actual"][:10],"2026-09-01")
        self.assertEqual(changed["groups"][1:],op["groups"][1:])

    def test_only_remote_business_data_is_used(self):
        self.store.put_document("cabinet_power","config:A",{"inventory":[],"history_ready":False})
        overview=self.service.overview("D")
        self.assertEqual(overview["source"],"local"); self.assertEqual(overview["record_count"],988)
        self.assertEqual(overview["inventory_only"],221)
        self.assertEqual(overview["counts"]["unknown"],0)
        rid=next(r["record_id"] for r in self.source_records if r["fields"]["楼栋"]=="D楼")
        self.remote.records.pop(rid); self.service.do_refresh("D",{}, {})
        self.assertEqual(self.service.overview("D")["record_count"],987)

    def test_local_snapshot_opens_offline_without_background_sync(self):
        self.service.overview("D")
        for scope in "ABCDE":
            saved=self.service.local.load(scope)
            self.assertEqual(saved["config"]["scope"],scope)
        self.assertIsNot(self.service._scope_locks["A"],self.service._scope_locks["B"])
        remote=FakeFeishu(self.source_records)
        remote.list_all=lambda path="records": (_ for _ in ()).throw(TimeoutError("offline"))
        cached=CabinetPowerService(self.store,remote,self.tmp.name)
        cached._directory=self.service._directory
        started=time.perf_counter(); overview=cached.overview("D")
        self.assertLess(time.perf_counter()-started,1)
        self.assertEqual(overview["record_count"],988)
        cached.pool.shutdown(wait=True)

    def test_first_missing_building_initializes_all_buildings_from_feishu(self):
        fresh=CabinetPowerService(MemoryStore(),FakeFeishu(self.source_records),Path(self.tmp.name)/"fresh")
        fresh._directory=FakeFeishu(self.directory_records)
        try:
            overview=fresh.overview("D")
            self.assertEqual(overview["record_count"],988)
            self.assertTrue(all(fresh.local.version(scope) for scope in "ABCDE"))
            self.assertEqual(fresh.remote.list_calls,1)
            self.assertEqual(fresh._directory.list_calls,1)
        finally:
            fresh.shutdown()

    def test_first_bootstrap_reports_all_building_progress(self):
        fresh=CabinetPowerService(MemoryStore(),FakeFeishu(),Path(self.tmp.name)/"progress")
        fresh._directory=FakeFeishu()
        release=threading.Event()
        def refresh(scope,payload,job):
            release.wait(5)
            fresh.local.replace(scope,{"scope":scope,"rooms":[],"inventory":[]},[],[])
            return {"count":0}
        fresh.do_refresh=refresh
        try:
            started=fresh.bootstrap("owner",start=True)
            self.assertEqual(started["total"],5)
            self.assertEqual(started["status"],"running")
            self.assertTrue(all(item["status"] in ("pending","running") for item in started["buildings"]))
            release.set()
            deadline=time.time()+10
            while time.time()<deadline and fresh.bootstrap("owner")["status"]=="running": time.sleep(.02)
            completed=fresh.bootstrap("owner")
            self.assertEqual((completed["status"],completed["ready"]),("succeeded",5))
        finally:
            release.set(); fresh.shutdown()

    def test_bootstrap_retries_only_failed_buildings(self):
        fresh=CabinetPowerService(MemoryStore(),FakeFeishu(),Path(self.tmp.name)/"retry")
        fresh._directory=FakeFeishu()
        attempts={scope:0 for scope in "ABCDE"}
        def refresh(scope,payload,job):
            attempts[scope]+=1
            if scope!="A" and attempts[scope]==1: raise TimeoutError("read timed out")
            fresh.local.replace(scope,{"scope":scope,"rooms":[],"inventory":[]},[],[])
            return {"count":0}
        fresh.do_refresh=refresh
        try:
            fresh.bootstrap("owner",start=True)
            deadline=time.time()+10
            while time.time()<deadline and fresh.bootstrap("owner")["status"]=="running": time.sleep(.02)
            self.assertEqual((fresh.bootstrap("owner")["ready"],fresh.bootstrap("owner")["failed"]),(1,4))
            fresh.bootstrap("owner",start=True)
            self.assertEqual(attempts,{scope:1 for scope in "ABCDE"})
            fresh.bootstrap("owner",start=True,retry_failed=True)
            deadline=time.time()+10
            while time.time()<deadline and fresh.bootstrap("owner")["status"]=="running": time.sleep(.02)
            self.assertEqual(attempts,{"A":1,"B":2,"C":2,"D":2,"E":2})
        finally: fresh.shutdown()

    def test_layout_does_not_open_xlsm_at_runtime(self):
        self.service.overview("D")
        layout=self.service.layout("D","201")
        self.assertEqual(layout["layout"]["sheet"],layout["room"]["sheet"])
        self.assertTrue(layout["layout"]["cells"])

    def test_packaged_template_replaces_compatible_old_feishu_hash_only(self):
        legacy=copy.deepcopy(self.configs["A"]); legacy["template_data"]["hash"]="0"*64
        current=self.service._packaged_config("A",legacy)
        self.assertNotEqual(current["template_data"]["hash"],legacy["template_data"]["hash"])
        self.assertTrue(current["map_values"])
        incompatible=copy.deepcopy(legacy); incompatible["inventory"][0]["positions"][0]["range"]="Z999:Z999"
        with self.assertRaisesRegex(CabinetError,"结构不一致"):
            self.service._packaged_config("A",incompatible)

    def test_a_new_rooms_extend_frozen_baseline_without_resetting_existing_records(self):
        from .lan_bitable_template_portal.cabinet_power import layout_identity
        from .lan_bitable_template_portal.cabinet_power_excel import inventory_state_baseline
        rooms={"203","303","403"}
        legacy=copy.deepcopy(self.configs["A"])
        legacy["rooms"]=[room for room in legacy["rooms"] if room["id"] not in rooms]
        legacy["inventory"]=[rack for rack in legacy["inventory"] if rack["room"] not in rooms]
        legacy["template_data"]["hash"]="previous-template"
        records=[record for record in self.source_records if record["fields"]["楼栋"]=="A楼"]
        baseline=inventory_state_baseline(legacy,[from_feishu(record) for record in records])[0]
        frozen={"version":1,"layout_identity":layout_identity(legacy),"racks":baseline,"created_at":"2026-09-21"}
        self.service.local.replace("A",legacy,records,[])
        self.service.local.document("A","power_baseline:frozen_v1",frozen)
        for rack in self.service._layout_data("A")["extension"]["inventory"]:
            rack["record_id"]="rackA"+rack["room"]+rack["rack"]
            self.service._directory.records[rack["record_id"]]["fields"]["机柜类型"]="网络机柜"
        snapshot=self.service._snapshot("A")
        self.assertEqual(len(snapshot["config"]["inventory"]),1072)
        self.assertEqual({key:value for key,value in snapshot["config"]["power_baseline"].items() if key in baseline},baseline)
        self.assertEqual({r["record_id"] for r in self.service.local.load("A")["records"]},{r["record_id"] for r in records})
        for room in rooms:
            layout=self.service.layout("A",room)
            self.assertEqual(layout["room"]["total"],28)
            self.assertEqual(layout["room"]["region"],"A1:T42")
            self.assertEqual(layout["summary"]["off"],28)
            self.assertTrue(all(r["rack_type"]=="网络机柜" for r in snapshot["config"]["inventory"] if r["room"]==room))
        saved=self.service.save_operation("A",{"operation_id":"new_room_first_power_on", "room":"203","rack":"A02",
            "rack_type":"服务器机柜","result":"成功","groups":[{"id":"new_room_event","action":"上正式电",
            "expected":"2026-09-21 09:00:00","actual":"2026-09-21 09:01:00","result":"成功"}]},"owner")
        fresh=CabinetPowerService(self.store,self.remote,self.tmp.name)
        try:
            current=fresh._snapshot("A")
            self.assertEqual(next(r for r in fresh.overview("A")["racks"] if (r["room"],r["rack"])==("203","A02"))["state"],"formal")
            content=(TEMPLATES/"A.xlsm").read_bytes()
            exported=Workbook(export_workbook(content,current["config"],current["operations"]))
            for room in rooms:
                cells=exported.cells(room+"-机柜平面图")
                self.assertEqual(exported.value(cells["K7"]),"B")
                self.assertEqual(exported.value(cells["E37"]),28)
                self.assertEqual(exported.value(cells["E38"]),1 if room=="203" else 0)
                self.assertEqual(exported.value(cells["E39"]),27 if room=="203" else 28)
                self.assertFalse(any(cell.get("t")=="e" or "#REF!" in cell.findtext(T("f"),"") for cell in cells.values()))
            summary=dict(exported.rows("机柜上电汇总表"))
            self.assertEqual(summary[13][2],1072)
            for row in (10,11,12):
                self.assertEqual(summary[row][5],27 if row==10 else 28)
                self.assertEqual(summary[row][8],1 if row==10 else 0)
            self.assertEqual(exported.archive.read("xl/vbaProject.bin"),Workbook(content).archive.read("xl/vbaProject.bin"))
            self.assertEqual({key:value for key,value in fresh.local.document("A","power_baseline:frozen_v1")["racks"].items() if key in baseline},baseline)
        finally: fresh.shutdown()

    def test_a_new_room_type_backfill_is_once_only_and_preserves_records(self):
        before=self.service.local.load("A")
        extension=self.service._layout_data("A")["extension"]
        with self.service.local.connect("A") as conn, conn:
            conn.execute("DELETE FROM documents WHERE key=?",("inventory_types:"+extension["types_revision"],))
            conn.execute("UPDATE inventory SET payload=json_set(payload,'$.rack_type','') WHERE room IN ('203','303','403')")
            self.service.local._version(conn)
        snapshot=self.service._snapshot("A")
        new=[r for r in snapshot["config"]["inventory"] if r["room"] in ("203","303","403")]
        self.assertEqual(len(new),84)
        self.assertTrue(all(r["rack_type"]=="网络机柜" for r in new))
        self.assertEqual(self.service.local.load("A")["records"],before["records"])
        with self.service.local.connect("A") as conn, conn:
            conn.execute("UPDATE inventory SET payload=json_set(payload,'$.rack_type','') WHERE room='203' AND rack='A01'")
            self.service.local._version(conn)
        fresh=CabinetPowerService(self.store,self.remote,self.tmp.name)
        try:
            current=fresh._snapshot("A")
            self.assertEqual(next(r["rack_type"] for r in current["config"]["inventory"] if (r["room"],r["rack"])==("203","A01")),"")
        finally: fresh.shutdown()

    def test_down_filter_includes_embedded_history_without_duplicate_records(self):
        down=self.service.operations("D",{"direction":"down","page_size":100})
        self.assertGreater(down["total"],0)
        self.assertTrue(all(any(e["action"].startswith("下") for e in r["events"]) for r in down["items"]))
        self.assertTrue(any(not r["action"].startswith("下") for r in down["items"]))
        self.assertEqual(len({r["record_id"] for r in down["items"]}),len(down["items"]))
        self.assertEqual(self.service.operations("E",{"direction":"empty"})["total"],870)

    def test_pc_table_columns_follow_each_building_workbook(self):
        for scope,lengths in (("A",[12,15]),("B",[12,13]),("C",[13,13]),("D",[31]),("E",[22])):
            formats=self.service.overview(scope)["sheet_formats"]
            self.assertEqual([len(f["columns"]) for f in formats],lengths)
            for fmt in formats:
                self.assertEqual([c["column"] for c in fmt["columns"]],list(range(1,len(fmt["columns"])+1)))
                result=self.service.operations(scope,{"sheet":fmt["sheet"],"page_size":100})
                self.assertEqual(result["total"],fmt["count"])
                self.assertTrue(all(o["display_sheet"]==fmt["sheet"] for o in result["items"]))

    def test_save_edits_groups_and_checks_versions(self):
        old=from_feishu(next(r for r in self.source_records if r["fields"]["楼栋"]=="E楼" and r["fields"]["来源行号"]==2))
        payload={"groups":copy.deepcopy(old["groups"]),"operation_id":"edit_test_123456789","expected_version":old["version"]}
        payload["groups"][0]["actual"]="2026-08-01 09:00:00"
        saved=self.service.save_operation("E",payload,"owner",old["record_id"])
        self.assertEqual(saved["groups"][0]["actual"],"2026-08-01 09:00:00")
        self.assertEqual(saved["groups"][1:],old["groups"][1:])
        with self.assertRaises(CabinetError):
            self.service.save_operation("E",{**payload,"operation_id":"edit_stale_12345678"},"owner",old["record_id"])

    def test_create_idempotence_response_loss(self):
        self.remote.fail_after_create=True
        payload={"room":"302","rack":"B03","rack_type":"网络机柜","power":16000,"result":"成功","groups":[{"action":"上正式电","actual":"2026-08-01 09:00:00","expected":""}],"operation_id":"create_test_123456789"}
        with self.assertRaises(TimeoutError): self.service.save_operation("A",payload,"owner")
        self.remote.fail_after_create=False
        saved=self.service.save_operation("A",payload,"owner")
        self.assertEqual(saved["record_id"],"recNew1"); self.assertEqual(self.remote.creates,1)
        self.assertEqual(self.service.save_operation("A",payload,"owner")["record_id"],"recNew1")
        with self.assertRaises(CabinetError): self.service.save_operation("A",payload,"someone_else")

    def test_invalid_new_history_rejected_before_write(self):
        payload={"room":"302","rack":"B03","rack_type":"网络机柜","result":"成功","groups":[{"action":"上正式电","actual":"","expected":""}],"operation_id":"invalid_test_123456"}
        with self.assertRaises(CabinetError): self.service.save_operation("A",payload,"owner")
        self.assertEqual(self.remote.creates,0)

    def test_de_cabinet_cannot_be_created_twice(self):
        payload={"room":"201","rack":"A01","rack_type":"网络机柜","result":"成功","groups":[{"action":"上正式电","actual":"2026-09-01 09:00:00","expected":""}],"operation_id":"duplicate_de_123456"}
        with self.assertRaises(CabinetError) as caught: self.service.save_operation("D",payload,"owner")
        self.assertEqual(caught.exception.status_code,409)
        self.assertEqual(self.remote.creates,0)

    def test_power_notice_batches_parse_samples_and_infer_only_from_prior_history(self):
        samples=[
            {
                "job_id":"job-notice-d","target_record_id":"rec-notice-d","notice_type":"上电通告","scope":"D",
                "title":"EA118机房D楼机柜上电通告","start_time":"2026-09-01 09:55","end_time":"2026-09-01 23:59","quantity":"40",
                "cabinet":"D-201包间B02、B04、B05、B07、B08、B10、B11、B13、B14、B16、C11、C13、C14、C16、D12、D14、D15、D17、E11、E13、E14、E16、F11、F13、F14、F16、G11、G13、G14、G16、H11、H13、H14、H16、I11、I13、I14、I16、J14、J16",
            },
            {
                "job_id":"job-notice-b-up","target_record_id":"rec-notice-b-up","notice_type":"上电通告","scope":"B",
                "title":"EA118机房B楼机柜上电通告","start_time":"2026-09-03 11:05","end_time":"2026-09-03 23:59","quantity":"4",
                "cabinet":"B-216运营商机房B04、B05，B-247运营商机房B04、B05",
            },
            {
                "job_id":"job-notice-b-down","target_record_id":"rec-notice-b-down","notice_type":"下电通告","scope":"B",
                "title":"EA118机房B楼机柜下电通告","start_time":"2026-09-03 14:55","end_time":"2026-09-03 19:00","quantity":"12",
                "cabinet":"B-402包间B15、B16、B17、B18，C13、C14、C15、C16，D15、D16、D17、D18",
            },
        ]
        batches=[self.service.batches.create_from_notice({**item,"owner_id":"owner"}) for item in samples]
        self.assertEqual([batch["stats"]["total"] for batch in batches],[40,4,12])
        self.assertEqual({(row["scope"],row["room"]) for row in batches[0]["rows"]},{("D","201")})
        self.assertEqual({(row["room"],row["rack_type"]) for row in batches[1]["rows"]},{("216","网络机柜"),("247","网络机柜")})
        self.assertEqual({row["action"] for row in batches[0]["rows"]},{"上正式电"})
        self.assertEqual({row["action"] for row in batches[1]["rows"]},{"上正式电"})
        self.assertEqual({row["action"] for row in batches[2]["rows"]},{"下测试电"})
        self.assertTrue(all(not row["expected"] for batch in batches for row in batch["rows"]))
        self.assertTrue(all(not row["actual"] and row["result"]=="成功" for batch in batches for row in batch["rows"]))
        self.assertTrue(all(row["current_power_state"]=="off" for row in batches[0]["rows"]))
        self.assertEqual(self.service.batches._infer_notice_action({"config":{"inventory":[]},"operations":[]},"201","A01","up",""),("","机柜未匹配当前目录，无法自动识别操作类型","unknown"))
        legacy=self.service.batches.store.get(batches[0]["batch_id"])
        legacy["rows"][0].update(action="",result="")
        legacy["rows"][0].pop("current_power_state",None)
        self.service.batches.store.save(legacy,legacy["version"])
        upgraded=self.service.batches.get(legacy["batch_id"])
        self.assertEqual((upgraded["rows"][0]["action"],upgraded["rows"][0]["result"],upgraded["rows"][0]["current_power_state"]),("上正式电","成功","off"))
        repeated=self.service.batches.create_from_notice({**samples[0],"owner_id":"owner"})
        self.assertEqual(repeated["batch_id"],batches[0]["batch_id"])

    def test_notice_batch_lifecycle_preserves_row_audit(self):
        source={"event_action":"start","target_record_id":"rec-notice-summary","notice_type":"上电通告",
                "scope":"B","cabinet":"B-216运营商机房B04、B05","quantity":"2","owner_id":"owner",
                "start_time":"2026-09-19 11:05","end_time":"2026-09-19 23:59","sent_at":"2026-09-19 11:06:00"}
        service=self.service.batches
        batch=service.apply_notice_event(source)
        batch=service.apply_notice_event({**source,"event_action":"update","cabinet":"B-247运营商机房B05",
                                          "quantity":"1"})
        self.assertEqual([(row["room"],row["rack"]) for row in batch["rows"] if not row.get("notice_removed")],[('247','B05')])
        self.assertEqual(sum(bool(row.get("notice_removed")) for row in batch["rows"]),2)
        service.apply_notice_event({**source,"event_action":"undo_end"})
        service.apply_notice_event({**source,"event_action":"end","sent_at":"2026-09-20 09:00:00"})
        service.apply_notice_event({**source,"event_action":"delete"})
        self.assertFalse(service.list("owner",["B"])["items"])
        self.assertEqual(service.get(batch["batch_id"])["stats"]["confirmable"],0)
        service.apply_notice_event({**source,"event_action":"undo_delete","prior_record_id":"rec-notice-summary",
                                    "target_record_id":"rec-notice-restored"})

        restored=service.get(batch["batch_id"])
        self.assertFalse(restored["source_notice"]["deleted_at"])
        self.assertEqual(restored["source_notice"]["target_record_id"],"rec-notice-restored")

    def test_notice_delete_before_batch_creation_is_terminal_and_restorable(self):
        service=self.service.batches
        source={"event_action":"delete","target_record_id":"rec-deleted-before-start",
                "notice_type":"上电通告","event_at":2.0,"idempotency_key":"early-delete"}
        with patch.object(service,"_notice_rows",side_effect=AssertionError("deletion must not read cabinet data")):
            deleted=service.apply_notice_event(source)
        self.assertTrue(deleted["source_notice"]["deleted_at"])
        self.assertEqual(deleted["rows"],[])
        self.assertFalse(deleted["source_notice"]["sent_at"])
        self.assertEqual(service.apply_notice_event(source)["batch_id"],deleted["batch_id"])
        for action in ("start","update","end","undo_end"):
            result=service.apply_notice_event({**source,"event_action":action,"event_at":1.0,
                "cabinet":"B-216运营商机房B04","quantity":"1"})
            self.assertTrue(result["source_notice"]["deleted_at"])
            self.assertFalse(result["rows"])
        self.assertEqual(service.list("system",list("ABCDE"),True)["pending_count"],0)
        with self.assertRaisesRegex(CabinetError,"尚未创建"):
            service.apply_notice_event({**source,"event_action":"update","target_record_id":"rec-not-deleted"})
        restored=service.apply_notice_event({**source,"event_action":"undo_delete","event_at":3.0,
            "prior_record_id":source["target_record_id"],"target_record_id":"rec-restored-before-start",
            "scope":"B","cabinet":"B-216运营商机房B04","quantity":"1"})
        self.assertEqual(restored["batch_id"],deleted["batch_id"])
        self.assertFalse(restored["source_notice"]["deleted_at"])
        self.assertEqual([(row["room"],row["rack"]) for row in restored["rows"]],[("216","B04")])
        self.assertEqual(len(service.store.notice_batches()),1)

    def test_batch_store_version_compare_is_atomic_across_instances(self):
        batch=self.service.batches.create_manual([self._manual_batch_row("A","2026-09-19 01:02:03")],"owner")
        first=self.service.batches.store.get(batch["batch_id"])
        other=type(self.service.batches.store)(self.service.batches.store.path)
        stale=other.get(batch["batch_id"])
        first["error"]="first"
        self.service.batches.store.save(first,first["version"])
        stale["error"]="stale"
        with self.assertRaisesRegex(CabinetError,"其他操作更新"):
            other.save(stale,stale["version"])

    def test_batch_store_migrates_legacy_embedded_rows_without_loss(self):
        store=self.service.batches.store
        batch=self.service.batches.create_manual([self._manual_batch_row("A","2026-09-19 01:02:03")],"owner")
        current=store.get(batch["batch_id"])
        with store._connect() as conn,conn:
            payload=json.loads(conn.execute("SELECT payload_json FROM batches WHERE batch_id=?",(batch["batch_id"],)).fetchone()[0])
            payload["rows"]=current["rows"]
            conn.execute("UPDATE batches SET payload_json=? WHERE batch_id=?",(json.dumps(payload,ensure_ascii=False),batch["batch_id"]))
            conn.execute("DELETE FROM batch_rows WHERE batch_id=?",(batch["batch_id"],))
            conn.execute("DELETE FROM batch_meta WHERE key='batch_payload_normalization_version'")
        migrated=type(store)(store.path)
        restored=migrated.get(batch["batch_id"])
        self.assertEqual([(row["room"],row["rack"]) for row in restored["rows"]],
                         [(current["rows"][0]["room"],current["rows"][0]["rack"])])
        with migrated._connect() as conn:
            self.assertNotIn("rows",json.loads(conn.execute("SELECT payload_json FROM batches WHERE batch_id=?",(batch["batch_id"],)).fetchone()[0]))

    def test_deleted_notice_rollback_failure_stays_visible_as_exception(self):
        service=self.service.batches
        source={"target_record_id":"rec-notice-delete","notice_type":"上电通告","scope":"B",
                "cabinet":"B-216运营商机房B04","quantity":"1","owner_id":"owner"}
        batch=service.create_from_notice(source)
        stored=service.store.get(batch["batch_id"])
        stored["rows"][0].update(status="completed",wrote_record=True,record_id="rec-written",
                                   operation_id="batch_written_operation")
        service.store.save(stored,stored["version"])
        self.service.rollback_batch_operation=lambda *_args: (_ for _ in ()).throw(CabinetError("已有后续操作",409))
        service.apply_notice_event({**source,"event_action":"delete"})
        result=self._wait_batch(batch["batch_id"])
        self.assertEqual(result["rows"][0]["status"],"rollback_blocked")
        self.assertEqual(service.list("owner",["B"])["total"],0)
        exceptions=service.list("owner",["B"],status="notice_rollback_error")
        self.assertEqual([item["batch_id"] for item in exceptions["items"]],[batch["batch_id"]])
        self.assertTrue(exceptions["items"][0]["source_notice_deleted"])

    def test_notice_delete_waits_for_inflight_write_then_resumes_rollback(self):
        service=self.service.batches
        source={"target_record_id":"rec-notice-race","notice_type":"上电通告","scope":"B",
                "cabinet":"B-216运营商机房B04","quantity":"1","owner_id":"owner"}
        batch=service.create_from_notice(source)
        stored=service.store.get(batch["batch_id"])
        stored["rows"][0].update(status="writing",operation_started=True,operation_id="batch_race_operation")
        service.store.save(stored,stored["version"])
        with self.assertRaisesRegex(CabinetError,"正在写入"):
            service.apply_notice_event({**source,"event_action":"delete"})
        self.assertEqual(service.list("owner",["B"],status="notice_rollback_error")["total"],1)
        stored=service.store.get(batch["batch_id"])
        stored["rows"][0].update(status="completed",record_id="rec-race-written",wrote_record=True)
        service.store.save(stored,stored["version"])
        calls=[]
        self.service.rollback_batch_operation=lambda *args: calls.append(args)
        service._resume_notice_rollback_after_confirm(batch["batch_id"])
        result=self._wait_batch(batch["batch_id"])
        self.assertEqual(result["rows"][0]["status"],"rolled_back")
        self.assertEqual(len(calls),1)
        self.assertEqual(service.list("owner",["B"],status="notice_rollback_error")["total"],0)

    def test_notice_update_removed_confirmed_rack_rolls_back_without_counting_it(self):
        service=self.service.batches
        source={"target_record_id":"rec-notice-update","notice_type":"上电通告","scope":"B",
                "cabinet":"B-216运营商机房B04、B05","quantity":"2","owner_id":"owner",
                "sent_at":"2026-09-18 10:00:00"}
        batch=service.create_from_notice(source)
        stored=service.store.get(batch["batch_id"])
        stored["rows"][0].update(status="completed",record_id="rec-written-update",wrote_record=True,
                                   operation_id="batch_update_operation")
        service.store.save(stored,stored["version"])
        calls=[]
        self.service.rollback_batch_operation=lambda *args: calls.append(args)
        service.apply_notice_event({**source,"event_action":"update",
            "cabinet":"B-247运营商机房B05","quantity":"1"})
        self._wait_batch(batch["batch_id"])
        self.assertEqual(len(calls),1)
        rows=service.get(batch["batch_id"])["rows"]
        self.assertEqual({(row["room"],row["rack"]) for row in rows if not row.get("notice_removed")},
                         {("247","B05")})
        service.apply_notice_event({**source,"event_action":"end","sent_at":"2026-09-05 11:00:00"})
        self.assertEqual([(item["room"],item["rack"]) for item in service.store.notice_batches()[0]["rows"] if not item.get("notice_removed")],
                         [("247","B05")])

    def test_notice_update_keeps_committed_identity_when_original_row_was_corrected(self):
        service=self.service.batches
        source={"target_record_id":"rec-notice-corrected","notice_type":"上电通告","scope":"B",
                "cabinet":"B-216运营商机房B04","quantity":"1","owner_id":"owner"}
        batch=service.create_from_notice(source)
        stored=service.store.get(batch["batch_id"])
        stored["rows"][0].update(rack="B05",status="completed",wrote_record=True,
                                   record_id="rec-written-B05",operation_id="batch_corrected_operation")
        service.store.save(stored,stored["version"])
        self.service.rollback_batch_operation=lambda *_args: None
        service.apply_notice_event({**source,"event_action":"update"})
        rows=service.get(batch["batch_id"])["rows"]
        old=next(row for row in rows if row.get("record_id")=="rec-written-B05")
        self.assertEqual(old["rack"],"B05")
        self.assertTrue(old["notice_removed"])
        self.assertTrue(any(row["rack"]=="B04" and not row.get("notice_removed") for row in rows))

    def test_notice_end_uses_final_cabinets_even_when_prior_update_failed(self):
        service=self.service.batches
        source={"target_record_id":"rec-notice-final","notice_type":"上电通告","scope":"B",
                "cabinet":"B-216运营商机房B04","quantity":"1","owner_id":"owner",
                "sent_at":"2026-09-18 10:00:00"}
        service.create_from_notice(source)
        with self.assertRaises(CabinetError):
            service.apply_notice_event({**source,"event_action":"update","cabinet":"无法识别"})
        service.apply_notice_event({**source,"event_action":"end",
            "cabinet":"B-247运营商机房B05","cabinet_verified":True,"sent_at":"2026-09-18 18:00:00"})
        self.assertEqual([(row["room"],row["rack"]) for row in service.store.notice_batches()[0]["rows"] if not row.get("notice_removed")],
                         [("247","B05")])

    def test_notice_batch_rejects_other_building_invalid_row(self):
        service=self.service.batches
        source={"target_record_id":"rec-notice-scope","notice_type":"上电通告","scope":"A",
                "cabinet":"A-202包间B01","quantity":"1","owner_id":"owner"}
        batch=service.create_from_notice(source)
        row=batch["rows"][0]
        batch=service.update(batch["batch_id"],{"version":batch["version"],"rows":[
            {"row_id":row["row_id"],"scope":"B","room":"202","rack":"B01"}]},"owner",["A","B"],True)
        self.assertIn("notice_scope",{issue["code"] for issue in batch["rows"][0]["issues"]})
        service.apply_notice_event({**source,"event_action":"end","sent_at":"2026-09-18 18:00:00"})
        self.assertFalse(service.get(batch["batch_id"])["stats"]["confirmable"])

    def test_notice_direction_change_resets_unconfirmed_action(self):
        service=self.service.batches
        source={"target_record_id":"rec-notice-direction","notice_type":"上电通告","scope":"D",
                "cabinet":"D-201包间B02","quantity":"1","owner_id":"owner"}
        batch=service.create_from_notice(source)
        self.assertEqual(batch["rows"][0]["action"],"上正式电")
        changed=service.apply_notice_event({**source,"event_action":"update","notice_type":"下电通告"})
        self.assertNotEqual(changed["rows"][0]["action"],"上正式电")

    def test_power_notice_count_warning_requires_fresh_acknowledgement(self):
        batch=self.service.batches.create_from_notice({
            "job_id":"job-notice-warning","target_record_id":"rec-notice-warning","notice_type":"上电通告","scope":"B",
            "title":"EA118机房B楼机柜上电通告","start_time":"2026-09-03 11:05","end_time":"2026-09-03 23:59","quantity":"3",
            "cabinet":"B-216运营商机房B04、B05","owner_id":"owner",
        })
        self.assertIn("quantity_mismatch",{item["code"] for item in batch["blocking_warnings"]})
        acknowledged=self.service.batches.update(batch["batch_id"],{
            "version":batch["version"],"rows":[],"acknowledge_warnings":True,
        },"owner",["B"])
        self.assertTrue(acknowledged["warnings_acknowledged"])
        row=acknowledged["rows"][0]
        changed=self.service.batches.update(batch["batch_id"],{
            "version":acknowledged["version"],"rows":[{"row_id":row["row_id"],"rack":"B11"}],
        },"owner",["B"])
        self.assertFalse(changed["warnings_acknowledged"])
        self.assertEqual((changed["rows"][0]["current_power_state"],changed["rows"][0]["action"]),("unknown",""))
        manual=self.service.batches.update(batch["batch_id"],{
            "version":changed["version"],"rows":[{"row_id":row["row_id"],"action":"上测试电"}],
        },"owner",["B"])
        moved=self.service.batches.update(batch["batch_id"],{
            "version":manual["version"],"rows":[{"row_id":row["row_id"],"rack":"B04","action":"上测试电"}],
        },"owner",["B"])
        self.assertEqual(moved["rows"][0]["action"],"上测试电")

    def test_power_notice_outbox_handoff_never_raises_into_notice_flow(self):
        bin_path=str(Path(__file__).parent)
        added_path=bin_path not in sys.path
        if added_path: sys.path.insert(0,bin_path)
        from .lan_bitable_template_portal.server import PortalRuntime
        original=PortalRuntime.state_store
        original_service=PortalRuntime.cabinet_power_service
        class BrokenStore:
            def enqueue_outbox_event(self,*_args,**_kwargs): raise TimeoutError("locked")
        PortalRuntime.state_store=BrokenStore()
        PortalRuntime.cabinet_power_service=None
        try:
            event_id=PortalRuntime.enqueue_cabinet_notice_batch({
                "work_type":"power","action":"start","notice_type":"上电通告","scope":"B",
                "title":"test","start_time":"2026-09-03 11:05","end_time":"2026-09-03 23:59",
                "cabinet":"B-216运营商机房B04","quantity":"1",
            },job_id="job-safe",target_record_id="rec-safe",request_payload={"_auth_open_id":"owner"})
            self.assertEqual(event_id,0)
        finally:
            PortalRuntime.state_store=original
            PortalRuntime.cabinet_power_service=original_service
            if added_path: sys.path.remove(bin_path)

    def test_power_notice_handoff_survives_primary_outbox_failure(self):
        bin_path=str(Path(__file__).parent)
        added_path=bin_path not in sys.path
        if added_path: sys.path.insert(0,bin_path)
        from .lan_bitable_template_portal.server import PortalRuntime
        original_store,original_service=PortalRuntime.state_store,PortalRuntime.cabinet_power_service
        class BrokenStore:
            def enqueue_outbox_event(self,*_args,**_kwargs): raise TimeoutError("locked")
            def lease_outbox_events(self,*_args,**_kwargs): raise TimeoutError("locked")
            def append_event_async(self,*_args,**_kwargs): raise TimeoutError("locked")
        PortalRuntime.state_store=BrokenStore()
        PortalRuntime.cabinet_power_service=self.service
        try:
            with patch.object(PortalRuntime,"ensure_cabinet_notice_worker",return_value=None):
                queued=PortalRuntime.enqueue_cabinet_notice_batch({
                    "work_type":"power","action":"start","notice_type":"上电通告","scope":"B",
                    "cabinet":"B-216运营商机房B04","quantity":"1",
                },job_id="fallback-job",target_record_id="rec-fallback")
            self.assertEqual(queued,-1)
            self.assertIsNotNone(self.service.batches.store.next_notice_handoff())
            result=PortalRuntime._process_cabinet_notice_queue_once()
            self.assertEqual(result["status"],"success")
            self.assertIsNone(self.service.batches.store.next_notice_handoff())
            self.assertEqual(self.service.batches.get(result["batch_id"])["stats"]["total"],1)
        finally:
            PortalRuntime.state_store=original_store
            PortalRuntime.cabinet_power_service=original_service
            if added_path: sys.path.remove(bin_path)

    def test_failed_notice_handoff_can_resume_after_restart(self):
        store=self.service.batches.store
        store.queue_notice_handoff({"idempotency_key":"resume-notice-handoff","event_action":"start"})
        for _ in range(5):
            self.assertIsNotNone(store.next_notice_handoff())
            store.finish_notice_handoff("resume-notice-handoff","temporary failure")
        self.assertIsNone(store.next_notice_handoff())
        store.requeue_failed_notice_handoffs()
        self.assertEqual(store.next_notice_handoff()["attempts"],0)

    def test_deleted_notice_handoff_recovery_clears_primary_and_fallback_errors(self):
        bin_path=str(Path(__file__).parent)
        added_path=bin_path not in sys.path
        if added_path: sys.path.insert(0,bin_path)
        from .lan_bitable_template_portal.server import PortalRuntime
        from .lan_bitable_template_portal.state_store import LanPortalStateStore
        state=LanPortalStateStore(Path(self.tmp.name)/"deleted-handoff.sqlite3")
        try:
            with patch.object(PortalRuntime,"state_store",state), patch.object(PortalRuntime,"cabinet_power_service",self.service), \
                    patch("bin.lan_bitable_template_portal.server.query_record_by_id",side_effect=AssertionError("deleted notice must not be queried")):
                channel=PortalRuntime.cabinet_notice_queue_channel
                source={"event_action":"delete","target_record_id":"rec-missing-batch","notice_type":"上电通告",
                        "idempotency_key":"missing-batch-delete","event_at":2.0}
                event_id=state.enqueue_outbox_event(channel,source)
                for _ in range(6):
                    state.mark_outbox_event(event_id,"pending",max_attempts=6,error="来源通告待办尚未创建，请稍后重试")
                state.requeue_failed_outbox_events(channel,max_attempts=5)
                self.assertTrue(state.list_outbox_events(channel,status="failed"))
                PortalRuntime._recover_cabinet_notice_rollbacks()
                self.assertFalse(state.list_outbox_events(channel,status="failed"))
                store=self.service.batches.store
                first=store.notice_by_target(source["target_record_id"])
                self.assertTrue(first["source_notice"]["deleted_at"])
                store.queue_notice_handoff(source)
                for _ in range(5): store.finish_notice_handoff(source["idempotency_key"],"来源通告待办尚未创建，请稍后重试")
                store.requeue_failed_notice_handoffs()
                second=PortalRuntime._process_cabinet_notice_queue_once()
                self.assertEqual((second["status"],second["batch_id"]),("success",first["batch_id"]))
                self.assertFalse(store.failed_notice_handoffs())
                state.enqueue_outbox_event(channel,{**source,"event_action":"end","event_at":1.0,"idempotency_key":"late-end"})
                self.assertEqual(PortalRuntime._process_cabinet_notice_queue_once()["status"],"success")
                self.assertFalse(PortalRuntime._process_cabinet_notice_queue_once()["processed"])
        finally:
            state.shutdown_write_worker(timeout=1)
            if added_path: sys.path.remove(bin_path)

    def test_stale_notice_update_does_not_overwrite_newer_end_snapshot(self):
        service=self.service.batches
        source={"target_record_id":"rec-ordered-notice","notice_type":"上电通告","scope":"B",
                "cabinet":"B-216运营商机房B04","quantity":"1","owner_id":"owner","event_at":1.0,
                "sent_at":"2026-09-18 10:00:00"}
        service.apply_notice_event(source)
        service.apply_notice_event({**source,"event_action":"end","event_at":3.0,
            "cabinet":"B-247运营商机房B05","cabinet_verified":True,"sent_at":"2026-09-18 18:00:00"})
        service.apply_notice_event({**source,"event_action":"update","event_at":2.0})
        self.assertEqual([(item["room"],item["rack"]) for item in service.store.notice_batches()[0]["rows"] if not item.get("notice_removed")],
                         [("247","B05")])

    def test_power_notice_outbox_creates_one_idempotent_batch(self):
        bin_path=str(Path(__file__).parent)
        added_path=bin_path not in sys.path
        if added_path: sys.path.insert(0,bin_path)
        from .lan_bitable_template_portal.server import PortalRuntime
        from .lan_bitable_template_portal.state_store import LanPortalStateStore
        original_store=PortalRuntime.state_store
        original_service=PortalRuntime.cabinet_power_service
        state=LanPortalStateStore(Path(self.tmp.name)/"notice-outbox.sqlite3")
        PortalRuntime.state_store=state
        PortalRuntime.cabinet_power_service=None
        try:
            prepared={
                "work_type":"power","action":"start","notice_type":"下电通告","scope":"B",
                "title":"EA118机房B楼机柜下电通告","start_time":"2026-09-03 14:55","end_time":"2026-09-03 19:00",
                "cabinet":"B-402包间B15、B16","quantity":"2",
            }
            first=PortalRuntime.enqueue_cabinet_notice_batch(prepared,job_id="job-outbox",target_record_id="rec-outbox",request_payload={"_auth_open_id":"owner"})
            second=PortalRuntime.enqueue_cabinet_notice_batch(prepared,job_id="job-outbox",target_record_id="rec-outbox",request_payload={"_auth_open_id":"owner"})
            self.assertEqual(first,second)
            PortalRuntime.cabinet_power_service=self.service
            result=PortalRuntime._process_cabinet_notice_queue_once()
            self.assertEqual(result["status"],"success")
            self.assertEqual(self.service.batches.get(result["batch_id"])["stats"]["total"],2)
            self.assertFalse(PortalRuntime._process_cabinet_notice_queue_once()["processed"])
        finally:
            PortalRuntime.stop_cabinet_notice_worker()
            PortalRuntime.state_store=original_store
            PortalRuntime.cabinet_power_service=original_service
            state.shutdown_write_worker(timeout=1)
            if added_path: sys.path.remove(bin_path)

    def test_power_notice_outbox_updates_then_records_successful_end(self):
        bin_path=str(Path(__file__).parent)
        added_path=bin_path not in sys.path
        if added_path: sys.path.insert(0,bin_path)
        from .lan_bitable_template_portal.server import PortalRuntime
        from .lan_bitable_template_portal.state_store import LanPortalStateStore
        old_store,old_service=PortalRuntime.state_store,PortalRuntime.cabinet_power_service
        state=LanPortalStateStore(Path(self.tmp.name)/"notice-lifecycle.sqlite3")
        PortalRuntime.state_store=state; PortalRuntime.cabinet_power_service=None
        try:
            base={"work_type":"power","notice_type":"上电通告","scope":"B",
                  "cabinet":"B-216运营商机房B04、B05","quantity":"2","response_time":"2026-09-19 11:05:00"}
            for action,job,changes in (
                ("start","start-job",{}),
                ("update","update-job",{"cabinet":"B-247运营商机房B05","quantity":"1"}),
                ("end","end-job",{"response_time":"2026-09-05 09:12:00"}),
            ):
                PortalRuntime.enqueue_cabinet_notice_batch({**base,**changes,"action":action},job_id=job,
                    target_record_id="rec-lifecycle",request_payload={"_auth_open_id":"owner"})
            PortalRuntime.cabinet_power_service=self.service
            with patch("bin.lan_bitable_template_portal.server.query_record_by_id",return_value=(True,{"fields":{
                "柜号":"B-247运营商机房B05","数量（个）":"1","上电状态":"结束","实际结束时间":"2026-09-05 09:12:00"}})):
                self.assertEqual([PortalRuntime._process_cabinet_notice_queue_once()["status"] for _ in range(3)],["success"]*3)
            batch=self.service.batches.store.notice_batches()[0]
            self.assertEqual([(item["room"],item["rack"]) for item in batch["rows"] if not item.get("notice_removed")],
                             [("247","B05")])
            self.assertEqual(batch["source_notice"]["sent_at"],"2026-09-19 11:05:00")
            self.assertEqual(batch["source_notice"]["ended_at"],"2026-09-05 09:12:00")
            self.assertFalse(PortalRuntime._process_cabinet_notice_queue_once()["processed"])
        finally:
            PortalRuntime.stop_cabinet_notice_worker()
            PortalRuntime.state_store=old_store; PortalRuntime.cabinet_power_service=old_service
            state.shutdown_write_worker(timeout=1)
            if added_path: sys.path.remove(bin_path)

    def test_notice_end_read_failure_keeps_start_cabinets_until_verified(self):
        bin_path=str(Path(__file__).parent)
        added_path=bin_path not in sys.path
        if added_path: sys.path.insert(0,bin_path)
        from .lan_bitable_template_portal.server import PortalRuntime
        from .lan_bitable_template_portal.state_store import LanPortalStateStore
        old_store,old_service=PortalRuntime.state_store,PortalRuntime.cabinet_power_service
        state=LanPortalStateStore(Path(self.tmp.name)/"notice-final-read.sqlite3")
        PortalRuntime.state_store=state; PortalRuntime.cabinet_power_service=self.service
        try:
            source={"target_record_id":"rec-notice-final-read","notice_type":"上电通告","scope":"B",
                    "cabinet":"B-216运营商机房B04","quantity":"1","owner_id":"owner",
                    "sent_at":"2026-09-18 10:00:00"}
            self.service.batches.create_from_notice(source)
            with patch.object(PortalRuntime,"ensure_cabinet_notice_worker",return_value=None):
                PortalRuntime.enqueue_cabinet_notice_batch({**source,"work_type":"power","action":"end",
                    "response_time":"2026-09-18 19:00:00"},job_id="final-read",target_record_id=source["target_record_id"])
            with patch("bin.lan_bitable_template_portal.server.query_record_by_id",return_value=(False,"timeout")):
                self.assertEqual(PortalRuntime._process_cabinet_notice_queue_once()["status"],"pending")
            self.assertEqual([(item["room"],item["rack"]) for item in self.service.batches.store.notice_batches()[0]["rows"] if not item.get("notice_removed")],[('216','B04')])
            with patch("bin.lan_bitable_template_portal.server.query_record_by_id",return_value=(True,{"fields":{
                "柜号":"B-216运营商机房B04","上电状态":"开始","实际结束时间":"2026-09-18 19:00:00"}})):
                self.assertEqual(PortalRuntime._process_cabinet_notice_queue_once()["status"],"pending")
            self.assertEqual([(item["room"],item["rack"]) for item in self.service.batches.store.notice_batches()[0]["rows"] if not item.get("notice_removed")],[('216','B04')])
            with patch("bin.lan_bitable_template_portal.server.query_record_by_id",return_value=(True,{"fields":{
                "柜号":"B-247运营商机房B05","数量（个）":"1","上电状态":"结束","实际结束时间":"2026-09-18 19:00:00"}})):
                self.assertEqual(PortalRuntime._process_cabinet_notice_queue_once()["status"],"success")
            self.assertEqual([(item["room"],item["rack"]) for item in self.service.batches.store.notice_batches()[0]["rows"] if not item.get("notice_removed")],[("247","B05")])
        finally:
            PortalRuntime.stop_cabinet_notice_worker()
            PortalRuntime.state_store=old_store; PortalRuntime.cabinet_power_service=old_service
            state.shutdown_write_worker(timeout=1)
            if added_path: sys.path.remove(bin_path)

    def test_power_notice_outbox_extracts_qt_notice_text(self):
        bin_path=str(Path(__file__).parent)
        added_path=bin_path not in sys.path
        if added_path: sys.path.insert(0,bin_path)
        from .lan_bitable_template_portal.server import PortalRuntime
        from .lan_bitable_template_portal.state_store import LanPortalStateStore
        old_store,old_service=PortalRuntime.state_store,PortalRuntime.cabinet_power_service
        state=LanPortalStateStore(Path(self.tmp.name)/"qt-notice.sqlite3")
        PortalRuntime.state_store=state; PortalRuntime.cabinet_power_service=None
        try:
            text="【下电通告】状态：开始\n【名称】EA118机房B楼机柜下电通告\n【时间】2026-09-03 14:55~2026-09-03 19:00\n【柜号】B-402包间B15、B16\n【数量】2\n【进度】准备工作已完成"
            PortalRuntime.enqueue_cabinet_notice_batch({"work_type":"power","action":"start",
                "notice_type":"下电通告","text":text},job_id="qt-start",target_record_id="rec-qt-power")
            PortalRuntime.cabinet_power_service=self.service
            result=PortalRuntime._process_cabinet_notice_queue_once()
            self.assertEqual(result["status"],"success")
            self.assertEqual(self.service.batches.get(result["batch_id"])["stats"]["total"],2)
        finally:
            PortalRuntime.stop_cabinet_notice_worker()
            PortalRuntime.state_store=old_store; PortalRuntime.cabinet_power_service=old_service
            state.shutdown_write_worker(timeout=1)
            if added_path: sys.path.remove(bin_path)

    def test_batch_pdf_overlap_is_local_until_confirmed(self):
        operation=next(
            op for record in self.source_records
            if record["fields"]["楼栋"]=="A楼"
            for op in [from_feishu(record)]
            if any(event["result"]=="成功" and event["actual"]<="2026-09-15 00:00:00" for event in op["events"])
        )
        event=next(event for event in operation["events"] if event["result"]=="成功" and event["actual"]<="2026-09-15 00:00:00")
        rack_type=next(item["rack_type"] for item in self.configs["A"]["inventory"] if (item["room"],item["rack"])==(operation["room"],operation["rack"]))
        text=f"""机柜{event['action']}确认单
申请时间： 2026-09-14 10:00:00
申请单号： [Z260914001234567890] 操作类型： {event['action']}
EA118  A{operation['room'][0]}-{int(operation['room'][1:])}.EA118  {operation['rack']}  {operation['rack']}  {rack_type}  2UR  {event['actual']}  成功  {event['actual']}
"""
        class Page:
            def extract_text(self,**_kwargs): return text
        class Reader:
            is_encrypted=False; pages=[Page()]
        self.service.batches._pdf_reader=lambda _path:Reader()
        before=self.service.overview("A")["record_count"]
        batch=self.service.batches.recognize([("sample.pdf",b"%PDF-test")],"owner")
        deadline=time.time()+5
        while time.time()<deadline and (batch:=self.service.batches.get(batch["batch_id"]))["status"]=="recognizing": time.sleep(.01)
        self.assertEqual(batch["stats"]["duplicate"],1)
        repeated=self.service.batches.recognize([("renamed.pdf",b"%PDF-test")],"owner")
        self.assertEqual(repeated["batch_id"],batch["batch_id"])
        cleared=self.service.batches.clear_overlaps(batch["batch_id"],batch["version"],"owner",["A"])
        self.assertEqual(cleared["rows"][0]["status"],"excluded_duplicate")
        self.assertEqual(self.service.overview("A")["record_count"],before)
        cancelled=self.service.batches.cancel(batch["batch_id"],"owner")
        self.assertEqual(cancelled["status"],"cancelled")
        cleaned=self.service.batches.cleanup_file(batch["batch_id"],cancelled["files"][0]["file_id"],"owner")
        self.assertTrue(cleaned["files"][0]["cleaned_at"])

    def test_pdf_duplicate_upload_authorizes_before_retry_and_preserves_business_history(self):
        service=self.service.batches
        files=[("sample.pdf",b"%PDF-retry-permissions")]
        with patch.object(service._pdf_pool,"submit") as submit:
            batch=service.recognize(files,"owner")
            def failed(current):
                current.update(status="failed",scopes=["A"],error="rollback failed")
                current["rows"]=[{**self._manual_batch_row("A","2026-09-14 01:02:03"),
                    "row_id":"r1","status":"rollback_failed","operation_id":"original-operation",
                    "record_id":"original-record","wrote_record":True,"attempts":[{"operation_id":"earlier-operation"}]}]
            saved=service._change(batch["batch_id"],failed)
            with self.assertRaises(CabinetError) as denied:
                service.recognize(files,"other",["B"])
            self.assertEqual(denied.exception.status_code,403)
            for owner,allowed,admin in (("owner",["A"],False),("viewer",["A"],False),("admin",[],True)):
                result=service.recognize(files,owner,allowed,admin)
                self.assertTrue(result["duplicate_upload"])
                self.assertEqual(service.get(batch["batch_id"]),saved)
            self.assertEqual(submit.call_count,1)

    def test_pdf_retry_only_restarts_empty_failed_recognition_with_current_version(self):
        service=self.service.batches
        files=[("sample.pdf",b"%PDF-retry-empty")]
        with patch.object(service._pdf_pool,"submit") as submit:
            batch=service.recognize(files,"owner")
            failed=service._change(batch["batch_id"],lambda current:current.update(status="failed",error="parser failed",scopes=["A"]))
            with self.assertRaises(CabinetError):
                service.recognize(files,"viewer",["A"])
            self.assertEqual(service.get(batch["batch_id"]),failed)
            retried=service.recognize(files,"admin",[],True)
            self.assertEqual(retried["status"],"recognizing")
            self.assertEqual(submit.call_count,2)
            service._change(batch["batch_id"],lambda current:current.update(status="failed"))
            original_change=service._change
            def concurrent_change(batch_id,callback,**kwargs):
                original_change(batch_id,lambda current:current.update(status="cancelled"))
                return original_change(batch_id,callback,**kwargs)
            with patch.object(service,"_change",side_effect=concurrent_change),self.assertRaises(CabinetError):
                service.recognize(files,"owner")
            self.assertEqual(service.get(batch["batch_id"])["status"],"cancelled")
            self.assertEqual(submit.call_count,2)

    def test_batch_confirm_creates_abc_and_appends_de(self):
        rows=[]
        for scope in ("A","D"):
            rack=self.configs[scope]["inventory"][0]
            rows.append({"scope":scope,"room":rack["room"],"rack":rack["rack"],"rack_type":rack["rack_type"],
                         "action":"上正式电","expected":"2026-09-14 01:02:03","actual":"2026-09-14 01:02:03","result":"成功"})
        before={scope:self.service.overview(scope)["record_count"] for scope in ("A","D")}
        batch=self.service.batches.create_manual(rows,"owner")
        with self.assertRaises(CabinetError) as denied:
            self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
        self.assertEqual(denied.exception.status_code,403)
        self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A","D"])
        deadline=time.time()+10
        while time.time()<deadline:
            batch=self.service.batches.get(batch["batch_id"])
            if batch["status"]!="running": break
            time.sleep(.01)
        self.assertEqual(batch["stats"]["completed"],2)
        self.assertEqual(self.service.overview("A")["record_count"],before["A"]+1)
        self.assertEqual(self.service.overview("D")["record_count"],before["D"])
        self.assertTrue(any(event["actual"]=="2026-09-14 01:02:03" for op in self.service._snapshot("D")["operations"] for event in op["events"]))
        for scope in ("A","D"):
            record_id=next(row["record_id"] for row in batch["rows"] if row["scope"]==scope)
            saved=next(op for op in self.service._snapshot(scope)["operations"] if op["record_id"]==record_id)
            self.assertTrue(any(item["batch_id"]==batch["batch_id"] for item in saved["meta"]["batch_rows"]))

    def _wait_batch(self,batch_id):
        deadline=time.time()+15
        while time.time()<deadline:
            batch=self.service.batches.get(batch_id)
            if batch["status"]!="running": return batch
            time.sleep(.01)
        self.fail("批次处理超时")

    def _manual_batch_row(self,scope,actual,action="上正式电"):
        rack=self.configs[scope]["inventory"][0]
        return {"scope":scope,"room":rack["room"],"rack":rack["rack"],"rack_type":rack["rack_type"],
                "action":action,"expected":actual,"actual":actual,"result":"成功"}

    def test_batch_rollback_restores_abc_create_and_de_history(self):
        before_d=copy.deepcopy(next(op for op in self.service._snapshot("D")["operations"] if (op["room"],op["rack"])==(self.configs["D"]["inventory"][0]["room"],self.configs["D"]["inventory"][0]["rack"])))
        count_a=self.service.overview("A")["record_count"]
        batch=self.service.batches.create_manual([self._manual_batch_row(scope,"2026-09-14 01:02:03") for scope in ("A","D")],"owner")
        self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A","D"])
        done=self._wait_batch(batch["batch_id"])
        self.assertEqual(done["stats"]["completed"],2)
        self.service.batches.rollback(batch["batch_id"],{"version":done["version"],"all":True},"owner",["A","D"])
        undone=self._wait_batch(batch["batch_id"])
        self.assertEqual(undone["stats"]["rolled_back"],2,undone["rows"])
        self.assertEqual(self.service.overview("A")["record_count"],count_a)
        self.assertEqual(next(op for op in self.service._snapshot("D")["operations"] if op["record_id"]==before_d["record_id"])["groups"],before_d["groups"])
        self.assertNotIn(next(row["record_id"] for row in undone["rows"] if row["scope"]=="A"),self.remote.records)

    def test_rolled_back_rows_can_be_confirmed_with_new_operation_ids(self):
        batch=self.service.batches.create_manual(
            [self._manual_batch_row(scope,"2026-09-14 01:02:03") for scope in ("A","D")],"owner"
        )
        self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A","D"])
        first=self._wait_batch(batch["batch_id"])
        old_ids={row["scope"]:row["operation_id"] for row in first["rows"]}
        self.service.batches.rollback(batch["batch_id"],{"version":first["version"],"all":True},"owner",["A","D"])
        undone=self._wait_batch(batch["batch_id"])
        self.assertEqual(undone["stats"]["confirmable"],2)
        self.service.batches.confirm(batch["batch_id"],{"version":undone["version"],"all":True},"owner",["A","D"])
        repeated=self._wait_batch(batch["batch_id"])
        self.assertEqual(repeated["stats"]["completed"],2,repeated["rows"])
        for row in repeated["rows"]:
            self.assertNotEqual(row["operation_id"],old_ids[row["scope"]])
            self.assertEqual(row["attempts"][0]["operation_id"],old_ids[row["scope"]])

    def test_batch_image_reason_rollback_and_reconfirm(self):
        from PIL import Image
        rack=self.configs["A"]["inventory"][0]
        actual="2026-09-14 01:02:03"
        row={**self._manual_batch_row("A",actual),"result":"失败","failure_reason":""}
        batch=self.service.batches.create_manual([row],"owner")
        self.assertIn("failure_reason",{issue["code"] for issue in batch["rows"][0]["issues"]})
        batch=self.service.batches.update(batch["batch_id"],{"version":batch["version"],"rows":[{"row_id":batch["rows"][0]["row_id"],"failure_reason":"现场核验未通过"}]},"owner",["A"])
        self.assertEqual(batch["stats"]["confirmable"],1)
        image=Image.new("RGB",(160,80),"white"); output=io.BytesIO(); image.save(output,format="PNG")
        candidate={"scope":"A","room":rack["room"],"rack":rack["rack"],"supplier_rack":"",
                   "action":"上正式电","expected":actual,"actual":actual}
        with patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",return_value=[candidate]):
            batch=self.service.batches.add_images(batch["batch_id"],[("proof.png",output.getvalue())],"owner",["A"])
            deadline=time.time()+5
            while time.time()<deadline and self.service.batches.get(batch["batch_id"])["images"][0]["status"]=="recognizing": time.sleep(.01)
        batch=self.service.batches.get(batch["batch_id"])
        self.assertEqual(batch["images"][0]["suggestions"][0]["status"],"applied")
        image_id=batch["images"][0]["image_id"]
        self.assertEqual(batch["rows"][0]["evidence_images"],[image_id])
        self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
        done=self._wait_batch(batch["batch_id"])
        self.assertEqual(done["stats"]["completed"],1,done["rows"])
        record_id=done["rows"][0]["record_id"]
        saved=self.remote.get(record_id)["fields"]
        self.assertEqual(saved["失败原因"],"现场核验未通过")
        self.assertEqual(len(saved["上下电确认截图"]),1)
        self.assertEqual(self.service.evidence_path("A",record_id,image_id)[0].read_bytes(),output.getvalue())
        thumbnail=self.service.evidence_path("A",record_id,image_id,True)[0]
        self.assertTrue(thumbnail.is_file())
        self.assertLess(thumbnail.stat().st_size,len(output.getvalue())+1024)
        local_image=self.service.evidence_path("A",record_id,image_id)[0]
        local_image.unlink()
        self.assertEqual(self.service.evidence_path("A",record_id,image_id)[0].read_bytes(),output.getvalue())
        self.assertEqual(self.service.batches.storage_status()["cloud_backed_files"],1)
        cleaned=self.service.batches.cleanup_evidence_cache()
        self.assertGreaterEqual(cleaned["deleted"],1)
        self.assertEqual(self.service.evidence_path("A",record_id,image_id)[0].read_bytes(),output.getvalue())
        self.service.batches.rollback(batch["batch_id"],{"version":done["version"],"all":True},"owner",["A"])
        undone=self._wait_batch(batch["batch_id"])
        self.assertEqual(undone["stats"]["rolled_back"],1,undone["rows"])
        self.service.batches.confirm(batch["batch_id"],{"version":undone["version"],"all":True},"owner",["A"])
        again=self._wait_batch(batch["batch_id"])
        self.assertEqual(again["stats"]["completed"],1,again["rows"])
        self.assertEqual(len(self.remote.attachments),1)
        self.assertEqual(self.remote.get(again["rows"][0]["record_id"])["fields"]["上下电确认截图"],saved["上下电确认截图"])

    def test_cross_scope_image_is_hidden_until_viewer_has_every_scope(self):
        from PIL import Image
        rows=[self._manual_batch_row(scope,"2026-09-14 01:02:03") for scope in ("A","B")]
        batch=self.service.batches.create_manual(rows,"owner")
        output=io.BytesIO();Image.new("RGB",(80,60),"white").save(output,format="PNG")
        candidates=[{"scope":row["scope"],"room":row["room"],"rack":row["rack"],"supplier_rack":"",
                     "action":row["action"],"expected":row["expected"],"actual":row["actual"]} for row in rows]
        with patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",return_value=candidates):
            self.service.batches.add_images(batch["batch_id"],[("mixed.png",output.getvalue())],"owner",["A","B"])
            deadline=time.time()+5
            while time.time()<deadline and self.service.batches.get(batch["batch_id"])["images"][0]["status"]=="recognizing":time.sleep(.01)
        batch=self.service.batches.get(batch["batch_id"]);image_id=batch["images"][0]["image_id"]
        limited=self.service.batches.visible(batch,"viewer",["A"])
        self.assertEqual(([row["scope"] for row in limited["rows"]],limited["images"],limited["files"]),(["A"],[],[]))
        with self.assertRaisesRegex(CabinetError,"无权查看截图"):
            self.service.batches.image_path(batch["batch_id"],image_id,"viewer",["A"])
        self.assertTrue(self.service.batches.image_path(batch["batch_id"],image_id,"viewer",["A","B"])[0].is_file())

    def test_batch_patch_delta_returns_only_changed_rows(self):
        rows=[self._manual_batch_row("A",f"2026-09-14 01:0{index}:03") for index in (2,3)]
        batch=self.service.batches.create_manual(rows,"owner")
        target=batch["rows"][0]
        changed=self.service.batches.update(batch["batch_id"],{
            "version":batch["version"],"response_mode":"delta",
            "rows":[{"row_id":target["row_id"],"expected":"2026-09-14 01:01:03"}],
        },"owner",["A"])
        self.assertTrue(changed["partial_rows"])
        self.assertEqual([row["row_id"] for row in changed["rows"]],[target["row_id"]])
        self.assertEqual(changed["rows"][0]["expected"],"2026-09-14 01:01:03")

    def test_existing_cloud_attachment_is_preserved_on_group_update(self):
        rack=self.configs["D"]["inventory"][0]
        existing=next(op for op in self.service._snapshot("D")["operations"] if (op["room"],op["rack"])==(rack["room"],rack["rack"]))
        operation=copy.deepcopy(existing)
        operation["raw_fields"]["上下电确认截图"]=[{"file_token":"fileOldEvidence123"}]
        operation["groups"].append({"id":"event_new","action":"上正式电","expected":"2026-09-14 01:02:03","actual":"2026-09-14 01:02:03","result":"成功","evidence_images":[{"image_id":"a"*64,"file_token":"fileNewEvidence123","extension":".png"}]})
        fields=to_fields(operation)
        self.assertEqual({item["file_token"] for item in fields["上下电确认截图"]},{"fileOldEvidence123","fileNewEvidence123"})

    def test_one_image_links_all_cabinets_and_overwrites_recognized_times(self):
        from PIL import Image
        racks=self.configs["A"]["inventory"][:2]
        rows=[{"scope":"A","room":rack["room"],"rack":rack["rack"],"rack_type":rack["rack_type"],
               "action":"上正式电","expected":"","actual":"","result":"成功"} for rack in racks]
        rows[1]["expected"]="2026-09-13 01:00:00"
        batch=self.service.batches.create_manual(rows,"owner")
        image=Image.new("RGB",(160,80),"white"); output=io.BytesIO(); image.save(output,format="PNG")
        candidates=[{"scope":"A","room":rack["room"],"rack":rack["rack"],"supplier_rack":"",
                     "action":"上正式电","expected":"2026-09-14 01:02:03","actual":"2026-09-14 01:03:04"} for rack in racks]
        with patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",return_value=candidates):
            self.service.batches.add_images(batch["batch_id"],[("both.png",output.getvalue())],"owner",["A"])
            deadline=time.time()+5
            while time.time()<deadline and self.service.batches.get(batch["batch_id"])["images"][0]["status"]=="recognizing": time.sleep(.01)
        result=self.service.batches.get(batch["batch_id"])
        self.assertEqual(result["rows"][0]["expected"],candidates[0]["expected"])
        self.assertEqual(result["rows"][0]["actual"],candidates[0]["actual"])
        self.assertEqual(result["rows"][0]["evidence_images"],[result["images"][0]["image_id"]])
        self.assertEqual(result["rows"][1]["expected"],candidates[1]["expected"])
        self.assertEqual(result["rows"][1]["actual"],candidates[1]["actual"])
        self.assertEqual(result["rows"][1]["evidence_images"],[result["images"][0]["image_id"]])
        self.assertEqual(result["images"][0]["suggestions"][1]["status"],"applied")
        self.assertTrue(any(edit.get("before")=="2026-09-13 01:00:00" and edit.get("image_id")
                            for edit in result["rows"][1]["edits"]))
        self.assertEqual(len(result["rows"]),2)

    def test_pending_image_recognition_resumes_after_restart(self):
        batch=self.service.batches.create_manual([self._manual_batch_row("A","2026-09-14 01:02:03")],"owner")
        image_id="a"*64
        self.service.batches._change(batch["batch_id"],lambda current:current.setdefault("images",[]).append(
            {"image_id":image_id,"name":"proof.png","extension":".png","status":"recognizing","suggestions":[]}
        ))
        pending=self.service.batches.get(batch["batch_id"])
        with self.assertRaisesRegex(CabinetError,"截图仍在识别"):
            self.service.batches.confirm(batch["batch_id"],{"version":pending["version"],"all":True},"owner",["A"])
        self.service.batches.shutdown(wait=True)
        with patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",return_value=[]):
            self.service._batches=CabinetBatchService(self.service,self.service.root)
            deadline=time.time()+5
            while time.time()<deadline and self.service.batches.get(batch["batch_id"])["images"][0]["status"]=="recognizing": time.sleep(.01)
        image=self.service.batches.get(batch["batch_id"])["images"][0]
        self.assertEqual(image["status"],"failed")

    def test_uncertain_batch_create_never_falls_back_to_another_create_token(self):
        service=self.service.batches
        racks=[rack for rack in self.service._snapshot('A')['config']['inventory'] if rack['room']=='203'][:2]
        rows=[dict(scope='A',room=rack['room'],rack=rack['rack'],rack_type=rack['rack_type'],action='上正式电',
                   expected='2026-09-20 10:00:00',actual='2026-09-20 10:00:00',result='成功') for rack in racks]
        batch=service.create_manual(rows,'owner')
        self.remote.fail_after_batch_create=True
        original_list=self.remote.list_all
        def delayed_search(path='records',filters=None):
            return [row for row in original_list(path,filters) if not row['fields'].get('数据标识','').startswith('manual_batch_')]
        with patch.object(self.remote,'list_all',side_effect=delayed_search):
            service.confirm(batch['batch_id'],dict(version=batch['version'],all=True),'owner',['A'])
            failed=self._wait_batch(batch['batch_id'])
            self.assertEqual(failed['stats']['completed'],0)
            for row in failed['rows']:
                self.assertTrue(row['operation_started'])
                with self.assertRaisesRegex(CabinetError,'结果仍未确认'):
                    self.service.reconcile_write('A',row['operation_id'],'owner')
            service.confirm(batch['batch_id'],dict(version=failed['version'],all=True),'owner',['A'])
            failed=self._wait_batch(batch['batch_id'])
        counts=Counter(row['fields'].get('数据标识') for row in self.remote.records.values()
                       if row['fields'].get('数据标识','').startswith('manual_batch_'))
        self.assertEqual(sorted(counts.values()),[1,1])
        service.confirm(batch['batch_id'],dict(version=failed['version'],all=True),'owner',['A'])
        self.assertEqual(self._wait_batch(batch['batch_id'])['stats']['completed'],2)
        self.assertEqual(self.remote.batch_create_calls,1)

    def test_preparation_failure_leaves_row_editable_without_a_write_journal(self):
        service=self.service.batches
        batch=service.create_manual([self._manual_batch_row('A','2026-09-20 10:00:00')],'owner')
        with patch.object(service,'_row_payload',side_effect=CabinetError('证明文件已丢失',409)):
            service.confirm(batch['batch_id'],dict(version=batch['version'],all=True),'owner',['A'])
            failed=self._wait_batch(batch['batch_id'])
        row=failed['rows'][0]
        self.assertFalse(row['operation_started'])
        self.assertIsNone(self.service.local.document('A','write:'+row['operation_id']))
        updated=service.update(batch['batch_id'],{'version':failed['version'],'rows':[{'row_id':row['row_id'],'type_detail':'更正'}]},'owner',['A'])
        self.assertEqual(updated['rows'][0]['type_detail'],'更正')

    def test_proof_business_conflict_requires_explicit_review_and_rechecks_changes(self):
        service=self.service.batches
        batch=service.create_manual([self._manual_batch_row('A','2026-09-20 10:00:00')],'owner')
        row=batch['rows'][0]
        def proof(current):
            current['rows'][0]['evidence_images']=['proof']
            current['images']=[{'image_id':'proof','name':'proof.png','status':'done','suggestions':[
                {**row,'row_id':row['row_id'],'action':'上测试电' if row['action']!='上测试电' else '上正式电','result':'失败'}]}]
        batch=service._change(batch['batch_id'],proof,validate=True)
        self.assertEqual(batch['stats']['confirmable'],0)
        self.assertTrue(batch['rows'][0]['evidence_business_conflict'])
        reviewed=service.apply_image(batch['batch_id'],'proof',{'version':batch['version'],'row_id':row['row_id'],'fields':{},'review_business':True},'owner',['A'])
        self.assertEqual(reviewed['stats']['confirmable'],1)
        changed=service.update(batch['batch_id'],{'version':reviewed['version'],'rows':[{'row_id':row['row_id'],'action':'下正式电'}]},'owner',['A'])
        self.assertTrue(changed['rows'][0]['evidence_business_conflict'])

    def test_image_correction_validates_directory_duplicates_and_preserves_proof(self):
        service=self.service.batches
        batch=service.create_image_batch('owner',['A'],'A')
        batch=service._change(batch['batch_id'],lambda current:current.update(images=[
            {'image_id':'proof','name':'proof.png','status':'failed','suggestions':[]}]))
        fields=self._manual_batch_row('A','2026-09-20 10:00:00')
        corrected=service.correct_image(batch['batch_id'],'proof',{'version':batch['version'],'fields':fields},'owner',['A'])
        self.assertEqual(corrected['rows'][0]['evidence_images'],['proof'])
        self.assertEqual(corrected['rows'][0]['edits'][0]['field'],'image_correction')
        with self.assertRaisesRegex(CabinetError,'已在本批'):
            service.correct_image(batch['batch_id'],'proof',{'version':corrected['version'],'fields':fields},'owner',['A'])
        with self.assertRaisesRegex(CabinetError,'不在当前楼栋目录'):
            service.correct_image(batch['batch_id'],'proof',{'version':corrected['version'],'fields':{**fields,'rack':'Z99'}},'owner',['A'])
        mixed=service._change(batch['batch_id'],lambda current:current['images'][0]['suggestions'].append({'scope':'B'}))
        with self.assertRaisesRegex(CabinetError,'无权使用其他楼栋截图'):
            service.correct_image(batch['batch_id'],'proof',{'version':mixed['version'],'fields':fields},'other',['A'])

    def test_image_retry_preserves_corrections_and_delete_restore_invalidates_old_worker(self):
        service=self.service.batches
        batch=service.create_manual([self._manual_batch_row('A','2026-09-20 10:00:00')],'owner')
        candidate={**batch['rows'][0],'actual':'2026-09-19 10:00:00'}
        batch=self._add_proof_images(batch,[[candidate]])
        row=batch['rows'][0]; image=batch['images'][0]
        batch=service.update(batch['batch_id'],{'version':batch['version'],'rows':[{'row_id':row['row_id'],'actual':'2026-09-20 10:01:02'}]},'owner',['A'])
        with patch('bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout',return_value=[candidate]):
            service.retry_image(batch['batch_id'],image['image_id'],{'version':batch['version']},'owner',['A'])
            deadline=time.time()+5
            while time.time()<deadline:
                batch=service.get(batch['batch_id'])
                if batch['images'][0]['status']!='recognizing': break
                time.sleep(.02)
        self.assertEqual(batch['rows'][0]['actual'],'2026-09-20 10:01:02')
        entered=threading.Event(); release=threading.Event(); calls=[]
        def slow_ocr(_content):
            calls.append(1)
            if len(calls)==1: entered.set();release.wait(5)
            return [candidate]
        with patch('bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout',side_effect=slow_ocr):
            try:
                service.retry_image(batch['batch_id'],image['image_id'],{'version':batch['version']},'owner',['A'])
                self.assertTrue(entered.wait(5))
                batch=service.get(batch['batch_id'])
                removed=service.delete_image(batch['batch_id'],image['image_id'],batch['version'],'owner',['A'])
                service.restore_image(batch['batch_id'],image['image_id'],removed['version'],'owner',['A'])
            finally: release.set()
            deadline=time.time()+5
            while time.time()<deadline:
                batch=service.get(batch['batch_id'])
                if batch['images'][0]['status']!='recognizing': break
                time.sleep(.02)
        self.assertEqual(batch['images'][0]['status'],'done')
        self.assertEqual(len(calls),2)
        self.assertEqual(batch['rows'][0]['evidence_images'],[image['image_id']])

    def test_single_export_is_local_and_downloadable_without_cloud_write(self):
        from concurrent.futures import ThreadPoolExecutor
        self.service._exports=ThreadPoolExecutor(max_workers=1)
        archive=FakeExportFeishu();self.service.export_remote=archive;self.service._export_schema_ready=False
        with patch('bin.lan_bitable_template_portal.cabinet_power.export_snapshot',return_value=b'test-export'):
            job=self.service.job('A','export','owner',{})
            deadline=time.time()+5
            while time.time()<deadline:
                status=self.service.job_status(job['job_id'],'A')
                if status['status'] in ('succeeded','failed'): break
                time.sleep(.02)
        self.assertEqual(status['status'],'succeeded',status.get('error'))
        export=self.service.local.document('A','export:'+status['result']['export_id'])
        self.assertEqual(Path(export['path']).read_bytes(),b'test-export')
        self.assertEqual(export['cloud_upload_status'],'local_only')
        self.assertEqual((archive.upload_calls,archive.creates),(0,0))
        with self.assertRaisesRegex(CabinetError,'单楼导出仅保存在本机'):
            self.service.upload_export('A',export['export_id'],'owner')
        self.assertEqual(status['result']['cloud_upload_status'],'local_only')

    def test_recovery_pages_old_unfinished_images_and_skips_live_worker(self):
        service=self.service.batches
        first=service.create_manual([self._manual_batch_row('A','2026-09-20 10:00:00')],'owner')
        service._change(first['batch_id'],lambda batch:batch.update(images=[
            {'image_id':'old','status':'recognizing','worker':None,'extension':'.png','name':'old.png'}]))
        with service.store._connect() as conn,conn:
            for index in range(1005):
                conn.execute("INSERT INTO batches VALUES(?,?,?,?,?,?,?,?,?)",(str(index),'owner','completed',None,'[]','{}',1,'2026-09-21','2026-09-21'))
        with patch.object(service,'_queue_image') as queue:
            service._recover_interrupted()
            self.assertEqual(queue.call_args.args[0],first['batch_id'])
        live=service._change(first['batch_id'],lambda batch:(batch['rows'][0].update(status='queued'),batch.update(worker=service._worker)))
        second=CabinetBatchService(self.service,self.service.root)
        try:
            self.assertEqual(second.store.get(live['batch_id'])['rows'][0]['status'],'queued')
        finally: second.shutdown(wait=True)

    def test_list_sql_pagination_filters_permissions_without_decoding_history(self):
        service=self.service.batches
        batch=service.create_manual([self._manual_batch_row('A','2026-09-20 10:00:00'),self._manual_batch_row('B','2026-09-20 10:00:00')],'owner')
        with patch.object(service.store,'runtime_list',side_effect=AssertionError('unbounded scan')):
            own=service.list('owner',[],scope='A',status='todo',page=99,page_size=1)
            other=service.list('other',['B'],status='todo')
            hidden=service.list('other',['E'],status='todo')
        self.assertEqual((own['total'],own['page'],own['pending_count']),(1,1,1))
        self.assertEqual(other['items'][0]['stats']['total'],1)
        self.assertEqual(other['items'][0]['scopes'],['B'])
        self.assertTrue(all(room.startswith('B楼') for room in other['items'][0]['rooms']))
        self.assertEqual(hidden['total'],0)

    def test_inventory_only_includes_new_directory_cabinets_without_operations(self):
        snap=self.service._snapshot('A')
        keys={(row['room'],row['rack']) for row in snap['config']['inventory']}
        operated={(row['room'],row['rack']) for row in snap['operations'] if row.get('events') and not row.get('meta',{}).get('baseline_correction')}
        self.assertEqual(self.service.overview('A')['inventory_only'],len(keys-operated))
        self.assertGreaterEqual(self.service.overview('A')['inventory_only'],84)

    def _add_proof_images(self, batch, candidates):
        from PIL import Image
        files=[]
        for index in range(len(candidates)):
            content=io.BytesIO()
            Image.new("RGB",(80,60),(40+index*30,80,100)).save(content,format="PNG")
            files.append((f"proof-{index}.png",content.getvalue()))
        with patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",
                   side_effect=candidates):
            self.service.batches.add_images(batch["batch_id"],files,"owner",list("ABCDE"),True)
            deadline=time.time()+10
            while time.time()<deadline:
                result=self.service.batches.get(batch["batch_id"])
                if all(image["status"]!="recognizing" for image in result["images"]): return result
                time.sleep(.01)
        self.fail("proof recognition timed out")

    def test_proof_linking_is_independent_of_missing_time_and_action_difference(self):
        for source in ("manual","notice","pdf"):
            batch=self.service.batches.create_manual([self._manual_batch_row("A","2026-09-14 01:02:03")],"owner")
            batch=self.service.batches._change(batch["batch_id"],lambda current:current.update(source=source))
            row=batch["rows"][0]
            candidate={key:row[key] for key in ("scope","room","rack")}
            candidate.update(action="下测试电",expected="",actual="2026-09-15 02:03:04",supplier_rack="")
            result=self._add_proof_images(batch,[[candidate]])
            self.assertEqual(result["images"][0]["status"],"done",(source,result["images"]))
            saved=result["rows"][0]
            self.assertEqual(saved["action"],"上正式电")
            self.assertEqual(saved["expected"],"2026-09-14 01:02:03")
            self.assertEqual(saved["actual"],"2026-09-15 02:03:04")
            self.assertEqual(saved["evidence_images"],[result["images"][0]["image_id"]])
            visible=self.service.batches.visible(result,"owner",["A"])
            self.assertEqual(visible["rows"][0]["evidence_images"],saved["evidence_images"])

    def test_proof_time_conflict_blocks_confirm_until_review_or_removal(self):
        service=self.service.batches
        batch=service.create_manual([self._manual_batch_row("A","2026-09-13 01:02:03")],"owner")
        row=batch["rows"][0]
        base={key:row[key] for key in ("scope","room","rack","action")}
        candidates=[{**base,"supplier_rack":"","expected":date,"actual":date}
                    for date in ("2026-09-14 01:02:03","2026-09-15 01:02:03")]
        result=self._add_proof_images(batch,[[item] for item in candidates])
        self.assertEqual((len(result["rows"]),result["stats"]["conflict"],result["stats"]["confirmable"]),(1,1,0))
        self.assertEqual(len(result["rows"][0]["evidence_images"]),2)
        with self.assertRaisesRegex(CabinetError,"没有可确认"):
            service.confirm(result["batch_id"],{"version":result["version"],"all":True},"owner",["A"])
        result=service.get(result["batch_id"])
        first,second=result["images"]
        removed=service.delete_image(result["batch_id"],first["image_id"],result["version"],"owner",["A"])
        self.assertFalse(removed["rows"][0]["evidence_time_conflict"])
        self.assertEqual(removed["rows"][0]["actual"],candidates[1]["actual"])
        restored=service.restore_image(result["batch_id"],first["image_id"],removed["version"],"owner",["A"])
        self.assertTrue(restored["rows"][0]["evidence_time_conflict"])
        reviewed=service.apply_image(result["batch_id"],first["image_id"],{
            "version":restored["version"],"row_id":row["row_id"],"candidate_index":0,
            "fields":{key:candidates[0][key] for key in ("expected","actual")}},"owner",["A"])
        self.assertFalse(reviewed["rows"][0]["evidence_time_conflict"])
        self.assertEqual(reviewed["stats"]["confirmable"],1)
        self.assertEqual(reviewed["rows"][0]["actual"],candidates[0]["actual"])
        self.assertEqual(service.get(result["batch_id"])["rows"][0]["evidence_time_review"]["owner"],"owner")
        delta=service.update(result["batch_id"],{"version":reviewed["version"],"response_mode":"delta",
            "rows":[{"row_id":row["row_id"],"actual":"2026-09-16 01:02:03"}]},"owner",["A"])
        self.assertTrue(delta["rows"][0]["evidence_time_conflict"])
        self.assertTrue(delta["images"])
        checked=service.apply_image(result["batch_id"],first["image_id"],{"version":delta["version"],
            "row_id":row["row_id"],"fields":{},"review_times":True},"owner",["A"])
        self.assertFalse(checked["rows"][0]["evidence_time_conflict"])
        self.assertEqual(checked["rows"][0]["actual"],"2026-09-16 01:02:03")

    def test_recognition_reads_without_triggering_notice_default_migration(self):
        service=self.service.batches
        batch=service.create_manual([self._manual_batch_row("A","2026-09-14 01:02:03")],"owner")
        image={"image_id":"a"*64,"extension":".png","name":"proof.png","status":"recognizing","suggestions":[]}
        service._atomic_write(self.service.root/"evidence"/("a"*64+".png"),b"test")
        service._change(batch["batch_id"],lambda current:current.update(source="notice",images=[image]))
        candidate={**batch["rows"][0],"actual":"2026-09-15 02:03:04"}
        with patch.object(service,"get",side_effect=AssertionError("OCR reads must not migrate notice defaults")), \
                patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",return_value=[candidate]):
            service._recognize_image(batch["batch_id"],image)
        saved=service.store.get(batch["batch_id"])
        self.assertEqual(saved["images"][0]["status"],"done",saved["images"])
        self.assertEqual(saved["rows"][0]["actual"],candidate["actual"])

    def test_proof_recognition_preserves_locked_rows_and_ambiguous_matches(self):
        rows=[self._manual_batch_row("A","2026-09-14 01:02:03") for _ in range(2)]
        batch=self.service.batches.create_manual(rows,"owner")
        candidate={key:rows[0][key] for key in ("scope","room","rack","action","expected","actual")}
        candidate["actual"]="2026-09-15 01:02:03"
        result=self._add_proof_images(batch,[[candidate]])
        self.assertFalse(any(row.get("evidence_images") for row in result["rows"]))
        self.assertEqual(result["images"][0]["suggestions"][0]["status"],"needs_review")
        for status in ("completed","queued","excluded_cancelled"):
            batch=self.service.batches.create_manual([rows[0]],"owner")
            stored=self.service.batches._change(batch["batch_id"],lambda current:current["rows"][0].update(status=status))
            image={"image_id":"a"*64,"extension":".png","name":"locked.png","status":"recognizing","suggestions":[]}
            self.service.batches._change(batch["batch_id"],lambda current:current.update(images=[image]))
            path=self.service.root/"evidence"/("a"*64+".png");path.parent.mkdir(exist_ok=True);path.write_bytes(b"test")
            with patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",return_value=[candidate]):
                self.service.batches._recognize_image(batch["batch_id"],image)
            result=self.service.batches.get(batch["batch_id"])
            self.assertEqual(result["rows"][0]["actual"],stored["rows"][0]["actual"])
            self.assertFalse(result["rows"][0].get("evidence_images"))

    def _pdf_proof_batch(self, scopes, date="2026-09-14 01:02:03", files_per_row=None):
        from pypdf import PdfWriter
        rows=[]
        for scope in scopes:
            row=self._manual_batch_row(scope,date)
            state=self.service.batches._current_state(self.service._snapshot(scope),row["room"],row["rack"])
            row["action"]={"formal":"下正式电","test":"下测试电"}.get(state,"上正式电")
            rows.append(row)
        batch=self.service.batches.create_manual(rows,"owner")
        indexes=files_per_row or [0]*len(rows); files=[]
        for index in sorted(set(indexes)):
            writer=PdfWriter();writer.add_blank_page(width=200+index,height=200)
            writer.add_metadata({"/Title":date})
            content=io.BytesIO();writer.write(content);data=content.getvalue()
            meta={"file_id":f"f{index}","name":f"confirmation-{index}.pdf","sha256":hashlib.sha256(data).hexdigest(),
                  "size":len(data),"status":"completed"}
            self.service.batches._atomic_write(self.service.batches.import_root/batch["batch_id"]/(meta["file_id"]+".pdf"),data)
            files.append(meta)
        def source(current):
            current.update(source="pdf",files=files)
            for row,index in zip(current["rows"],indexes):
                row.update(file_id=f"f{index}",file_name=f"confirmation-{index}.pdf")
        return self.service.batches._change(batch["batch_id"],source,validate=True)

    def test_pdf_proofs_upload_once_and_follow_rows_across_buildings(self):
        service=self.service.batches
        batch=self._pdf_proof_batch(["A","D","E"],files_per_row=[0,0,1])
        visible=service.visible(batch,"owner",list("ADE"))
        self.assertEqual([row["proof_files"][0]["name"] for row in visible["rows"]],
                         ["confirmation-0.pdf","confirmation-0.pdf","confirmation-1.pdf"])
        row=batch["rows"][0]
        delta=service.update(batch["batch_id"],{"version":batch["version"],"response_mode":"delta",
            "rows":[{"row_id":row["row_id"],"expected":"2026-09-14 02:00:00"}]},"owner",list("ADE"))
        self.assertEqual(service.visible(delta,"owner",list("ADE"))["rows"][0]["proof_files"][0]["file_id"],"f0")
        batch=service.get(batch["batch_id"])
        service.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",list("ADE"))
        done=self._wait_batch(batch["batch_id"])
        self.assertEqual(done["stats"]["completed"],3,done["rows"])
        self.assertEqual(len(self.remote.attachments),2)
        for row in done["rows"]:
            saved=from_feishu(self.remote.get(row["record_id"]))
            document=next(group["evidence_files"][0] for group in saved["groups"] if group.get("evidence_files"))
            expected=next(file for file in done["files"] if file["file_id"]==row["file_id"])
            self.assertEqual(document["file_token"],expected["cloud_file_token"])
            self.assertIn({"file_token":document["file_token"]},saved["raw_fields"]["上下电确认截图"])
            self.assertTrue(any(event.get("evidence_files") for event in saved["events"]))
            path,name=self.service.document_path(row["scope"],row["record_id"],document["file_id"],list("ADE"))
            self.assertEqual(name,expected["name"])
            path.unlink()
            self.assertEqual(hashlib.sha256(self.service.document_path(row["scope"],row["record_id"],document["file_id"],list("ADE"))[0].read_bytes()).hexdigest(),expected["sha256"])
            if row["scope"]=="A":
                with self.assertRaisesRegex(CabinetError,"其他楼栋"):
                    self.service.document_path("A",row["record_id"],document["file_id"],["A"])
        service.shutdown(wait=True)
        restored=CabinetBatchService(self.service,self.service.root)
        self.service._batches=restored
        self.assertTrue(all(file.get("cloud_file_token") for file in restored.get(batch["batch_id"])["files"]))

    def test_pdf_proofs_survive_edit_and_rollback_preserves_previous_attachment(self):
        service=self.service.batches
        first=self._pdf_proof_batch(["D"])
        service.confirm(first["batch_id"],{"version":first["version"],"all":True},"owner",["D"])
        first=self._wait_batch(first["batch_id"]); rid=first["rows"][0]["record_id"]
        previous=copy.deepcopy(self.remote.get(rid)["fields"]["上下电确认截图"])
        second=self._pdf_proof_batch(["D"],"2026-09-15 01:02:03")
        service.confirm(second["batch_id"],{"version":second["version"],"all":True},"owner",["D"])
        done=self._wait_batch(second["batch_id"])
        self.assertEqual(done["stats"]["completed"],1,done["rows"])
        operation=from_feishu(self.remote.get(rid))
        self.assertEqual(len(operation["raw_fields"]["上下电确认截图"]),2)
        groups=copy.deepcopy(operation["groups"])
        for group in groups: group.pop("evidence_files",None)
        edited=self.service.validate_op("D",{"groups":groups},operation)
        self.assertEqual(to_fields(edited)["上下电确认截图"],operation["raw_fields"]["上下电确认截图"])
        self.assertEqual([group.get("evidence_files") for group in edited["groups"]],
                         [group.get("evidence_files") for group in operation["groups"]])
        book=Workbook(export_workbook((TEMPLATES/"D.xlsm").read_bytes(),self.service.config("D"),self.service._snapshot("D")["operations"]))
        self.assertFalse(any(".pdf" in name for name in book.archive.namelist()))
        self.assertFalse(any("confirmation-0.pdf" in str(value) for name in book.sheets for _,row in book.rows(name) for value in row.values()))
        service.rollback(done["batch_id"],{"version":done["version"],"all":True},"owner",["D"])
        undone=self._wait_batch(done["batch_id"])
        self.assertEqual(undone["stats"]["rolled_back"],1,undone["rows"])
        self.assertEqual(self.remote.get(rid)["fields"]["上下电确认截图"],previous)
        service.confirm(undone["batch_id"],{"version":undone["version"],"all":True},"owner",["D"])
        repeated=self._wait_batch(done["batch_id"])
        self.assertEqual(repeated["stats"]["completed"],1,repeated["rows"])
        self.assertEqual(len(self.remote.attachments),2)

    def test_identical_pdf_bytes_reuse_upload_across_source_file_ids(self):
        service=self.service.batches
        batch=self._pdf_proof_batch(["A","E"],files_per_row=[0,1])
        content=(service.import_root/batch["batch_id"]/"f0.pdf").read_bytes()
        service._atomic_write(service.import_root/batch["batch_id"]/"f1.pdf",content)
        service._change(batch["batch_id"],lambda current:current["files"][1].update(sha256=current["files"][0]["sha256"]))
        first=service._upload_proof(batch["batch_id"],"files","f0","A")
        second=service._upload_proof(batch["batch_id"],"files","f1","E")
        self.assertEqual(first,second)
        self.assertEqual(len(self.remote.attachments),1)
        self.assertTrue(all(item.get("cloud_file_token")==first for item in service.get(batch["batch_id"])["files"]))

    def test_pdf_upload_failure_retries_without_creating_duplicate_records(self):
        service=self.service.batches
        batch=self._pdf_proof_batch(["A"])
        operation_id=batch["rows"][0]["operation_id"]
        with patch.object(self.remote,"upload_attachment",side_effect=TimeoutError("upload failed")):
            service.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
            failed=self._wait_batch(batch["batch_id"])
        self.assertEqual(failed["stats"]["failed"],1)
        self.assertEqual(failed["stats"]["completed"],0)
        self.assertEqual(self.remote.creates,0)
        self.remote.fail_after_create=True
        service.confirm(failed["batch_id"],{"version":failed["version"],"all":True},"owner",["A"])
        unconfirmed=self._wait_batch(batch["batch_id"])
        self.assertEqual(unconfirmed["stats"]["failed"],1)
        self.assertEqual(self.remote.creates,1)
        self.remote.fail_after_create=False
        service.confirm(unconfirmed["batch_id"],{"version":unconfirmed["version"],"all":True},"owner",["A"])
        done=self._wait_batch(batch["batch_id"])
        self.assertEqual(done["stats"]["completed"],1,done["rows"])
        self.assertEqual(done["rows"][0]["operation_id"],operation_id)
        self.assertEqual((self.remote.creates,len(self.remote.attachments)),(1,1))

    def test_de_history_keeps_evidence_group_and_rollback_removes_new_attachment(self):
        from PIL import Image
        rack=self.configs["D"]["inventory"][0]
        original=copy.deepcopy(next(op for op in self.service._snapshot("D")["operations"] if (op["room"],op["rack"])==(rack["room"],rack["rack"])))
        row=self._manual_batch_row("D","2026-09-14 01:02:03")
        batch=self.service.batches.create_manual([row],"owner")
        image=Image.new("RGB",(80,60),"white"); output=io.BytesIO();image.save(output,format="PNG")
        candidate={"scope":"D","room":rack["room"],"rack":rack["rack"],"supplier_rack":"",
                   "action":"上正式电","expected":row["expected"],"actual":row["actual"]}
        with patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",return_value=[candidate]):
            self.service.batches.add_images(batch["batch_id"],[("proof.png",output.getvalue())],"owner",["D"])
            deadline=time.time()+5
            while time.time()<deadline and self.service.batches.get(batch["batch_id"])["images"][0]["status"]=="recognizing":time.sleep(.01)
        batch=self.service.batches.get(batch["batch_id"])
        self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["D"])
        done=self._wait_batch(batch["batch_id"])
        self.assertEqual(done["stats"]["completed"],1,done["rows"])
        saved=next(op for op in self.service._snapshot("D")["operations"] if op["record_id"]==original["record_id"])
        self.assertEqual(len(saved["groups"]),len(original["groups"])+1)
        self.assertTrue(any(group.get("evidence_images") for group in saved["groups"]))
        self.service.batches.rollback(batch["batch_id"],{"version":done["version"],"all":True},"owner",["D"])
        undone=self._wait_batch(batch["batch_id"])
        self.assertEqual(undone["stats"]["rolled_back"],1,undone["rows"])
        restored=next(op for op in self.service._snapshot("D")["operations"] if op["record_id"]==original["record_id"])
        self.assertEqual(restored["groups"],original["groups"])
        self.assertFalse(self.remote.get(original["record_id"])["fields"].get("上下电确认截图"))

    def test_pasted_image_can_be_removed_before_write_or_after_rollback(self):
        from PIL import Image
        rack=self.configs["A"]["inventory"][0]
        row=self._manual_batch_row("A","2026-09-14 01:02:03")
        batch=self.service.batches.create_manual([row],"owner")
        output=io.BytesIO();Image.new("RGB",(80,60),"white").save(output,format="PNG")
        candidate={"scope":"A","room":rack["room"],"rack":rack["rack"],"supplier_rack":"",
                   "action":row["action"],"expected":row["expected"],"actual":row["actual"]}
        with patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",return_value=[candidate]):
            for attempt in range(2):
                self.service.batches.add_images(batch["batch_id"],[("pasted.png",output.getvalue())],"owner",["A"])
                deadline=time.time()+5
                while time.time()<deadline and self.service.batches.get(batch["batch_id"])["images"][0]["status"]=="recognizing":time.sleep(.01)
                batch=self.service.batches.get(batch["batch_id"])
                image_id=batch["images"][0]["image_id"]
                self.assertEqual(batch["rows"][0]["evidence_images"],[image_id])
                if attempt==0:
                    with self.assertRaises(CabinetError):
                        self.service.batches.delete_image(batch["batch_id"],image_id,batch["version"],"other",["A"])
                    with self.assertRaises(CabinetError):
                        self.service.batches.delete_image(batch["batch_id"],image_id,batch["version"]-1,"owner",["A"])
                    removed=self.service.batches.delete_image(batch["batch_id"],image_id,batch["version"],"owner",["A"])
                    self.assertTrue(removed["images"][0]["deleted_at"])
                    self.assertEqual(removed["rows"][0]["evidence_images"],[])
                    self.assertTrue((self.service.root/"evidence"/(image_id+".png")).is_file())
                    restored=self.service.batches.restore_image(batch["batch_id"],image_id,removed["version"],"owner",["A"])
                    self.assertFalse(restored["images"][0]["deleted_at"])
                    self.assertEqual(restored["rows"][0]["evidence_images"],[image_id])
        self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
        done=self._wait_batch(batch["batch_id"])
        self.assertEqual(done["stats"]["completed"],1,done["rows"])
        with self.assertRaisesRegex(CabinetError,"先回退"):
            self.service.batches.delete_image(batch["batch_id"],image_id,done["version"],"owner",["A"])
        self.service.batches.rollback(batch["batch_id"],{"version":done["version"],"all":True},"owner",["A"])
        undone=self._wait_batch(batch["batch_id"])
        removed=self.service.batches.delete_image(batch["batch_id"],image_id,undone["version"],"owner",["A"])
        self.assertEqual(removed["rows"][0]["evidence_images"],[])
        restored=self.service.batches.restore_image(batch["batch_id"],image_id,removed["version"],"owner",["A"])
        self.assertEqual(restored["rows"][0]["evidence_images"],[image_id])

    def test_image_registration_creates_rows_then_uses_existing_confirm_flow(self):
        from PIL import Image
        racks=self.configs["A"]["inventory"][:2]
        batch=self.service.batches.create_image_batch("owner",["A"])
        self.assertEqual((batch["source"],batch["rows"]),("image",[]))
        content=io.BytesIO();Image.new("RGB",(160,80),"white").save(content,format="PNG")
        states={(item["room"],item["rack"]):item["state"] for item in self.service.overview("A")["racks"]}
        candidates=[{"scope":"A","room":rack["room"],"rack":rack["rack"],"supplier_rack":"",
                     "action":{"formal":"下正式电","test":"下测试电"}.get(states.get((rack["room"],rack["rack"])),"上正式电"),"expected":"2026-09-14 01:02:03",
                     "actual":"2026-09-14 01:02:03","result":""} for rack in racks]
        with patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",return_value=candidates):
            self.service.batches.add_images(batch["batch_id"],[("mail.png",content.getvalue())],"owner",["A"])
            deadline=time.time()+5
            while time.time()<deadline and self.service.batches.get(batch["batch_id"])["images"][0]["status"]=="recognizing":time.sleep(.01)
        batch=self.service.batches.get(batch["batch_id"])
        self.assertEqual(len(batch["rows"]),2)
        image_id=batch["images"][0]["image_id"]
        self.assertEqual({row["evidence_images"][0] for row in batch["rows"]},{image_id})
        self.assertEqual([row["rack_type"] for row in batch["rows"]],[rack["rack_type"] for rack in racks])
        self.assertEqual(batch["stats"]["confirmable"],0)
        removed=self.service.batches.delete_image(batch["batch_id"],image_id,batch["version"],"owner",["A"])
        self.assertTrue(all(row["status"]=="excluded_image" for row in removed["rows"]))
        batch=self.service.batches.restore_image(batch["batch_id"],image_id,removed["version"],"owner",["A"])
        self.assertEqual({row["evidence_images"][0] for row in batch["rows"]},{image_id})
        with self.assertRaises(CabinetError):
            self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
        batch=self.service.batches.get(batch["batch_id"])
        batch=self.service.batches.update(batch["batch_id"],{"version":batch["version"],"rows":[
            {"row_id":row["row_id"],"result":"成功"} for row in batch["rows"]]},"owner",["A"])
        self.assertEqual(batch["stats"]["confirmable"],2,batch["rows"])
        self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
        done=self._wait_batch(batch["batch_id"])
        self.assertEqual(done["stats"]["completed"],2,done["rows"])
        self.assertEqual(len(self.remote.attachments),1)
        for row in done["rows"]:
            self.assertEqual(len(self.remote.get(row["record_id"])["fields"]["上下电确认截图"]),1)

    def test_empty_image_batch_can_be_permanently_deleted(self):
        batch=self.service.batches.create_image_batch("owner",["A"],"A")
        result=self.service.batches.delete_empty(batch["batch_id"],"owner",expected_version=batch["version"])
        self.assertTrue(result["deleted"])
        with self.assertRaisesRegex(CabinetError,"批次不存在"):
            self.service.batches.get(batch["batch_id"])

    def test_image_registration_keeps_out_of_scope_candidates_unsubmitted(self):
        from PIL import Image
        batch=self.service.batches.create_image_batch("owner",["A"])
        content=io.BytesIO();Image.new("RGB",(80,60),"white").save(content,format="PNG")
        rack=self.configs["B"]["inventory"][0]
        candidate={"scope":"B","room":rack["room"],"rack":rack["rack"],"supplier_rack":"",
                   "action":"上正式电","expected":"2026-09-14 01:02:03","actual":"2026-09-14 01:02:03","result":"成功"}
        with patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",return_value=[candidate]):
            self.service.batches.add_images(batch["batch_id"],[("wrong-building.png",content.getvalue())],"owner",["A"])
            deadline=time.time()+5
            while time.time()<deadline and self.service.batches.get(batch["batch_id"])["images"][0]["status"]=="recognizing":time.sleep(.01)
        batch=self.service.batches.get(batch["batch_id"])
        self.assertEqual(batch["rows"],[])
        self.assertEqual(batch["images"][0]["suggestions"][0]["status"],"unauthorized")

    def test_image_registration_flags_different_times_without_creating_duplicate_cabinet(self):
        from PIL import Image
        rack=self.configs["A"]["inventory"][0]
        batch=self.service.batches.create_image_batch("owner",["A"])
        files=[]
        for color in ("white","gray"):
            content=io.BytesIO();Image.new("RGB",(80,60),color).save(content,format="PNG")
            files.append((color+".png",content.getvalue()))
        candidates=[{"scope":"A","room":rack["room"],"rack":rack["rack"],"supplier_rack":"",
                     "action":"上正式电","expected":stamp,"actual":stamp,"result":"成功"}
                    for stamp in ("2026-09-14 01:02:03","2026-09-15 01:02:03")]
        with patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",side_effect=[[item] for item in candidates]):
            self.service.batches.add_images(batch["batch_id"],files,"owner",["A"])
            deadline=time.time()+5
            while time.time()<deadline and any(item["status"]=="recognizing" for item in self.service.batches.get(batch["batch_id"])["images"]):time.sleep(.01)
        result=self.service.batches.get(batch["batch_id"])
        self.assertEqual(len(result["rows"]),1)
        self.assertTrue(result["rows"][0]["evidence_time_conflict"])
        self.assertEqual(result["stats"]["confirmable"],0)
        self.assertEqual(set(result["rows"][0]["evidence_images"]),{item["image_id"] for item in result["images"]})

    def test_todo_batch_list_and_badge_are_scoped_to_building(self):
        a=self.service.batches.create_manual([self._manual_batch_row("A","2026-09-14 01:02:03")],"owner")
        b=self.service.batches.create_manual([self._manual_batch_row("B","2026-09-14 01:02:03")],"owner")
        pending_a=self.service.batches.list("owner",["A","B"],scope="A",status="todo")
        pending_b=self.service.batches.list("owner",["A","B"],scope="B",status="todo")
        self.assertEqual(([item["batch_id"] for item in pending_a["items"]],pending_a["pending_count"]),([a["batch_id"]],1))
        self.assertEqual(([item["batch_id"] for item in pending_b["items"]],pending_b["pending_count"]),([b["batch_id"]],1))
        image=self.service.batches.create_image_batch("owner",["A","B"],"A")
        self.assertEqual(self.service.batches.list("owner",["A","B"],scope="A",status="todo")["pending_count"],2)
        self.assertNotIn(image["batch_id"],[item["batch_id"] for item in self.service.batches.list("owner",["A","B"],scope="B")["items"]])
        mixed=self.service.batches.create_manual([self._manual_batch_row(scope,"2026-09-14 01:05:03") for scope in ("A","B")],"owner")
        limited=next(item for item in self.service.batches.list("viewer",["A"],scope="A")["items"] if item["batch_id"]==mixed["batch_id"])
        self.assertEqual((limited["scopes"],limited["stats"]["total"],limited["pending_rows"]),(["A"],1,1))
        self.service.batches.confirm(a["batch_id"],{"version":a["version"],"all":True},"owner",["A"])
        self._wait_batch(a["batch_id"])
        self.assertEqual(self.service.batches.list("owner",["A","B"],scope="A",status="todo")["pending_count"],2)

    def test_notice_rollback_error_filter_does_not_leak_other_building(self):
        service=self.service.batches
        batch=service.create_from_notice({"target_record_id":"rec-mixed-notice","notice_type":"上电通告",
            "scope":"CAMPUS","cabinet":"A-202包间B01、B-202包间B01","quantity":"2","owner_id":"owner"})
        stored=service.store.get(batch["batch_id"])
        target=next(row for row in stored["rows"] if row["scope"]=="B")
        target.update(status="rollback_blocked",error="后续操作")
        stored["source_notice"]["deleted_at"]="2026-09-18 19:00:00"
        service.store.save(stored,stored["version"])
        self.assertEqual(service.list("viewer",["A"],status="notice_rollback_error")["total"],0)
        self.assertEqual(service.list("owner",["A","B"],status="notice_rollback_error")["total"],1)

    def test_cancelled_batch_detail_and_selected_restore_stay_consistent(self):
        racks=self.configs["A"]["inventory"][:2]
        rows=[{"scope":"A","room":rack["room"],"rack":rack["rack"],"rack_type":rack["rack_type"],
               "action":"上正式电","expected":"2026-09-14 01:02:03","actual":"2026-09-14 01:02:03","result":"成功"}
              for rack in racks]
        batch=self.service.batches.create_manual(rows,"owner")
        cancelled=self.service.batches.cancel(batch["batch_id"],"owner",expected_version=batch["version"])
        self.assertEqual(cancelled["status"],"cancelled")
        self.assertEqual([row["status"] for row in self.service.batches.get(batch["batch_id"])["rows"]],["excluded_cancelled"]*2)
        listing=self.service.batches.list("owner",["A"],scope="A")
        self.assertEqual(next(item for item in listing["items"] if item["batch_id"]==batch["batch_id"])["status"],"cancelled")
        self.assertTrue(all(row["restorable"] for row in self.service.batches.visible(cancelled,"owner",["A"])["rows"]))
        first=cancelled["rows"][0]["row_id"]
        restored=self.service.batches.restore_rows(batch["batch_id"],{"version":cancelled["version"],"row_ids":[first]},"owner",["A"])
        self.assertEqual(restored["status"],"pending")
        self.assertEqual([row["status"] for row in restored["rows"]],["ready","excluded_cancelled"])
        self.assertEqual(self.service.batches.list("owner",["A"],scope="A",status="todo")["pending_count"],1)
        restored=self.service.batches.restore_rows(batch["batch_id"],{"version":restored["version"],"row_ids":[restored["rows"][1]["row_id"]]},"owner",["A"])
        self.assertEqual([row["status"] for row in restored["rows"]],["ready","ready"])

    def test_cancel_during_image_ocr_does_not_recreate_rows(self):
        from PIL import Image
        batch=self.service.batches.create_image_batch("owner",["A"])
        content=io.BytesIO();Image.new("RGB",(80,60),"white").save(content,format="PNG")
        rack=self.configs["A"]["inventory"][0]
        candidate={"scope":"A","room":rack["room"],"rack":rack["rack"],"supplier_rack":"",
                   "action":"上正式电","expected":"2026-09-14 01:02:03","actual":"2026-09-14 01:02:03","result":"成功"}
        entered=threading.Event();release=threading.Event()
        def slow_ocr(_content):
            entered.set();release.wait(5);return [candidate]
        with patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",side_effect=slow_ocr):
            self.service.batches.add_images(batch["batch_id"],[("mail.png",content.getvalue())],"owner",["A"])
            self.assertTrue(entered.wait(5))
            current=self.service.batches.get(batch["batch_id"])
            cancelled=self.service.batches.cancel(batch["batch_id"],"owner",expected_version=current["version"])
            release.set()
            self.service.batches._ocr_pool.shutdown(wait=True)
        final=self.service.batches.get(batch["batch_id"])
        self.assertEqual((cancelled["status"],final["status"],final["rows"],final["images"][0]["status"]),
                         ("cancelled","cancelled",[],"cancelled"))

    def test_cancel_rejects_unverified_write(self):
        batch=self.service.batches.create_manual([self._manual_batch_row("A","2026-09-14 01:02:03")],"owner")
        def pending(current):current["rows"][0].update(status="failed",operation_started=True)
        batch=self.service.batches._change(batch["batch_id"],pending)
        with self.assertRaisesRegex(CabinetError,"未核验"):
            self.service.batches.cancel(batch["batch_id"],"owner",expected_version=batch["version"])

    def test_cancel_then_restore_preserves_rolled_back_attempt(self):
        batch=self.service.batches.create_manual([self._manual_batch_row("A","2026-09-14 01:02:03")],"owner")
        self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
        done=self._wait_batch(batch["batch_id"])
        self.service.batches.rollback(batch["batch_id"],{"version":done["version"],"all":True},"owner",["A"])
        undone=self._wait_batch(batch["batch_id"])
        old_id=undone["rows"][0]["operation_id"]
        cancelled=self.service.batches.cancel(batch["batch_id"],"owner",expected_version=undone["version"])
        self.assertEqual(cancelled["rows"][0]["status"],"excluded_cancelled")
        restored=self.service.batches.restore_rows(batch["batch_id"],{"version":cancelled["version"],"row_ids":[cancelled["rows"][0]["row_id"]]},"owner",["A"])
        self.assertEqual(restored["rows"][0]["status"],"rolled_back")
        self.service.batches.confirm(batch["batch_id"],{"version":restored["version"],"all":True},"owner",["A"])
        again=self._wait_batch(batch["batch_id"])
        self.assertEqual(again["stats"]["completed"],1,again["rows"])
        self.assertNotEqual(again["rows"][0]["operation_id"],old_id)

    def test_exclude_restore_reconfirms_rolled_back_row_with_new_attempt(self):
        service=self.service.batches
        batch=service.create_manual([self._manual_batch_row("A","2026-09-14 01:02:03")],"owner")
        for restore_mode in ("bulk","patch","legacy"):
            with self.subTest(restore=restore_mode):
                service.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
                done=self._wait_batch(batch["batch_id"])
                self.assertEqual(done["stats"]["completed"],1,done["rows"])
                old_id=done["rows"][0]["operation_id"]
                service.rollback(batch["batch_id"],{"version":done["version"],"all":True},"owner",["A"])
                undone=self._wait_batch(batch["batch_id"])
                row_id=undone["rows"][0]["row_id"]
                excluded=service.update(batch["batch_id"],{"version":undone["version"],"rows":[{"row_id":row_id,"excluded":True}]},"owner",["A"])
                if restore_mode=="patch":
                    batch=service.update(batch["batch_id"],{"version":excluded["version"],"rows":[{"row_id":row_id,"excluded":False}]},"owner",["A"])
                else:
                    if restore_mode=="legacy":
                        excluded=service._change(batch["batch_id"],lambda current:current["rows"][0].pop("excluded_from_status",None))
                    else:
                        excluded=service.cancel(batch["batch_id"],"owner",expected_version=excluded["version"])
                    batch=service.restore_rows(batch["batch_id"],{"version":excluded["version"],"row_ids":[row_id]},"owner",["A"])
                self.assertEqual(batch["rows"][0]["status"],"rolled_back")
                service.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
                done=self._wait_batch(batch["batch_id"])
                self.assertEqual(done["stats"]["completed"],1,done["rows"])
                self.assertNotEqual(done["rows"][0]["operation_id"],old_id)
                self.assertEqual(done["rows"][0]["attempts"][-1]["operation_id"],old_id)
                service.rollback(batch["batch_id"],{"version":done["version"],"all":True},"owner",["A"])
                batch=self._wait_batch(batch["batch_id"])

    def test_image_registration_does_not_infer_sequential_events_from_conflicting_screenshots(self):
        from PIL import Image
        rack=self.configs["A"]["inventory"][0]
        current=next(item for item in self.service.overview("A")["racks"] if (item["room"],item["rack"])==(rack["room"],rack["rack"]))
        first,second={"formal":("下正式电","上正式电"),"test":("下测试电","上测试电"),"off":("上正式电","下正式电")}[current["state"]]
        candidates=[{"scope":"A","room":rack["room"],"rack":rack["rack"],"supplier_rack":"",
                     "action":action,"expected":stamp,"actual":stamp,"result":"成功"}
                    for action,stamp in ((first,"2026-09-18 00:01:00"),(second,"2026-09-18 00:02:00"))]
        files=[]
        for color in ("white","gray"):
            content=io.BytesIO();Image.new("RGB",(80,60),color).save(content,format="PNG")
            files.append((color+".png",content.getvalue()))
        batch=self.service.batches.create_image_batch("owner",["A"])
        with patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",side_effect=[[item] for item in candidates]):
            self.service.batches.add_images(batch["batch_id"],files,"owner",["A"])
            deadline=time.time()+5
            while time.time()<deadline and any(item["status"]=="recognizing" for item in self.service.batches.get(batch["batch_id"])["images"]):time.sleep(.01)
        result=self.service.batches.get(batch["batch_id"])
        self.assertEqual(len(result["rows"]),1)
        self.assertEqual(result["stats"]["confirmable"],0,result["rows"])
        self.assertEqual(result["rows"][0]["action"],first)
        self.assertTrue(result["rows"][0]["evidence_time_conflict"])
        self.assertEqual(len(result["rows"][0]["evidence_images"]),2)

    def test_deleted_photo_cannot_be_restored_into_completed_record(self):
        from PIL import Image
        rack=self.configs["A"]["inventory"][0]
        row=self._manual_batch_row("A","2026-09-14 01:02:03")
        batch=self.service.batches.create_manual([row],"owner")
        content=io.BytesIO();Image.new("RGB",(80,60),"white").save(content,format="PNG")
        candidate={"scope":"A","room":rack["room"],"rack":rack["rack"],"supplier_rack":"",
                   "action":row["action"],"expected":row["expected"],"actual":row["actual"]}
        with patch("bin.lan_bitable_template_portal.cabinet_power_evidence.recognize_image_with_timeout",return_value=[candidate]):
            self.service.batches.add_images(batch["batch_id"],[("proof.png",content.getvalue())],"owner",["A"])
            deadline=time.time()+5
            while time.time()<deadline and self.service.batches.get(batch["batch_id"])["images"][0]["status"]=="recognizing":time.sleep(.01)
        batch=self.service.batches.get(batch["batch_id"])
        image_id=batch["images"][0]["image_id"]
        removed=self.service.batches.delete_image(batch["batch_id"],image_id,batch["version"],"owner",["A"])
        self.service.batches.confirm(batch["batch_id"],{"version":removed["version"],"all":True},"owner",["A"])
        done=self._wait_batch(batch["batch_id"])
        self.assertEqual(done["stats"]["completed"],1,done["rows"])
        with self.assertRaisesRegex(CabinetError,"先回退"):
            self.service.batches.restore_image(batch["batch_id"],image_id,done["version"],"owner",["A"])

    def test_batch_confirm_uses_one_cloud_write_per_building_and_rolls_back(self):
        rows=[]
        for scope in ("A","E"):
            for rack in self.configs[scope]["inventory"][:2]:
                rows.append({"scope":scope,"room":rack["room"],"rack":rack["rack"],
                             "rack_type":rack["rack_type"],"action":"上正式电",
                             "expected":"2026-09-14 01:02:03","actual":"2026-09-14 01:02:03","result":"成功"})
        batch=self.service.batches.create_manual(rows,"owner")
        self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A","E"])
        done=self._wait_batch(batch["batch_id"])
        self.assertEqual(done["stats"]["completed"],4,done["rows"])
        self.assertEqual((self.remote.batch_create_calls,self.remote.batch_update_calls),(1,1))
        original_update=self.remote.update
        def checked_update(record_id,fields):
            for field in ("机柜功率（W）","来源行号"):
                if fields.get(field) is not None:
                    self.assertIsInstance(fields[field],(int,float),field)
            return original_update(record_id,fields)
        self.remote.update=checked_update
        self.service.batches.rollback(batch["batch_id"],{"version":done["version"],"all":True},"owner",["A","E"])
        undone=self._wait_batch(batch["batch_id"])
        self.assertEqual(undone["stats"]["rolled_back"],4,undone["rows"])
        self.assertEqual(undone["status"],"rolled_back")

    def test_batch_create_lost_response_reconciles_without_duplicate(self):
        rows=[]
        for rack in self.configs["A"]["inventory"][:2]:
            rows.append({"scope":"A","room":rack["room"],"rack":rack["rack"],
                         "rack_type":rack["rack_type"],"action":"上正式电",
                         "expected":"2026-09-14 01:02:03","actual":"2026-09-14 01:02:03","result":"成功"})
        batch=self.service.batches.create_manual(rows,"owner")
        self.remote.fail_after_batch_create=True
        self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
        done=self._wait_batch(batch["batch_id"])
        self.assertEqual(done["stats"]["completed"],2,done["rows"])
        self.assertEqual(self.remote.batch_create_calls,1)
        self.assertEqual(self.remote.creates,2)

    def test_batch_update_lost_response_reconciles_each_cabinet(self):
        rows=[]
        for rack in self.configs["E"]["inventory"][:2]:
            rows.append({"scope":"E","room":rack["room"],"rack":rack["rack"],
                         "rack_type":rack["rack_type"],"action":"下正式电",
                         "expected":"2026-09-14 01:02:03","actual":"2026-09-14 01:02:03","result":"成功"})
        batch=self.service.batches.create_manual(rows,"owner")
        before_count=len(self.remote.records)
        self.remote.fail_after_batch_update=True
        self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["E"])
        done=self._wait_batch(batch["batch_id"])
        self.assertEqual(done["stats"]["completed"],2,done["rows"])
        self.assertEqual(self.remote.batch_update_calls,1)
        self.assertEqual(len(self.remote.records),before_count)

    def test_all_failed_rollbacks_are_not_reported_as_partial(self):
        batch={"status":"partial","rows":[{"status":"rollback_failed","scope":"E"},
                                          {"status":"rollback_blocked","scope":"E"}]}
        self.service.batches._refresh_summary(batch)
        self.assertEqual(batch["status"],"failed")
        self.assertEqual(batch["stats"]["rolled_back"],0)

    def test_one_failed_rollback_does_not_block_other_cabinet(self):
        rows=[]
        for rack in self.configs["E"]["inventory"][:2]:
            rows.append({"scope":"E","room":rack["room"],"rack":rack["rack"],
                         "rack_type":rack["rack_type"],"action":"下正式电",
                         "expected":"2026-09-14 01:02:03","actual":"2026-09-14 01:02:03","result":"成功"})
        batch=self.service.batches.create_manual(rows,"owner")
        self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["E"])
        done=self._wait_batch(batch["batch_id"])
        self.assertEqual(done["stats"]["completed"],2)
        blocked_id=done["rows"][-1]["record_id"]
        original_update=self.remote.update
        def fail_one(record_id,fields):
            if record_id==blocked_id: raise CabinetError("模拟数字字段写入失败")
            return original_update(record_id,fields)
        self.remote.update=fail_one
        self.service.batches.rollback(batch["batch_id"],{"version":done["version"],"all":True},"owner",["E"])
        partial=self._wait_batch(batch["batch_id"])
        self.assertEqual(partial["stats"]["rollback_failed"],1,partial["rows"])
        self.assertEqual(partial["stats"]["rolled_back"],1,partial["rows"])
        self.remote.update=original_update
        self.service.batches.rollback(batch["batch_id"],{"version":partial["version"],"all":True},"owner",["E"])
        self.assertEqual(self._wait_batch(batch["batch_id"])["stats"]["rolled_back"],2)

    def test_batch_rollback_skips_cabinet_with_later_batch(self):
        first=self.service.batches.create_manual([self._manual_batch_row("A","2026-09-14 01:02:03")],"owner")
        self.service.batches.confirm(first["batch_id"],{"version":first["version"],"all":True},"owner",["A"])
        first=self._wait_batch(first["batch_id"])
        second=self.service.batches.create_manual([self._manual_batch_row("A","2026-09-14 01:03:03", "下正式电")],"owner")
        self.service.batches.confirm(second["batch_id"],{"version":second["version"],"all":True},"owner",["A"])
        second=self._wait_batch(second["batch_id"])
        self.service.batches.rollback(first["batch_id"],{"version":first["version"],"all":True},"owner",["A"])
        first=self._wait_batch(first["batch_id"])
        self.assertEqual(first["rows"][0]["status"],"rollback_blocked")
        self.assertIn(first["rows"][0]["record_id"],self.remote.records)
        self.service.batches.rollback(second["batch_id"],{"version":second["version"],"all":True},"owner",["A"])
        self.assertEqual(self._wait_batch(second["batch_id"])["rows"][0]["status"],"rolled_back")
        self.service.batches.rollback(first["batch_id"],{"version":first["version"],"all":True},"owner",["A"])
        self.assertEqual(self._wait_batch(first["batch_id"])["rows"][0]["status"],"rolled_back")

    def test_rollback_uses_commit_order_not_batch_creation_or_business_time(self):
        service=self.service.batches
        for scope in ("A","D"):
            with self.subTest(scope=scope):
                first=service.create_manual([self._manual_batch_row(scope,"2026-09-14 01:02:03")],"owner")
                second=service.create_manual([self._manual_batch_row(scope,"2026-09-14 01:03:03","下正式电")],"owner")
                completed=[]
                for batch in (second,first):
                    service.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",[scope])
                    batch=self._wait_batch(batch["batch_id"])
                    self.assertEqual(batch["stats"]["completed"],1,batch["rows"])
                    row=batch["rows"][0]
                    journal=self.service.local.document(scope,"write:"+row["operation_id"])
                    journal["completed_at"]=1000
                    self.service.local.document(scope,"write:"+row["operation_id"],journal)
                    completed.append(batch)
                early,late=completed
                self.assertGreater(service._write_order(scope,late["rows"][0]),service._write_order(scope,early["rows"][0]))
                service.rollback(early["batch_id"],{"version":early["version"],"all":True},"owner",[scope])
                self.assertEqual(self._wait_batch(early["batch_id"])["rows"][0]["status"],"rollback_blocked")
                for batch in (late,early):
                    batch=service.get(batch["batch_id"])
                    service.rollback(batch["batch_id"],{"version":batch["version"],"all":True},"owner",[scope])
                    self.assertEqual(self._wait_batch(batch["batch_id"])["rows"][0]["status"],"rolled_back")
                early=service.get(early["batch_id"])
                service.confirm(early["batch_id"],{"version":early["version"],"all":True},"owner",[scope])
                self.assertEqual(self._wait_batch(early["batch_id"])["stats"]["completed"],1)
                late=service.get(late["batch_id"])
                with self.assertRaisesRegex(CabinetError,"没有可确认"):
                    service.confirm(late["batch_id"],{"version":late["version"],"all":True},"owner",[scope])
                self.assertIn("later_batch",{issue["code"] for issue in service.get(late["batch_id"])["rows"][0]["issues"]})

    def test_same_batch_rollback_reverses_actual_commit_order(self):
        service=self.service.batches
        batch=service.create_manual([self._manual_batch_row("A","2026-09-14 01:02:03"),
                                    self._manual_batch_row("A","2026-09-14 01:03:03","下正式电")],"owner")
        for row_id in [row["row_id"] for row in reversed(batch["rows"])]:
            service.confirm(batch["batch_id"],{"version":batch["version"],"row_ids":[row_id]},"owner",["A"])
            batch=self._wait_batch(batch["batch_id"])
        self.assertEqual(batch["stats"]["completed"],2,batch["rows"])
        service.rollback(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
        batch=self._wait_batch(batch["batch_id"])
        self.assertEqual(batch["stats"]["rolled_back"],2)
        service.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
        batch=self._wait_batch(batch["batch_id"])
        self.assertEqual(batch["stats"]["completed"],2,batch["rows"])

    def test_previous_release_write_order_uses_journal_completion_time(self):
        service=self.service.batches
        first=service.create_manual([self._manual_batch_row("A","2026-09-14 01:02:03")],"owner")
        second=service.create_manual([self._manual_batch_row("A","2026-09-14 01:03:03")],"owner")
        for batch,completed_at in ((first,2000.2),(second,2000.1)):
            def complete(current):
                current["rows"][0].update(status="completed",wrote_record=True,completed_at="2026-09-14 02:00:00")
            service._change(batch["batch_id"],complete)
            self.service.local.document("A","write:"+batch["rows"][0]["operation_id"],
                                        {"status":"completed","completed_at":completed_at})
        first,second=service.get(first["batch_id"]),service.get(second["batch_id"])
        self.assertFalse(service._later_completed(first["batch_id"],first["rows"][0]))
        self.assertTrue(service._later_completed(second["batch_id"],second["rows"][0]))
        self.service.local.document("A","write:"+second["rows"][0]["operation_id"],{"status":"completed"})
        service._change(second["batch_id"],lambda current:current["rows"][0].pop("completed_at"))
        self.assertTrue(service._later_completed(first["batch_id"],first["rows"][0]),"unknown order must not permit rollback")

    def test_batch_rollback_recovers_lost_delete_response(self):
        batch=self.service.batches.create_manual([self._manual_batch_row("A","2026-09-14 01:02:03")],"owner")
        lists_before=self.remote.list_calls
        self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
        batch=self._wait_batch(batch["batch_id"])
        self.assertEqual(self.remote.list_calls,lists_before,"首次创建不应先扫描飞书")
        self.remote.fail_after_delete=True
        self.service.batches.rollback(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
        batch=self._wait_batch(batch["batch_id"])
        self.assertEqual(batch["rows"][0]["status"],"rollback_failed")
        self.remote.fail_after_delete=False
        self.service.batches.rollback(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["A"])
        batch=self._wait_batch(batch["batch_id"])
        self.assertEqual(batch["rows"][0]["status"],"rolled_back",batch["rows"][0])

    def test_batch_rollback_restores_directory_type_and_rejects_cloud_edits(self):
        row=self._manual_batch_row("D","2026-09-14 01:02:03")
        original_type=row["rack_type"]
        row["rack_type"]="网络机柜" if original_type=="服务器机柜" else "服务器机柜"
        row["type_resolution"]="sync_current"
        batch=self.service.batches.create_manual([row],"owner")
        self.assertEqual(batch["rows"][0]["status"],"ready",batch["rows"][0])
        self.service.batches.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["D"])
        batch=self._wait_batch(batch["batch_id"])
        self.assertEqual(batch["rows"][0]["status"],"completed",batch["rows"][0])
        self.assertEqual(next(item["rack_type"] for item in self.service._snapshot("D")["config"]["inventory"] if (item["room"],item["rack"])==(row["room"],row["rack"])),row["rack_type"])
        self.service.batches.rollback(batch["batch_id"],{"version":batch["version"],"all":True},"owner",["D"])
        batch=self._wait_batch(batch["batch_id"])
        self.assertEqual(batch["rows"][0]["status"],"rolled_back",batch["rows"][0])
        self.assertEqual(next(item["rack_type"] for item in self.service._snapshot("D")["config"]["inventory"] if (item["room"],item["rack"])==(row["room"],row["rack"])),original_type)

        other=self.service.batches.create_manual([self._manual_batch_row("A","2026-09-14 01:02:03")],"owner")
        self.service.batches.confirm(other["batch_id"],{"version":other["version"],"all":True},"owner",["A"])
        other=self._wait_batch(other["batch_id"])
        rid=other["rows"][0]["record_id"]
        self.remote.records[rid]["fields"]["操作类型"]="下正式电"
        self.service.batches.rollback(other["batch_id"],{"version":other["version"],"all":True},"owner",["A"])
        other=self._wait_batch(other["batch_id"])
        self.assertEqual(other["rows"][0]["status"],"rollback_blocked")
        self.assertIn(rid,self.remote.records)

    def test_de_newest_group_stays_in_current_operation_columns(self):
        old=from_feishu(next(r for r in self.source_records if r["fields"]["楼栋"]=="E楼" and r["fields"]["来源行号"]==2))
        groups=old["groups"]+[{"action":"下正式电","actual":"2026-09-09 10:00:00","expected":""}]
        edited=self.service.validate_op("E",{"groups":groups},old)
        self.assertEqual(edited["groups"][0]["action"],"下正式电")
        self.assertEqual(to_fields(edited)["操作类型"],"下正式电")

    def test_de_repeated_saves_expand_history_headers_dates_and_styles(self):
        from .lan_bitable_template_portal.cabinet_power_excel import col_name, coord
        from openpyxl.styles.numbers import BUILTIN_FORMATS
        for scope in 'DE':
            snapshot=self.service._snapshot(scope)
            if scope=='D':
                old=next(op for op in snapshot['operations'] if (op['room'],op['rack'])==('202','F12'))
            else:
                old=max(snapshot['operations'],key=lambda op:sum(bool(g.get('action')) for g in op['groups']))
            count=len(snapshot['operations']); initial=copy.deepcopy(old['groups'])
            for index,action in enumerate(('正式电转测试电','测试电转正式电','下正式电')):
                prior_groups=copy.deepcopy(old['groups'])
                group={'id':f'export_repeat_{scope}_{index}','action':action,'result':'成功',
                       'expected':f'2026-09-21 {10+index:02}:00:00','actual':f'2026-09-21 {10+index:02}:01:00'}
                old=self.service.save_operation(scope,{'operation_id':f'export_repeat_operation_{scope}_{index}',
                    'expected_version':old['version'],'groups':[*old['groups'],group]},'owner',old['record_id'])
                self.assertEqual(old['groups'][0],group)
                self.assertEqual(old['groups'][1:],prior_groups)
            snapshot=self.service._snapshot(scope)
            self.assertEqual(len(snapshot['operations']),count)
            before=Workbook((TEMPLATES/(scope+'.xlsm')).read_bytes())
            after=Workbook(export_workbook((TEMPLATES/(scope+'.xlsm')).read_bytes(),snapshot['config'],snapshot['operations']))
            fmt=snapshot['config']['template_data']['formats'][0]; sheet=fmt['sheet']
            cells=after.cells(sheet); row=old['source_row']
            styles=ET.fromstring(after.archive.read('xl/styles.xml')); xfs=styles.find(T('cellXfs'))
            num_formats={**BUILTIN_FORMATS,**{int(node.get('numFmtId')):node.get('formatCode') for node in styles.iter(T('numFmt'))}}
            self.assertEqual(after.value(cells[f'H{row}']),old['rack_type'])
            self.assertEqual(after.value(cells[f'I{row}']),old['power'])
            self.assertEqual(after.value(cells[f'J{row}']),'成功')
            columns=sorted(coord(ref)[0] for ref,cell in cells.items() if coord(ref)[1]==fmt['header'] and after.value(cell)=='操作类型')
            self.assertGreaterEqual(len(columns),len(initial)+3)
            exported=[{'action':after.value(cells[f'{col_name(col)}{row}']),
                       'expected':dates(after.value(cells[f'{col_name(col+1)}{row}'])),
                       'actual':dates(after.value(cells[f'{col_name(col+2)}{row}']))} for col in columns[:len(old['groups'])]]
            self.assertEqual(exported,[{'action':g.get('action',''),'expected':dates(g.get('expected')),'actual':dates(g.get('actual'))} for g in old['groups']])
            for col in columns[len(fmt['groups']):]:
                for offset,label in enumerate(('操作类型','期望完成时间','实际完成时间')):
                    self.assertEqual(after.value(cells[f'{col_name(col+offset)}1']),label)
                    key=('action','expected','actual')[offset]
                    sample=f"{col_name(fmt['groups'][-1][key])}2"
                    style=int(cells[f'{col_name(col+offset)}{row}'].get('s'))
                    self.assertEqual(after.styles[style],before.styles[int(before.cells(sheet)[sample].get('s'))])
                    if offset: self.assertEqual(num_formats[int(xfs[style].get('numFmtId'))],'yyyy-mm-dd hh:mm')
                    width=next(node.get('width') for node in after.sheet(sheet).find(T('cols')) if int(node.get('min'))<=col+offset<=int(node.get('max')))
                    source_col=fmt['groups'][-1][key]
                    expected_width=next(node.get('width') for node in before.sheet(sheet).find(T('cols')) if int(node.get('min'))<=source_col<=int(node.get('max')))
                    self.assertEqual(width,expected_width)
            definitions=sorted(after.sheet(sheet).find(T('cols')),key=lambda node:int(node.get('min')))
            self.assertTrue(all(int(left.get('max'))<int(right.get('min')) for left,right in zip(definitions,definitions[1:])))
            self.assertEqual(after.archive.read('xl/vbaProject.bin'),before.archive.read('xl/vbaProject.bin'))

    def test_abc_history_compression_preserves_empty_times_and_grows_rows(self):
        for scope in 'ABC':
            snapshot=self.service._snapshot(scope)
            rack=snapshot['config']['inventory'][0]
            for category in ('up','down'):
                groups=[{'action':'上正式电','expected':'2026-09-20 09:00:00','actual':'2026-09-20 09:01:00'},
                        {'action':'正式电转测试电','expected':'','actual':''},
                        {'action':'测试电转正式电','expected':'','actual':'2026-09-20 11:01:00'},
                        {'action':'正式电转测试电','expected':'2026-09-20 12:00:00','actual':'2026-09-20 12:01:00'},
                        {'action':'测试电转正式电','expected':'2026-09-20 13:00:00','actual':'2026-09-20 13:01:00'}]
                if category=='down': groups.append({'action':'下正式电','expected':'','actual':'2026-09-20 14:01:00'})
                else: groups=groups[:3]
                source=next(fmt['sheet'] for fmt in snapshot['config']['template_data']['formats']
                            if ('下电' in fmt['sheet'] and '上下电' not in fmt['sheet'])==(category=='down'))
                op={'record_id':f'recExport{scope}{category}','scope':scope,'room':rack['room'],'rack':rack['rack'],
                    'system_name':system_name(scope,rack['room']),'source':source,'source_row':None,'rack_type':rack['rack_type'],
                    'result':'成功','groups':groups,'meta':{'schema':3,'category':category},'category':category}
                op=from_feishu({'record_id':op['record_id'],'fields':{**to_fields(op),'来源工作表':source}})
                book=Workbook(export_workbook((TEMPLATES/(scope+'.xlsm')).read_bytes(),snapshot['config'],[op]))
                row_number,row=next((n,r) for n,r in book.rows(source) if n>1 and r.get(4)==rack['rack'])
                fmt=next(fmt for fmt in snapshot['config']['template_data']['formats'] if fmt['sheet']==source)
                actions='\n'.join(str(row.get(g['action'],'')) for g in fmt['groups'])
                self.assertEqual(Counter(re.findall('上正式电|正式电转测试电|测试电转正式电|下正式电',actions)),Counter(g['action'] for g in groups),(scope,category))
                merged=str(row.get(fmt['groups'][1]['actual'],''))
                self.assertIn('2、2026-09-20 11:01',merged,(scope,category))
                self.assertNotIn('11:01:00',merged)
                self.assertEqual([line.split('、',1)[0] for line in str(row[fmt['groups'][1]['action']]).splitlines()],
                                 [line.split('、',1)[0] for line in merged.splitlines()])
                sheet_row=next(r for r in book.sheet(source).find(T('sheetData')) if r.get('r')==str(row_number))
                self.assertGreater(float(sheet_row.get('ht',0)),30,(scope,category))
                if scope=='B' and category=='down': self.assertGreaterEqual(float(sheet_row.get('ht',0)),110)

    def test_original_export_preserves_all_five_formats_and_raw_rows(self):
        from openpyxl import load_workbook
        for scope in "ABCDE":
            original=(TEMPLATES/(scope+".xlsm")).read_bytes(); before=Workbook(original)
            ops=[from_feishu(r) for r in self.source_records if r["fields"]["楼栋"]==scope+"楼"]
            result=export_workbook(original,self.configs[scope],ops); after=Workbook(result)
            self.assertEqual(list(after.sheets),list(before.sheets))
            self.assertEqual(before.archive.read("xl/vbaProject.bin"),after.archive.read("xl/vbaProject.bin"))
            for name in before.sheets:
                original_merges=[x.get("ref") for x in before.sheet(name).iter(T("mergeCell"))]
                exported_merges=[x.get("ref") for x in after.sheet(name).iter(T("mergeCell"))]
                self.assertTrue(set(original_merges)<=set(exported_merges),name)
                if "平面图" not in name: self.assertEqual(original_merges,exported_merges)
            for fmt in self.models[scope]["formats"]:
                name=fmt["sheet"]; raw=dict(before.rows(name)); actual=dict(after.rows(name))
                cells=after.cells(name)
                date_cols={group[key] for group in fmt['groups'] for key in ('actual','expected') if group.get(key)}
                styles=ET.fromstring(after.archive.read('xl/styles.xml')); xfs=styles.find(T('cellXfs'))
                num_formats={int(node.get('numFmtId')):node.get('formatCode') for node in styles.iter(T('numFmt'))}
                for rn,row in raw.items():
                    if rn<=fmt["header"] or not row.get(fmt["rack"]): continue
                    for col,value in row.items():
                        exported=actual[rn].get(col,"")
                        if col in date_cols and dates(value):
                            from .lan_bitable_template_portal.cabinet_power_excel import col_name,export_time_text
                            style=int(cells[f'{col_name(col)}{rn}'].get('s','0'))
                            self.assertEqual(num_formats[int(xfs[style].get('numFmtId'))],'yyyy-mm-dd hh:mm')
                            if isinstance(value,str):
                                if isinstance(exported,str): self.assertEqual(exported,export_time_text(value))
                                else: self.assertEqual(dates(exported),dates(value))
                                continue
                        self.assertEqual(exported,value,(scope,name,rn,col))
            checked=load_workbook(io.BytesIO(result),keep_vba=True,data_only=True); checked.close()

    def test_notice_expected_time_stays_blank_through_confirmation_all_buildings(self):
        for scope in "ABCDE":
            with self.subTest(scope=scope):
                snapshot=self.service._snapshot(scope)
                rack=next(item for item in derive_records(snapshot["config"],snapshot["operations"])["racks"]
                          if item["state"]=="off")
                service=self.service.batches
                batch=service.create_from_notice({"target_record_id":"recEmptyExpected"+scope,
                    "notice_type":"上电通告","scope":scope,"cabinet":f"{scope}-{rack['room']}包间{rack['rack']}",
                    "quantity":"1","owner_id":"owner","end_time":"2026-09-20 19:00:00"})
                row=batch["rows"][0]
                self.assertEqual(row["expected"],"")
                self.assertEqual(batch["source_notice"]["end_time"],"2026-09-20 19:00:00")
                self.assertNotIn("expected",{issue["code"] for issue in row["issues"]})
                batch=service.update(batch["batch_id"],{"version":batch["version"],"rows":[
                    {"row_id":row["row_id"],"action":"上正式电","actual":"2026-09-20 18:00:00"}]},"owner",[scope])
                self.assertEqual(batch["stats"]["confirmable"],1,batch["rows"])
                service.confirm(batch["batch_id"],{"version":batch["version"],"all":True},"owner",[scope])
                done=self._wait_batch(batch["batch_id"])
                self.assertEqual(done["stats"]["completed"],1,done["rows"])
                saved=from_feishu(self.remote.get(done["rows"][0]["record_id"]))
                group=next(group for group in saved["groups"] if group.get("actual")=="2026-09-20 18:00:00")
                self.assertEqual(group["expected"],"")

    def test_notice_old_default_is_cleared_without_erasing_user_or_evidence_times(self):
        service=self.service.batches
        batch=service.create_from_notice({"target_record_id":"recOldExpected","notice_type":"上电通告",
            "scope":"D","cabinet":"D-201包间B02、B04、B05、B07、B08","quantity":"5","owner_id":"owner",
            "end_time":"2026-09-20 19:00:00"})
        for row in batch["rows"]:
            row["expected"]="2026-09-20 19:00:00"
            row["original"]["expected"]=row["expected"]
        batch["rows"][1]["edits"].append({"field":"expected","before":"","after":"2026-09-20 19:00:00"})
        batch["rows"][2]["evidence_images"]=["proof"]
        batch["rows"][3]["status"]="completed"
        batch["rows"][4]["operation_started"]=True
        service.store.save(batch,batch["version"])
        upgraded=service.get(batch["batch_id"])
        self.assertEqual([row["expected"] for row in upgraded["rows"]],[""]+["2026-09-20 19:00:00"]*4)
        self.assertEqual(upgraded["rows"][0]["original"]["expected"],"2026-09-20 19:00:00")
        self.assertEqual(upgraded["rows"][0]["edits"][-1]["owner"],"system")
        self.assertEqual(service.get(batch["batch_id"])["version"],upgraded["version"])
        manual=service.create_manual([{**self._manual_batch_row("A","2026-09-20 18:00:00"),"expected":""}],"owner")
        manual=service.update(manual["batch_id"],{"version":manual["version"],"rows":[
            {"row_id":manual["rows"][0]["row_id"],"expected":""}]},"owner",["A"])
        self.assertIn("expected",{issue["code"] for issue in manual["rows"][0]["issues"]})
        self.assertFalse(service._valid_date("2026-02-30 12:00:00",allow_future=True))

    def test_removed_summary_projection_cleanup_preserves_batches(self):
        store=self.service.batches.store
        batch=self.service.batches.create_manual([self._manual_batch_row("A","2026-09-20 18:00:00")],"owner")
        before=store.get(batch["batch_id"])
        with store._connect() as conn,conn:
            conn.execute("CREATE TABLE notice_summary_rows(id TEXT)")
            conn.execute("CREATE TABLE notice_summary_baselines(id TEXT)")
            conn.execute("INSERT INTO batch_meta VALUES('notice_summary_projection_version','3')")
            conn.execute("DELETE FROM batch_meta WHERE key='batch_runtime_projection_version'")
        migrated=type(store)(store.path)
        self.assertEqual(migrated.get(batch["batch_id"]),before)
        self.assertEqual(migrated.runtime_status(batch["batch_id"])["stats"],before["stats"])
        with migrated._connect() as conn:
            self.assertFalse(conn.execute("SELECT name FROM sqlite_master WHERE name LIKE 'notice_summary_%'").fetchall())

    def test_mail_daily_and_monthly_statistics_all_buildings(self):
        for scope in "ABCDE":
            with self.subTest(scope=scope):
                snapshot=self.service._snapshot(scope); config=snapshot["config"]
                rack=next(item for item in derive_records(config,snapshot["operations"])["racks"] if item["state"]=="off")
                operations=copy.deepcopy(snapshot["operations"])
                for index,(action,category) in enumerate((("上正式电","up"),("下正式电","down")),2):
                    actual=f"2027-01-{index:02d} 10:00:00"
                    with patch("bin.lan_bitable_template_portal.cabinet_power.stamp",return_value="2027-01-04 12:00:00"):
                        operation=self.service.validate_op(scope,{"room":rack["room"],"rack":rack["rack"],
                            "rack_type":rack["rack_type"],"result":"成功","category":category,
                            "groups":[{"id":f"new_event_{index}","action":action,"expected":"","actual":actual,"result":"成功"}]})
                    fields=to_fields(operation); fields["来源工作表"]=operation["source"]
                    operations.append(from_feishu({"record_id":f"recExport{scope}{index}","fields":fields}))
                book=Workbook(export_workbook((TEMPLATES/(scope+".xlsm")).read_bytes(),config,operations))
                rows=dict(book.rows("机柜上电汇总表"))
                date_col,up_col,down_col={"A":(2,3,5),"B":(13,14,15),"C":(2,3,4),"D":(2,3,4),"E":(2,3,4)}[scope]
                up=next(row for row in rows.values() if row.get(date_col)=="2027.1.2")
                down=next(row for row in rows.values() if row.get(4 if scope=="A" else date_col)=="2027.1.3")
                self.assertEqual((up[up_col],down[down_col]),(1.0,1.0))
                month_col={"A":9,"B":19,"C":33,"D":8,"E":8}[scope]
                month=next(row for row in rows.values() if str(row.get(month_col))=="2027年1月"
                           or dates(row.get(month_col)) and dates(row[month_col])[0].startswith("2027-01"))
                self.assertEqual((month[month_col+1],month[month_col+2]),(1.0,1.0))
                self.assertFalse(any("每月阿里" in name or "（通告）" in name for name in book.sheets))

    def test_new_cycle_rows_use_real_table_end_and_fill_mail_periods(self):
        scope="B"; snapshot=self.service._snapshot(scope); config=snapshot["config"]
        rack=next(item for item in derive_records(config,snapshot["operations"])["racks"] if item["state"]=="off")
        group={"id":"event_new_floorplan_up","action":"上正式电","expected":"2026-09-20 11:00:00",
               "actual":"2026-09-20 10:00:00","result":"成功"}
        created=self.service.validate_op(scope,{"operation_id":"floorplan_new_up_1234","room":rack["room"],
            "rack":rack["rack"],"rack_type":rack["rack_type"],"groups":[group],"result":"成功","category":"up"})
        fields=to_fields(created); fields.update({"来源工作表":created["source"],"数据标识":"manual_floorplan_new_up"})
        operation=from_feishu({"record_id":"recFloorplanNewUp","fields":fields})
        workbook=Workbook(export_workbook((TEMPLATES/"B.xlsm").read_bytes(),config,[*snapshot["operations"],operation]))
        rows=dict(workbook.rows("机柜上电时间统计"))
        added=[(rn,row) for rn,row in rows.items() if dates(row.get(7))==["2026-09-20 10:00:00"]]
        self.assertEqual(len(added),1)
        self.assertEqual(added[0][0],1024)
        self.assertEqual(added[0][1][1],1023.0)
        self.assertEqual(max(rows),1024)
        mail=dict(workbook.rows("机柜上电汇总表"))
        day=next(row for row in mail.values() if str(row.get(13))=="2026.9.20")
        self.assertEqual(day[14],1.0)
        self.assertFalse(any(str(value).startswith("系统新增上下电") for row in mail.values() for value in row.values()))

        operation["meta"]["batch_rows"]=[{"source":"notice","source_notice":{"target_record_id":"recNotice"}}]
        from_notice=Workbook(export_workbook((TEMPLATES/"B.xlsm").read_bytes(),config,
            [*snapshot["operations"],operation]))
        day=next(row for row in dict(from_notice.rows("机柜上电汇总表")).values()
                 if str(row.get(13))=="2026.9.20")
        self.assertEqual(day[14],1.0)

    def test_manual_conversion_updates_mail_stock_without_counting_an_event(self):
        scope="C"; snapshot=self.service._snapshot(scope); config=snapshot["config"]
        config=copy.deepcopy(config)
        config["power_baseline"]=map_state_baseline((TEMPLATES/"C.xlsm").read_bytes(),config,snapshot["operations"])[0]
        baseline=config["power_baseline"]
        current={(item["room"],item["rack"]):item["state"] for item in derive_records(config,snapshot["operations"])["racks"]}
        room_rack=next(tuple(key.split("/",1)) for key,value in baseline.items()
                       if value.get("state")=="formal" and current.get(tuple(key.split("/",1)))=="formal")
        rack=next(item for item in config["inventory"] if (item["room"],item["rack"])==room_rack)
        original=Workbook(export_workbook((TEMPLATES/"C.xlsm").read_bytes(),config,snapshot["operations"]))
        item=self.service.validate_op(scope,{"operation_id":"manual_transition_1234","room":rack["room"],
            "rack":rack["rack"],"rack_type":rack["rack_type"],"groups":[{"id":"transition_event",
            "action":"正式电转测试电","expected":"2026-09-20 12:00:00","actual":"2026-09-20 12:01:00",
            "result":"成功"}],"result":"成功","category":"up"})
        fields=to_fields(item); fields.update({"来源工作表":item["source"],"数据标识":"manual_transition_record"})
        operation=from_feishu({"record_id":"recManualTransition","fields":fields})
        workbook=Workbook(export_workbook((TEMPLATES/"C.xlsm").read_bytes(),config,[*snapshot["operations"],operation]))

        def metric(book,sheet,label):
            rows=dict(book.rows(sheet))
            row=next(value for value in rows.values() if any(label in str(cell) for cell in value.values()))
            column=next(column for column,value in row.items() if label in str(value))
            return next(float(row[index]) for index in range(column+1,column+5) if isinstance(row.get(index),(int,float)))

        for sheet in ("机柜上电汇总表",):
            self.assertEqual(metric(workbook,sheet,"测试电总数"),metric(original,sheet,"测试电总数")+1)
            self.assertEqual(metric(workbook,sheet,"正式电总数"),metric(original,sheet,"正式电总数")-1)
            before=sum("2026.9.20" in str(value) for row in original.rows(sheet) for value in row[1].values())
            after=sum("2026.9.20" in str(value) for row in workbook.rows(sheet) for value in row[1].values())
            self.assertEqual(after,before)

    def test_abc_cycle_history_fits_original_up_and_down_columns(self):
        scope="A"; snapshot=self.service._snapshot(scope); config=snapshot["config"]
        rack=next(item for item in derive_records(config,snapshot["operations"])["racks"] if item["state"]=="off")
        groups=[
            {"id":"cycle_up","action":"上正式电","expected":"2026-09-20 09:00:00","actual":"2026-09-20 09:01:00","result":"成功"},
            {"id":"cycle_to_test","action":"正式电转测试电","expected":"2026-09-20 10:00:00","actual":"2026-09-20 10:01:00","result":"成功"},
            {"id":"cycle_to_formal","action":"测试电转正式电","expected":"2026-09-20 11:00:00","actual":"2026-09-20 11:01:00","result":"成功"},
        ]
        def operation(record_id,category,items):
            payload={"operation_id":record_id+"_operation","room":rack["room"],"rack":rack["rack"],
                     "rack_type":rack["rack_type"],"groups":items,"result":"成功","category":category}
            item=self.service.validate_op(scope,payload); fields=to_fields(item)
            fields.update({"来源工作表":item["source"],"数据标识":"manual_"+record_id})
            return from_feishu({"record_id":record_id,"fields":fields})
        up=operation("recCycleUp","up",groups)
        down_group={"id":"cycle_down","action":"下正式电","expected":"2026-09-20 12:00:00","actual":"2026-09-20 12:01:00","result":"成功"}
        down=operation("recCycleDown","down",[*groups,down_group])
        workbook=Workbook(export_workbook((TEMPLATES/"A.xlsm").read_bytes(),config,[*snapshot["operations"],up,down]))
        up_row=next(row for row in dict(workbook.rows("机柜上电时间统计")).values() if dates(row.get(7))==["2026-09-20 09:01:00"])
        self.assertEqual(up_row[5],"上正式电")
        self.assertIn("正式电转测试电",str(up_row[8])); self.assertIn("测试电转正式电",str(up_row[8]))
        down_row=next(row for row in dict(workbook.rows("机柜下电时间统计")).values() if dates(row.get(7))==["2026-09-20 09:01:00"])
        self.assertEqual(down_row[5],"上正式电")
        self.assertIn("正式电转测试电",str(down_row[8])); self.assertIn("测试电转正式电",str(down_row[8]))
        self.assertEqual(down_row[10],"下正式电")

    def test_saving_an_empty_existing_record_deletes_cloud_and_local_rows(self):
        rack=self.configs["A"]["inventory"][0]
        created=self.service.save_operation("A",{"operation_id":"create_then_delete_01","room":rack["room"],
            "rack":rack["rack"],"rack_type":rack["rack_type"],"result":"成功","groups":[{
            "id":"created_event","action":"上正式电","expected":"2026-09-20 09:00:00",
            "actual":"2026-09-20 09:01:00","result":"成功"}]},"owner")
        record_id=created["record_id"]
        deleted=self.service.save_operation("A",{"operation_id":"delete_existing_record_01",
            "expected_version":created["version"],"groups":[{"id":"created_event","action":"","expected":"","actual":"","result":"成功"}]},
            "owner",record_id)
        self.assertTrue(deleted["deleted"])
        self.assertNotIn(record_id,self.remote.records)
        self.assertFalse(any(item["record_id"]==record_id for item in self.service._snapshot("A")["operations"]))

    def test_abc_batch_conversion_updates_active_row_and_down_copies_cycle(self):
        snapshot=self.service._snapshot("A"); state=derive_records(snapshot["config"],snapshot["operations"])
        rack=next(item for item in state["racks"] if item["state"] in ("formal","test") and item["last_operation"])
        latest=max(((event,operation) for operation in snapshot["operations"]
                    if (operation["room"],operation["rack"])==(rack["room"],rack["rack"])
                    for event in operation["events"] if completed_state_event(event)),
                   key=lambda item:(item[0]["actual"],item[0]["id"]))
        batch={"batch_id":"batch_cycle_test","source":"manual","images":[]}
        base={"scope":"A","room":rack["room"],"rack":rack["rack"],"rack_type":rack["rack_type"],
              "supplier_rack":"","type_resolution":"keep_current","expected":"2026-09-20 20:00:00",
              "actual":"2026-09-20 20:01:00","result":"成功","failure_reason":"","attempts":[],
              "file_name":"","file_sha256":"","application_ids":[],"page":0,"source_row":1,"edits":[]}
        conversion={**base,"row_id":"row_conversion","operation_id":"operation_conversion_01",
                    "action":"正式电转测试电" if rack["state"]=="formal" else "测试电转正式电"}
        (request,record_id),duplicate=self.service.batches._row_payload(batch,conversion)
        self.assertEqual(duplicate,"")
        self.assertEqual(record_id,latest[1]["record_id"])
        self.assertEqual(request["groups"][-1]["action"],conversion["action"])
        down={**base,"row_id":"row_down","operation_id":"operation_down_cycle_01",
              "action":"下正式电" if rack["state"]=="formal" else "下测试电"}
        (request,record_id),duplicate=self.service.batches._row_payload(batch,down)
        self.assertEqual((record_id,duplicate),("",""))
        self.assertEqual(request["groups"][-1]["action"],down["action"])
        self.assertGreater(len(request["groups"]),1)

    def test_exported_summary_formula_inputs_keep_empty_counts_numeric_blank(self):
        scope="A"; original=(TEMPLATES/(scope+".xlsm")).read_bytes()
        ops=[from_feishu(record) for record in self.source_records if record["fields"]["楼栋"]==scope+"楼"]
        exported=Workbook(export_workbook(original,self.configs[scope],ops))
        for sheet in ("机柜上电汇总表",):
            cells=exported.cells(sheet)
            row=next(number for number,values in exported.rows(sheet) if values.get(2)=="2025.5.17")
            self.assertIsNone(cells[f"E{row}"].get("t"),(sheet,f"E{row}"))
            self.assertIsNone(cells[f"E{row}"].find(T("is")),(sheet,f"E{row}"))
            self.assertEqual(cells[f"F{row}"].findtext(T("f")),f"F{row-1}+C{row}-E{row}")

    def test_notice_batch_does_not_enter_export_until_confirmed(self):
        archive=FakeExportFeishu()
        self.service.export_remote=archive
        source={"target_record_id":"rec-export-chain","notice_type":"上电通告","scope":"B",
                "cabinet":"B-216运营商机房B04、B-247运营商机房B05","quantity":"2","owner_id":"owner",
                "sent_at":"2026-09-18 18:05:00"}
        self.service.batches.apply_notice_event(source)
        self.service.batches.apply_notice_event({**source,"event_action":"end","sent_at":"2026-09-18 19:05:00"})
        job=self.service.job("B","export","owner",{"batch_id":"all_"+"d"*32})
        deadline=time.time()+30
        while time.time()<deadline:
            state=self.service.job_status(job["job_id"])
            if state["status"] in ("succeeded","failed"): break
            time.sleep(.02)
        self.assertEqual(state["status"],"succeeded",state.get("error"))
        result=state["result"]
        self.assertEqual(result["cloud_upload_status"],"local_only")
        saved=self.service.local.document("B","export:"+result["export_id"])
        workbook=Workbook(Path(saved["path"]).read_bytes())
        self.assertIn("机柜上电汇总表",workbook.sheets)
        self.assertFalse(any("每月阿里统计" in name or "（通告）" in name for name in workbook.sheets))
        mail=dict(workbook.rows("机柜上电汇总表"))
        self.assertFalse(any(values.get(13)=="2026.9.18" for values in mail.values()))
        self.assertEqual(len(archive.records),0)

    def _fake_export_file(self,scope,payload,job):
        eid=uuid.uuid4().hex
        content=(scope+str(job.get("batch_id") or "")).encode()
        path=self.service.atomic_file(Path("exports")/eid/(scope+".xlsm"),content)
        export={"export_id":eid,"scope":scope,"path":path,"filename":scope+".xlsm",
                 "version":payload["snapshot"]["version"],"export_format_version":EXPORT_FORMAT_VERSION,
                 "created_at":job["created_at"],"sha256":hashlib.sha256(content).hexdigest(),
                 "owner":job.get("owner",""),"batch_id":job.get("batch_id",""),"cloud_upload_status":"local_only"}
        self.service.write("export:"+eid,export)
        return self.service._public_export(export)

    def test_export_format_upgrade_invalidates_old_job_but_notice_changes_do_not(self):
        version=self.service._snapshot("B")["version"]
        self.service.write("export:old-stale",{"export_id":"old-stale","scope":"B","path":str(Path(self.tmp.name)/"old.xlsm"),
            "filename":"old.xlsm","version":version,"created_at":"2026-09-19 08:31:00"})
        calls=[]
        self.service.do_export=lambda scope,payload,job: calls.append(payload["snapshot"]["version"]) or self._fake_export_file(scope,payload,job)
        request={"batch_id":"single_"+"f"*32}
        def wait(job):
            deadline=time.time()+10
            while time.time()<deadline:
                state=self.service.job_status(job["job_id"])
                if state["status"] in ("succeeded","failed"): break
                time.sleep(.01)
            self.assertEqual(state["status"],"succeeded",state.get("error"))
        first=self.service.job("B","export","owner",request); wait(first)
        old=self.service.local.document("B","job:"+first["job_id"])
        old.pop("export_format_version")
        self.service.write("job:"+first["job_id"],old)
        second=self.service.job("B","export","owner",request); wait(second)
        self.assertNotEqual(first["job_id"],second["job_id"])
        self.service.batches.create_from_notice({"target_record_id":"rec-no-export-effect","notice_type":"上电通告",
            "scope":"B","cabinet":"B-216运营商机房B04","quantity":"1","owner_id":"owner"})
        retry=self.service.job("B","export","owner",request)
        self.assertEqual(retry["job_id"],second["job_id"])
        self.assertEqual(len(calls),2)
        history=self.service.export_history("B")
        self.assertEqual(next(item for item in history["items"] if item["export_id"]=="old-stale")["stale_reason"],"导出格式已更新")

    def test_export_uses_local_snapshot_and_all_buildings_can_run_together(self):
        self.service.overview("D")
        self.remote.list_all=lambda path="records": (_ for _ in ()).throw(AssertionError("export must not refresh Feishu"))
        result=self.service.do_export("D",{},{}); saved=self.service.read("export:"+result["export_id"])
        self.assertTrue(Path(saved["path"]).is_file())

        barrier=threading.Barrier(5)
        self.service.do_export=lambda scope,payload,job: {"scope":scope,"passed":barrier.wait(timeout=15)}
        jobs=[self.service.job(scope,"export","owner",{}) for scope in "ABCDE"]
        deadline=time.time()+20
        while time.time()<deadline and any(self.service.job_status(job["job_id"])["status"] not in ("succeeded","failed") for job in jobs): time.sleep(.01)
        states=[self.service.job_status(job["job_id"])["status"] for job in jobs]
        self.assertEqual(states,["succeeded"]*5,[self.service.job_status(job['job_id']) for job in jobs])

    def test_completed_all_export_request_reuses_same_batch_job(self):
        calls=[]
        def generate(scope,payload,job):
            calls.append(payload["batch_id"])
            return self._fake_export_file(scope,payload,job)
        self.service.do_export=generate
        batch_id="all_"+"a"*32
        first=self.service.job("D","export","owner",{"batch_id":batch_id})
        deadline=time.time()+5
        while time.time()<deadline and self.service.job_status(first["job_id"])["status"]!="succeeded": time.sleep(.01)
        self.assertEqual(self.service.job_status(first["job_id"])["status"],"succeeded")
        retry=self.service.job("D","export","owner",{"batch_id":batch_id})
        self.assertEqual(retry["job_id"],first["job_id"])
        self.assertEqual(calls,[batch_id])
        for missing in ("cleaned","unlinked","metadata"):
            with self.subTest(missing=missing):
                result=self.service.job_status(first["job_id"])["result"]
                export=self.service.local.document("D","export:"+result["export_id"])
                if missing=="cleaned": self.service.cleanup_export("D",result["export_id"])
                elif missing=="unlinked": Path(export["path"]).unlink()
                else:
                    with self.service.local.connect("D") as conn,conn:
                        conn.execute("DELETE FROM documents WHERE key=?",("export:"+result["export_id"],))
                retry=self.service.job("D","export","owner",{"batch_id":batch_id})
                self.assertNotEqual(retry["job_id"],first["job_id"])
                deadline=time.time()+5
                while time.time()<deadline and self.service.job_status(retry["job_id"])["status"]!="succeeded": time.sleep(.01)
                self.assertEqual(self.service.job_status(retry["job_id"])["status"],"succeeded")
                first=retry
        self.assertEqual(calls,[batch_id]*4)
        second=self.service.job("D","export","owner",{"batch_id":"all_"+"b"*32})
        self.assertNotEqual(second["job_id"],first["job_id"])

    def test_export_start_response_is_immutable_when_worker_runs_immediately(self):
        def start_immediately(_func,job):
            job.update(status="running",started_at="2026-09-18 20:00:00")
        with patch.object(self.service,"snapshot",side_effect=AssertionError("export preparation must run in the worker")), \
                patch.object(self.service.local,"documents",side_effect=AssertionError("export start must not scan historical jobs")), \
                patch.object(self.service._pools["D"],"submit",side_effect=start_immediately):
            response=self.service.job("D","export","owner",{"batch_id":"all_"+"c"*32})
        self.assertEqual(response["status"],"pending")
        self.assertNotIn("started_at",response)
        self.assertNotIn("payload",response)
        self.service._running.pop(response["job_id"],None)

    def test_running_all_export_retry_keeps_original_job_after_ledger_changes(self):
        started=threading.Event(); release=threading.Event()
        def slow_export(scope,payload,job):
            started.set(); self.assertTrue(release.wait(5)); return {"scope":scope}
        self.service.do_export=slow_export
        batch_id="all_"+"e"*32
        first=self.service.job("D","export","owner",{"batch_id":batch_id})
        self.assertTrue(started.wait(5))
        try:
            with patch.object(self.service,"snapshot",side_effect=AssertionError("running job must keep its frozen snapshot")):
                retry=self.service.job("D","export","owner",{"batch_id":batch_id})
            self.assertEqual(retry["job_id"],first["job_id"])
        finally: release.set()

    def test_job_status_tolerates_short_worker_registration_gap(self):
        jid=uuid.uuid4().hex
        job={"job_id":jid,"scope":"D","kind":"export","status":"pending","pid":os.getpid(),
             "created_at":dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),"version":1,"payload":{}}
        self.service.write("job:"+jid,job)
        self.assertEqual(self.service.job_status(jid)["status"],"pending")
        job["created_at"]="2020-01-01 00:00:00"
        self.service.write("job:"+jid,job)
        self.assertEqual(self.service.job_status(jid)["status"],"failed")

    def _wait_export_batch(self,batch_id,service=None):
        service=service or self.service
        deadline=time.time()+20
        while time.time()<deadline:
            status=service.export_batch_status(batch_id,"owner")
            if status["status"] in ("succeeded","failed"): return status
            time.sleep(.03)
        self.fail("五楼导出批次未完成")

    def test_export_archive_schema_renames_attachment_without_losing_history(self):
        archive=FakeExportFeishu()
        archive.records["recLegacy"]={"record_id":"recLegacy","fields":{"导出文件":[{"file_token":"old-file"}]}}
        self.service.export_remote=archive
        self.service.ensure_export_archive_fields("2026")
        self.assertNotIn("导出文件",archive.fields)
        self.assertEqual(archive.records["recLegacy"]["fields"]["上传文件"],[{"file_token":"old-file"}])
        self.assertTrue(set(EXPORT_ARCHIVE_FIELDS)<=set(archive.fields))
        self.assertEqual([option["name"] for option in archive.fields["月份"]["property"]["options"]],
                         [f"{i:02d}" for i in range(1,13)])

    def test_five_exports_use_one_monthly_record_and_replace_it(self):
        archive=FakeExportFeishu(); self.service.export_remote=archive
        self.service.do_export=self._fake_export_file
        first="all_"+"a"*32; second="all_"+"b"*32
        self.service.start_export_batch(first,"owner",list("ABCDE"))
        saved=self._wait_export_batch(first)
        self.assertEqual(saved["status"],"succeeded",saved.get("error"))
        self.assertEqual((len(archive.records),archive.upload_calls),(1,5))
        before=next(iter(archive.records.values()))
        self.assertEqual(before["fields"]["子分类"],"机柜上下电记录")
        self.assertEqual(len(before["fields"]["上传文件"]),5)
        self.assertEqual(before["fields"]["链接"],saved["archive_url"])
        self.service.start_export_batch(second,"owner",list("ABCDE"))
        updated=self._wait_export_batch(second)
        self.assertEqual(updated["status"],"succeeded",updated.get("error"))
        self.assertEqual((len(archive.records),archive.upload_calls),(1,10))
        after=next(iter(archive.records.values()))
        self.assertEqual(after["record_id"],before["record_id"])
        self.assertEqual(after["fields"]["批次标识"],second)
        self.assertEqual(len(after["fields"]["上传文件"]),5)
        old_eid=saved["items"]["A"]["result"]["export_id"]
        self.assertEqual(self.service.local.document("A","export:"+old_eid)["cloud_upload_status"],"replaced")

    def test_export_batch_creates_new_month_without_changing_prior_archive(self):
        archive=FakeExportFeishu(); self.service.export_remote=archive
        self.service.do_export=self._fake_export_file
        now=dt.datetime.now(dt.timezone(dt.timedelta(hours=8)))
        previous_month=now.replace(day=1)-dt.timedelta(days=1)
        old={"record_id":"recPriorMonth","fields":{"子分类":"机柜上下电记录",
             "年度":f"{previous_month.year:04d}","月份":f"{previous_month.month:02d}",
             "上传文件":[{"file_token":"old-file"}]}}
        archive.records[old["record_id"]]=copy.deepcopy(old)
        batch_id="all_"+"9"*32
        self.service.start_export_batch(batch_id,"owner",list("ABCDE"))
        result=self._wait_export_batch(batch_id)
        self.assertEqual(result["status"],"succeeded",result.get("error"))
        self.assertEqual(len(archive.records),2)
        self.assertEqual(archive.records[old["record_id"]],old)

    def test_export_batch_retries_link_without_reupload_and_reconciles_create(self):
        archive=FakeExportFeishu(); archive.fail_after_create=True; archive.fail_share=True
        self.service.export_remote=archive; self.service.do_export=self._fake_export_file
        batch_id="all_"+"c"*32
        self.service.start_export_batch(batch_id,"owner",list("ABCDE"))
        partial=self._wait_export_batch(batch_id)
        self.assertEqual((partial["status"],partial["phase"]),("failed","linking"))
        self.assertEqual((archive.upload_calls,archive.creates,len(archive.records)),(5,1,1))
        archive.fail_share=False; archive.fail_after_create=False
        self.service.start_export_batch(batch_id,"owner",list("ABCDE"))
        recovered=self._wait_export_batch(batch_id)
        self.assertEqual(recovered["status"],"succeeded",recovered.get("error"))
        self.assertEqual((archive.upload_calls,archive.creates,len(archive.records)),(5,1,1))

    def test_export_batch_waits_for_all_buildings_and_retries_only_failure(self):
        archive=FakeExportFeishu(); self.service.export_remote=archive
        failed=[]
        def export(scope,payload,job):
            if scope=="C" and not failed:
                failed.append(scope); raise CabinetError("C楼暂时不可用")
            return self._fake_export_file(scope,payload,job)
        self.service.do_export=export
        batch_id="all_"+"d"*32
        self.service.start_export_batch(batch_id,"owner",list("ABCDE"))
        first=self._wait_export_batch(batch_id)
        self.assertEqual(first["status"],"failed")
        self.assertEqual((archive.upload_calls,len(archive.records)),(0,0))
        self.assertNotIn("导出文件",archive.fields)
        self.assertTrue({"上传文件","子分类","年度","月份","链接"}<=set(archive.fields))
        old_a=first["items"]["A"]["result"]["export_id"]
        self.service.start_export_batch(batch_id,"owner",list("ABCDE"))
        second=self._wait_export_batch(batch_id)
        self.assertEqual(second["status"],"succeeded",second.get("error"))
        self.assertEqual(second["items"]["A"]["result"]["export_id"],old_a)
        self.assertEqual((archive.upload_calls,len(archive.records)),(5,1))

    def test_export_batch_rejects_duplicate_month_rows_and_preserves_files(self):
        archive=FakeExportFeishu(); self.service.export_remote=archive; self.service.do_export=self._fake_export_file
        now=dt.datetime.now(dt.timezone(dt.timedelta(hours=8)))
        for index in range(2):
            archive.records[f"recDuplicate{index}"]={"record_id":f"recDuplicate{index}","fields":{
                "子分类":"机柜上下电记录","年度":f"{now.year:04d}","月份":f"{now.month:02d}"}}
        batch_id="all_"+"f"*32
        self.service.start_export_batch(batch_id,"owner",list("ABCDE"))
        result=self._wait_export_batch(batch_id)
        self.assertEqual(result["status"],"failed")
        self.assertIn("多条",result["error"])
        self.assertEqual(len(archive.records),2)
        self.assertTrue(all(item["status"]=="succeeded" for item in result["items"].values()))

    def test_export_batch_resume_on_new_service_reuses_uploaded_tokens(self):
        archive=FakeExportFeishu(); archive.fail_share=True
        self.service.export_remote=archive; self.service.do_export=self._fake_export_file
        batch_id="all_"+"0"*32
        self.service.start_export_batch(batch_id,"owner",list("ABCDE"))
        self.assertEqual(self._wait_export_batch(batch_id)["phase"],"linking")
        fresh=CabinetPowerService(self.store,self.remote,self.tmp.name,export_remote=archive)
        try:
            archive.fail_share=False
            fresh.start_export_batch(batch_id,"owner",list("ABCDE"))
            recovered=self._wait_export_batch(batch_id,fresh)
            self.assertEqual(recovered["status"],"succeeded",recovered.get("error"))
            self.assertEqual((archive.upload_calls,len(archive.records)),(5,1))
        finally: fresh.shutdown()

    def test_export_cleanup_refuses_running_batch(self):
        eid="cleanup-during-batch"; batch_id="all_"+"e"*32
        path=self.service.atomic_file(Path("exports")/eid/"test.xlsm",b"test")
        self.service.write("export:"+eid,{"export_id":eid,"scope":"D","path":path,"filename":"test.xlsm","batch_id":batch_id})
        self.store.put_document("cabinet_export_batches",batch_id,{"status":"running"})
        with self.assertRaisesRegex(CabinetError,"批次仍在处理"):
            self.service.cleanup_export("D",eid)
        self.store.put_document("cabinet_export_batches",batch_id,{"status":"succeeded"})
        self.service.cleanup_export("D",eid)
        self.assertFalse(Path(path).exists())
        history=next(item for item in self.service.export_history("D")["items"] if item["export_id"]==eid)
        self.assertFalse(history["file_available"])

    def test_export_history_is_paged_in_sqlite_order(self):
        version=self.service.local.version("D")
        for index in range(25):
            eid=f"paged-{index:02d}"
            self.service.write("export:"+eid,{"export_id":eid,"scope":"D","path":str(Path(self.tmp.name)/(eid+".xlsm")),
                "filename":eid+".xlsm","version":version,"export_format_version":EXPORT_FORMAT_VERSION,
                "created_at":f"2026-09-{index+1:02d} 10:00:00"})
        page=self.service.export_history("D",2,10)
        self.assertEqual((page["total"],page["page"],page["page_size"]),(25,2,10))
        self.assertEqual([item["export_id"] for item in page["items"]],
                         [f"paged-{index:02d}" for index in range(14,4,-1)])
        self.assertEqual(page["current"]["export_id"],"paged-24")

    def test_export_remote_edit_overrides_original_cells_and_keeps_conversion(self):
        original=(TEMPLATES/"E.xlsm").read_bytes()
        rows=[from_feishu(r) for r in self.source_records if r["fields"]["楼栋"]=="E楼"]
        row=next(o for o in rows if o["source_row"]==2)
        row["groups"][0]["actual"]="2026-08-01 09:00:00"; row["groups"][0]["action"]="下正式电"
        row=from_feishu({"record_id":row["record_id"],"fields":{**row["raw_fields"],**to_fields(row)}})
        rows=[row if o["record_id"]==row["record_id"] else o for o in rows]
        exported=Workbook(export_workbook(original,self.configs["E"],rows))
        cells=exported.cells("机柜上下电时间统计")
        self.assertEqual(exported.value(cells["E2"]),"下正式电")
        self.assertEqual(dates(exported.value(cells["G2"])),["2026-08-01 09:00:00"])
        self.assertEqual(exported.value(cells["K2"]),"上正式电")
        self.assertEqual(exported.value(cells["N2"]),"正式电转测试电")

    def test_removed_import_routes_and_scope_permissions(self):
        from fastapi import FastAPI
        from fastapi.responses import JSONResponse
        from fastapi.testclient import TestClient
        from .lan_bitable_template_portal.cabinet_power_routes import install_cabinet_power_routes
        class Controller:
            def _current_session(self,request): return {"open_id":"test"} if request.headers.get("x-test-login") else None
            def _auth_required_response(self): return JSONResponse({},status_code=401)
            def _json_ok(self,request,session,data): return JSONResponse({"ok":True,"data":data})
            def _portal_error_response(self,exc,default_status): return JSONResponse({"error":str(exc)},status_code=default_status)
            async def _read_json_request(self,request,max_bytes): return await request.json()
        app=FastAPI(); controller=Controller()
        runtime=SimpleNamespace(state_store=self.store,auth_manager=SimpleNamespace(is_admin=lambda s:False,scope_allowed=lambda s,scope:scope=="A"))
        install_cabinet_power_routes(app,controller,runtime)
        self.assertIn("/api/cabinet-power/exports/{export_id}/upload",{route.path for route in app.routes})
        controller._cabinet_power.remote=self.remote; controller._cabinet_power._directory=self.service._directory
        controller._cabinet_power.root=Path(self.tmp.name)
        controller._cabinet_power.local=self.service.local
        with TestClient(app) as client:
            self.assertEqual(client.get("/api/cabinet-power/buildings").status_code,401)
            self.assertEqual(client.get("/api/cabinet-power/overview?scope=B",headers={"x-test-login":"1"}).status_code,403)
            self.assertEqual(client.post("/api/cabinet-power/imports/commit",headers={"x-test-login":"1"}).status_code,404)
            bootstrap=client.post("/api/cabinet-power/bootstrap",headers={"x-test-login":"1"},json={"scope":"A"}).json()["data"]
            self.assertEqual((bootstrap["status"],bootstrap["ready"]),("succeeded",5))
            data=client.get("/api/cabinet-power/overview?scope=A",headers={"x-test-login":"1"}).json()["data"]
            self.assertEqual(data["record_count"],1031)
            self.assertEqual(client.get("/api/cabinet-power/batches",headers={"x-test-login":"1"}).status_code,200)
            rack=self.configs["A"]["inventory"][0]
            created=client.post("/api/cabinet-power/batches",headers={"x-test-login":"1"},json={"rows":[{"scope":"A","room":rack["room"],"rack":rack["rack"],"rack_type":rack["rack_type"],"action":"上正式电","expected":"2026-09-14 12:00:00","actual":"2026-09-14 12:00:00","result":"成功"}]}).json()["data"]
            self.assertEqual(created["stats"]["total"],1)
            self.assertEqual(client.get("/api/cabinet-power/batches/"+created["batch_id"],headers={"x-test-login":"1"}).status_code,200)
            compact=client.get("/api/cabinet-power/batches/"+created["batch_id"]+"/status",headers={"x-test-login":"1"}).json()["data"]
            self.assertEqual((compact["batch_id"],compact["version"],compact["status"]),(created["batch_id"],created["version"],created["status"]))
            self.assertNotIn("rows",compact)
            text=f"机柜上测试电确认单\n申请时间： 2026-09-14 10:00:00\n申请单号：[Z260914001234567890] 操作类型： 上测试电\nEA118  A{rack['room'][0]}-{int(rack['room'][1:])}.EA118  {rack['rack']}  {rack['rack']}  {rack['rack_type']}  2UR  2026-09-14 10:00:00  成功  2026-09-14 09:59:00"
            class Page:
                def extract_text(self,**_kwargs): return text
            controller._cabinet_power.batches._pdf_reader=lambda _path:SimpleNamespace(is_encrypted=False,pages=[Page()])
            recognized=client.post("/api/cabinet-power/batches/recognize",headers={"x-test-login":"1"},files=[("files",("sample.pdf",b"%PDF-route","application/pdf"))])
            self.assertEqual(recognized.status_code,202)
            recognized_id=recognized.json()["data"]["batch_id"]
            deadline=time.time()+5
            while time.time()<deadline:
                detail=client.get("/api/cabinet-power/batches/"+recognized_id,headers={"x-test-login":"1"}).json()["data"]
                if detail["status"]!="recognizing": break
                time.sleep(.01)
            self.assertEqual(detail["stats"]["total"],1)
            self.assertEqual(detail["rows"][0]["supplier_rack"],rack["rack"])

if __name__=="__main__": unittest.main()
