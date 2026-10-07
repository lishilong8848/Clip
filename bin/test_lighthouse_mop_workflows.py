"""Isolated MOP fill workflow test (engineer maintenance sheet, native_mop form).

Exercises the real PortalAgent MOP prepare -> field_options -> amend -> confirm
(twice, high risk) -> _execute flow through fake FastAPI routes, driving the
native MaintenancePortalService._parse_xlsx_preview / _extract_mop_sheet_targets /
fill_engineer_mop_file (real Excel fill + signature anchors).

No real network / credentials / business data.  Only the directory, signature
people sources and cloud transport are stubbed on an isolated service subclass.
"""
from __future__ import annotations

import asyncio
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import Workbook

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse

_TOP = Path(__file__).resolve().parent
sys.path.insert(0, str(_TOP))

from lan_bitable_template_portal.lighthouse_agent import PortalAgent  # noqa: E402
from lan_bitable_template_portal.lighthouse_ai import AssistantError, LighthouseAssistant  # noqa: E402
from lan_bitable_template_portal.lighthouse_api import PortalAPICatalog  # noqa: E402
from lan_bitable_template_portal.lighthouse_files import LighthouseFiles  # noqa: E402
from lan_bitable_template_portal.portal_service import (  # noqa: E402
    MaintenancePortalService,
    PortalError,
)
from lan_bitable_template_portal.state_store import LanPortalStateStore  # noqa: E402
from test_drill_management import _signature_png  # noqa: E402
from test_lighthouse_stream import Store  # noqa: E402


def _public_contains(public, private: str) -> bool:
    """Recursively check whether any string inside public contains the private path."""
    needle = private.replace("\\", "/")
    if isinstance(public, str):
        return needle in public.replace("\\", "/")
    if isinstance(public, dict):
        return any(_public_contains(value, private) for value in public.values())
    if isinstance(public, (list, tuple)):
        return any(_public_contains(value, private) for value in public)
    return False


STAFF = [
    {
        "source": "staff", "record_id": "person-impl", "name": "王实施", "open_id": "ou_impl",
        "employee_no": "P001", "building": "A楼", "can_receive_message": True, "has_signature": True,
    },
    {
        "source": "staff", "record_id": "person-aud", "name": "李审核", "open_id": "ou_aud",
        "employee_no": "P002", "building": "A楼", "can_receive_message": True, "has_signature": True,
    },
]
TEMP = [
    {
        "source": "temporary", "temp_id": "temp-1", "name": "临时工甲", "role": "implementer",
        "building": "A楼", "has_signature": True,
    },
]


def _build_workbook_bytes() -> bytes:
    wb = Workbook()
    cover = wb.active
    cover.title = "封面"
    cover["A1"] = "维保单封面"
    ws = wb.create_sheet("维护单")
    ws["A1"] = "维护实施人："
    ws.merge_cells("A1:B1")
    ws["D1"] = "维护审核人："
    ws.merge_cells("D1:E1")
    # Signature value areas are the untouched single cells C1 (implementer) / F1 (auditor).
    ws["A2"] = "设备：A栋1号电梯"
    ws["A3"] = "维护开始时间："
    ws.merge_cells("A3:B3")
    ws["A4"] = "维护完成情况：□正常 □异常"
    ws["A5"] = "维护完成时间："
    ws.merge_cells("A5:B5")
    ws["A6"] = "备注"
    blob = io.BytesIO()
    wb.save(blob)
    return blob.getvalue()


class IsolatedMopService(MaintenancePortalService):
    def __init__(self, tmpdir: Path):
        super().__init__()
        self._state_store = LanPortalStateStore(str(tmpdir / "state.sqlite3"))
        self._signature_management = SimpleNamespace(references=lambda items: items)
        self._fill_calls = 0
        self._upload_calls = 0

    def _load_signature_people(self, *, force: bool = False) -> list[dict]:
        return list(STAFF)

    def _load_external_signature_people(self, *, force: bool = False) -> list[dict]:
        return []

    def signature_image_bytes(self, *, record_id: str):
        return _signature_png(), "image/png"

    def temporary_signature_image_bytes(self, *, temp_id: str):
        return _signature_png(), "image/png"

    def external_signature_image_bytes(self, *, record_id: str):
        return _signature_png(), "image/png"

    def _normalize_scope(self, scope: str) -> str:
        return str(scope or "").strip().upper()

    def _upload_bitable_file(self, **kwargs):
        self._upload_calls += 1
        raise AssertionError("MOP fill must not upload to cloud transport")

    def _patch_record_fields(self, **kwargs):
        raise AssertionError("MOP fill must not patch bitable records")


class MopWorkflowTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmpdir_obj = tempfile.TemporaryDirectory()
        self.tmp = Path(self.tmpdir_obj.name)
        self.store = Store(self.tmp / "assistant_state.sqlite3")
        self.model = Mock()
        self.model.settings.return_value = {}
        self.model.profile.return_value = {}
        self.assistant = LighthouseAssistant(self.store, Mock(return_value=([], [])), model=self.model)
        self.service = IsolatedMopService(self.tmp)
        self._data_patch = patch(
            "lan_bitable_template_portal.portal_service.get_data_file_path",
            side_effect=lambda name: str(self.tmp / (name or "")),
        )
        self._data_patch.start()
        self.addCleanup(self._data_patch.stop)

        self.content = _build_workbook_bytes()
        self.source_path = self.tmp / "mop_source.xlsx"
        self.source_path.write_bytes(self.content)
        parsed = self.service._parse_xlsx_preview(self.content)
        self.snapshot = dict(parsed)
        self.snapshot.update(
            {
                "local_file": {"path": str(self.source_path), "file_name": "mop_source.xlsx", "size": len(self.content)},
                "attachment": {"file_token": "mop-file-1", "name": "mop_test.xlsx"},
                "mop_record_id": "mop-1",
                "mop_title": "10月维保",
                "mop_file_name": "mop_test.xlsx",
            }
        )

        self.app = FastAPI()
        self.sig_calls: list[tuple] = []
        self.binary_response = False

        @self.app.get("/api/signatures/people")
        async def sig_people(scope: str, q: str = "", limit: int = 100, notice_key: str = ""):
            self.sig_calls.append(("staff", scope, notice_key))
            return {"ok": True, "data": {"people": STAFF}}

        @self.app.get("/api/signatures/temporary/people")
        async def sig_temp(scope: str, q: str = "", limit: int = 100, notice_key: str = ""):
            self.sig_calls.append(("temporary", scope, notice_key))
            return {"ok": True, "data": {"people": TEMP}}

        @self.app.post("/api/engineer/mop/fill")
        async def mop_fill(payload: dict, request: Request):
            if payload.get("scope") not in ("A", "ALL"):
                raise HTTPException(status_code=403, detail="无权填写该范围维护单。")
            kwargs = {
                "scope": payload.get("scope"),
                "notice_key": payload.get("notice_key", ""),
                "signature_context_key": payload.get("signature_context_key", ""),
                "operator_open_id": "ou_operator",
                "local_file_path": payload.get("local_file_path"),
                "mop_record_id": payload.get("mop_record_id", ""),
                "mop_title": payload.get("mop_title", ""),
                "mop_file_name": payload.get("mop_file_name", ""),
                "sheet_name": payload.get("sheet_name"),
                "fields": payload["fields"],
                "checkboxes": payload["checkboxes"],
                "cell_edits": payload["cell_edits"],
                "signatures": payload["signatures"],
            }
            try:
                result = self.service.fill_engineer_mop_file(**kwargs)
                self.service._fill_calls += 1
                if self.binary_response and request.query_params.get("download") == "1":
                    return FileResponse(result["path"], filename=result["file_name"])
                return {"ok": True, "data": result}
            except PortalError as exc:
                raise HTTPException(status_code=400, detail=str(exc)) from None

        self.catalog = PortalAPICatalog(self.app)
        self.portal = PortalAgent(self.assistant, self.catalog, LighthouseFiles(self.store))
        self.actor = {"id": "actor-a", "scopes": ["A"], "is_admin": False}
        self.request = Request(
            scope={
                "type": "http",
                "method": "NOOP",
                "path": "/",
                "scheme": "http",
                "server": ("testserver", 80),
                "headers": [],
                "query_string": b"",
                "client": ("127.0.0.1", 12345),
            },
            receive=None,
            send=None,
        )

    async def asyncTearDown(self):
        self.tmpdir_obj.cleanup()

    def _decision(self, *, scope: str = "A", mop_record_id: str = "mop-1") -> dict:
        return {
            "title": "维护单填写",
            "explanation": "填写10月维保维护单",
            "notice_key": "notice-1",
            "scopes": ["A"],
            "operations": [
                {
                    "id": "mop_fill",
                    "api_id": "POST /api/engineer/mop/fill",
                    "body": {
                        "scope": scope,
                        "mop_record_id": mop_record_id,
                        "mop_title": "10月维保",
                        "mop_file_name": "mop_test.xlsx",
                        "notice_key": "notice-1",
                        "fields": [],
                        "checkboxes": [],
                        "cell_edits": [],
                        "signatures": [],
                    },
                }
            ],
        }

    def _find_ref(self, plan, record_id: str, role: str) -> str:
        for ref, binding in (plan.get("_references") or {}).items():
            value = binding.get("value") or {}
            if binding.get("field") == "signatures" and value.get("record_id") == record_id and value.get("role") == role:
                return ref
        raise AssertionError(f"missing signer ref {record_id}/{role}")

    def _confirm_staff_usage(self, context_key: str):
        for role, rid, oid, name in (
            ("implementer", "person-impl", "ou_impl", "王实施"),
            ("auditor", "person-aud", "ou_aud", "李审核"),
        ):
            record = self.service._state_store.create_mop_signature_usage_confirmation(
                scope="A",
                notice_key=context_key,
                role=role,
                signer_record_id=rid,
                signer_open_id=oid,
                signer_name=name,
                requested_by_openid="ou_operator",
                requested_by_name="操作人",
            )
            self.service._state_store.confirm_mop_signature_usage(token=record["token"])

    async def _run_happy_flow(self, confirm_signatures: bool = True):
        self.sig_calls.clear()
        plan = self.portal.prepare(
            self.actor, self._decision(), "turn-mop", file_ids=[], references={}, queries={"q_mop": self.snapshot}
        )
        self.assertEqual(plan["status"], "needs_input")
        self.assertEqual(plan["scopes"], ["A"])
        identity = plan["id"]
        op = plan["operations"][0]
        self.assertEqual(op["body"]["scope"], "A")
        self.assertIn("$reference", op["body"].get("local_file_path", {}))
        self.assertIn("$reference", op["body"].get("signature_context_key", {}))

        # field_options loads staff + temporary people via native invoke, keeps scope.
        pub = await self.portal.field_options(self.actor, identity, "step0.mop", self.request)
        self.assertEqual(pub["status"], "needs_input")
        version = pub["version"]
        full = self.portal.get_plan(self.actor, identity)
        self.assertEqual([c[:2] for c in self.sig_calls], [("staff", "A"), ("temporary", "A")])
        impl_ref = self._find_ref(full, "person-impl", "implementer")
        aud_ref = self._find_ref(full, "person-aud", "auditor")
        context_key = op["body"]["signature_context_key"]["$reference"]
        context_value = full["_references"][context_key]["value"]
        self.assertEqual(context_value, "notice-1|mop:mop-1|attachment:mop-file-1")

        values = {
            "sheet_name": "维护单",
            "sheet_0": {
                "fields": {"field_0": "2026-10-02T10:30", "field_1": "2026-10-02T16:00"},
                "checkboxes": {"check_0": "正常"},
                "cell_edits": [{"row": 6, "column": "C", "value": "已按时完成维护"}],
                "implementer": [impl_ref],
                "auditor": [aud_ref],
            },
        }
        amend_pub = self.portal.amend(
            self.actor, identity, {"version": version, "values": {"step0.mop": values}}
        )
        self.assertEqual(amend_pub["status"], "awaiting_confirmation")
        ver_after_amend = amend_pub["version"]

        plan_after = self.portal.get_plan(self.actor, identity)
        self.assertEqual(plan_after["risk"], "high")
        if confirm_signatures:
            self._confirm_staff_usage(context_value)

        rev = await self.portal.confirm(
            self.actor, identity, {"version": ver_after_amend, "stage": "review"}, self.request
        )
        self.assertEqual(rev["status"], "awaiting_second_confirmation")
        ver2 = rev["version"]
        exc = await self.portal.confirm(
            self.actor, identity, {"version": ver2, "stage": "execute"}, self.request
        )
        self.assertEqual(exc["status"], "running")
        return identity

    async def test_mop_fill_happy_path(self):
        identity = await self._run_happy_flow()
        plan = await self._wait_completed(identity)
        self.assertEqual(plan["status"], "completed")
        self.assertEqual(self.service._fill_calls, 1)
        self.assertEqual(self.service._upload_calls, 0)

        result = plan["results"][0]["_raw"]
        output_path = Path(result["path"])
        self.assertTrue(output_path.exists())
        self.assertTrue(output_path.name.endswith(".xlsx"))
        self.assertGreaterEqual(result["inserted"], 2)

        from openpyxl import load_workbook

        wb = load_workbook(str(output_path))
        ws = wb["维护单"]
        self.assertEqual(ws["C3"].value, "2026年10月2日10时30分")
        self.assertEqual(ws["A5"].value, "2026年10月2日16时00分")
        self.assertIn("☑正常", str(ws["A4"].value))
        self.assertEqual(ws["C6"].value, "已按时完成维护")
        self.assertGreaterEqual(len(ws._images), 2)

        # Public plan must not leak private file path / open_id / signature pixels.
        public_plan = self.portal.public_plan(plan)
        public_json = json.dumps(public_plan, ensure_ascii=False)
        for private in (str(self.source_path), result["path"], result.get("relative_path", "")):
            self.assertFalse(_public_contains(public_plan, private), "private path leaked into public plan")
        self.assertNotIn("ou_impl", public_json)
        self.assertNotIn("ou_aud", public_json)
        self.assertNotIn("iVBORw0KGgo", public_json)  # PNG magic -> base64 pixels

    async def test_native_fill_download_returns_authorized_assistant_file(self):
        self.binary_response = True
        identity = await self._run_happy_flow()
        plan = await self._wait_completed(identity)
        self.assertEqual(plan["status"], "completed", plan.get("error"))
        file = plan["results"][0]["data"]
        self.assertTrue(file["url"].startswith("/api/assistant/files/"))
        self.assertTrue(file["name"].endswith(".xlsx"))
        saved = self.portal.files.get(self.actor, file["id"])
        self.assertTrue(Path(saved["path"]).is_file())
        with self.assertRaises(AssistantError):
            self.portal.files.get({**self.actor, "id": "different-owner"}, file["id"])
        public = self.portal.public_plan(plan)
        self.assertFalse(_public_contains(public, saved["path"]))
        self.assertEqual(public["results"][0]["data"]["url"], file["url"])

    async def _wait_completed(self, identity: str):
        for _ in range(200):
            plan = self.portal.get_plan(self.actor, identity)
            if plan["status"] in {"completed", "failed"}:
                return plan
            await asyncio.sleep(0.05)
        raise AssertionError("workflow did not finish")

    async def test_prepare_requires_matching_binding(self):
        # No snapshot bound -> prepare must refuse.
        with self.assertRaises(AssistantError) as cm:
            self.portal.prepare(self.actor, self._decision(), "turn-mop", file_ids=[], references={}, queries={})
        self.assertIn("维护单", str(cm.exception))

        # Foreign scope -> 403.
        with self.assertRaises(AssistantError) as cm:
            self.portal.prepare(
                self.actor, self._decision(scope="B"), "turn-mop", file_ids=[], references={}, queries={"q_mop": self.snapshot}
            )
        self.assertIn("无权", str(cm.exception))
        self.assertEqual(cm.exception.status, 403)

    async def test_fail_invalid_edits_and_forged_signer(self):
        plan = self.portal.prepare(
            self.actor, self._decision(), "turn-neg", file_ids=[], references={}, queries={"q_mop": self.snapshot}
        )
        identity = plan["id"]
        pub = await self.portal.field_options(self.actor, identity, "step0.mop", self.request)
        version = pub["version"]
        full = self.portal.get_plan(self.actor, identity)
        impl_ref = self._find_ref(full, "person-impl", "implementer")
        aud_ref = self._find_ref(full, "person-aud", "auditor")
        base = {
            "sheet_name": "维护单",
            "sheet_0": {
                "fields": {"field_0": "2026-10-02T10:30", "field_1": "2026-10-02T16:00"},
                "checkboxes": {"check_0": "正常"},
                "cell_edits": [{"row": 6, "column": "C", "value": "x"}],
                "implementer": [impl_ref],
                "auditor": [aud_ref],
            },
        }

        def amend_with(vals_overrides, ver):
            values = json.loads(json.dumps(base))
            values.update(vals_overrides)
            self.portal.amend(self.actor, identity, {"version": ver, "values": {"step0.mop": values}})

        with self.assertRaises(AssistantError) as cm:
            amend_with({"sheet_name": "不存在的表"}, version)
        self.assertIn("工作表", str(cm.exception))

        with self.assertRaises(AssistantError) as cm:
            amend_with({"sheet_0": {**base["sheet_0"], "implementer": ["forged-ref"]}}, version)
        self.assertIn("签署人员", str(cm.exception))

        with self.assertRaises(AssistantError) as cm:
            amend_with({"sheet_0": {**base["sheet_0"], "cell_edits": [{"row": 99, "column": "C", "value": "x"}]}}, version)
        self.assertIn("单元格位置", str(cm.exception))

        with self.assertRaises(AssistantError) as cm:
            amend_with({"sheet_0": {**base["sheet_0"], "cell_edits": [{"row": 1, "column": "C", "value": "x"}]}}, version)
        self.assertIn("普通单元格", str(cm.exception))

        # readonly scope cannot be tampered through the form.
        with self.assertRaises(AssistantError) as cm:
            amend_with({"scope": "B"}, version)
        self.assertIn("只读", str(cm.exception))

    async def test_unconfirmed_signer_stops_then_local_consent_allows(self):
        identity = await self._run_happy_flow(confirm_signatures=False)
        plan = await self._wait_completed(identity)
        self.assertEqual(plan["status"], "failed")

        # Local native consent now permits the same workflow on a fresh plan.
        self._confirm_staff_usage("notice-1|mop:mop-1|attachment:mop-file-1")

        identity2 = await self._run_happy_flow(confirm_signatures=True)
        plan2 = await self._wait_completed(identity2)
        self.assertEqual(plan2["status"], "completed")


if __name__ == "__main__":
    unittest.main(verbosity=2)
