"""Offline extraction checks: render scanned PDFs locally, never request a model."""
import io
import unittest
import zipfile
from unittest.mock import Mock, patch

from bin.test_lighthouse_knowledge_extract import make_pdf_text, make_docx, make_xlsx
from openclaw_service.assistant.lighthouse_ai import AssistantError
from openclaw_service.assistant.lighthouse_knowledge_extract import extract_sections, SUPPORTED_EXTENSIONS
from openclaw_service.assistant.lighthouse_knowledge import EXTENSIONS


def scanned_pdf():
    from PIL import Image, ImageDraw
    with Image.new('RGB', (400, 300), 'white') as image:
        ImageDraw.Draw(image).text((30, 30), 'SCAN PAGE TEST 1234', fill='black')
        result = io.BytesIO()
        image.save(result, format='PDF')
    return result.getvalue()


class DocumentFormatTests(unittest.TestCase):
    def setUp(self):
        self.enterContext(patch('socket.socket.connect', side_effect=AssertionError('No network')))

    def test_extension_registry_is_shared(self):
        self.assertEqual(EXTENSIONS, SUPPORTED_EXTENSIONS)

    def test_scanned_pdf_is_rendered_as_an_image_for_ocr(self):
        seen = []
        def ocr(content):
            from PIL import Image
            with Image.open(io.BytesIO(content)) as image:
                self.assertLessEqual(max(image.size), 2400)
                self.assertLess(image.convert('L').getextrema()[0], 240)
            seen.append(content)
            return '设备检查时间 08:30'
        sections = extract_sections(scanned_pdf(), '扫描.pdf', ocr=ocr)
        self.assertEqual(len(seen), 1)
        self.assertEqual(sections[0]['text'], '设备检查时间 08:30')
        self.assertIn('第1页（OCR）', sections[0]['location'])

    def test_text_and_scan_pdf_keeps_all_page_numbers(self):
        from pypdf import PdfReader, PdfWriter
        writer = PdfWriter()
        writer.append(PdfReader(io.BytesIO(make_pdf_text('Printed text'))))
        writer.add_blank_page(width=300, height=300)
        writer.append(PdfReader(io.BytesIO(scanned_pdf())))
        result = io.BytesIO()
        writer.write(result)
        ocr = Mock(return_value='扫描正文')
        sections = extract_sections(result.getvalue(), '混合.pdf', ocr=ocr)
        self.assertEqual([s['text'] for s in sections], ['Printed text', '扫描正文'])
        self.assertIn('第3页', sections[1]['location'])
        ocr.assert_called_once()

    def test_ocr_failure_and_empty_result_are_not_partial_success(self):
        for ocr in (Mock(side_effect=RuntimeError('failure')), Mock(return_value='')):
            with self.subTest(ocr=ocr), self.assertRaisesRegex(AssistantError, '第1页OCR'):
                extract_sections(scanned_pdf(), '扫描.pdf', ocr=ocr)

    def test_existing_text_pdf_never_uses_ocr(self):
        ocr = Mock(side_effect=AssertionError('No OCR for text PDF'))
        self.assertEqual(extract_sections(make_pdf_text('Text content'), '原文.pdf', ocr=ocr)[0]['text'], 'Text content')

    def test_office_variants_use_text_parsers_without_executing_macros(self):
        for suffix in ('.docm', '.dotx', '.dotm'):
            with self.subTest(suffix=suffix):
                self.assertIn('维保说明', extract_sections(make_docx(['维保说明']), '模板' + suffix)[0]['text'])
        for suffix in ('.xltx', '.xltm'):
            with self.subTest(suffix=suffix):
                self.assertIn('设备', extract_sections(make_xlsx([['设备'], ['AHU-1']]), '模板' + suffix)[0]['text'])

    def test_presentation_variant_retains_slide_location(self):
        from pptx import Presentation
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = '专项检查'
        deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = '设备要求'
        first = deck.slides._sldIdLst[0]
        deck.slides._sldIdLst.remove(first)
        deck.slides._sldIdLst.append(first)
        result = io.BytesIO()
        deck.save(result)
        for suffix in ('.pptm', '.ppsx', '.potx'):
            with self.subTest(suffix=suffix):
                types = {'.pptm': b'application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml',
                    '.ppsx': b'application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml',
                    '.potx': b'application/vnd.openxmlformats-officedocument.presentationml.template.main+xml'}
                variant = io.BytesIO()
                with zipfile.ZipFile(io.BytesIO(result.getvalue())) as source, zipfile.ZipFile(variant, 'w') as target:
                    for item in source.infolist():
                        body = source.read(item)
                        if item.filename == '[Content_Types].xml':
                            body = body.replace(b'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml', types[suffix])
                        target.writestr(item, body)
                sections = extract_sections(variant.getvalue(), '材料' + suffix)
                self.assertIn('设备要求', sections[0]['text'])
                self.assertIn('专项检查', sections[1]['text'])
                self.assertIn('幻灯片1', sections[0]['location'])

    def test_utf16_and_text_formats(self):
        for suffix in ('.txt', '.md', '.log', '.json', '.yaml', '.tsv'):
            self.assertEqual(extract_sections('设备正常'.encode('utf-16'), '说明' + suffix)[0]['text'], '设备正常')

    def test_html_no_active_content_or_remote_fetch(self):
        result = extract_sections(b'<p>Visible text</p><script>hidden()</script><img src="https://fixture.invalid/a.png">', '页面.html')
        self.assertIn('Visible text', result[0]['text'])
        self.assertNotIn('hidden', result[0]['text'])

    def test_opendocument_text(self):
        result = io.BytesIO()
        with zipfile.ZipFile(result, 'w') as archive:
            archive.writestr('content.xml', '<root xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><text:p>原始内容</text:p></root>')
        for suffix in ('.odt', '.ods', '.odp'):
            self.assertIn('原始内容', extract_sections(result.getvalue(), '文档' + suffix)[0]['text'])

    def test_epub_uses_existing_converter_without_remote_fetch(self):
        result = io.BytesIO()
        with zipfile.ZipFile(result, 'w') as archive:
            archive.writestr('META-INF/container.xml', '<container><rootfiles><rootfile full-path="content.opf"/></rootfiles></container>')
            archive.writestr('content.opf', '<package xmlns:dc="http://purl.org/dc/elements/1.1/"><metadata><dc:title>技术手册</dc:title></metadata><manifest><item id="p1" href="chapter.html"/></manifest><spine><itemref idref="p1"/></spine></package>')
            archive.writestr('chapter.html', '<p>设备操作方法</p>')
        self.assertIn('设备操作方法', '\n'.join(s['text'] for s in extract_sections(result.getvalue(), '资料.epub')))

    def test_tiff_all_pages_are_read(self):
        from PIL import Image
        output = io.BytesIO()
        with Image.new('RGB', (40, 40), 'white') as first, Image.new('RGB', (40, 40), 'black') as second:
            first.save(output, format='TIFF', save_all=True, append_images=[second])
        ocr = Mock(side_effect=['第一页', '第二页'])
        result = extract_sections(output.getvalue(), '扫描.tiff', ocr=ocr)
        self.assertEqual([s['text'] for s in result], ['第一页', '第二页'])
        self.assertEqual(ocr.call_count, 2)

    def test_unsupported_binary_is_not_decoded_as_text(self):
        with self.assertRaisesRegex(AssistantError, '不支持的文件类型'):
            extract_sections(b'\x00\xff\x01binary', 'old.doc')
