"""Synthetic attachment and chart checks; no model/network/production business writes."""
import asyncio
import copy
import io
import os
import socket
import sys
import tempfile
import unittest
from contextlib import asynccontextmanager
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent))
from openclaw_service.assistant.lighthouse_ai import AssistantError
from openclaw_service.assistant.lighthouse_files import LighthouseFiles
from openclaw_service.assistant import lighthouse_document_tools as tools

ACTOR = {'id': 'document-fixture', 'scopes': ['D']}


class Store:
    def __init__(self, path):
        self.db_path, self.docs = path, {}

    def get_document(self, namespace, key):
        return copy.deepcopy(self.docs.get((namespace, key)))

    def put_document(self, namespace, key, value):
        self.docs[namespace, key] = copy.deepcopy(value)


class DocumentToolsTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.files = LighthouseFiles(Store(self.root / 'state.sqlite3'))

    def test_word_table_structure_is_preserved_offline(self):
        from docx import Document
        document = Document()
        document.add_heading('设备检查', 1)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = '设备编号', '检查结果'
        table.cell(1, 0).text, table.cell(1, 1).text = 'D-101', '正常'
        output = io.BytesIO()
        document.save(output)
        file = self.files.upload(ACTOR, '检查.docx', output.getvalue(), extract=False)
        result = tools.read_attachment(self.files, ACTOR, file['id'])
        self.assertIn('设备检查', result['text'])
        self.assertIn('|', result['text'])
        self.assertIn('D-101', result['text'])
        self.assertFalse(result['empty'])
        self.assertFalse(result['truncated'])

    def test_cache_is_private_rechecked_and_paginated(self):
        file = self.files.upload(ACTOR, '检查.pdf', b'fixture', extract=False, source_scopes=['D'])
        with patch.object(tools, 'run_local', return_value={'text': 'x' * 7000, 'note': 'fixture'}) as run:
            first = tools.read_attachment(self.files, ACTOR, file['id'])
            second = tools.read_attachment(self.files, ACTOR, file['id'], offset=6000)
            self.assertEqual(first['next_offset'], 6000)
            self.assertEqual(second['text'], 'x' * 1000)
            self.assertIsNone(second['next_offset'])
            run.assert_called_once()
            for actor in ({'id': 'other', 'scopes': ['D'], 'is_admin': True}, {**ACTOR, 'scopes': ['A']}):
                with self.assertRaises(AssistantError):
                    tools.read_attachment(self.files, actor, file['id'])

    def test_excel_and_slides_are_read_with_their_structure(self):
        from openpyxl import Workbook
        from pptx import Presentation
        book = Workbook()
        book.active.append(['设备', '数量'])
        book.active.append(['D-101', 2])
        output = io.BytesIO()
        book.save(output)
        book.close()
        file = self.files.upload(ACTOR, '核对.xlsx', output.getvalue(), extract=False)
        result = tools.read_attachment(self.files, ACTOR, file['id'])
        self.assertIn('D-101', result['text'])
        self.assertIn('|', result['text'])
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = '设备检查'
        slide.placeholders[1].text = '运行正常'
        output = io.BytesIO()
        deck.save(output)
        file = self.files.upload(ACTOR, '检查.pptx', output.getvalue(), extract=False)
        result = tools.read_attachment(self.files, ACTOR, file['id'])
        self.assertIn('设备检查', result['text'])
        self.assertIn('运行正常', result['text'])

    def test_chart_shapes_render_and_invalid_values_fail(self):
        from PIL import Image
        for kind in ('bar', 'line', 'pie'):
            result = tools.run_local('chart', {'kind': kind, 'title': '各楼栋维修项目统计（测试数据）',
                'labels': ['A楼', 'B楼', 'C楼', 'D楼', 'E楼', 'H楼'],
                'series': [{'name': '维修中', 'values': [2, 4, 1, 3, 2, 1]}], 'unit': '项'})
            with Image.open(io.BytesIO(result)) as image:
                self.assertEqual(image.size, (1350, 810))
                self.assertIsNotNone(image.convert('RGB').getextrema())
            if os.environ.get('LIGHTHOUSE_TEST_ARTIFACTS'):
                target = Path(os.environ['LIGHTHOUSE_TEST_ARTIFACTS'])
                target.mkdir(parents=True, exist_ok=True)
                (target / (kind + '.png')).write_bytes(result)
        for value in (float('nan'), float('inf'), None, True):
            with self.assertRaises(AssistantError):
                tools.chart_png('bar', '测试', ['A'], [{'name': '项目', 'values': [value]}])
        for values in ([-1], [0]):
            with self.assertRaises(AssistantError):
                tools.chart_png('pie', '测试', ['A'], [{'name': '项目', 'values': values}])

    def test_local_ocr_reads_image_and_scanned_pdf_one_page_at_a_time(self):
        from PIL import Image, ImageDraw, ImageFont
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.lib.utils import ImageReader
        image = Image.new('RGB', (1000, 260), 'white')
        font = ImageFont.truetype(str(Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts/msyh.ttc'), 38)
        ImageDraw.Draw(image).text((30, 40), '设备编号 D-101\n运行正常 2026', fill='black', font=font)
        output = io.BytesIO()
        image.save(output, format='PNG')
        file = self.files.upload(ACTOR, '识别.png', output.getvalue(), extract=False)
        result = tools.read_attachment(self.files, ACTOR, file['id'], mode='ocr')
        self.assertIn('D-101', result['text'])
        self.assertIn('运行正常', result['text'])
        pdf = io.BytesIO()
        canvas = Canvas(pdf, pagesize=(600, 300))
        for _ in range(2):
            canvas.drawImage(ImageReader(image), 20, 70, width=550, height=143)
            canvas.showPage()
        canvas.save()
        image.close()
        file = self.files.upload(ACTOR, '扫描件.pdf', pdf.getvalue(), extract=False)
        result = tools.read_attachment(self.files, ACTOR, file['id'], mode='ocr', page=1)
        self.assertIn('2026', result['text'])
        self.assertEqual(result['total_pages'], 2)
        self.assertEqual(result['next_page'], 2)

    def test_timeout_releases_worker_slot_and_offline_calls_are_blocked(self):
        with self.assertRaises(AssistantError) as raised:
            tools.run_local('chart', {'kind': 'bar'}, timeout=.001)
        self.assertEqual(raised.exception.status, 408)
        self.assertTrue(tools._SLOT.acquire(blocking=False))
        tools._SLOT.release()
        with patch.object(socket, 'getaddrinfo', tools._offline), self.assertRaises(OSError):
            socket.getaddrinfo('example.com', 443)


class ConversationToolsTests(unittest.IsolatedAsyncioTestCase):
    async def test_tools_are_registered_and_chart_is_delivered_to_current_account(self):
        from bin import test_lighthouse_stream as fixtures
        from pydantic_ai.models.function import DeltaToolCall, FunctionModel
        from openclaw_service.assistant.lighthouse_model import LighthouseModel
        from PIL import Image
        await fixtures.StreamTests.asyncSetUp(self)
        calls, artifacts = [], []
        output = io.BytesIO()
        Image.new('RGB', (20, 20), 'blue').save(output, format='PNG')

        async def stream(messages, info):
            self.assertTrue({'create_chart', 'read_attachment_document', 'feishu_document'} <= {tool.name for tool in info.function_tools})
            calls.append(1)
            if len(calls) == 1:
                yield {0: DeltaToolCall(name='create_chart', json_args='{"kind":"bar","title":"统计","labels":["A","B"],"series":[{"name":"数量","values":[1,2]}]}')}
            else:
                yield '图表已生成，可在对话中查看。'

        @asynccontextmanager
        async def model(*_):
            yield FunctionModel(stream_function=stream)

        async def authorize():
            return self.actor

        async def emit(event, data):
            if event == 'artifact':
                artifacts.append(data['file'])

        engine = LighthouseModel(self.portal, model_factory=model)
        try:
            with patch.object(tools, 'run_local', return_value=output.getvalue()):
                await engine.answer(self.actor, {'question': '把用户提供的A为1、B为2画成柱状图', '_profile': self.model.profile()}, [], self.request, emit, authorize, {})
            self.assertEqual(len(artifacts), 1)
            self.assertTrue(artifacts[0]['is_image'])
            self.assertEqual(self.reads, [])
            with self.assertRaises(AssistantError):
                self.portal.files.get({**self.actor, 'id': 'someone-else'}, artifacts[0]['id'])
        finally:
            await engine.close()
            await self.runtime.close()


if __name__ == '__main__':
    unittest.main()
