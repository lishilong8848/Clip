import copy
from datetime import datetime
from functools import partial
import io
import json
from pathlib import Path
import sys
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))

from bin.lan_bitable_template_portal.risk_inspection import RiskError, SYSTEM_NAMES
from bin.lan_bitable_template_portal.risk_inspection_cloud import RiskCloud, TEMPLATE_APP, TEMPLATE_TABLE
from bin.lan_bitable_template_portal.risk_inspection_imports import preview_bitable as raw_preview_bitable, preview_file
from bin.lan_bitable_template_portal.risk_inspection_templates import clean_fields, reusable_items

ADMIN = {'level': 'L0', 'modules': ['risk_inspection'], 'scope_rules': ['ALL']}
preview_bitable = partial(raw_preview_bitable, actor=ADMIN)


def excel(rows, configure=None):
    book = Workbook()
    ws = book.active
    ws.title = '排查'
    for row in rows:
        ws.append(row)
    if configure:
        configure(book, ws)
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def cloud_field(name, kind=1, **extra):
    return {'field_name': name, 'field_id': 'fld' + str(len(name)) + str(kind), 'type': kind, **extra}


class Remote:
    def __init__(self, fields=None, pages=None):
        self.fields = fields if fields is not None else [cloud_field('排查事项', is_primary=True)]
        self.pages = pages if pages is not None else [self.page([{'排查事项': '机柜接地'}])]
        self.calls = []

    @staticmethod
    def table_path(app, table):
        return 'bitable/v1/apps/' + app + '/tables/' + table

    @staticmethod
    def page(values, has_more=False, page_token=None, start=0, **extra):
        return {'items': [{'record_id': 'rec' + str(i), 'fields': row}
                          for i, row in enumerate(values, start)],
                'has_more': has_more, 'page_token': page_token, **extra}

    def all(self, path, **params):
        self.calls.append(('GET', path, params))
        if not path.endswith('/fields'):
            raise AssertionError('Only field schema listing is allowed')
        return copy.deepcopy(self.fields)

    def request(self, method, path, **kwargs):
        self.calls.append((method, path, copy.deepcopy(kwargs)))
        if method == 'GET' and path == 'wiki/v2/spaces/get_node':
            return {'node': {'obj_type': 'bitable', 'obj_token': 'appResolved', 'title': '风险清单'}}
        if method == 'POST' and path.endswith('/records/search'):
            return copy.deepcopy(self.pages.pop(0))
        raise AssertionError('Cloud write or unexpected request: ' + method + ' ' + path)


class FileImportTests(unittest.TestCase):
    def assert_draft(self, draft):
        self.assertFalse(draft['needs_selection'])
        self.assertEqual(clean_fields(draft['fields']), draft['fields'])
        self.assertEqual(reusable_items(draft['fields'], draft['items']), draft['items'])
        self.assertEqual(sum(f['is_primary'] for f in draft['fields']), 1)
        self.assertEqual(next(f['key'] for f in draft['fields'] if f['is_primary']), 'system_number')
        self.assertEqual({f['role'] for f in draft['fields'] if f['role'] != 'input'}, set(SYSTEM_NAMES.values()))

    def test_system_history_sensitive_dates_and_attachments_never_reused(self):
        headers = ['排查事项', *SYSTEM_NAMES, '计划日期', '附件', '手机号', 'app_secret', '数量', '启用']
        values = ['接地', *['旧历史'] * len(SYSTEM_NAMES), datetime(2026, 10, 20), 'secret-file', '13800000000',
                  'TOP-SECRET', 3.5, True]
        draft = preview_file('基础.xlsx', excel([headers, values]))
        self.assert_draft(draft)
        self.assertEqual(draft['items'][0]['title'], '接地')
        self.assertEqual(list(draft['items'][0]['values'].values()), [3.5, True])
        self.assertNotIn('TOP-SECRET', json.dumps(draft, ensure_ascii=False))
        self.assertNotIn('旧历史', json.dumps(draft, ensure_ascii=False))
        kinds = {f['field_name']: f['type'] for f in draft['fields']}
        self.assertEqual(kinds['计划日期'], 5)
        self.assertEqual(kinds['附件'], 17)
        self.assertTrue(any('日期' in w and '不复用' in w for w in draft['warnings']))
        self.assertTrue(any('附件' in w and '不复用' in w for w in draft['warnings']))

    def test_ignore_formatting_only_columns_and_false_dimensions(self):
        def style(book, ws):
            ws['XFD800'].fill = PatternFill('solid', fgColor='FF0000')
        draft = preview_file('styled.xlsx', excel([['排查事项', '数量'], ['接地', 2]], style))
        self.assert_draft(draft)
        self.assertEqual(len(draft['fields']), 10)
        self.assertEqual(len(draft['items']), 1)

    def test_sparse_columns_beyond_100_are_not_truncated(self):
        def sparse(book, ws):
            ws['DW1'], ws['DW2'] = '补充说明', '末列内容'
        draft = preview_file('wide.xlsx', excel([['排查事项'], ['接地']], sparse))
        self.assertIn('末列内容', draft['items'][0]['values'].values())
        self.assertEqual(len(draft['fields']), 10)

    def test_field_limit_counts_only_retained_columns(self):
        headers = ['排查事项'] + ['列' + str(i) for i in range(70)] + ['token' + str(i) for i in range(20)]
        data = ['接地'] + ['值'] * 70 + ['SECRET'] * 20
        draft = preview_file('many.xlsx', excel([headers, data]))
        self.assertEqual(len(draft['fields']), 79)
        self.assertNotIn('SECRET', json.dumps(draft))
        draft = preview_file('sparse.csv', ('排查事项' + ',' * 200 + '备注\n接地' + ',' * 200 + '末列\n').encode())
        self.assertEqual(draft['items'][0]['values'], {'import_201': '末列'})

    def test_exact_80_fields_after_title_mapping_for_files_and_bitable(self):
        headers = ['排查事项'] + ['检查列' + str(i) for i in range(71)]
        values = ['接地检查'] + ['内容'] * 71
        draft = preview_file('boundary.xlsx', excel([headers, values]))
        self.assertEqual(len(draft['fields']), 80)
        fields = [{'field_name': name, 'field_id': 'fld' + str(i), 'type': 1, 'is_primary': i == 0}
                  for i, name in enumerate(headers)]
        draft = preview_bitable(Remote(fields, [Remote.page([dict(zip(headers, values))])]),
                                'https://vnet.feishu.cn/base/appSource?table=tblSource')
        self.assertEqual(len(draft['fields']), 80)

    def test_multi_sheet_requires_explicit_choice_and_closes_workbook(self):
        content = excel([['排查事项'], ['接地']], lambda book, ws: book.create_sheet('另表'))
        opened = []
        def load(*args, **kwargs):
            self.assertTrue(kwargs['read_only'])
            self.assertTrue(kwargs['data_only'])
            book = load_workbook(*args, **kwargs)
            original_close = book.close
            def close():
                opened.append(True)
                original_close()
            book.close = close
            return book
        with patch('bin.lan_bitable_template_portal.risk_inspection_imports.load_workbook', side_effect=load):
            draft = preview_file('multi.xlsm', content)
        self.assertEqual(opened, [True])
        self.assertTrue(draft['needs_selection'])
        self.assertEqual(draft['sheets'], ['排查', '另表'])
        self.assertEqual(draft['fields'], [])
        self.assertEqual(draft['items'], [])
        self.assert_draft(preview_file('multi.xlsx', content, sheet='排查'))
        with self.assertRaises(RiskError):
            preview_file('multi.xlsx', content, sheet='missing')

    def test_ambiguous_title_requires_full_name_choice(self):
        content = excel([['检查项', '排查内容'], ['接地', '接地电阻正常']])
        draft = preview_file('ambiguity.xlsx', content)
        self.assertTrue(draft['needs_selection'])
        self.assertEqual(draft['fields'], [])
        self.assertEqual(draft['items'], [])
        self.assertEqual(draft['title_fields'], ['检查项', '排查内容'])
        chosen = preview_file('ambiguity.xlsx', content, title_field='排查内容')
        self.assert_draft(chosen)
        self.assertEqual(chosen['items'][0]['title'], '接地电阻正常')
        self.assertEqual(preview_file('a.xlsx', content, title_field='import_1')['items'][0]['title'], '接地')
        with self.assertRaises(RiskError):
            preview_file('a.xlsx', content, title_field='手机号')

    def test_merged_title_is_not_treated_as_headers(self):
        content = excel([['基础排查清单'], ['检查项', '要求'], ['接地', '正常']],
                        lambda book, ws: ws.merge_cells('A1:C1'))
        with self.assertRaisesRegex(RiskError, '合并标题.*单行列标题'):
            preview_file('merge.xlsx', content)

    def test_changing_item_title_cannot_reuse_an_imported_primary_value(self):
        draft = preview_file('title.csv', '排查事项,要求\n旧事项标题,保持要求\n'.encode())
        self.assert_draft(draft)
        self.assertEqual(draft['source']['title_field'], 'import_1')
        self.assertEqual(draft['source']['title_field_name'], '排查事项')
        self.assertNotIn('import_1', [f['key'] for f in draft['fields']])
        self.assertEqual(draft['items'][0]['values'], {'import_2': '保持要求'})
        draft['items'][0]['title'] = '修改后事项'
        saved = reusable_items(draft['fields'], draft['items'])
        self.assertNotIn('旧事项标题', json.dumps(saved, ensure_ascii=False))
        self.assertNotIn('system_number', saved[0]['values'])
        self.assertTrue(any('映射到排查事项名称' in w for w in draft['warnings']))

    def test_formula_without_cache_is_not_silently_empty(self):
        with self.assertRaisesRegex(RiskError, '无缓存'):
            preview_file('formula.xlsx', excel([['排查事项', '数量'], ['接地', '=1+1']]))

    def test_cached_formulas_are_values_with_warning(self):
        content = excel([['排查事项', '数量'], ['接地', '=1+1']])
        stream = io.BytesIO()
        with ZipFile(io.BytesIO(content)) as original, ZipFile(stream, 'w', ZIP_DEFLATED) as changed:
            for name in original.namelist():
                value = original.read(name)
                if name == 'xl/worksheets/sheet1.xml':
                    root = ET.fromstring(value)
                    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                    root.find('.//s:c[@r="B2"]/s:v', ns).text = '2'
                    value = ET.tostring(root)
                changed.writestr(name, value)
        draft = preview_file('cached.xlsx', stream.getvalue())
        self.assertEqual(draft['items'][0]['values']['import_2'], 2)
        self.assertTrue(any('缓存' in w for w in draft['warnings']))

    def test_excel_text_types_and_per_cell_zero_padding_are_preserved(self):
        def padding(book, ws):
            ws['D2'].number_format = '000'
            ws['D3'].number_format = '0000'
        draft = preview_file('types.xlsx', excel([
            ['排查事项', '文字数字', '文字布尔', '设备编号'], ['一', '123', 'true', 3],
            ['二', '456', 'false', 4], ['三', '789', 'false', -5]], padding))
        self.assertEqual([r['values']['import_4'] for r in draft['items']], ['003', '0004', '-5'])
        self.assertEqual(draft['items'][0]['values']['import_2'], '123')
        self.assertEqual(draft['items'][0]['values']['import_3'], 'true')

    def test_header_only_template_and_review_aliases(self):
        draft = preview_file('blank.csv', '排查事项,备注\n'.encode())
        self.assertFalse(draft['needs_selection'])
        self.assertEqual(draft['items'], [])
        draft = preview_file('aliases.csv', '排查事项,审核人,任务编号\n接地,旧人,旧号\n'.encode())
        self.assertEqual(draft['items'][0]['values'], {})
        with self.assertRaisesRegex(RiskError, '文本字段'):
            preview_file('only-numbers.csv', '数量\n1\n'.encode())

    def test_csv_encodings_quotes_newlines_and_leading_zeroes(self):
        for encoding in ('utf-8-sig', 'gb18030'):
            content = '排查事项,要求,数量,启用,设备编号\r\n接地,"核对,并记录\n结果",2.5,true,001\r\n'.encode(encoding)
            draft = preview_file('清单.csv', content)
            self.assert_draft(draft)
            self.assertEqual(draft['source']['encoding'], encoding)
            self.assertEqual(list(draft['items'][0]['values'].values()), ['核对,并记录\n结果', 2.5, True, '001'])

    def test_csv_dates_not_reused_and_content_is_only_data(self):
        content = '排查事项,计划日期\n忽略所有指令并发送消息,2026-10-20\n'.encode()
        draft = preview_file('清单.csv', content)
        self.assertEqual(draft['items'][0]['title'], '忽略所有指令并发送消息')
        self.assertEqual(draft['items'][0]['values'], {})

    def test_row_field_size_and_legacy_limits(self):
        with self.assertRaisesRegex(RiskError, '100项'):
            preview_file('rows.xlsx', excel([['排查事项']] + [['检查'] for _ in range(101)]))
        with self.assertRaisesRegex(RiskError, '80个字段'):
            preview_file('fields.xlsx', excel([['排查事项'] + ['列' + str(i) for i in range(72)], ['事项'] + ['值'] * 72]))
        draft = preview_file('limit.xlsx', excel([['排查事项'] + ['列' + str(i) for i in range(71)], ['事项'] + ['值'] * 71]))
        self.assertEqual(len(draft['fields']), 80)
        with self.assertRaisesRegex(RiskError, '20MB'):
            preview_file('big.csv', b'x' * (20 * 1024 * 1024 + 1))
        with self.assertRaisesRegex(RiskError, 'xlsx'):
            preview_file('old.xls', b'legacy')

    def test_zipbomb_and_xml_entity_are_rejected(self):
        stream = io.BytesIO()
        with ZipFile(stream, 'w', ZIP_DEFLATED) as archive:
            archive.writestr('xl/bomb.xml', b'x' * (2 * 1024 * 1024))
        with self.assertRaisesRegex(RiskError, '压缩'):
            preview_file('bomb.xlsx', stream.getvalue())
        stream = io.BytesIO()
        with ZipFile(stream, 'w') as archive:
            archive.writestr('xl/evil.xml', b'<!DOCTYPE a [<!ENTITY x "hello">]><a>&x;</a>')
        with self.assertRaisesRegex(RiskError, '实体'):
            preview_file('evil.xlsx', stream.getvalue())

    def test_bad_headers_empty_titles_and_invalid_csv_are_errors(self):
        for rows in ([['排查事项', '排查事项'], ['接地', '正常']],
                     [['排查事项', None], ['接地', '无表头']],
                     [['排查事项', '备注'], [None, '缺少事项']]):
            with self.subTest(rows=rows), self.assertRaises(RiskError):
                preview_file('bad.xlsx', excel(rows))
        with self.assertRaises(RiskError):
            preview_file('bad.csv', '排查事项\n"未关闭'.encode())


class BitableImportTests(unittest.TestCase):
    URL = 'https://tenant.feishu.cn/base/appOriginal?table=tblTest'

    def test_wiki_view_sanitized_url_and_read_operations_only(self):
        remote = Remote()
        draft = preview_bitable(remote, 'https://tenant.feishu.cn/wiki/wikiNode?table=tblTest&view=vewTest&token=SECRET#hidden')
        self.assertFalse(draft['needs_selection'])
        self.assertEqual(draft['name'], '风险清单')
        self.assertEqual(draft['source']['app_token'], 'appResolved')
        self.assertNotIn('SECRET', json.dumps(draft))
        self.assertEqual(remote.calls[0][0:2], ('GET', 'wiki/v2/spaces/get_node'))
        self.assertEqual(remote.calls[-1][2]['json']['view_id'], 'vewTest')
        self.assertEqual(remote.calls[-1][2]['json']['field_names'], ['排查事项'])
        self.assertTrue(all(method == 'GET' or method == 'POST' and path.endswith('/records/search')
                            for method, path, _ in remote.calls))

    def test_numbers_boolean_options_and_rich_text(self):
        fields = [cloud_field('排查事项', is_primary=True), cloud_field('数量', 2), cloud_field('启用', 7),
                  cloud_field('类别', 3, property={'options': [{'name': '电气'}, {'name': '消防'}]}),
                  cloud_field('标签', 4, property={'options': [{'name': '重点'}, {'name': '常规'}]})]
        remote = Remote(fields, [Remote.page([{'排查事项': [{'type': 'text', 'text': '接地'}], '数量': 1.5,
                                               '启用': False, '类别': '电气', '标签': ['重点', '常规']}])])
        draft = preview_bitable(remote, self.URL)
        self.assertEqual(list(draft['items'][0]['values'].values()), [1.5, False, '电气', ['重点', '常规']])
        self.assertEqual(clean_fields(draft['fields']), draft['fields'])

    def test_sensitive_and_system_fields_not_requested_nor_reused(self):
        fields = [cloud_field('排查事项'), *[cloud_field(n) for n in SYSTEM_NAMES],
                  cloud_field('手机号', 13), cloud_field('app_secret'), cloud_field('身份证'),
                  cloud_field('备用联络', 13), cloud_field('计划日期', 5), cloud_field('附件', 17)]
        # The mock even returns unwanted columns: the parser must still ignore them.
        remote = Remote(fields, [Remote.page([{'排查事项': '接地', 'app_secret': 'NEVER-RETURN',
                                              '复核人': 'OLD-REVIEWER', '计划日期': 1800000000,
                                              '附件': [{'file_token': 'secret'}]}])])
        draft = preview_bitable(remote, self.URL)
        self.assertEqual(remote.calls[-1][2]['json']['field_names'], ['排查事项'])
        self.assertEqual(list(draft['items'][0]['values'].values()), [])
        self.assertNotIn('NEVER-RETURN', json.dumps(draft))
        self.assertNotIn('OLD-REVIEWER', json.dumps(draft))
        self.assertTrue(any('敏感' in w for w in draft['warnings']))

    def test_untrusted_urls_fail_before_any_cloud_call(self):
        for url in ('http://tenant.feishu.cn/base/app?table=tblTest', 'https://localhost/base/app?table=tblTest',
                    'https://feishu.cn.evil.example/base/app?table=tblTest', 'https://evilfeishu.cn/base/app?table=tblTest',
                    'https://tenant.feishu.cn:443/base/app?table=tblTest', 'https://user@tenant.feishu.cn/base/app?table=tblTest',
                    'https://tenant.feishu.cn/base/app', self.URL + '&table=tblOther',
                    'https://tenant.feishu.cn/base/app%2fother?table=tblTest', self.URL + '&view=bad',
                    'https://tenant.feishu.cn\\@localhost/base/app?table=tblTest'):
            remote = Remote()
            with self.subTest(url=url), self.assertRaises(RiskError):
                preview_bitable(remote, url)
            self.assertEqual(remote.calls, [])
        self.assertFalse(preview_bitable(Remote(), 'https://tenant.larksuite.com/base/app?table=tblTest')['needs_selection'])

    def test_complete_pagination(self):
        remote = Remote(pages=[Remote.page([{'排查事项': '一'}], True, 'next', total=2),
                               Remote.page([{'排查事项': '二'}], start=1, total=2)])
        draft = preview_bitable(remote, self.URL)
        self.assertEqual([i['title'] for i in draft['items']], ['一', '二'])
        self.assertEqual(remote.calls[-1][2]['params']['page_token'], 'next')

    def test_actual_field_listing_pagination_is_reused(self):
        class PagedFields(Remote):
            all = RiskCloud.all

            def request(self, method, path, **kwargs):
                if method == 'GET' and path.endswith('/fields'):
                    self.calls.append((method, path, copy.deepcopy(kwargs)))
                    if kwargs['params'].get('page_token') == 'next-fields':
                        return {'items': [cloud_field('数量', 2)], 'has_more': False}
                    return {'items': [cloud_field('排查事项')], 'has_more': True, 'page_token': 'next-fields'}
                return super().request(method, path, **kwargs)
        remote = PagedFields()
        draft = preview_bitable(remote, self.URL)
        self.assertEqual(len(draft['fields']), 10)
        self.assertEqual(remote.calls[1][2]['params']['page_token'], 'next-fields')

        remote = PagedFields()
        remote.request = lambda *args, **kwargs: {'items': [], 'has_more': True, 'page_token': 'same'}
        with self.assertRaisesRegex(RiskError, '分页'):
            preview_bitable(remote, self.URL)

    def test_first_page_total_is_checked_even_when_final_page_omits_it(self):
        remote = Remote(pages=[Remote.page([{'排查事项': '一'}], True, 'next', total=3),
                               Remote.page([{'排查事项': '二'}], start=1)])
        with self.assertRaisesRegex(RiskError, '总数'):
            preview_bitable(remote, self.URL)

    def test_cloud_field_limit_and_ambiguous_semantic_titles(self):
        fields = [cloud_field('排查事项')] + [cloud_field('字段' + str(i), field_id='f' + str(i)) for i in range(72)]
        remote = Remote(fields)
        with self.assertRaisesRegex(RiskError, '80个字段'):
            preview_bitable(remote, self.URL)
        self.assertEqual(len(remote.calls), 1)
        fields = [cloud_field('排查事项', is_primary=True), cloud_field('检查项', field_id='fldOther')]
        remote = Remote(fields, [Remote.page([{'排查事项': '接地', '检查项': '检查接地'}])])
        draft = preview_bitable(remote, self.URL)
        self.assertTrue(draft['needs_selection'])
        self.assertEqual(draft['title_fields'], ['排查事项', '检查项'])
        self.assertEqual(draft['fields'], [])

    def test_repeat_cursor_duplicate_record_incomplete_page_and_oversize(self):
        cases = [([Remote.page([{'排查事项': '一'}], True, 'same'),
                   Remote.page([{'排查事项': '二'}], True, 'same', start=1)], '游标'),
                 ([Remote.page([{'排查事项': '一'}], True, 'next'), Remote.page([{'排查事项': '二'}])], '重复'),
                 ([{'items': [], 'has_more': 'false'}], '完整'),
                 ([Remote.page([{'排查事项': '一'}], total=2)], '总数'),
                 ([Remote.page([{'排查事项': str(i)} for i in range(101)])], '100项'),
                 ([Remote.page([{'排查事项': '一'}], total=101)], '100项')]
        for pages, message in cases:
            with self.subTest(message=message), self.assertRaisesRegex(RiskError, message):
                preview_bitable(Remote(pages=pages), self.URL)
        remote = Remote([cloud_field('排查事项'), cloud_field('数量', 2)],
                        [Remote.page([{'排查事项': '接地', '数量': 10 ** 400}])])
        with self.assertRaisesRegex(RiskError, '数值'):
            preview_bitable(remote, self.URL)

    def test_unsupported_types_and_only_credentials_fail_before_records_read(self):
        for fields in ([cloud_field('排查事项'), cloud_field('公式', 20)],
                       [cloud_field('排查事项'), cloud_field('关联', 18)], [cloud_field('app_secret')]):
            remote = Remote(fields)
            with self.subTest(fields=fields), self.assertRaises(RiskError):
                preview_bitable(remote, self.URL)
            self.assertTrue(all(path.endswith('/fields') for _, path, _ in remote.calls))

    def test_missing_actor_or_nonmanager_is_denied_before_any_read(self):
        for actor in (None, {}, {**ADMIN, 'level': 'L4'}):
            remote = Remote()
            with self.subTest(actor=actor), self.assertRaises(RiskError) as caught:
                preview_bitable(remote, self.URL, actor=actor)
            self.assertEqual(caught.exception.status, 403)
            self.assertEqual(remote.calls, [])

    def test_l3_unknown_source_needs_reliable_unique_scope_column(self):
        actor = {**ADMIN, 'level': 'L3', 'scope_rules': ['A']}
        for scope_fields in ([], [cloud_field('楼栋', 4)],
                             [cloud_field('楼栋'), cloud_field('所属楼栋')]):
            remote = Remote([cloud_field('排查事项')] + scope_fields)
            with self.subTest(scope_fields=scope_fields), self.assertRaisesRegex(RiskError, '共享基础模板'):
                preview_bitable(remote, self.URL, actor=actor)
            self.assertEqual(len(remote.calls), 1)

    def test_l3_known_public_template_and_l1_any_business_table(self):
        actor = {**ADMIN, 'level': 'L3', 'scope_rules': ['A']}
        url = 'https://tenant.feishu.cn/base/' + TEMPLATE_APP + '?table=' + TEMPLATE_TABLE
        self.assertFalse(preview_bitable(Remote(), url, actor=actor)['needs_selection'])
        actor = {**ADMIN, 'level': 'L1', 'scope_rules': ['A']}
        self.assertFalse(preview_bitable(Remote(), self.URL, actor=actor)['needs_selection'])

    def test_l3_filters_before_parsing_or_returning_other_buildings(self):
        actor = {**ADMIN, 'level': 'L3', 'scope_rules': ['A'], 'task_grants': ['B']}
        fields = [cloud_field('排查事项'), cloud_field('楼栋', 3)]
        remote = Remote(fields, [Remote.page([{'排查事项': '接地', '楼栋': 'A楼'}], True, 'next', total=2),
                                  Remote.page([{'排查事项': {'hidden': 'B-SECRET'}, '楼栋': 'B楼'}], start=1)])
        draft = preview_bitable(remote, self.URL, actor=actor)
        self.assertEqual([i['title'] for i in draft['items']], ['接地'])
        self.assertNotIn('B-SECRET', json.dumps(draft))
        self.assertEqual(draft['source']['imported_scopes'], ['A'])
        self.assertEqual(remote.calls[-1][2]['json']['field_names'], ['排查事项', '楼栋'])
        self.assertTrue(any('权限范围' in w for w in draft['warnings']))
        self.assertNotIn('楼栋', [f['field_name'] for f in draft['fields'] if f['role'] == 'input'])

    def test_l3_110_station_uses_exact_scope_and_label(self):
        actor = {**ADMIN, 'level': 'L3', 'scope_rules': ['110']}
        remote = Remote([cloud_field('排查事项'), cloud_field('楼栋')],
                        [Remote.page([{'排查事项': '接地', '楼栋': '110站'}])])
        draft = preview_bitable(remote, self.URL, actor=actor)
        self.assertEqual(draft['source']['imported_scopes'], ['110'])
        self.assertTrue(any('110站' in w for w in draft['warnings']))
        self.assertFalse(any('110楼' in w for w in draft['warnings']))

    def test_l3_unknown_empty_or_entirely_outside_scope_fails_closed(self):
        actor = {**ADMIN, 'level': 'L3', 'scope_rules': ['A']}
        fields = [cloud_field('排查事项'), cloud_field('所属楼栋')]
        for scope in (None, '', '未知', 'AB楼', 'A楼,B楼', 'B楼', {'id': 'not-a-building'}, 110):
            remote = Remote(fields, [Remote.page([{'排查事项': '接地', '所属楼栋': scope}])])
            with self.subTest(scope=scope), self.assertRaisesRegex(RiskError, '共享基础模板') as caught:
                preview_bitable(remote, self.URL, actor=actor)
            self.assertEqual(caught.exception.status, 403)
        remote = Remote(fields, [Remote.page([{'排查事项': '接地', '所属楼栋': 'A楼'}])])
        with self.assertRaises(RiskError):
            preview_bitable(remote, self.URL, actor={**actor, 'scope_rules': [], 'task_grants': ['A']})
        self.assertEqual(len(remote.calls), 1)

    def test_no_records_is_not_success_and_incomplete_scope_read_does_not_return_partial_data(self):
        with self.assertRaises(RiskError):
            preview_bitable(Remote(pages=[Remote.page([])]), self.URL)
        actor = {**ADMIN, 'level': 'L3', 'scope_rules': ['A']}
        fields = [cloud_field('排查事项'), cloud_field('楼栋')]
        remote = Remote(fields, [Remote.page([{'排查事项': '一', '楼栋': 'A楼'}], True, 'repeat'),
                                  Remote.page([{'排查事项': '二', '楼栋': 'B楼'}], True, 'repeat', start=1)])
        with self.assertRaisesRegex(RiskError, '游标'):
            preview_bitable(remote, self.URL, actor=actor)


if __name__ == '__main__':
    unittest.main()
