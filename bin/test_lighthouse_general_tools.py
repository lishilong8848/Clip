"""Shared web/Feishu general-tool flows with synthetic models and private temp files."""
import io
import json
import asyncio
import unittest
from contextlib import asynccontextmanager
from pathlib import Path
from unittest.mock import AsyncMock, patch

from bin import test_lighthouse_stream as fixtures
from pydantic_ai.models.function import DeltaToolCall, FunctionModel
from openclaw_service.assistant.lighthouse_model import LighthouseModel
from openclaw_service.assistant.lighthouse_files import extract_text


class GeneralToolTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        await fixtures.StreamTests.asyncSetUp(self)
        self.addAsyncCleanup(self.runtime.close)

    async def run_tools(self, question, steps, *, file_ids=(), channel='', history=(), public_sources=None):
        calls = []
        async def stream(messages, info):
            index = len(calls)
            calls.append(messages)
            if index < len(steps):
                if isinstance(steps[index], str):
                    yield steps[index]
                else:
                    name, args = steps[index]
                    yield {0: DeltaToolCall(name=name, json_args=json.dumps(args, ensure_ascii=False))}
            else:
                yield '已根据提供的资料完成整理。'
        @asynccontextmanager
        async def model(*_):
            yield FunctionModel(stream_function=stream)
        engine = LighthouseModel(self.portal, model_factory=model, public_sources=public_sources)
        self.addAsyncCleanup(engine.close)
        actor = {**self.actor, 'channel': channel}
        result = await engine.answer(actor, {'question': question, 'file_ids': list(file_ids), '_profile': self.model.profile()},
                                     list(history), self.request, AsyncMock(), AsyncMock(return_value=actor), {})
        return result, calls

    async def test_chinese_public_lookup_is_shared_by_web_and_feishu(self):
        import httpx
        from openclaw_service.assistant.lighthouse_public import PublicSources
        question = '你上网搜一下黑尾大勾鱼是什么?'
        requests = []
        def handler(request):
            requests.append(request)
            if request.url.host == 'api.mwmbl.org':
                return httpx.Response(200, json=[])
            if request.url.host == 'www.bing.com':
                return httpx.Response(503)
            return httpx.Response(200, text='<ul class="result"><li class="res-list">'
                '<h3 class="res-title"><a data-mdurl="https://example.com/fish" href="/link?m=tracking">黑尾大勾鱼</a></h3>'
                '<p data-result-summary>一种小型观赏鱼。</p></li></ul>')
        source = PublicSources(client_factory=lambda: httpx.AsyncClient(transport=httpx.MockTransport(handler)))
        for channel in ('', 'feishu:oc_fixture'):
            with self.subTest(channel=channel):
                result, calls = await self.run_tools(question,
                    [('public_search', {'query': question}), '黑尾大勾鱼是一种小型观赏鱼。[1]'],
                    channel=channel, public_sources=source)
                self.assertIn('小型观赏鱼', result['answer'])
                self.assertEqual(result['sources'][0]['url'], 'https://example.com/fish')
                response = next(part.content for message in calls[-1] for part in message.parts
                    if getattr(part, 'part_kind', '') == 'tool-return')
                self.assertEqual(response['actual_query'], '黑尾大勾鱼')
                self.assertTrue(response['ok'])
        self.assertEqual(len(requests), 3, 'The second channel can reuse safe public evidence')
        self.assertEqual(self.reads, [])

    async def test_attachment_business_words_use_file_not_live_business_data(self):
        file = self.portal.files.upload(self.actor, '维修记录.txt', 'A楼维修项目：测试设备一、测试设备二。'.encode())
        for channel in ('', 'feishu:oc_fixture'):
            with self.subTest(channel=channel):
                result, _ = await self.run_tools('统计我上传的附件里A楼有多少个维修项目，并注明来源',
                    [('read_file', {'file_id': file['id']})], file_ids=[file['id']], channel=channel)
                self.assertEqual(result['answer'], '已根据提供的资料完成整理。')
                self.assertTrue(any(file['name'] in source['title'] for source in result['sources']))
                self.assertEqual(self.reads, [])
                self.assertNotIn('plan', result)

    async def test_generated_file_defaults_to_conversation_in_both_channels(self):
        for channel in ('', 'feishu:oc_fixture'):
            with self.subTest(channel=channel):
                result, _ = await self.run_tools('帮我写个py文件，冒泡排序的示例，然后发给我',
                    [('create_text_file', {'name': '冒泡排序.py', 'content': 'def bubble_sort(a):\n    for i in range(len(a)):\n        for j in range(len(a)-i-1):\n            if a[j] > a[j+1]: a[j], a[j+1] = a[j+1], a[j]\n    return a\n'})], channel=channel)
                self.assertEqual(len(result['output_files']), 1)
                self.assertNotIn('plan', result)
                self.assertEqual(self.reads, [])

    async def test_previous_file_can_be_read_again_without_live_business_queries(self):
        file=self.portal.files.upload(self.actor,'维修记录.txt','测试设备甲、测试设备乙。'.encode())
        result,calls=await self.run_tools('统计上个附件里有多少个维修项目', [('read_file',{'file_id':file['id']})],
            history=[{'question':'保存这份资料','answer':'已收到','file_ids':[file['id']],'scopes':self.actor['scopes']}])
        self.assertEqual(len(calls),2)
        self.assertEqual(self.reads,[])
        self.assertIn(file['name'],result['sources'][0]['title'])

    async def test_vision_attachment_is_evidence_without_forcing_ocr_or_public_search(self):
        from PIL import Image
        from openclaw_service.assistant import lighthouse_files
        data=io.BytesIO()
        Image.new('RGB',(10,10),'red').save(data,format='PNG')
        self.model.profile.return_value['image_input']=True
        with patch.object(lighthouse_files,'recognize_text',return_value=''):
            file=self.portal.files.upload(self.actor,'测试.png',data.getvalue())
            result,calls=await self.run_tools('分析这个附件内容',[],file_ids=[file['id']])
        self.assertEqual(len(calls),1)
        self.assertIn('图片附件',result['sources'][0]['title'])
        self.assertEqual(self.reads,[])

    async def test_file_claim_without_attachment_is_retried_into_real_office_file(self):
        for extension, data, magic in (
            ('docx', {'title':'交接说明','paragraphs':['已核对现场信息。']}, b'PK'),
            ('xlsx', {'sheets':[{'name':'汇总','headers':['项目','数量'],'rows':[['测试项',2]]}]}, b'PK'),
            ('pdf', {'title':'交接说明','paragraphs':['已核对现场信息。']}, b'%PDF'),
            ('pptx', {'slides':[{'title':'交接说明','bullets':['已核对现场信息。']}]}, b'PK'),
        ):
            with self.subTest(extension=extension):
                result, calls = await self.run_tools('生成一个'+extension+'文件发给我',
                    ['文件已生成，请下载。', ('create_document_file', {'name':'测试.'+extension,'data':data})])
                self.assertEqual(len(calls),3)
                self.assertEqual(len(result['output_files']),1)
                item=self.portal.files.get(self.actor,result['output_files'][0]['id'])
                self.assertTrue(Path(item['path']).read_bytes().startswith(magic))
                self.assertNotIn('plan',result)

    async def test_reading_business_attachment_cannot_prepare_unrelated_business_write(self):
        file = self.portal.files.upload(self.actor, '通告.txt', '示例通告内容，仅供文档分析。'.encode())
        result, calls = await self.run_tools('分析这个附件中的通告，不办理业务',
            [('query', {'operation': {'api_id':'GET /api/repair-management/records','params':{'scope':'D'}}}), ('read_file', {'file_id': file['id']})], file_ids=[file['id']])
        self.assertEqual(self.reads, [])
        self.assertNotIn('plan', result)


class FileReadTests(unittest.TestCase):
    def test_excel_ignores_empty_formatting_and_keeps_real_rows(self):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill
        book = Workbook()
        book.active.append(['设备', '数量'])
        book.active.append(['测试设备', 0])
        book.active.cell(7000, 120).fill = PatternFill('solid', fgColor='FFFFFF')
        data = io.BytesIO()
        book.save(data)
        book.close()
        text = extract_text(data.getvalue(), '.xlsx')
        self.assertIn('测试设备\t0', text)
        self.assertLess(len(text), 100)


class PublicSearchReuseTests(unittest.IsolatedAsyncioTestCase):
    async def test_new_caller_during_cancel_cleanup_starts_a_new_search(self):
        from openclaw_service.assistant.lighthouse_public import PublicSources, _now_iso
        source=PublicSources()
        self.addAsyncCleanup(source.close)
        started,cancelling,cleanup=asyncio.Event(),asyncio.Event(),asyncio.Event()
        count=0
        async def read(*_):
            nonlocal count
            count+=1
            if count==1:
                started.set()
                try:
                    await asyncio.Event().wait()
                finally:
                    cancelling.set()
                    await cleanup.wait()
            return {'ok':True,'queried_at':_now_iso(),'results':[{'title':'Python'}]}
        source._search_uncached=read
        first=asyncio.create_task(source.search('Python','Python'))
        await started.wait()
        first.cancel()
        await cancelling.wait()
        second=asyncio.create_task(source.search('Python','Python'))
        await asyncio.sleep(.01)
        cleanup.set()
        results=await asyncio.gather(first,second,return_exceptions=True)
        self.assertIsInstance(results[0],asyncio.CancelledError)
        self.assertIsInstance(results[1],dict)
        self.assertEqual(count,2)
        self.assertEqual(source._search_tasks,{})
        self.assertEqual(source._search_waiters,{})

    async def test_concurrent_public_queries_share_request_but_not_mutable_results(self):
        from openclaw_service.assistant.lighthouse_public import PublicSources, _now_iso, PublicQueryError
        source=PublicSources()
        self.addAsyncCleanup(source.close)
        gate=asyncio.Event()
        async def read(*_):
            await gate.wait()
            return {'ok':True,'queried_at':_now_iso(),'results':[{'title':'Python文档','url':'https://docs.python.org/3/'}]}
        source._search_uncached=AsyncMock(side_effect=read)
        tasks=[asyncio.create_task(source.search('Python documentation','搜索Python文档')) for _ in range(20)]
        await asyncio.sleep(.02)
        tasks[0].cancel()
        await asyncio.gather(tasks[0],return_exceptions=True)
        gate.set()
        results=await asyncio.gather(*tasks[1:])
        source._search_uncached.assert_awaited_once()
        results[0]['results'][0]['title']='changed'
        cached=await source.search('Python documentation','搜索Python文档')
        self.assertTrue(cached['cached'])
        self.assertEqual(cached['results'][0]['title'],'Python文档')
        with self.assertRaises(PublicQueryError):
            await source.search('内部台账','内部台账')
        source._search_uncached.assert_awaited_once()
        key=next(key for key in source._search_cache if key[:2] == ('Python documentation','搜索Python文档'))
        source._search_cache[key]=(asyncio.get_running_loop().time()-121,cached)
        await source.search(*key[:2])
        self.assertEqual(source._search_uncached.await_count,2)

    async def test_failed_search_is_not_cached(self):
        from openclaw_service.assistant.lighthouse_public import PublicSources, _now_iso
        source=PublicSources()
        self.addAsyncCleanup(source.close)
        source._search_uncached=AsyncMock(return_value={'ok':False,'queried_at':_now_iso(),'results':[]})
        await source.search('Python文档','搜索Python文档')
        await source.search('Python文档','搜索Python文档')
        self.assertEqual(source._search_uncached.await_count,2)


if __name__ == '__main__':
    unittest.main()
