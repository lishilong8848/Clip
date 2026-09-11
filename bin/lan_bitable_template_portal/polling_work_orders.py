# -*- coding: utf-8 -*-
from __future__ import annotations

import base64
import copy
import datetime as dt
import hashlib
import hmac
import io
import json
import math
import os
import re
import secrets
import shutil
import threading
import time
import uuid
import zipfile
from pathlib import Path
from typing import Any
from urllib.parse import quote, urlsplit

from upload_event_module.utils import get_data_file_path
from .polling_work_order_relay import public_link_with_configured_port

from .portal_service import (
    BUILDING_OPEN_ID_MAP,
    PortalConflictError,
    PortalError,
    PortalNotFoundError,
)


POLLING_SOP_NAMESPACE = "polling_sop"
POLLING_WORK_ORDER_NAMESPACE = "polling_work_order"
POLLING_WORK_ORDER_SECRET_NAMESPACE = "polling_work_order_secret"
POLLING_WORK_ORDER_SECRET_KEY = "hmac"
POLLING_UNITS = tuple(f"{index}#" for index in range(1, 7))
POLLING_UNIT_GROUPS = (POLLING_UNITS[:3], POLLING_UNITS[3:])
POLLING_SOP_SCOPES = frozenset({"110", "A", "B", "C", "D", "E", "H"})
POLLING_H_DUTY_RECORD_ID = "h_duty_account"
POLLING_SOP_MAX_FILE_BYTES = 20 * 1024 * 1024
POLLING_SOP_MAX_FILES = 10
POLLING_SOP_MAX_TOTAL_BYTES = 100 * 1024 * 1024
POLLING_SOP_MAX_STEPS = 30
POLLING_STEP_MAX_SECONDS = 24 * 60 * 60
POLLING_STEP_PHOTO_MAX_BYTES = 8 * 1024 * 1024
POLLING_STEP_MAX_PHOTOS = 5
POLLING_WORK_ORDER_MAX_PHOTOS = 100
POLLING_WORK_ORDER_MAX_PHOTO_BYTES = 200 * 1024 * 1024
POLLING_STEP_PHOTO_MAX_PIXELS = 40_000_000
POLLING_STEP_PHOTO_MAX_DIMENSION = 12_000
POLLING_STEP_PHOTO_MIME_TYPES = frozenset(
    {"image/jpeg", "image/png", "image/webp"}
)
POLLING_WORK_ORDER_TEMPLATE_NAME = "轮巡操作流程.xlsx"
POLLING_WORK_ORDER_CACHE_NAME = "轮巡操作流程.v3.xlsx"
POLLING_WORK_ORDER_MAX_BYTES = 20 * 1024 * 1024
POLLING_SOP_APP_TOKEN = "HU38bc1vnamMK9sCeOgclUvXnFc"
POLLING_SOP_TABLE_ID = "tbl5KlqJAhpqWbK4"
POLLING_SOP_CLOUD_FIELDS = {
    "SOP标识": 1,
    "楼栋": 1,
    "工单类型": 1,
    "步骤数据": 1,
    "附件": 17,
    "附件元数据": 1,
    "记录版本": 2,
    "创建时间": 1,
    "更新时间": 1,
    "创建人": 1,
    "更新人": 1,
    "内容哈希": 1,
    "状态": 1,
}
WORK_ORDER_TYPES = frozenset({"polling", "maintenance", "adjust"})
ADJUST_COOLING_MODES = {
    "1#": "停机状态",
    "2#": "板换模式",
    "3#": "预冷模式",
    "4#": "制冷模式",
}


class PollingWorkOrderTokenError(PortalError):
    status_code = 403


def _flag(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "yes", "on", "是"}


def _safe_file_name(value: Any) -> str:
    name = Path(str(value or "").strip()).name
    name = re.sub(r"[\x00-\x1f<>:\"/\\|?*]+", "_", name).strip(" .")
    return name[:160] or "attachment.bin"


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _excel_text(value: Any) -> str:
    text = str(value or "")
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def _polling_unit_group(unit: str) -> tuple[str, ...]:
    return next((group for group in POLLING_UNIT_GROUPS if unit in group), ())


def _work_order_type(value: Any) -> str:
    normalized = str(value or "polling").strip().lower()
    if normalized not in WORK_ORDER_TYPES:
        raise PortalError("工单类型必须是 maintenance、polling 或 adjust。")
    return normalized


def _stored_work_order_type(payload: dict | None) -> str:
    normalized = str((payload or {}).get("work_type") or "polling").strip().lower()
    return normalized if normalized in WORK_ORDER_TYPES else "polling"


def _run_label(work_type: str, run: dict) -> str:
    if work_type == "adjust":
        return str(run.get("label") or "设备调整作业")
    if work_type == "maintenance":
        return "维保作业"
    return f"{run.get('from_unit') or ''}→{run.get('to_unit') or ''}"


def _adjust_cooling_sop_tokens(steps: list[dict]) -> set[str]:
    content = "\n".join(str(step.get("content") or "") for step in steps)
    return {
        token
        for token in ("{{from}}", "{{to}}", "{{other}}")
        if token in content
    }


class PollingSopCloudStore:
    _schema_ready = False
    _schema_lock = threading.Lock()

    def __init__(self) -> None:
        from upload_event_module.services.http_client import FeishuHttpClient

        self.client = FeishuHttpClient(retries=1)
        self.root = (
            "https://open.feishu.cn/open-apis/bitable/v1/apps/"
            f"{POLLING_SOP_APP_TOKEN}/tables/{POLLING_SOP_TABLE_ID}"
        )
        self._cache_lock = threading.RLock()
        self._cache: list[dict] = []
        self._cache_until = 0.0

    @staticmethod
    def _token() -> str:
        from upload_event_module.services.feishu_token_manager import token_manager

        return token_manager.get_tenant_token()

    def _request(self, method: str, path: str, body=None, params=None) -> dict:
        if method != "GET":
            from .portal_service import external_real_write_guard

            guard = external_real_write_guard()
            if not guard.get("real_write_allowed"):
                raise PortalError(str(guard.get("reason") or "真实飞书写入未确认。"))
        payload = self.client.request_json(
            method,
            self.root + "/" + path.lstrip("/"),
            headers={"Authorization": "Bearer " + self._token()},
            params=params or {},
            json_payload=body,
            retries=1 if method == "GET" else 0,
        )
        if int(payload.get("code") or 0):
            raise PortalError(
                f"SOP多维表请求失败：code={payload.get('code')}，{payload.get('msg') or ''}"
            )
        return payload.get("data") if isinstance(payload.get("data"), dict) else {}

    def ensure_schema(self) -> None:
        if type(self)._schema_ready:
            return
        with type(self)._schema_lock:
            if type(self)._schema_ready:
                return
            fields = {
                str(item.get("field_name") or ""): item
                for item in self._request("GET", "fields", params={"page_size": 100}).get("items", [])
            }
            for name, field_type in POLLING_SOP_CLOUD_FIELDS.items():
                current = fields.get(name)
                if current and int(current.get("type") or 0) != field_type:
                    raise PortalError(f"SOP多维表字段类型错误：{name}")
                if not current:
                    self._request("POST", "fields", {"field_name": name, "type": field_type})
            type(self)._schema_ready = True

    def list_records(self) -> list[dict]:
        self.ensure_schema()
        records: list[dict] = []
        page_token = ""
        while True:
            params = {"page_size": 500}
            if page_token:
                params["page_token"] = page_token
            page = self._request("GET", "records", params=params)
            records.extend(item for item in page.get("items", []) if isinstance(item, dict))
            if not page.get("has_more"):
                return records
            page_token = str(page.get("page_token") or "")
            if not page_token:
                raise PortalError("SOP多维表分页游标无效。")

    @staticmethod
    def _decode(record: dict) -> dict | None:
        fields = record.get("fields") if isinstance(record.get("fields"), dict) else {}
        sop_id = str(fields.get("SOP标识") or "").strip()
        if not sop_id or str(fields.get("状态") or "启用").strip() == "已删除":
            return None
        try:
            steps = json.loads(str(fields.get("步骤数据") or "[]"))
            attachment_meta = json.loads(str(fields.get("附件元数据") or "[]"))
        except (TypeError, ValueError) as exc:
            raise PortalError(f"SOP云端数据格式错误：{sop_id}") from exc
        raw_attachments = fields.get("附件") if isinstance(fields.get("附件"), list) else []
        tokens = {
            str(item.get("file_token") or item.get("token") or ""): item
            for item in raw_attachments
            if isinstance(item, dict)
        }
        attachments = []
        for index, item in enumerate(attachment_meta if isinstance(attachment_meta, list) else []):
            if not isinstance(item, dict):
                continue
            token = str(item.get("file_token") or "")
            remote = tokens.get(token) or (raw_attachments[index] if index < len(raw_attachments) else {})
            attachments.append(
                {
                    **item,
                    "file_token": token or str((remote or {}).get("file_token") or (remote or {}).get("token") or ""),
                    "name": str(item.get("name") or (remote or {}).get("name") or "SOP附件"),
                    "_cloud_download_url": str((remote or {}).get("url") or ""),
                }
            )
        return {
            "sop_id": sop_id,
            "work_type": str(fields.get("工单类型") or "polling"),
            "scope": str(fields.get("楼栋") or "").replace("楼", "").upper(),
            "name": str(fields.get("文本") or ""),
            "steps": steps if isinstance(steps, list) else [],
            "attachments": attachments,
            "version": int(float(fields.get("记录版本") or 0)),
            "created_at": str(fields.get("创建时间") or ""),
            "updated_at": str(fields.get("更新时间") or ""),
            "created_by": str(fields.get("创建人") or ""),
            "updated_by": str(fields.get("更新人") or ""),
            "_cloud_record_id": str(record.get("record_id") or ""),
        }

    def list_sops(self, *, force: bool = False) -> list[dict]:
        with self._cache_lock:
            if not force and time.monotonic() < self._cache_until:
                return copy.deepcopy(self._cache)
        items = [sop for record in self.list_records() if (sop := self._decode(record))]
        with self._cache_lock:
            self._cache = copy.deepcopy(items)
            self._cache_until = time.monotonic() + 15.0
        return items

    def get_sop(self, sop_id: str, *, force: bool = False) -> dict | None:
        return next((sop for sop in self.list_sops(force=force) if sop["sop_id"] == sop_id), None)

    def invalidate(self) -> None:
        with self._cache_lock:
            self._cache_until = 0.0

    def close(self) -> None:
        self.client.close()

    @staticmethod
    def _fields(sop: dict) -> dict:
        attachments = [
            {key: value for key, value in item.items() if key in {"attachment_id", "name", "size", "sha256", "file_token", "created_at", "created_by"}}
            for item in sop.get("attachments") or []
        ]
        fields = {
            "文本": str(sop.get("name") or ""),
            "SOP标识": str(sop.get("sop_id") or ""),
            "楼栋": str(sop.get("scope") or "") + "楼",
            "工单类型": str(sop.get("work_type") or "polling"),
            "步骤数据": json.dumps(sop.get("steps") or [], ensure_ascii=False, separators=(",", ":")),
            "附件": [{"file_token": item["file_token"]} for item in attachments if item.get("file_token")],
            "附件元数据": json.dumps(attachments, ensure_ascii=False, separators=(",", ":")),
            "记录版本": int(sop.get("version") or 0),
            "创建时间": str(sop.get("created_at") or ""),
            "更新时间": str(sop.get("updated_at") or ""),
            "创建人": str(sop.get("created_by") or ""),
            "更新人": str(sop.get("updated_by") or ""),
            "内容哈希": hashlib.sha256(json.dumps([sop.get("steps") or [], attachments], ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest(),
            "状态": "启用",
        }
        return fields

    def save_sop(self, sop: dict, *, expected_version: int, allow_create: bool = False) -> dict:
        existing = self.get_sop(str(sop.get("sop_id") or ""), force=True)
        if existing and int(existing.get("version") or 0) != int(expected_version or 0):
            raise PortalConflictError("SOP 已被其他电脑修改，请刷新后重试。")
        if not existing and expected_version and not allow_create:
            raise PortalConflictError("SOP 云端版本不存在，请刷新后重试。")
        body = {"fields": self._fields(sop)}
        if existing:
            record = self._request("PUT", f"records/{existing['_cloud_record_id']}", body).get("record") or {}
        else:
            record = self._request("POST", "records", body).get("record") or {}
        result = copy.deepcopy(sop)
        result["_cloud_record_id"] = str(record.get("record_id") or (existing or {}).get("_cloud_record_id") or "")
        self.invalidate()
        return result

    def upload_attachment(self, path: Path, file_name: str) -> str:
        payload = self.client.request_file_json(
            "POST",
            "https://open.feishu.cn/open-apis/drive/v1/medias/upload_all",
            headers={"Authorization": "Bearer " + self._token()},
            data={"file_name": file_name, "parent_type": "bitable_file", "parent_node": POLLING_SOP_APP_TOKEN, "size": str(path.stat().st_size)},
            file_path=str(path),
            file_name=file_name,
            retries=1,
        )
        if int(payload.get("code") or 0) or not str((payload.get("data") or {}).get("file_token") or ""):
            raise PortalError(f"SOP附件上传失败：{payload.get('msg') or '未返回文件标识'}")
        return str(payload["data"]["file_token"])

    def download_attachment(self, attachment: dict) -> bytes:
        file_token = str(attachment.get("file_token") or "").strip()
        if not file_token:
            raise PortalError("SOP附件缺少有效的飞书文件标识。")
        download_url = str(attachment.get("_cloud_download_url") or "").strip()
        if download_url:
            parsed = urlsplit(download_url)
            if (
                parsed.scheme != "https"
                or parsed.hostname != "open.feishu.cn"
                or not parsed.path.startswith("/open-apis/drive/v1/medias/")
            ):
                raise PortalError("SOP附件下载地址无效。")
        content, _ = self.client.request_bytes(
            "GET",
            download_url
            or "https://open.feishu.cn/open-apis/drive/v1/medias/" + quote(file_token, safe="") + "/download",
            headers={"Authorization": "Bearer " + self._token()},
            retries=1,
            max_bytes=POLLING_SOP_MAX_FILE_BYTES,
        )
        return content

    def delete_sop(self, sop: dict) -> None:
        record_id = str(sop.get("_cloud_record_id") or "")
        if record_id:
            self._request("DELETE", f"records/{record_id}")
            self.invalidate()


class PollingWorkOrderService:
    _lock = threading.RLock()

    def __init__(self, state_store, cloud: PollingSopCloudStore | None = None) -> None:
        self.state_store = state_store
        self.cloud = cloud
        self.sop_root = Path(get_data_file_path("polling_sop")).resolve()
        self.work_order_root = Path(
            get_data_file_path("polling_work_orders")
        ).resolve()
        self.work_order_template_path = (
            Path(__file__).resolve().parent
            / "templates"
            / POLLING_WORK_ORDER_TEMPLATE_NAME
        )
        self._legacy_cloud_sync_lock = threading.Lock()
        self._legacy_cloud_sync_done = False

    @staticmethod
    def _now_text() -> str:
        return time.strftime("%Y-%m-%d %H:%M:%S", time.localtime())

    @staticmethod
    def _public_sop(sop: dict) -> dict:
        result = copy.deepcopy(sop or {})
        result.pop("_cloud_record_id", None)
        sop_id = str(result.get("sop_id") or "")
        result["attachments"] = [
            {
                **{key: value for key, value in item.items() if key not in {"path", "file_token", "_cloud_download_url"}},
                "download_url": (
                    f"/api/polling-sops/{quote(sop_id, safe='')}/attachments/"
                    f"{quote(str(item.get('attachment_id') or ''), safe='')}"
                ),
            }
            for item in result.get("attachments") or []
            if isinstance(item, dict)
        ]
        result["ready"] = bool(result.get("steps") and result.get("attachments"))
        return result

    def _cache_cloud_sop(self, remote: dict) -> dict:
        sop_id = str(remote.get("sop_id") or "")
        scope = str(remote.get("scope") or "").strip().upper()
        work_type = _work_order_type(remote.get("work_type"))
        if not re.fullmatch(r"[A-Za-z0-9_-]{8,80}", sop_id) or scope not in POLLING_SOP_SCOPES:
            raise PortalError("SOP 云端标识或楼栋无效。")
        remote = {
            **copy.deepcopy(remote),
            "scope": scope,
            "work_type": work_type,
            "steps": self._normalized_steps(remote.get("steps"), work_type=work_type),
        }
        name = str(remote.get("name") or "").strip()
        if not name or len(name) > 160:
            raise PortalError("SOP 云端名称不能为空且不能超过 160 个字符。")
        remote["name"] = name
        raw_attachments = remote.get("attachments") or []
        if not isinstance(raw_attachments, list) or len(raw_attachments) > POLLING_SOP_MAX_FILES:
            raise PortalError(f"SOP 云端附件不能超过 {POLLING_SOP_MAX_FILES} 个。")
        local = self.state_store.get_document(POLLING_SOP_NAMESPACE, sop_id) or {}
        local_attachments = {
            str(item.get("attachment_id") or ""): item
            for item in local.get("attachments") or []
            if isinstance(item, dict)
        }
        attachments = []
        total_size = 0
        for item in raw_attachments:
            if not isinstance(item, dict):
                raise PortalError("SOP 云端附件格式无效。")
            attachment = copy.deepcopy(item)
            if not re.fullmatch(r"[A-Za-z0-9_-]{8,80}", str(attachment.get("attachment_id") or "")):
                raise PortalError("SOP 云端附件标识无效。")
            if not re.fullmatch(r"[A-Za-z0-9_-]{8,256}", str(attachment.get("file_token") or "")):
                raise PortalError("SOP 云端附件缺少有效的飞书文件标识。")
            try:
                size = int(attachment.get("size") or 0)
            except (TypeError, ValueError) as exc:
                raise PortalError("SOP 云端附件大小无效。") from exc
            if size < 0 or size > POLLING_SOP_MAX_FILE_BYTES:
                raise PortalError("SOP 云端单个附件不能超过 20MB。")
            total_size += size
            if total_size > POLLING_SOP_MAX_TOTAL_BYTES:
                raise PortalError("SOP 云端附件总大小不能超过 100MB。")
            digest = str(attachment.get("sha256") or "").strip().lower()
            if digest and not re.fullmatch(r"[0-9a-f]{64}", digest):
                raise PortalError("SOP 云端附件校验值无效。")
            attachment["size"] = size
            attachment["sha256"] = digest
            attachment["name"] = _safe_file_name(attachment.get("name"))
            existing = local_attachments.get(str(attachment.get("attachment_id") or "")) or {}
            path = str(existing.get("path") or "")
            if not path:
                directory = self._sop_directory(sop_id)
                path = str((directory / f"{attachment.get('attachment_id')}_{_safe_file_name(attachment.get('name'))}").resolve())
            attachment["path"] = path
            attachments.append(attachment)
        cached = {**copy.deepcopy(remote), "attachments": attachments}
        if cached != local:
            self.state_store.put_document(POLLING_SOP_NAMESPACE, sop_id, cached)
        return cached

    def _sync_legacy_local_sops(self, cloud_items: list[dict]) -> tuple[list[dict], list[dict]]:
        if not self.cloud or self._legacy_cloud_sync_done:
            return cloud_items, []
        with self._legacy_cloud_sync_lock:
            if self._legacy_cloud_sync_done:
                return cloud_items, []
            known_ids = {str(item.get("sop_id") or "") for item in cloud_items}
            pending: list[dict] = []
            changed = False
            documents = self.state_store.list_documents(POLLING_SOP_NAMESPACE)
            for document in documents:
                local = copy.deepcopy(document.get("payload") or {})
                sop_id = str(local.get("sop_id") or "").strip()
                if not sop_id or sop_id in known_ids:
                    continue
                try:
                    scope = str(local.get("scope") or "").strip().upper()
                    work_type = _work_order_type(local.get("work_type"))
                    name = str(local.get("name") or "").strip()
                    if not re.fullmatch(r"[A-Za-z0-9_-]{8,80}", sop_id) or scope not in POLLING_SOP_SCOPES:
                        raise PortalError("本地旧 SOP 标识或楼栋无效。")
                    if not name or len(name) > 160:
                        raise PortalError("本地旧 SOP 名称无效。")
                    local.update(
                        scope=scope,
                        work_type=work_type,
                        name=name,
                        steps=self._normalized_steps(local.get("steps"), work_type=work_type),
                    )
                    attachments = copy.deepcopy(local.get("attachments") or [])
                    if not isinstance(attachments, list) or len(attachments) > POLLING_SOP_MAX_FILES:
                        raise PortalError("本地旧 SOP 附件数量无效。")
                    total_size = 0
                    for attachment in attachments:
                        if not isinstance(attachment, dict):
                            raise PortalError("本地旧 SOP 附件格式无效。")
                        attachment_id = str(attachment.get("attachment_id") or "")
                        if not re.fullmatch(r"[A-Za-z0-9_-]{8,80}", attachment_id):
                            raise PortalError("本地旧 SOP 附件标识无效。")
                        path = Path(str(attachment.get("path") or "")).resolve()
                        directory = self._sop_directory(sop_id)
                        if not path.is_file() or not path.is_relative_to(directory):
                            raise PortalError(f"本地旧 SOP 附件不存在：{_safe_file_name(attachment.get('name'))}")
                        size = path.stat().st_size
                        total_size += size
                        if size > POLLING_SOP_MAX_FILE_BYTES or total_size > POLLING_SOP_MAX_TOTAL_BYTES:
                            raise PortalError("本地旧 SOP 附件大小超限。")
                        digest_value = _file_sha256(path)
                        expected_hash = str(attachment.get("sha256") or "").strip().lower()
                        if expected_hash and digest_value != expected_hash:
                            raise PortalError(f"本地旧 SOP 附件校验失败：{path.name}")
                        attachment.update(size=size, sha256=digest_value, name=_safe_file_name(attachment.get("name")))
                        if not str(attachment.get("file_token") or "").strip():
                            attachment["file_token"] = self.cloud.upload_attachment(path, attachment["name"])
                            local["attachments"] = attachments
                            self.state_store.put_document(POLLING_SOP_NAMESPACE, sop_id, local)
                    local["attachments"] = attachments
                    self.cloud.save_sop(local, expected_version=0, allow_create=True)
                    known_ids.add(sop_id)
                    changed = True
                except PortalConflictError:
                    existing = self.cloud.get_sop(sop_id, force=True)
                    if existing:
                        known_ids.add(sop_id)
                        changed = True
                    else:
                        pending.append(local)
                except Exception as exc:
                    print(f"[ClipFlow] 本地旧 SOP 云端迁移待重试: sop_id={sop_id}, error={exc}")
                    pending.append(local)
            if changed:
                cloud_items = self.cloud.list_sops(force=True)
            self._legacy_cloud_sync_done = not pending
            return cloud_items, pending

    def _cloud_sops(self, *, force: bool = False) -> list[dict] | None:
        if not self.cloud:
            return None
        try:
            cloud_items, pending = self._sync_legacy_local_sops(
                self.cloud.list_sops(force=force)
            )
            return [self._cache_cloud_sop(item) for item in cloud_items] + pending
        except Exception as exc:
            print(f"[ClipFlow] SOP cloud read failed, using local cache: {exc}")
            return None

    def list_sops(self, scope: str, work_type: str = "polling") -> list[dict]:
        scope = str(scope or "").strip().upper()
        work_type = _work_order_type(work_type)
        if scope not in POLLING_SOP_SCOPES:
            raise PortalError("请在明确的单楼页面读取 SOP。")
        cloud_items = self._cloud_sops()
        source = cloud_items if cloud_items is not None else [
            document.get("payload") or {}
            for document in self.state_store.list_documents(POLLING_SOP_NAMESPACE)
            if isinstance(document.get("payload"), dict)
        ]
        items = [self._public_sop(item) for item in source if str(item.get("scope") or "").strip().upper() == scope and _stored_work_order_type(item) == work_type]
        return sorted(items, key=lambda item: str(item.get("name") or "").casefold())

    def get_sop(self, sop_id: str, *, public: bool = True, refresh: bool = False) -> dict:
        sop_id = str(sop_id or "").strip()
        if self.cloud:
            try:
                remote = self.cloud.get_sop(sop_id, force=refresh)
            except Exception as exc:
                print(f"[ClipFlow] SOP cloud read failed, using local cache: {exc}")
            else:
                if not remote:
                    self.state_store.delete_document(POLLING_SOP_NAMESPACE, sop_id)
                    raise PortalNotFoundError("轮巡 SOP 不存在。")
                sop = self._cache_cloud_sop(remote)
                return self._public_sop(sop) if public else copy.deepcopy(sop)
        sop = self.state_store.get_document(POLLING_SOP_NAMESPACE, sop_id)
        if not isinstance(sop, dict):
            raise PortalNotFoundError("轮巡 SOP 不存在。")
        return self._public_sop(sop) if public else copy.deepcopy(sop)

    @staticmethod
    def _normalized_steps(value: Any, *, work_type: str = "polling") -> list[dict]:
        work_type = _work_order_type(work_type)
        raw_steps = value if isinstance(value, list) else []
        if len(raw_steps) > POLLING_SOP_MAX_STEPS:
            raise PortalError(f"SOP 步骤不能超过 {POLLING_SOP_MAX_STEPS} 条。")
        steps: list[dict] = []
        for index, raw in enumerate(raw_steps):
            if not isinstance(raw, dict):
                raise PortalError(f"第 {index + 1} 个 SOP 步骤格式无效。")
            content = str(raw.get("content") or "").strip()
            if not content:
                raise PortalError(f"第 {index + 1} 个 SOP 步骤缺少操作内容。")
            if re.search(r"(?<!\{)\{(?:from|to|other)\}(?!\})", content):
                raise PortalError(
                    f"第 {index + 1} 个 SOP 步骤占位符必须使用 "
                    "{{from}}、{{to}}、{{other}}。"
                )
            if work_type == "maintenance" and re.search(
                r"\{\{(?:from|to|other)\}\}", content
            ):
                raise PortalError(
                    f"第 {index + 1} 个维保 SOP 步骤不能使用设备指向占位符。"
                )
            if len(content) > 5000:
                raise PortalError(f"第 {index + 1} 个 SOP 步骤内容过长。")
            operator_required = _flag(raw.get("operator_required"))
            reviewer_required = _flag(raw.get("reviewer_required"))
            photo_required = (
                True if "photo_required" not in raw else _flag(raw.get("photo_required"))
            )
            if not operator_required and not reviewer_required:
                raise PortalError(
                    f"第 {index + 1} 个 SOP 步骤至少需要操作人或现场审核人确认。"
                )
            try:
                time_limit_seconds = int(raw.get("time_limit_seconds") or 0)
            except (TypeError, ValueError) as exc:
                raise PortalError(
                    f"第 {index + 1} 个 SOP 步骤时间限制必须为整数秒。"
                ) from exc
            if time_limit_seconds not in range(POLLING_STEP_MAX_SECONDS + 1):
                raise PortalError(
                    f"第 {index + 1} 个 SOP 步骤时间限制必须在 0–{POLLING_STEP_MAX_SECONDS} 秒之间。"
                )
            steps.append(
                {
                    "step_id": str(raw.get("step_id") or uuid.uuid4().hex),
                    "order": index + 1,
                    "content": content,
                    "operator_required": operator_required,
                    "reviewer_required": reviewer_required,
                    "photo_required": photo_required,
                    "time_limit_seconds": time_limit_seconds,
                }
            )
        return steps

    def save_sop(self, payload: dict, *, actor_open_id: str = "") -> dict:
        payload = payload if isinstance(payload, dict) else {}
        work_type = _work_order_type(payload.get("work_type"))
        name = str(payload.get("name") or "").strip()
        scope = str(payload.get("scope") or "").strip().upper()
        if scope not in POLLING_SOP_SCOPES:
            raise PortalError("请在明确的单楼页面维护 SOP。")
        if not name or len(name) > 160:
            raise PortalError("SOP 名称不能为空且不能超过 160 个字符。")
        steps = self._normalized_steps(payload.get("steps"), work_type=work_type)
        sop_id = str(payload.get("sop_id") or "").strip() or uuid.uuid4().hex
        expected_version = int(payload.get("expected_version") or 0)
        with self._lock:
            existing = self.state_store.get_document(POLLING_SOP_NAMESPACE, sop_id)
            if existing and expected_version != int(existing.get("version") or 0):
                raise PortalConflictError("SOP 已被其他用户修改，请刷新后重试。")
            if not existing and expected_version:
                raise PortalConflictError("SOP 版本已失效，请刷新后重试。")
            if existing and str(existing.get("scope") or "").strip().upper() not in {"", scope}:
                raise PortalConflictError("SOP 不能跨楼栋修改。")
            if existing and _stored_work_order_type(existing) != work_type:
                raise PortalConflictError("SOP 不能跨通告类型修改。")
            for item in self.list_sops(scope, work_type):
                if (
                    str(item.get("sop_id") or "") != sop_id
                    and str(item.get("scope") or "").strip().upper() == scope
                    and str(item.get("name") or "").strip().casefold() == name.casefold()
                ):
                    raise PortalConflictError("当前楼栋已存在同名 SOP。")
            now = self._now_text()
            sop = {
                **(existing or {}),
                "sop_id": sop_id,
                "work_type": work_type,
                "scope": scope,
                "name": name,
                "steps": steps,
                "attachments": copy.deepcopy((existing or {}).get("attachments") or []),
                "version": int((existing or {}).get("version") or 0) + 1,
                "created_at": str((existing or {}).get("created_at") or now),
                "created_by": str((existing or {}).get("created_by") or actor_open_id),
                "updated_at": now,
                "updated_by": str(actor_open_id or ""),
            }
            if self.cloud:
                sop = self.cloud.save_sop(
                    sop,
                    expected_version=expected_version,
                    allow_create=bool(existing),
                )
            self.state_store.put_document(POLLING_SOP_NAMESPACE, sop_id, sop)
            self._sync_unstarted_work_order_limits(sop)
        return self._public_sop(sop)

    def _sync_unstarted_work_order_limits(self, sop: dict) -> None:
        settings_by_id = {
            str(step.get("step_id") or ""): (
                int(step.get("time_limit_seconds") or 0),
                self._step_photo_required(step),
            )
            for step in sop.get("steps") or []
            if str(step.get("step_id") or "")
        }
        settings_by_index = {
            index: (
                int(step.get("time_limit_seconds") or 0),
                self._step_photo_required(step),
            )
            for index, step in enumerate(sop.get("steps") or [], start=1)
        }
        for document in self.state_store.list_documents(POLLING_WORK_ORDER_NAMESPACE):
            group = document.get("payload") or {}
            if (
                not isinstance(group, dict)
                or str(group.get("state") or "") != "active"
                or str(group.get("sop_id") or "") != str(sop.get("sop_id") or "")
            ):
                continue
            steps = list(group.get("steps") or [])
            started_runs = {
                int(step.get("run_index") or 0)
                for step in steps
                if float(step.get("activated_at_ts") or 0) > 0
                or bool(step.get("operator_confirmation"))
                or bool(step.get("reviewer_confirmation"))
            }
            changed = False
            for step in steps:
                if int(step.get("run_index") or 0) in started_runs:
                    continue
                sop_step_id = str(step.get("sop_step_id") or "")
                settings = settings_by_id.get(sop_step_id) if sop_step_id else None
                if not sop_step_id:
                    settings = settings_by_index.get(int(step.get("step_index") or 0))
                if settings is not None:
                    limit, photo_required = settings
                    if int(step.get("time_limit_seconds") or 0) != limit:
                        step["time_limit_seconds"] = limit
                        changed = True
                    if self._step_photo_required(step) != photo_required:
                        step["photo_required"] = photo_required
                        changed = True
            if changed:
                group["steps"] = steps
                group["updated_at"] = self._now_text()
                self.state_store.put_document(
                    POLLING_WORK_ORDER_NAMESPACE,
                    str(group.get("target_record_id") or ""),
                    group,
                )

    def _sop_directory(self, sop_id: str) -> Path:
        if not re.fullmatch(r"[A-Za-z0-9_-]{8,80}", str(sop_id or "")):
            raise PortalError("SOP ID 无效。")
        path = (self.sop_root / sop_id).resolve()
        if path == self.sop_root or not path.is_relative_to(self.sop_root):
            raise PortalError("SOP 附件路径无效。")
        return path

    def _ensure_local_attachment(self, sop: dict, attachment: dict) -> Path:
        directory = self._sop_directory(str(sop.get("sop_id") or ""))
        path = Path(str(attachment.get("path") or "")).resolve()
        if path.is_file() and path.is_relative_to(directory):
            return path
        if not self.cloud or not str(attachment.get("file_token") or "").strip():
            raise PortalNotFoundError("SOP 附件文件不存在。")
        content = self.cloud.download_attachment(attachment)
        if len(content) > POLLING_SOP_MAX_FILE_BYTES:
            raise PortalError("SOP 单个附件不能超过 20MB。")
        expected_hash = str(attachment.get("sha256") or "")
        if expected_hash and hashlib.sha256(content).hexdigest() != expected_hash:
            raise PortalError("SOP 附件下载后校验失败。")
        directory.mkdir(parents=True, exist_ok=True)
        path = (directory / f"{attachment.get('attachment_id')}_{_safe_file_name(attachment.get('name'))}").resolve()
        temporary = path.with_suffix(path.suffix + ".tmp")
        temporary.write_bytes(content)
        os.replace(temporary, path)
        attachment["path"] = str(path)
        self.state_store.put_document(POLLING_SOP_NAMESPACE, str(sop.get("sop_id") or ""), sop)
        return path

    def add_sop_attachment(
        self,
        sop_id: str,
        *,
        file_name: str,
        content: bytes,
        expected_version: int,
        actor_open_id: str = "",
    ) -> dict:
        content = bytes(content or b"")
        if not content:
            raise PortalError("SOP 附件内容为空。")
        if len(content) > POLLING_SOP_MAX_FILE_BYTES:
            raise PortalError("SOP 单个附件不能超过 20MB。")
        with self._lock:
            sop = self.get_sop(sop_id, public=False, refresh=True)
            if int(expected_version or 0) != int(sop.get("version") or 0):
                raise PortalConflictError("SOP 已被修改，请刷新后重试。")
            attachments = list(sop.get("attachments") or [])
            if len(attachments) >= POLLING_SOP_MAX_FILES:
                raise PortalError(f"每个 SOP 最多上传 {POLLING_SOP_MAX_FILES} 个附件。")
            if sum(int(item.get("size") or 0) for item in attachments) + len(content) > POLLING_SOP_MAX_TOTAL_BYTES:
                raise PortalError("SOP 附件总大小不能超过 100MB。")
            attachment_id = uuid.uuid4().hex
            safe_name = _safe_file_name(file_name)
            directory = self._sop_directory(sop_id)
            directory.mkdir(parents=True, exist_ok=True)
            path = (directory / f"{attachment_id}_{safe_name}").resolve()
            if not path.is_relative_to(directory):
                raise PortalError("SOP 附件路径无效。")
            temporary = path.with_suffix(path.suffix + ".tmp")
            temporary.write_bytes(content)
            os.replace(temporary, path)
            attachment = {
                "attachment_id": attachment_id,
                "name": safe_name,
                "size": len(content),
                "sha256": hashlib.sha256(content).hexdigest(),
                "path": str(path),
                "created_at": self._now_text(),
                "created_by": str(actor_open_id or ""),
            }
            if self.cloud:
                try:
                    attachment["file_token"] = self.cloud.upload_attachment(path, safe_name)
                except Exception:
                    path.unlink(missing_ok=True)
                    raise
            attachments.append(attachment)
            sop.update(
                {
                    "attachments": attachments,
                    "version": int(sop.get("version") or 0) + 1,
                    "updated_at": self._now_text(),
                    "updated_by": str(actor_open_id or ""),
                }
            )
            if self.cloud:
                try:
                    sop = self.cloud.save_sop(sop, expected_version=expected_version)
                except Exception:
                    path.unlink(missing_ok=True)
                    raise
            self.state_store.put_document(POLLING_SOP_NAMESPACE, sop_id, sop)
        return self._public_sop(sop)

    def get_sop_attachment(self, sop_id: str, attachment_id: str) -> tuple[bytes, str]:
        sop = self.get_sop(sop_id, public=False)
        attachment = next(
            (
                item
                for item in sop.get("attachments") or []
                if str(item.get("attachment_id") or "") == str(attachment_id or "")
            ),
            None,
        )
        if not attachment:
            raise PortalNotFoundError("SOP 附件不存在。")
        path = self._ensure_local_attachment(sop, attachment)
        return path.read_bytes(), str(attachment.get("name") or path.name)

    def delete_sop_attachment(
        self,
        sop_id: str,
        attachment_id: str,
        *,
        expected_version: int,
        actor_open_id: str = "",
    ) -> dict:
        with self._lock:
            sop = self.get_sop(sop_id, public=False, refresh=True)
            if int(expected_version or 0) != int(sop.get("version") or 0):
                raise PortalConflictError("SOP 已被修改，请刷新后重试。")
            attachments = list(sop.get("attachments") or [])
            attachment = next(
                (item for item in attachments if str(item.get("attachment_id") or "") == str(attachment_id or "")),
                None,
            )
            if not attachment:
                raise PortalNotFoundError("SOP 附件不存在。")
            path = Path(str(attachment.get("path") or "")).resolve()
            directory = self._sop_directory(sop_id)
            sop.update(
                {
                    "attachments": [item for item in attachments if item is not attachment],
                    "version": int(sop.get("version") or 0) + 1,
                    "updated_at": self._now_text(),
                    "updated_by": str(actor_open_id or ""),
                }
            )
            if self.cloud:
                sop = self.cloud.save_sop(sop, expected_version=expected_version)
            self.state_store.put_document(POLLING_SOP_NAMESPACE, sop_id, sop)
            if path.is_relative_to(directory) and path.is_file():
                path.unlink()
        return self._public_sop(sop)

    def delete_sop(self, sop_id: str, *, expected_version: int) -> dict:
        with self._lock:
            sop = self.get_sop(sop_id, public=False, refresh=True)
            if int(expected_version or 0) != int(sop.get("version") or 0):
                raise PortalConflictError("SOP 已被修改，请刷新后重试。")
            if self.cloud:
                self.cloud.delete_sop(sop)
            directory = self._sop_directory(sop_id)
            if directory.exists():
                shutil.rmtree(directory)
            self.state_store.delete_document(POLLING_SOP_NAMESPACE, sop_id)
        return {"deleted": True, "sop_id": sop_id}

    @staticmethod
    def _person_by_id(
        people: list[dict], record_id: str, *, allow_h_duty: bool = False
    ) -> dict:
        if allow_h_duty and str(record_id or "").strip() == POLLING_H_DUTY_RECORD_ID:
            return {
                "record_id": POLLING_H_DUTY_RECORD_ID,
                "name": "H楼值班账号",
                "open_id": str(BUILDING_OPEN_ID_MAP.get("H") or "").strip(),
                "employee_no": "",
                "building": "H楼",
                "position": "值班账号",
                "shift": "",
            }
        person = next(
            (
                item
                for item in people or []
                if str(item.get("record_id") or "").strip() == str(record_id or "").strip()
            ),
            None,
        )
        if not person:
            raise PortalError("所选人员已不在当前人员表，请重新选择。")
        result = {
            key: str(person.get(key) or "").strip()
            for key in ("record_id", "name", "open_id", "employee_no", "building", "position", "shift", "account_nature")
        }
        if person.get("can_receive_message") is False:
            result["open_id"] = ""
        return result

    def prepare_start(
        self,
        request_payload: dict,
        *,
        job_id: str,
        people: list[dict],
    ) -> dict:
        work_type = str(request_payload.get("work_type") or "").strip().lower()
        if (
            work_type not in WORK_ORDER_TYPES
            or str(request_payload.get("action") or "").strip().lower() != "start"
            or not _flag(request_payload.get("_web_action_request"))
        ):
            return {}
        if _flag(request_payload.get("polling_work_order_exempt")):
            return {"polling_work_order_exempt": True}
        sop = self.get_sop(
            str(request_payload.get("polling_sop_id") or ""),
            public=False,
            refresh=bool(self.cloud),
        )
        request_scope = str(request_payload.get("scope") or "").strip().upper()
        if request_scope not in POLLING_SOP_SCOPES or str(
            sop.get("scope") or ""
        ).strip().upper() != request_scope:
            raise PortalError("所选 SOP 不属于当前楼栋，请重新选择。")
        if _stored_work_order_type(sop) != work_type:
            raise PortalError("所选 SOP 不适用于当前工单类型，请重新选择。")
        sop_steps = [step for step in sop.get("steps") or [] if isinstance(step, dict)]
        adjust_tokens = _adjust_cooling_sop_tokens(sop_steps)
        has_adjust_cooling_step = work_type == "adjust" and adjust_tokens == {"{{from}}"}
        has_any_placeholder = bool(adjust_tokens)
        has_unsupported_placeholder = has_any_placeholder and not has_adjust_cooling_step
        if work_type in {"maintenance", "adjust"} and has_unsupported_placeholder:
            if work_type == "maintenance":
                raise PortalError("所选 SOP 包含轮巡设备指向占位符，不能用于维保工单。")
            raise PortalError("制冷单元模式切换 SOP 只允许使用 {{from}} 表示本次选择的制冷单元。")
        expected_version = int(request_payload.get("polling_sop_version") or 0)
        if expected_version != int(sop.get("version") or 0):
            raise PortalConflictError("所选 SOP 已修改，请重新选择。")
        if not sop.get("steps") or not sop.get("attachments"):
            raise PortalError("所选 SOP 缺少步骤或附件，不能用于发送开始。")
        normalized_runs: list[dict] = [{"run_index": 1, "label": _run_label(work_type, {})}]
        if has_adjust_cooling_step:
            runs = request_payload.get("polling_runs")
            runs = runs if isinstance(runs, list) else []
            run_count = int(request_payload.get("polling_run_count") or 0)
            run = runs[0] if len(runs) == 1 and isinstance(runs[0], dict) else {}
            from_mode = str(run.get("from_unit") or "").strip()
            to_mode = str(run.get("to_unit") or "").strip()
            cooling_unit = str(run.get("other_unit") or "").strip()
            if run_count != 1 or len(runs) != 1 or cooling_unit not in POLLING_UNITS:
                raise PortalError("请选择本次需要调整的 1#–6# 制冷单元。")
            if from_mode not in ADJUST_COOLING_MODES or to_mode not in ADJUST_COOLING_MODES:
                raise PortalError("请选择制冷单元当前运行模式和切换后运行模式。")
            if from_mode == to_mode:
                raise PortalError("制冷单元当前运行模式和切换后运行模式不能相同。")
            from_label = ADJUST_COOLING_MODES[from_mode]
            to_label = ADJUST_COOLING_MODES[to_mode]
            normalized_runs = [{
                "run_index": 1,
                "label": f"{cooling_unit}制冷单元由{from_label}指向{to_label}",
                "from_unit": from_mode,
                "to_unit": to_mode,
                "other_unit": cooling_unit,
            }]
        if work_type == "polling":
            runs = request_payload.get("polling_runs")
            runs = runs if isinstance(runs, list) else []
            run_count = int(request_payload.get("polling_run_count") or 0)
            if run_count not in range(1, 3) or len(runs) != run_count:
                raise PortalError("轮巡次数必须为 1–2 且与轮巡组合数量一致。")
            normalized_runs = []
            seen_pairs: set[tuple[str, str]] = set()
            used_units: set[str] = set()
            for index, item in enumerate(runs):
                item = item if isinstance(item, dict) else {}
                from_unit = str(item.get("from_unit") or "").strip()
                to_unit = str(item.get("to_unit") or "").strip()
                if from_unit not in POLLING_UNITS or to_unit not in POLLING_UNITS:
                    raise PortalError(f"第 {index + 1} 次轮巡设备必须从 1#–6# 中选择。")
                if from_unit == to_unit:
                    raise PortalError(f"第 {index + 1} 次轮巡的起点和终点不能相同。")
                group = _polling_unit_group(from_unit)
                if not group or to_unit not in group:
                    raise PortalError(
                        f"第 {index + 1} 次轮巡不能跨越 1#–3# 与 4#–6# 分组。"
                    )
                if from_unit in used_units or to_unit in used_units:
                    raise PortalError(
                        f"第 {index + 1} 次轮巡使用了前序工单已选择的设备编号。"
                    )
                pair = (from_unit, to_unit)
                if pair in seen_pairs:
                    raise PortalError("轮巡组合不能重复。")
                seen_pairs.add(pair)
                used_units.update((from_unit, to_unit))
                normalized_runs.append(
                    {
                        "run_index": index + 1,
                        "from_unit": from_unit,
                        "to_unit": to_unit,
                        "other_unit": next(
                            unit for unit in group if unit not in {from_unit, to_unit}
                        ),
                    }
                )
        operator = self._person_by_id(
            people,
            str(request_payload.get("polling_operator_record_id") or ""),
        )
        reviewer = self._person_by_id(
            people,
            str(request_payload.get("polling_reviewer_record_id") or ""),
            allow_h_duty=True,
        )
        if (
            operator["record_id"] == reviewer["record_id"]
            or (
                operator.get("open_id")
                and operator.get("open_id") == reviewer.get("open_id")
            )
        ):
            raise PortalError("操作人和现场审核人不能是同一人。")
        pending_root = (self.work_order_root / "pending").resolve()
        staging = (pending_root / re.sub(r"[^A-Za-z0-9_-]", "_", job_id)).resolve()
        if staging == pending_root or not staging.is_relative_to(pending_root):
            raise PortalError("工单附件暂存路径无效。")
        staging.mkdir(parents=True, exist_ok=True)
        snapshot_attachments: list[dict] = []
        for attachment in sop.get("attachments") or []:
            source = self._ensure_local_attachment(sop, attachment)
            target = (staging / f"{attachment.get('attachment_id')}_{_safe_file_name(attachment.get('name'))}").resolve()
            if not target.is_relative_to(staging):
                raise PortalError("工单附件暂存路径无效。")
            if not target.is_file() or _file_sha256(target) != str(attachment.get("sha256") or ""):
                shutil.copy2(source, target)
            snapshot_attachments.append(
                {
                    "attachment_id": str(attachment.get("attachment_id") or ""),
                    "name": str(attachment.get("name") or target.name),
                    "size": int(attachment.get("size") or target.stat().st_size),
                    "sha256": str(attachment.get("sha256") or _file_sha256(target)),
                    "staged_path": str(target),
                }
            )
        prepared_steps = copy.deepcopy(sop.get("steps") or [])
        if has_adjust_cooling_step:
            unit_text = f"{normalized_runs[0]['other_unit']}制冷单元"
            for step in prepared_steps:
                step["content"] = str(step.get("content") or "").replace(
                    "{{from}}", unit_text
                )
        return {
            "polling_work_order_required": True,
            "polling_work_order_spec": {
                "work_type": work_type,
                "sop_id": str(sop.get("sop_id") or ""),
                "sop_version": int(sop.get("version") or 0),
                "sop_name": str(sop.get("name") or ""),
                "scope": request_scope,
                "steps": prepared_steps,
                "attachments": snapshot_attachments,
                "runs": normalized_runs,
                "operator": operator,
                "reviewer": reviewer,
                "initiator_open_id": str(request_payload.get("_auth_open_id") or ""),
                "initiator_name": str(request_payload.get("_auth_user_name") or ""),
            },
            "polling_operator_name": operator["name"],
            "polling_reviewer_name": reviewer["name"],
        }

    def _secret(self) -> bytes:
        with self._lock:
            document = self.state_store.get_document(
                POLLING_WORK_ORDER_SECRET_NAMESPACE,
                POLLING_WORK_ORDER_SECRET_KEY,
            ) or {}
            encoded = str(document.get("secret") or "").strip()
            if not encoded:
                encoded = secrets.token_urlsafe(48)
                self.state_store.put_document(
                    POLLING_WORK_ORDER_SECRET_NAMESPACE,
                    POLLING_WORK_ORDER_SECRET_KEY,
                    {"secret": encoded, "created_at": self._now_text()},
                )
            return encoded.encode("utf-8")

    def role_token(self, target_record_id: str, role: str) -> str:
        payload = f"{target_record_id}.{role}"
        signature = base64.urlsafe_b64encode(
            hmac.new(self._secret(), payload.encode("utf-8"), hashlib.sha256).digest()
        ).decode("ascii").rstrip("=")
        return f"{payload}.{signature}"

    def _resolve_token(self, token: str) -> tuple[dict, str]:
        parts = str(token or "").strip().rsplit(".", 2)
        if len(parts) != 3 or parts[1] not in {"operator", "reviewer"}:
            raise PollingWorkOrderTokenError("工单链接无效或已失效。")
        target_record_id, role, _signature = parts
        expected = self.role_token(target_record_id, role)
        if not hmac.compare_digest(expected, str(token or "").strip()):
            raise PollingWorkOrderTokenError("工单链接无效或已失效。")
        group = self.get_group(target_record_id)
        if str(group.get("state") or "") in {"cancelled", "stopped"}:
            raise PollingWorkOrderTokenError("该工单链接已失效。")
        token_hash = hashlib.sha256(str(token).encode("utf-8")).hexdigest()
        if not hmac.compare_digest(
            token_hash,
            str((group.get("token_hashes") or {}).get(role) or ""),
        ):
            raise PollingWorkOrderTokenError("工单链接无效或已失效。")
        return group, role

    def get_group(self, target_record_id: str) -> dict:
        group = self.state_store.get_document(
            POLLING_WORK_ORDER_NAMESPACE,
            str(target_record_id or "").strip(),
        )
        if not isinstance(group, dict):
            raise PortalNotFoundError("工单不存在。")
        return copy.deepcopy(group)

    def validate_group_token(self, token: str, target_record_id: str) -> None:
        group, _role = self._resolve_token(token)
        if str(group.get("target_record_id") or "") != str(
            target_record_id or ""
        ).strip():
            raise PollingWorkOrderTokenError("工单链接无效或已失效。")

    def _group_directory(self, target_record_id: str) -> Path:
        directory_name = hashlib.sha256(
            str(target_record_id or "").strip().encode("utf-8")
        ).hexdigest()[:24]
        groups_root = (self.work_order_root / "groups").resolve()
        directory = (groups_root / directory_name).resolve()
        if directory == groups_root or not directory.is_relative_to(groups_root):
            raise PortalError("工单附件路径无效。")
        return directory

    def create_group(
        self,
        prepared: dict,
        *,
        target_record_id: str,
        title: str,
        public_base_url: str,
        public_relay: bool = False,
    ) -> dict:
        target_record_id = str(target_record_id or "").strip()
        spec = prepared.get("polling_work_order_spec")
        if not prepared.get("polling_work_order_required") or not isinstance(spec, dict):
            return {}
        work_type = _work_order_type(spec.get("work_type"))
        with self._lock:
            existing = self.state_store.get_document(
                POLLING_WORK_ORDER_NAMESPACE, target_record_id
            )
            if isinstance(existing, dict):
                return self.group_with_links(existing, public_base_url)
            directory = self._group_directory(target_record_id)
            directory.mkdir(parents=True, exist_ok=True)
            attachments: list[dict] = []
            seen_hashes: set[str] = set()
            pending_root = (self.work_order_root / "pending").resolve()
            staging_directories: set[Path] = set()
            for item in spec.get("attachments") or []:
                digest = str(item.get("sha256") or "").strip()
                if digest in seen_hashes:
                    continue
                source = Path(str(item.get("staged_path") or "")).resolve()
                if not source.is_file() or not source.is_relative_to(pending_root):
                    raise PortalError(f"工单附件暂存文件不存在：{item.get('name') or '-'}")
                staging_directories.add(source.parent)
                target = (directory / f"{item.get('attachment_id')}_{_safe_file_name(item.get('name'))}").resolve()
                if not target.is_relative_to(directory):
                    raise PortalError("工单附件路径无效。")
                if not target.is_file() or _file_sha256(target) != digest:
                    shutil.copy2(source, target)
                attachments.append(
                    {
                        "attachment_id": str(item.get("attachment_id") or ""),
                        "name": str(item.get("name") or target.name),
                        "size": int(item.get("size") or target.stat().st_size),
                        "sha256": digest or _file_sha256(target),
                        "path": str(target),
                    }
                )
                seen_hashes.add(digest)
            flattened: list[dict] = []
            runs = list(spec.get("runs") or [])
            for run_index, run in enumerate(runs, start=1):
                run_label = _run_label(work_type, run)
                for step_index, template_step in enumerate(spec.get("steps") or [], start=1):
                    content = str(template_step.get("content") or "")
                    content = content.replace("{{from}}", str(run.get("from_unit") or ""))
                    content = content.replace("{{to}}", str(run.get("to_unit") or ""))
                    content = content.replace("{{other}}", str(run.get("other_unit") or ""))
                    flattened.append(
                        {
                            "step_key": f"{run_index}:{step_index}",
                            "sop_step_id": str(template_step.get("step_id") or ""),
                            "global_index": len(flattened),
                            "run_index": run_index,
                            "run_count": len(runs),
                            "run_label": run_label,
                            "step_index": step_index,
                            "step_count": len(spec.get("steps") or []),
                            "content": content,
                            "operator_required": bool(template_step.get("operator_required")),
                            "reviewer_required": bool(template_step.get("reviewer_required")),
                            "photo_required": self._step_photo_required(template_step),
                            "time_limit_seconds": int(
                                template_step.get("time_limit_seconds") or 0
                            ),
                            "activated_at_ts": 0.0,
                            "photos": [],
                            "operator_confirmation": {},
                            "reviewer_confirmation": {},
                        }
                    )
            operator_token = self.role_token(target_record_id, "operator")
            reviewer_token = self.role_token(target_record_id, "reviewer")
            now = self._now_text()
            group = {
                "group_id": target_record_id,
                "target_record_id": target_record_id,
                "work_type": work_type,
                "notice_type": {"maintenance": "维保通告", "adjust": "设备调整", "polling": "设备轮巡"}[work_type],
                "title": str(title or target_record_id),
                "sop_id": str(spec.get("sop_id") or ""),
                "sop_version": int(spec.get("sop_version") or 0),
                "sop_name": str(spec.get("sop_name") or ""),
                "scope": str(spec.get("scope") or "").strip().upper(),
                "runs": runs,
                "steps": flattened,
                "attachments": attachments,
                "operator": copy.deepcopy(spec.get("operator") or {}),
                "reviewer": copy.deepcopy(spec.get("reviewer") or {}),
                "initiator_open_id": str(spec.get("initiator_open_id") or ""),
                "initiator_name": str(spec.get("initiator_name") or ""),
                "token_hashes": {
                    "operator": hashlib.sha256(operator_token.encode("utf-8")).hexdigest(),
                    "reviewer": hashlib.sha256(reviewer_token.encode("utf-8")).hexdigest(),
                },
                "current_index": 0,
                "selected_run_index": 0,
                "selected_by_role": "",
                "version": 1,
                "state": "active",
                "uploaded_file_tokens": [],
                "last_error": "",
                "notifications": {},
                "created_at": now,
                "updated_at": now,
            }
            relay_mode = str(
                prepared.get("polling_work_order_mode") or ""
            ).strip()
            if relay_mode not in {"public_service", "public_relay", "local", "local_fallback"}:
                relay_mode = "public_relay" if public_relay else "local"
            group["relay"] = {
                "mode": relay_mode,
                "fallback_reason": str(
                    prepared.get("polling_work_order_fallback_reason") or ""
                ),
                "relay_checked_at": float(
                    prepared.get("polling_work_order_relay_checked_at") or 0
                ),
                "relay_url": str(
                    prepared.get("polling_work_order_relay_url") or ""
                ),
            }
            if public_relay:
                group["relay"]["registration_state"] = "registration_pending"
            self.state_store.put_document(
                POLLING_WORK_ORDER_NAMESPACE, target_record_id, group
            )
            for staging in staging_directories:
                if staging.parent == pending_root and staging.is_dir():
                    shutil.rmtree(staging)
        return self.group_with_links(group, public_base_url)

    def group_with_links(self, group: dict, public_base_url: str) -> dict:
        result = copy.deepcopy(group or {})
        relay = result.get("relay") if isinstance(result.get("relay"), dict) else {}
        relay.pop("management_token", None)
        relay.pop("link_credentials", None)
        result["relay"] = relay
        if str(relay.get("mode") or "") in {"public_service", "public_relay"}:
            if str(relay.get("registration_state") or "") == "registered":
                for key in ('operator_link','reviewer_link'):
                    result[key]=relay[key]=public_link_with_configured_port(relay.get(key),str(relay.get('relay_url') or ''))
            else:
                result["operator_link"] = ""
                result["reviewer_link"] = ""
            return result
        base = str(public_base_url or "").rstrip("/")
        if base:
            result["operator_link"] = (
                f"{base}/polling-work-order?token={quote(self.role_token(str(group.get('target_record_id') or ''), 'operator'), safe='')}"
            )
            result["reviewer_link"] = (
                f"{base}/polling-work-order?token={quote(self.role_token(str(group.get('target_record_id') or ''), 'reviewer'), safe='')}"
            )
        return result

    @staticmethod
    def _step_public(step: dict, current_index: int) -> dict:
        result = copy.deepcopy(step)
        result["photo_required"] = PollingWorkOrderService._step_photo_required(result)
        index = int(result.get("global_index") or 0)
        result["position"] = "current" if index == current_index else "previous" if index < current_index else "next"
        result["operator_confirmed"] = bool(result.pop("operator_confirmation", {}))
        result["reviewer_confirmed"] = bool(result.pop("reviewer_confirmation", {}))
        activated_at = float(result.get("activated_at_ts") or 0)
        time_limit = int(result.get("time_limit_seconds") or 0)
        timer_started = bool(activated_at > 0 or time_limit <= 0)
        available_at = activated_at + time_limit if timer_started else 0
        result["timer_started"] = timer_started
        result["confirm_available_at"] = available_at
        result["remaining_seconds"] = max(
            0,
            time_limit
            if not timer_started
            else int(math.ceil(available_at - time.time())),
        )
        result["photos"] = [
            {
                key: value
                for key, value in photo.items()
                if key not in {"path"}
            }
            for photo in result.get("photos") or []
            if isinstance(photo, dict)
        ]
        return result

    @staticmethod
    def _step_done(step: dict) -> bool:
        return (
            not step.get("operator_required")
            or bool(step.get("operator_confirmation"))
        ) and (
            not step.get("reviewer_required")
            or bool(step.get("reviewer_confirmation"))
        )

    @staticmethod
    def _step_photo_required(step: dict) -> bool:
        return True if "photo_required" not in step else _flag(step.get("photo_required"))

    def _step_has_local_photo(self, group: dict, step: dict) -> bool:
        photo_root = (
            self._group_directory(str(group.get("target_record_id") or ""))
            / "photos"
        ).resolve()
        return any(
            path.is_file() and path.is_relative_to(photo_root)
            for photo in step.get("photos") or []
            if isinstance(photo, dict)
            for path in [Path(str(photo.get("path") or "")).resolve()]
        )

    @staticmethod
    def _workbook_output_name(group: dict) -> str:
        work_type = _stored_work_order_type(group)
        scope = str(group.get("scope") or "").strip().upper()
        if scope == "110":
            building = "110站"
        elif scope in POLLING_SOP_SCOPES:
            building = f"{scope}楼"
        else:
            match = re.search(r"(110站|[ABCDEH]楼)", str(group.get("title") or ""))
            building = match.group(1) if match else "园区"
        confirmed_times = [
            str((step.get(key) or {}).get("confirmed_at") or "").strip()
            for step in group.get("steps") or []
            if isinstance(step, dict)
            for key in ("operator_confirmation", "reviewer_confirmation")
            if str((step.get(key) or {}).get("confirmed_at") or "").strip()
        ]
        date_text = (max(confirmed_times) if confirmed_times else str(group.get("updated_at") or ""))[:10]
        try:
            completed_date = dt.date.fromisoformat(date_text)
        except ValueError:
            completed_date = dt.date.today()
        prefix = (
            f"{building}{completed_date.year}年{completed_date.month}月"
            f"{completed_date.day}日-"
        )
        direction = ""
        if work_type == "polling":
            runs = [item for item in group.get("runs") or [] if isinstance(item, dict)]
            if len(runs) == 1:
                direction = f"{runs[0].get('from_unit') or ''}轮巡至{runs[0].get('to_unit') or ''}"
            else:
                direction = (
                    "".join(str(item.get("from_unit") or "") for item in runs)
                    + "轮巡至"
                    + "".join(str(item.get("to_unit") or "") for item in runs)
                )
        tail = (
            (f"-{direction}" if direction else "")
            + "-操作人-"
            f"{_safe_file_name((group.get('operator') or {}).get('name') or '未填写')[:32]}-"
            f"审核人-{_safe_file_name((group.get('reviewer') or {}).get('name') or '未填写')[:32]}-"
            "操作记录"
        )
        sop_name = _safe_file_name(
            group.get("sop_name")
            or ({"maintenance": "维保SOP", "adjust": "设备调整SOP"}.get(work_type, "轮巡SOP"))
        )
        sop_name = sop_name[: max(1, 155 - len(prefix) - len(tail))]
        return f"{prefix}{sop_name}{tail}.xlsx"

    @classmethod
    def _selected_run_index(cls, group: dict, steps: list[dict]) -> int:
        selected = int(group.get("selected_run_index") or 0)
        if selected > 0 or "selected_run_index" in group:
            return selected
        current_index = int(group.get("current_index") or 0)
        if str(group.get("state") or "") != "active" or not 0 <= current_index < len(steps):
            return 0
        step = steps[current_index]
        if (
            float(step.get("activated_at_ts") or 0) > 0
            or bool(step.get("operator_confirmation"))
            or bool(step.get("reviewer_confirmation"))
        ):
            return int(step.get("run_index") or 0)
        return 0

    def session(self, token: str) -> dict:
        group, role = self._resolve_token(token)
        steps = list(group.get("steps") or [])
        current_index = int(group.get("current_index") or 0)
        selected_run_index = self._selected_run_index(group, steps)
        run_indexes = [
            index
            for index, step in enumerate(steps)
            if int(step.get("run_index") or 0) == selected_run_index
        ]
        incomplete_indexes = [index for index in run_indexes if not self._step_done(steps[index])]
        if selected_run_index and not incomplete_indexes:
            selected_run_index = 0
        current = (
            current_index
            if current_index in incomplete_indexes
            else incomplete_indexes[0]
            if incomplete_indexes
            else 0
        )
        position = run_indexes.index(current) if current in run_indexes else 0
        visible_indexes = run_indexes[max(0, position - 1) : position + 2] if selected_run_index else []
        public_steps = [
            self._step_public(
                steps[index], int(group.get("current_index") or 0)
            )
            for index in visible_indexes
        ]
        for step in public_steps:
            for photo in step.get("photos") or []:
                photo["preview_url"] = (
                    "/api/polling-work-orders/photos/"
                    f"{quote(str(photo.get('photo_id') or ''), safe='')}"
                    f"?token={quote(str(token or ''), safe='')}"
                )
        work_orders = []
        for run_index, run in enumerate(group.get("runs") or [], start=1):
            run_steps = [
                step
                for step in steps
                if int(step.get("run_index") or 0) == run_index
            ]
            completed_steps = sum(1 for step in run_steps if self._step_done(step))
            state = (
                "completed"
                if run_steps and completed_steps == len(run_steps)
                else "active"
                if run_index == selected_run_index
                else "available"
                if str(group.get("state") or "") == "active" and not selected_run_index
                else "locked"
            )
            work_orders.append(
                {
                    "run_index": run_index,
                    "from_unit": str(run.get("from_unit") or ""),
                    "to_unit": str(run.get("to_unit") or ""),
                    "label": _run_label(_stored_work_order_type(group), run),
                    "step_count": len(run_steps),
                    "completed_steps": completed_steps,
                    "state": state,
                    "selectable": state in {"active", "available"},
                }
            )
        selected_steps = [
            step
            for step in steps
            if int(step.get("run_index") or 0) == selected_run_index
        ]
        selection_owner = str(group.get("selected_by_role") or "")
        can_release_selection = bool(
            selected_run_index
            and (not selection_owner or selection_owner == role)
            and not any(
                step.get("operator_confirmation")
                or step.get("reviewer_confirmation")
                or step.get("photos")
                for step in selected_steps
            )
        )
        can_rollback_previous = bool(
            role == "reviewer"
            and str(group.get("state") or "") == "active"
            and selected_run_index
            and current_index > 0
            and current_index < len(steps)
            and int(steps[current_index - 1].get("run_index") or 0)
            == selected_run_index
            and self._step_done(steps[current_index - 1])
        )
        return {
            "group_id": str(group.get("group_id") or ""),
            "title": str(group.get("title") or ""),
            "work_type": _stored_work_order_type(group),
            "sop_name": str(group.get("sop_name") or ""),
            "role": role,
            "role_label": "操作人" if role == "operator" else "现场审核人",
            "assigned_person": copy.deepcopy(group.get(role) or {}),
            "state": str(group.get("state") or ""),
            "version": int(group.get("version") or 0),
            "current_index": current_index,
            "total_steps": len(steps),
            "current_run_index": selected_run_index,
            "can_release_selection": can_release_selection,
            "can_rollback_previous": can_rollback_previous,
            "work_orders": work_orders,
            "steps": public_steps,
            "last_error": str(group.get("last_error") or ""),
        }

    def activate(
        self,
        token: str,
        *,
        run_index: int,
        expected_version: int,
    ) -> dict:
        with self._lock:
            group, _role = self._resolve_token(token)
            if str(group.get("state") or "") != "active":
                raise PortalConflictError("当前工单已不能开始倒计时。")
            steps = list(group.get("steps") or [])
            selected_run_index = self._selected_run_index(group, steps)
            if selected_run_index and selected_run_index != int(run_index or 0):
                raise PortalConflictError("另一角色已选择其他工单，请进入已选择的工单。")
            run_indexes = [
                index
                for index, item in enumerate(steps)
                if int(item.get("run_index") or 0) == int(run_index or 0)
                and not self._step_done(item)
            ]
            if not run_indexes:
                raise PortalConflictError("所选工单已完成。")
            current_index = int(group.get("current_index") or 0)
            if current_index not in run_indexes:
                current_index = run_indexes[0]
            step = steps[current_index]
            if float(step.get("activated_at_ts") or 0) > 0:
                return self.session(token)
            if int(expected_version or 0) != int(group.get("version") or 0):
                raise PortalConflictError("工单状态已更新，请刷新后重试。")
            step["activated_at_ts"] = time.time()
            selected_by_role = (
                str(group.get("selected_by_role") or "")
                if selected_run_index
                else _role
            )
            group.update(
                {
                    "steps": steps,
                    "current_index": current_index,
                    "selected_run_index": int(run_index or 0),
                    "selected_by_role": selected_by_role,
                    "version": int(group.get("version") or 0) + 1,
                    "updated_at": self._now_text(),
                    "last_error": "",
                }
            )
            self.state_store.put_document(
                POLLING_WORK_ORDER_NAMESPACE,
                str(group.get("target_record_id") or ""),
                group,
            )
        return self.session(token)

    def release_selection(
        self,
        token: str,
        *,
        run_index: int,
        expected_version: int,
    ) -> dict:
        with self._lock:
            group, role = self._resolve_token(token)
            if str(group.get("state") or "") != "active":
                raise PortalConflictError("当前工单已不能退出选择。")
            steps = list(group.get("steps") or [])
            selected_run_index = self._selected_run_index(group, steps)
            if not selected_run_index:
                return self.session(token)
            if selected_run_index != int(run_index or 0):
                raise PortalConflictError("当前选择已变化，请刷新后重试。")
            selection_owner = str(group.get("selected_by_role") or "")
            if selection_owner and selection_owner != role:
                raise PortalConflictError("当前工单由另一角色选择，不能代为退出。")
            if int(expected_version or 0) != int(group.get("version") or 0):
                raise PortalConflictError("工单状态已更新，请刷新后重试。")
            selected_steps = [
                step
                for step in steps
                if int(step.get("run_index") or 0) == selected_run_index
            ]
            if any(
                step.get("operator_confirmation")
                or step.get("reviewer_confirmation")
                or step.get("photos")
                for step in selected_steps
            ):
                raise PortalConflictError("当前工单已有操作记录，不能退出后改选其他工单。")
            for step in selected_steps:
                step["activated_at_ts"] = 0.0
            remaining_indexes = [
                index
                for index, step in enumerate(steps)
                if not self._step_done(step)
            ]
            group.update(
                {
                    "steps": steps,
                    "current_index": remaining_indexes[0] if remaining_indexes else len(steps),
                    "selected_run_index": 0,
                    "selected_by_role": "",
                    "version": int(group.get("version") or 0) + 1,
                    "updated_at": self._now_text(),
                    "last_error": "",
                }
            )
            self.state_store.put_document(
                POLLING_WORK_ORDER_NAMESPACE,
                str(group.get("target_record_id") or ""),
                group,
            )
        return self.session(token)

    def confirm(
        self,
        token: str,
        *,
        step_key: str,
        expected_version: int,
        actual_open_id: str = "",
        actual_name: str = "",
    ) -> dict:
        with self._lock:
            group, role = self._resolve_token(token)
            if int(expected_version or 0) != int(group.get("version") or 0):
                steps = list(group.get("steps") or [])
                prior = next((item for item in steps if str(item.get("step_key") or "") == str(step_key or "")), None)
                if prior and prior.get(f"{role}_confirmation"):
                    return self.session(token)
                raise PortalConflictError("工单状态已更新，请刷新后重试。")
            steps = list(group.get("steps") or [])
            current_index = int(group.get("current_index") or 0)
            selected_run_index = self._selected_run_index(group, steps)
            if not selected_run_index:
                raise PortalConflictError("请先从工单总览选择要执行的工单。")
            if current_index >= len(steps):
                return self.session(token)
            step = steps[current_index]
            if int(step.get("run_index") or 0) != selected_run_index:
                raise PortalConflictError("当前工单选择已变化，请刷新后重试。")
            if str(step.get("step_key") or "") != str(step_key or ""):
                raise PortalConflictError("只能确认当前步骤。")
            time_limit = int(step.get("time_limit_seconds") or 0)
            activated_at = float(step.get("activated_at_ts") or 0)
            if time_limit > 0 and activated_at <= 0:
                raise PortalConflictError("请先进入当前工单并启动步骤倒计时。")
            available_at = activated_at + time_limit
            if time.time() < available_at:
                raise PortalConflictError(
                    f"当前步骤还需等待 {max(1, int(math.ceil(available_at - time.time())))} 秒。"
                )
            if not bool(step.get(f"{role}_required")):
                raise PortalConflictError("当前步骤不需要该角色确认。")
            if role == "reviewer" and step.get("operator_required") and not step.get("operator_confirmation"):
                raise PortalConflictError("请先等待操作人确认。")
            if self._step_photo_required(step) and not self._step_has_local_photo(group, step):
                raise PortalConflictError("请先拍摄并上传当前步骤照片。")
            if not step.get(f"{role}_confirmation"):
                step[f"{role}_confirmation"] = {
                    "assigned_record_id": str((group.get(role) or {}).get("record_id") or ""),
                    "assigned_name": str((group.get(role) or {}).get("name") or ""),
                    "actual_open_id": str(actual_open_id or ""),
                    "actual_name": str(actual_name or ""),
                    "confirmed_at": self._now_text(),
                }
            required_done = (
                (not step.get("operator_required") or bool(step.get("operator_confirmation")))
                and (not step.get("reviewer_required") or bool(step.get("reviewer_confirmation")))
            )
            if required_done:
                next_indexes = [
                    index
                    for index, item in enumerate(steps)
                    if int(item.get("run_index") or 0) == selected_run_index
                    and not self._step_done(item)
                ]
                remaining_indexes: list[int] = []
                if next_indexes:
                    group["current_index"] = next_indexes[0]
                    steps[next_indexes[0]]["activated_at_ts"] = time.time()
                else:
                    group["selected_run_index"] = 0
                    group["selected_by_role"] = ""
                    remaining_indexes = [
                        index
                        for index, item in enumerate(steps)
                        if not self._step_done(item)
                    ]
                    group["current_index"] = remaining_indexes[0] if remaining_indexes else len(steps)
                if not next_indexes and not remaining_indexes:
                    group["state"] = "upload_pending"
            group.update(
                {
                    "steps": steps,
                    "version": int(group.get("version") or 0) + 1,
                    "updated_at": self._now_text(),
                    "last_error": "",
                }
            )
            self.state_store.put_document(
                POLLING_WORK_ORDER_NAMESPACE,
                str(group.get("target_record_id") or ""),
                group,
            )
        return self.session(token)

    def rollback_previous(
        self,
        token: str,
        *,
        step_key: str,
        expected_version: int,
        actual_open_id: str = "",
        actual_name: str = "",
    ) -> dict:
        with self._lock:
            group, role = self._resolve_token(token)
            if role != "reviewer":
                raise PortalConflictError("只有现场审核人可以回退上一步。")
            if str(group.get("state") or "") != "active":
                raise PortalConflictError("当前工单已不能回退步骤。")
            if int(expected_version or 0) != int(group.get("version") or 0):
                raise PortalConflictError("工单状态已更新，请刷新后重试。")
            steps = list(group.get("steps") or [])
            current_index = int(group.get("current_index") or 0)
            selected_run_index = self._selected_run_index(group, steps)
            if not selected_run_index or not 0 < current_index < len(steps):
                raise PortalConflictError("当前没有可回退的上一步。")
            current = steps[current_index]
            previous = steps[current_index - 1]
            if str(current.get("step_key") or "") != str(step_key or ""):
                raise PortalConflictError("当前步骤已变化，请刷新后重试。")
            if int(previous.get("run_index") or 0) != selected_run_index:
                raise PortalConflictError("当前工单没有可回退的上一步。")
            if not self._step_done(previous):
                raise PortalConflictError("上一步尚未完成，无需回退。")
            photo_root = (
                self._group_directory(str(group.get("target_record_id") or ""))
                / "photos"
            ).resolve()
            discarded_paths: list[Path] = []
            for index in (current_index - 1, current_index):
                item = steps[index]
                discarded_paths.extend(
                    Path(str(photo.get("path") or "")).resolve()
                    for photo in item.get("photos") or []
                    if isinstance(photo, dict)
                )
                item["photos"] = []
                item["operator_confirmation"] = {}
                item["reviewer_confirmation"] = {}
                item["activated_at_ts"] = time.time() if index == current_index - 1 else 0.0
            group.update(
                {
                    "steps": steps,
                    "current_index": current_index - 1,
                    "version": int(group.get("version") or 0) + 1,
                    "updated_at": self._now_text(),
                    "last_error": "",
                    "last_rollback": {
                        "from_step_key": str(current.get("step_key") or ""),
                        "to_step_key": str(previous.get("step_key") or ""),
                        "actual_open_id": str(actual_open_id or ""),
                        "actual_name": str(actual_name or ""),
                        "rolled_back_at": self._now_text(),
                    },
                }
            )
            self.state_store.put_document(
                POLLING_WORK_ORDER_NAMESPACE,
                str(group.get("target_record_id") or ""),
                group,
            )
            for path in discarded_paths:
                if path.is_file() and path.is_relative_to(photo_root):
                    try:
                        path.unlink()
                    except OSError:
                        pass
        return self.session(token)

    def add_step_photo(
        self,
        token: str,
        *,
        step_key: str,
        expected_version: int,
        file_name: str,
        mime_type: str,
        content: bytes,
    ) -> dict:
        content = bytes(content or b"")
        if not content:
            raise PortalError("操作照片内容为空。")
        if len(content) > POLLING_STEP_PHOTO_MAX_BYTES:
            raise PortalError("单张操作照片不能超过 8MB。")
        if not str(mime_type or "").startswith("image/"):
            raise PortalError("只能上传图片作为操作照片。")
        try:
            from PIL import Image

            with Image.open(io.BytesIO(content)) as image:
                detected_mime_type = str(Image.MIME.get(image.format) or "")
                width, height = image.size
                if (
                    width <= 0
                    or height <= 0
                    or width > POLLING_STEP_PHOTO_MAX_DIMENSION
                    or height > POLLING_STEP_PHOTO_MAX_DIMENSION
                    or width * height > POLLING_STEP_PHOTO_MAX_PIXELS
                ):
                    raise PortalError("操作照片像素或尺寸超过限制。")
                image.verify()
            if detected_mime_type not in POLLING_STEP_PHOTO_MIME_TYPES:
                raise ValueError("unsupported image format")
        except PortalError:
            raise
        except Exception as exc:
            raise PortalError("操作照片内容损坏，无法保存。") from exc
        mime_type = detected_mime_type
        with self._lock:
            group, role = self._resolve_token(token)
            if str(group.get("state") or "") != "active":
                raise PortalConflictError("当前工单已不能上传步骤照片。")
            if int(expected_version or 0) != int(group.get("version") or 0):
                raise PortalConflictError("工单状态已更新，请刷新后重试。")
            steps = list(group.get("steps") or [])
            selected_run_index = self._selected_run_index(group, steps)
            if not selected_run_index:
                raise PortalConflictError("请先从工单总览选择要执行的工单。")
            current_index = int(group.get("current_index") or 0)
            if current_index >= len(steps):
                raise PortalConflictError("当前工单步骤已完成。")
            step = steps[current_index]
            if int(step.get("run_index") or 0) != selected_run_index:
                raise PortalConflictError("当前工单选择已变化，请刷新后重试。")
            if str(step.get("step_key") or "") != str(step_key or ""):
                raise PortalConflictError("只能给当前步骤拍照。")
            if step.get("reviewer_confirmation") or (
                role == "operator" and step.get("operator_confirmation")
            ):
                raise PortalConflictError("当前步骤已有确认，不能再添加操作照片。")
            digest = hashlib.sha256(content).hexdigest()
            existing_photos = list(step.get("photos") or [])
            if any(
                str(photo.get("sha256") or "") == digest
                for photo in existing_photos
                if isinstance(photo, dict)
            ):
                return self.session(token)
            if len(existing_photos) >= POLLING_STEP_MAX_PHOTOS:
                raise PortalError(
                    f"每个步骤最多上传 {POLLING_STEP_MAX_PHOTOS} 张操作照片。"
                )
            total_photos = sum(
                len(item.get("photos") or [])
                for item in steps
                if isinstance(item, dict)
            )
            if total_photos >= POLLING_WORK_ORDER_MAX_PHOTOS:
                raise PortalError(
                    f"整个工单组最多上传 {POLLING_WORK_ORDER_MAX_PHOTOS} 张操作照片。"
                )
            total_photo_bytes = sum(
                int(photo.get("size") or 0)
                for item in steps
                if isinstance(item, dict)
                for photo in item.get("photos") or []
                if isinstance(photo, dict)
            )
            if total_photo_bytes + len(content) > POLLING_WORK_ORDER_MAX_PHOTO_BYTES:
                raise PortalError("整个工单组操作照片总大小不能超过 200MB。")
            directory = (self._group_directory(str(group.get("target_record_id") or "")) / "photos").resolve()
            group_directory = self._group_directory(
                str(group.get("target_record_id") or "")
            )
            if not directory.is_relative_to(group_directory):
                raise PortalError("操作照片路径无效。")
            directory.mkdir(parents=True, exist_ok=True)
            photo_id = uuid.uuid4().hex
            safe_name = _safe_file_name(file_name or "step_photo.png")
            path = (directory / f"{photo_id}_{safe_name}").resolve()
            if not path.is_relative_to(directory):
                raise PortalError("操作照片路径无效。")
            temporary = path.with_suffix(path.suffix + ".tmp")
            temporary.write_bytes(content)
            os.replace(temporary, path)
            step["photos"] = existing_photos + [
                {
                    "photo_id": photo_id,
                    "name": safe_name,
                    "mime_type": str(mime_type or "image/png"),
                    "size": len(content),
                    "sha256": digest,
                    "path": str(path),
                    "uploaded_role": role,
                    "uploaded_at": self._now_text(),
                }
            ]
            group.update(
                {
                    "steps": steps,
                    "version": int(group.get("version") or 0) + 1,
                    "updated_at": self._now_text(),
                    "last_error": "",
                }
            )
            try:
                self.state_store.put_document(
                    POLLING_WORK_ORDER_NAMESPACE,
                    str(group.get("target_record_id") or ""),
                    group,
                )
            except Exception:
                try:
                    path.unlink()
                except OSError:
                    pass
                raise
        return self.session(token)

    def step_photo_content(
        self, token: str, *, photo_id: str
    ) -> tuple[bytes, str, str]:
        group, _role = self._resolve_token(token)
        photo = next(
            (
                item
                for step in group.get("steps") or []
                for item in step.get("photos") or []
                if str(item.get("photo_id") or "") == str(photo_id or "")
            ),
            None,
        )
        if not photo:
            raise PortalNotFoundError("操作照片不存在。")
        directory = (self._group_directory(str(group.get("target_record_id") or "")) / "photos").resolve()
        path = Path(str(photo.get("path") or "")).resolve()
        if not path.is_file() or not path.is_relative_to(directory):
            raise PortalNotFoundError("操作照片文件不存在。")
        return (
            path.read_bytes(),
            str(photo.get("mime_type") or "image/png"),
            str(photo.get("name") or path.name),
        )

    def build_execution_workbook(self, target_record_id: str) -> dict:
        group = self.get_group(target_record_id)
        work_type = _stored_work_order_type(group)
        if str(group.get("state") or "") not in {"upload_pending", "completed"}:
            raise PortalConflictError("工单步骤尚未全部完成。")
        try:
            from openpyxl import load_workbook
            from openpyxl.drawing.image import Image as ExcelImage
            from openpyxl.utils import get_column_letter
            from PIL import Image as PillowImage
            from PIL import ImageOps
        except Exception as exc:
            raise PortalError("缺少 openpyxl/Pillow，无法生成工单表格。") from exc
        template_path = Path(self.work_order_template_path).resolve()
        if not template_path.is_file():
            raise PortalError("工单模板不存在。")
        try:
            with zipfile.ZipFile(template_path) as archive:
                logo_bytes = archive.read("xl/media/image1.png")
        except Exception as exc:
            raise PortalError("工单模板中的Logo无法读取。") from exc
        directory = self._group_directory(target_record_id)
        directory.mkdir(parents=True, exist_ok=True)
        path = (directory / POLLING_WORK_ORDER_CACHE_NAME).resolve()
        if path.is_file():
            size = path.stat().st_size
            if size > POLLING_WORK_ORDER_MAX_BYTES:
                raise PortalError("工单Excel超过允许大小，已停止上传。")
            return {
                "name": self._workbook_output_name(group),
                "path": str(path),
                "size": size,
                "sha256": _file_sha256(path),
            }
        runs = [item for item in group.get("runs") or [] if isinstance(item, dict)]
        if not runs:
            raise PortalError("工单缺少执行项。")
        try:
            workbook = load_workbook(
                template_path,
                data_only=False,
                read_only=False,
                keep_links=False,
            )
        except Exception as exc:
            raise PortalError("工单模板无法打开。") from exc
        template_sheet = workbook.active
        sheets = [template_sheet]
        for _ in runs[1:]:
            sheets.append(workbook.copy_worksheet(template_sheet))
        steps = [item for item in group.get("steps") or [] if isinstance(item, dict)]
        photo_root = (directory / "photos").resolve()
        image_handles: list[Any] = []
        try:
            for run_index, (run, sheet) in enumerate(zip(runs, sheets), start=1):
                run_steps = [
                    step
                    for step in steps
                    if int(step.get("run_index") or 0) == run_index
                ]
                if not run_steps:
                    raise PortalError(f"工单{run_index}缺少操作步骤。")
                run_label = _run_label(work_type, run)
                sheet.title = f"工单{run_index} {run_label}"[:31]
                sheet["A1"] = None
                sheet["C3"] = _excel_text((group.get("operator") or {}).get("name"))
                sheet["D3"] = _excel_text((group.get("reviewer") or {}).get("name"))
                confirmed_times = [
                    str((step.get(key) or {}).get("confirmed_at") or "").strip()
                    for step in run_steps
                    for key in ("operator_confirmation", "reviewer_confirmation")
                    if str((step.get(key) or {}).get("confirmed_at") or "").strip()
                ]
                completion_text = max(confirmed_times) if confirmed_times else str(
                    group.get("updated_at") or self._now_text()
                )
                try:
                    sheet["E3"] = dt.datetime.strptime(
                        completion_text, "%Y-%m-%d %H:%M:%S"
                    )
                    sheet["E3"].number_format = "yyyy-mm-dd hh:mm:ss"
                except ValueError:
                    sheet["E3"] = _excel_text(completion_text)
                sheet["B7"] = _excel_text(
                    f"{group.get('sop_name') or (({'maintenance': '维保', 'adjust': '设备调整'}.get(work_type, '轮巡')) + '操作流程')} · {run_label}"
                )

                source_styles = [
                    copy.copy(sheet.cell(row=10, column=column)._style)
                    for column in range(1, 6)
                ]
                source_height = float(sheet.row_dimensions[10].height or 156.5)
                for merged in list(sheet.merged_cells.ranges):
                    if int(merged.min_row) >= 9:
                        sheet.unmerge_cells(str(merged))
                if len(run_steps) > 2:
                    sheet.insert_rows(11, amount=len(run_steps) - 2)
                elif len(run_steps) == 1:
                    sheet.delete_rows(10, amount=1)

                logo_buffer = io.BytesIO(logo_bytes)
                logo = ExcelImage(logo_buffer)
                logo.width = 163
                logo.height = 59
                sheet.add_image(logo, "A1")
                image_handles.extend((logo_buffer, logo))
                max_photo_count = max(
                    len(step.get("photos") or []) for step in run_steps
                )
                last_column = 5 + max(0, max_photo_count - 1)
                for column in range(6, last_column + 1):
                    sheet.column_dimensions[get_column_letter(column)].width = 34

                for step_offset, step in enumerate(run_steps):
                    row_number = 9 + step_offset
                    sheet.row_dimensions[row_number].height = source_height
                    for column, style in enumerate(source_styles, start=1):
                        sheet.cell(row=row_number, column=column)._style = copy.copy(
                            style
                        )
                    for column in range(6, last_column + 1):
                        sheet.cell(row=row_number, column=column)._style = copy.copy(
                            source_styles[2]
                        )
                    sheet.merge_cells(
                        start_row=row_number,
                        start_column=1,
                        end_row=row_number,
                        end_column=2,
                    )
                    sheet.merge_cells(
                        start_row=row_number,
                        start_column=4,
                        end_row=row_number,
                        end_column=5,
                    )
                    sheet.cell(row=row_number, column=1).value = None
                    sheet.cell(row=row_number, column=3).value = int(
                        step.get("step_index") or step_offset + 1
                    )
                    sheet.cell(row=row_number, column=4).value = _excel_text(
                        step.get("content")
                    )
                    photos = [
                        item
                        for item in step.get("photos") or []
                        if isinstance(item, dict)
                    ]
                    if not photos:
                        raise PortalError(
                            f"工单{run_index}第{step_offset + 1}步缺少操作照片。"
                        )
                    for photo_index, photo in enumerate(photos):
                        photo_path = Path(str(photo.get("path") or "")).resolve()
                        if (
                            not photo_path.is_file()
                            or not photo_path.is_relative_to(photo_root)
                        ):
                            raise PortalError(
                                f"工单{run_index}第{step_offset + 1}步的本地照片不存在。"
                            )
                        try:
                            with PillowImage.open(photo_path) as source:
                                normalized = ImageOps.exif_transpose(source)
                                if "A" in normalized.getbands():
                                    prepared = PillowImage.new(
                                        "RGB", normalized.size, "white"
                                    )
                                    prepared.paste(
                                        normalized,
                                        mask=normalized.getchannel("A"),
                                    )
                                else:
                                    prepared = normalized.convert("RGB")
                                width, height = prepared.size
                                scale = min(
                                    1.0,
                                    230 / max(1, width),
                                    190 / max(1, height),
                                )
                                target_width = max(1, int(round(width * scale)))
                                target_height = max(1, int(round(height * scale)))
                                if (target_width, target_height) != prepared.size:
                                    prepared = prepared.resize(
                                        (target_width, target_height),
                                        getattr(
                                            PillowImage, "Resampling", PillowImage
                                        ).LANCZOS,
                                    )
                                photo_buffer = io.BytesIO()
                                prepared.save(photo_buffer, format="PNG", optimize=True)
                        except Exception as exc:
                            raise PortalError(
                                f"工单{run_index}第{step_offset + 1}步的照片无法写入Excel。"
                            ) from exc
                        photo_buffer.seek(0)
                        excel_photo = ExcelImage(photo_buffer)
                        excel_photo.width = target_width
                        excel_photo.height = target_height
                        anchor_column = "A" if photo_index == 0 else get_column_letter(5 + photo_index)
                        sheet.add_image(excel_photo, f"{anchor_column}{row_number}")
                        image_handles.extend((photo_buffer, excel_photo))

                last_row = 8 + len(run_steps)
                sheet.print_area = f"A1:{get_column_letter(last_column)}{last_row}"
                sheet.print_title_rows = "8:8"
                sheet.page_setup.orientation = "portrait"
                sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
                sheet.page_setup.scale = None
                sheet.page_setup.fitToWidth = 1
                sheet.page_setup.fitToHeight = 0
                sheet.sheet_properties.pageSetUpPr.fitToPage = True
                sheet.print_options.horizontalCentered = True
            temporary = path.with_suffix(path.suffix + ".tmp")
            if temporary.is_file():
                temporary.unlink()
            workbook.save(temporary)
            os.replace(temporary, path)
            if path.stat().st_size > POLLING_WORK_ORDER_MAX_BYTES:
                raise PortalError("工单Excel超过允许大小，已停止上传。")
        except Exception:
            temporary = path.with_suffix(path.suffix + ".tmp")
            if temporary.is_file():
                temporary.unlink()
            raise
        finally:
            workbook.close()
        return {
            "name": self._workbook_output_name(group),
            "path": str(path),
            "size": path.stat().st_size,
            "sha256": _file_sha256(path),
        }

    def pending_upload_groups(self) -> list[dict]:
        return [
            copy.deepcopy(document.get("payload") or {})
            for document in self.state_store.list_documents(POLLING_WORK_ORDER_NAMESPACE)
            if isinstance(document.get("payload"), dict)
            and str((document.get("payload") or {}).get("state") or "") == "upload_pending"
        ]

    def open_groups(self) -> list[dict]:
        return [
            copy.deepcopy(document.get("payload") or {})
            for document in self.state_store.list_documents(POLLING_WORK_ORDER_NAMESPACE)
            if isinstance(document.get("payload"), dict)
            and str((document.get("payload") or {}).get("state") or "")
            in {"active", "upload_pending", "completed"}
        ]

    def update_relay(self, target_record_id: str, changes: dict) -> dict:
        with self._lock:
            group = self.get_group(target_record_id)
            group["relay"] = {**dict(group.get("relay") or {}), **copy.deepcopy(changes or {})}
            group["updated_at"] = self._now_text()
            self.state_store.put_document(POLLING_WORK_ORDER_NAMESPACE, target_record_id, group)
        return group

    def mark_public_artifact_ready(self, target_record_id: str) -> dict:
        with self._lock:
            group = self.get_group(target_record_id)
            if str(group.get("state") or "") not in {"completed", "cancelled", "stopped"}:
                group["state"] = "upload_pending"
                group["last_error"] = ""
                group["updated_at"] = self._now_text()
                self.state_store.put_document(POLLING_WORK_ORDER_NAMESPACE, target_record_id, group)
        return group

    def save_public_artifact(
        self,
        target_record_id: str,
        *,
        content: bytes,
        name: str,
        expected_sha256: str = "",
    ) -> dict:
        if not content or len(content) > POLLING_WORK_ORDER_MAX_BYTES:
            raise PortalError("公网工单Excel为空或超过允许大小。")
        digest = hashlib.sha256(content).hexdigest()
        expected = str(expected_sha256 or "").strip().lower()
        if expected and not hmac.compare_digest(digest, expected):
            raise PortalError("公网工单Excel校验失败。")
        safe_name = _safe_file_name(name)
        if not safe_name.lower().endswith(".xlsx"):
            safe_name += ".xlsx"
        directory = self._group_directory(target_record_id)
        directory.mkdir(parents=True, exist_ok=True)
        path = (directory / safe_name).resolve()
        if not path.is_relative_to(directory):
            raise PortalError("公网工单Excel路径无效。")
        temporary = path.with_suffix(path.suffix + ".tmp")
        with temporary.open("wb") as stream:
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
        return {"name": safe_name, "path": str(path), "size": len(content), "sha256": digest}

    def cancel_and_get(self, target_record_id: str, *, reason: str) -> dict:
        self.cancel_group(target_record_id, reason=reason)
        return self.get_group(target_record_id)

    def mark_upload_result(
        self,
        target_record_id: str,
        *,
        success: bool,
        file_tokens: list[str] | None = None,
        photo_file_tokens: list[str] | None = None,
        error: str = "",
    ) -> dict:
        with self._lock:
            group = self.get_group(target_record_id)
            group.update(
                {
                    "state": "completed" if success else "upload_pending",
                    "uploaded_file_tokens": list(dict.fromkeys(file_tokens or group.get("uploaded_file_tokens") or [])),
                    "uploaded_photo_file_tokens": list(
                        dict.fromkeys(
                            photo_file_tokens
                            or group.get("uploaded_photo_file_tokens")
                            or []
                        )
                    ),
                    "last_error": "" if success else str(error or "工单附件上传失败。"),
                    "upload_completed_at": self._now_text() if success else "",
                    "updated_at": self._now_text(),
                    "version": int(group.get("version") or 0) + 1,
                }
            )
            self.state_store.put_document(POLLING_WORK_ORDER_NAMESPACE, target_record_id, group)
        return group

    def mark_upload_progress(
        self,
        target_record_id: str,
        *,
        token_by_sha256: dict[str, str],
        workbook_name: str = "",
        photo_token_by_sha256: dict[str, str] | None = None,
        error: str = "",
    ) -> dict:
        with self._lock:
            group = self.get_group(target_record_id)
            group["uploaded_by_sha256"] = {
                str(key): str(value)
                for key, value in (token_by_sha256 or {}).items()
                if str(key).strip() and str(value).strip()
            }
            if workbook_name:
                group["uploaded_workbook_name"] = str(workbook_name)
            if photo_token_by_sha256 is not None:
                group["uploaded_photo_by_sha256"] = {
                    str(key): str(value)
                    for key, value in photo_token_by_sha256.items()
                    if str(key).strip() and str(value).strip()
                }
            group["last_error"] = str(error or "")
            group["updated_at"] = self._now_text()
            self.state_store.put_document(
                POLLING_WORK_ORDER_NAMESPACE, target_record_id, group
            )
        return group

    def update_notifications(self, target_record_id: str, notifications: dict) -> dict:
        with self._lock:
            group = self.get_group(target_record_id)
            group["notifications"] = copy.deepcopy(notifications or {})
            group["updated_at"] = self._now_text()
            self.state_store.put_document(POLLING_WORK_ORDER_NAMESPACE, target_record_id, group)
        return group

    def cancel_group(self, target_record_id: str, *, reason: str) -> None:
        with self._lock:
            try:
                group = self.get_group(target_record_id)
            except PortalNotFoundError:
                return
            if str(group.get("state") or "") in {"cancelled", "stopped"}:
                return
            group.update(
                {
                    "state": "cancelled",
                    "cancel_reason": str(reason or "target_terminal"),
                    "updated_at": self._now_text(),
                    "version": int(group.get("version") or 0) + 1,
                }
            )
            self.state_store.put_document(POLLING_WORK_ORDER_NAMESPACE, target_record_id, group)
