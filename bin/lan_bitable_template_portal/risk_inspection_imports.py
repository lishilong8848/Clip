"""Read-only, in-memory candidates for new risk inspection base templates.

Selection responses contain no items. Source metadata is descriptive, never an
authorization decision; callers must enforce manager permissions before reading.
"""
import csv
from datetime import date, datetime
import io
import math
from pathlib import PurePosixPath
import re
from urllib.parse import parse_qs, urlencode, urlsplit, urlunsplit
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile
from zlib import error as ZlibError

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

from .risk_inspection import FIELD_TYPES, SCOPES, SYSTEM_NAMES, RiskError, can_scope, manager, text
from .risk_inspection_templates import clean_fields, reusable_items


MAX_BYTES = 20 * 1024 * 1024
MAX_ITEMS = 100
MAX_FIELDS = 80
TITLE_NAMES = {'排查事项', '检查项', '检查事项', '排查内容', '事项', '任务名称', '检查内容', '事项名称'}
SENSITIVE = re.compile(r'身份证|证件号|手机号|手机号码|联系电话|电话号码|密码|口令|密钥|凭据|'
                       r'token|secret|password|credential|authorization|phone|mobile|id[_\s-]*card|'
                       r'api[_\s-]*key|access[_\s-]*key', re.I)
SYSTEM_ALIASES = {'任务编号', '审核人', '审核结果', '审核日期', '审核时间', '排查时间', '所属楼栋'}
SCOPE_NAMES = {'楼栋', '所属楼栋'}
SCOPE_LABELS = {label: scope for scope in SCOPES for label in (scope, scope + '楼')}
SCOPE_LABELS['110站'] = '110'
SOURCE_DENIED = '请全楼管理员导入为共享基础模板，或改上传本人有权的表格。'
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def _empty(value):
    return value is None or value == [] or isinstance(value, str) and not value.strip()


def _excluded(name, warnings):
    if name in SYSTEM_NAMES or name in SYSTEM_ALIASES or name.startswith('灯塔'):
        message = '系统列不导入历史值，使用新的系统字段：' + name
    elif SENSITIVE.search(name):
        message = '已排除敏感凭据或人员信息列：' + name
    else:
        return False
    if message not in warnings:
        warnings.append(message)
    return True


def _field(key, name, kind, **extra):
    return {'key': key, 'field_name': name, 'type': kind, 'role': 'input', **extra}


def _draft(name, fields, rows, source, warnings, title_field):
    """Clean both schemas and prefills through the existing template contract."""
    choices = [f for f in fields if f['type'] == 1]
    if not choices:
        raise RiskError('源表没有可用于事项标题的非敏感文本字段，请先提供排查事项文本列。')
    if title_field is not None:
        if not isinstance(title_field, str):
            raise RiskError('请选择有效的事项标题字段。')
        matches = [f for f in choices if title_field in {f['key'], f['field_name']}]
        if len(matches) != 1:
            raise RiskError('事项标题字段必须是唯一的非系统、非敏感文本列。')
    else:
        matches = [f for f in choices if f['field_name'] in TITLE_NAMES]
        if not matches:
            matches = [f for f in choices if f.get('is_primary')]
        if not matches and len(choices) == 1:
            matches = choices
    selected = matches[0] if len(matches) == 1 else None
    result = {'name': text(name, 120, True), 'fields': [], 'items': [], 'source': source,
              'warnings': warnings, 'needs_selection': selected is None}
    if selected is None:
        result.update(requires_selection='title_field', title_fields=[f['field_name'] for f in choices])
        warnings.append('无法唯一确定事项标题，请显式选择 title_field；当前没有生成事项。')
        return result
    fields = clean_fields([_field('system_number', '排查任务编号', 1, is_primary=True)] + [
        {**f, 'is_primary': False} for f in fields if f is not selected])
    result['fields'] = fields
    for f in fields:
        if f['role'] == 'input' and f['type'] in {5, 17}:
            warnings.append(('日期' if f['type'] == 5 else '附件') + '字段仅保留定义，不复用实际日期或附件：' + f['field_name'])
    source['title_field'] = selected['key']
    source['title_field_name'] = selected['field_name']
    warnings.append('标题列已完整映射到排查事项名称，不另存为填写列，避免修改事项名称后残留旧值：' + selected['field_name'])
    items = []
    for row in rows:
        if all(_empty(v) for v in row.values()):
            continue
        title = row.get(selected['key'])
        if not isinstance(title, str) or not title.strip():
            raise RiskError('存在事项标题为空的源记录，请补齐后重新预览。')
        items.append({'title': title, 'values': row})
    try:
        result['items'] = reusable_items(fields, items) if items else []
    except OverflowError:
        raise RiskError('源记录数值超出有效范围。') from None
    return result


def _check_zip(archive):
    entries = archive.infolist()
    if len(entries) > 2000 or len({e.filename for e in entries}) != len(entries):
        raise RiskError('Excel 压缩包条目过多或重复，已拒绝读取。')
    if sum(e.file_size for e in entries) > 80 * 1024 * 1024:
        raise RiskError('Excel 解压后超过80MB，已拒绝读取以防压缩炸弹。')
    for entry in entries:
        if entry.flag_bits & 1 or entry.file_size > max(1024 * 1024, entry.compress_size * 200):
            raise RiskError('Excel 压缩比例异常或已加密，已拒绝读取。')
        if entry.filename.endswith(('.xml', '.rels')):
            raw = archive.read(entry).replace(b'\x00', b'').upper()
            if b'<!DOCTYPE' in raw or b'<!ENTITY' in raw:
                raise RiskError('Excel XML 包含不安全的实体声明，已拒绝读取。')


def _sheet_layout(archive, path):
    active_rows, columns, merged = set(), set(), []
    header_row, has_formula = None, False
    for _, element in ET.iterparse(io.BytesIO(archive.read(path)), events=('end',)):
        if element.tag == NS + 'c':
            formula = element.find(NS + 'f')
            has_formula = has_formula or formula is not None
            cached = element.find(NS + 'v')
            if formula is not None and (cached is None or cached.text is None):
                raise RiskError('Excel 含无缓存结果的公式；请用 Excel 重新计算并保存，或粘贴为值后重试。')
            if element.get('t') == 'e':
                raise RiskError('Excel 包含计算错误单元格，请修正源文件后重试。')
            if formula is not None or any((child.text or '').strip() for child in element.iter()
                                          if child.tag in {NS + 'v', NS + 't'}):
                row, column = coordinate_to_tuple(element.attrib['r'])
                if row > 1048576 or column > 16384:
                    raise RiskError('Excel 单元格位置超出有效范围。')
                active_rows.add(row)
                columns.add(column)
                header_row = row if header_row is None else min(header_row, row)
            element.clear()
        elif element.tag == NS + 'mergeCell':
            merged.append(element.attrib['ref'])
        elif element.tag == NS + 'row':
            element.clear()
    if not active_rows:
        raise RiskError('工作表没有有效表头或排查事项。')
    return active_rows, sorted(columns), header_row, merged, has_formula


def _scalar(value):
    if not isinstance(value, str):
        return value
    stripped = value.strip()
    if stripped.lower() in {'true', 'false'}:
        return stripped.lower() == 'true'
    if re.fullmatch(r'-?(?:0|[1-9]\d*)(?:\.\d+)?(?:[eE][+-]?\d+)?', stripped):
        try:
            number = float(stripped) if any(c in stripped for c in '.eE') else int(stripped)
            valid = math.isfinite(number)
        except (ValueError, OverflowError):
            valid = False
        if not valid:
            raise RiskError('数值超出有效范围。')
        return number
    return value


def _tabular(name, headers, data, source, warnings, title_field, formats=None, positions=None):
    fields, rows, names = [], [], set()
    columns = []
    scalar = _scalar if source.get('encoding') else lambda value: value
    for index, header in enumerate(headers):
        if _empty(header) and all(_empty(row[index]) for row in data):
            continue
        if not isinstance(header, str) or not header.strip():
            raise RiskError('存在无有效名称的表头列，请补齐第一行有效表头。')
        header = text(header, 80, True)
        if header in names:
            raise RiskError('表头名称重复：' + header)
        names.add(header)
        if _excluded(header, warnings):
            continue
        values = [scalar(row[index]) for row in data if not _empty(row[index])]
        column_formats = [fmt for (_, col), fmt in (formats or {}).items() if col == index]
        if any(word in header for word in ('附件', '照片', '图片')):
            kind = 17
        elif values and all(isinstance(v, (date, datetime)) for v in values):
            kind = 5
        elif any(word in header for word in ('日期', '时间')) and values and all(
                isinstance(v, str) and re.fullmatch(r'\d{4}[-/]\d{1,2}[-/]\d{1,2}(?:[ T].*)?', v.strip()) for v in values):
            kind = 5
        elif values and all(type(v) is bool for v in values):
            kind = 7
        elif values and all(type(v) in {int, float} for v in values):
            # Identifiers with a zero-padded Excel format must stay text.
            kind = 1 if any(re.fullmatch(r'0{2,}', fmt) for fmt in column_formats) else 2
        else:
            kind = 1
        if header in TITLE_NAMES:
            kind = 1
        key = 'import_' + str((positions[index] if positions is not None else index) + 1)
        fields.append(_field(key, header, kind))
        if len(fields) - 1 + len(SYSTEM_NAMES) > MAX_FIELDS:
            raise RiskError('包含9个系统字段后超过80个字段，请精简源表。')
        columns.append((index, key, kind))
    for row_index, raw in enumerate(data):
        values = {}
        for index, key, kind in columns:
            value = scalar(raw[index])
            if _empty(value):
                continue
            if kind == 1 and not isinstance(value, str):
                fmt = (formats or {}).get((row_index, index), '')
                value = (format(value, '0' + str(len(fmt)) + 'd') if re.fullmatch(r'0{2,}', fmt)
                         and type(value) is int else str(value))
            values[key] = value
        rows.append(values)
    return _draft(name, fields, rows, source, warnings, title_field)


def preview_file(name, content, *, sheet=None, title_field=None):
    """Preview bytes only; no uploaded originals are saved anywhere."""
    name = text(name, 255, True).replace('\\', '/').rsplit('/', 1)[-1]
    suffix = PurePosixPath(name).suffix.lower()
    if not isinstance(content, bytes) or not content:
        raise RiskError('请上传非空文件。')
    if len(content) > MAX_BYTES:
        raise RiskError('上传文件不得超过20MB。', 413)
    if suffix == '.xls':
        raise RiskError('不支持旧二进制 xls 文件，请另存为 xlsx 后导入。')
    if suffix not in {'.xlsx', '.xlsm', '.csv'}:
        raise RiskError('仅支持 xlsx、xlsm 或 csv 文件。')
    source = {'kind': 'file', 'filename': name}
    draft_name = text(PurePosixPath(name).stem, 120, True)
    warnings = []
    if suffix == '.csv':
        for encoding in ('utf-8-sig', 'gb18030'):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise RiskError('CSV 编码无效，请使用 UTF-8 BOM 或 GB18030。')
        source['encoding'] = encoding
        try:
            rows = []
            for row_number, row in enumerate(csv.reader(io.StringIO(decoded, newline=''), strict=True), 1):
                if any(not _empty(v) for v in row):
                    if not rows:
                        source['header_row'] = row_number
                    rows.append(row)
                if len(rows) > MAX_ITEMS + 1:
                    raise RiskError('源表超过100项，请拆分后导入；不会截断数据。')
        except csv.Error:
            raise RiskError('CSV 格式无效，请检查引号、列分隔和单元格长度。') from None
        if not rows:
            raise RiskError('CSV 没有有效表头。')
        columns = sorted({i for row in rows for i, v in enumerate(row) if not _empty(v)})
        rows = [[row[i] if i < len(row) else None for i in columns] for row in rows]
        return _tabular(draft_name, rows[0], rows[1:], source, warnings, title_field, positions=columns)
    workbook = None
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            _check_zip(archive)
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
            sheets = workbook.sheetnames
            if sheet is None and len(sheets) > 1:
                return {'name': draft_name, 'fields': [], 'items': [], 'source': source,
                        'warnings': ['存在多个工作表，请显式选择 sheet 后重新预览。'],
                        'needs_selection': True, 'requires_selection': 'sheet', 'sheets': sheets}
            if sheet is not None and (not isinstance(sheet, str) or sheet not in sheets):
                raise RiskError('请选择存在的工作表名称。')
            ws = workbook[sheet or sheets[0]]
            source['sheet'] = ws.title
            active, columns, first_row, merges, has_formula = _sheet_layout(archive, ws._worksheet_path.lstrip('/'))
            source['header_row'] = first_row
            merged = []
            for ref in merges:
                left, top, right, bottom = range_boundaries(ref)
                if left != right and top <= first_row <= bottom:
                    merged.append(ref)
            if merged:
                raise RiskError('检测到跨列合并标题，请整理为单行列标题后重试。')
            if len(active) - 1 > MAX_ITEMS:
                raise RiskError('源表超过100项，请拆分后导入；不会截断数据。')
            ws.reset_dimensions()
            rows, formats = [], {}
            for index, cells in enumerate(ws.iter_rows(min_row=first_row, max_row=max(active),
                                                       max_col=max(columns)), first_row):
                if index not in active:
                    continue
                rows.append([cells[i - 1].value for i in columns])
                if index == first_row:
                    continue
                for i, col in enumerate(columns):
                    cell = cells[col - 1]
                    if not _empty(cell.value) and cell.number_format != 'General':
                        formats[(len(rows) - 2, i)] = cell.number_format
            if formats:
                warnings.append('Excel 显示样式（颜色、百分号等）不复用；保留字段类型和数值，纯零补位编号按文本保留。')
            if has_formula:
                warnings.append('Excel 公式不执行；有缓存的公式只读取已保存的结果，请确认源文件已重新计算。')
            return _tabular(draft_name, rows[0], rows[1:], source, warnings, title_field, formats,
                            positions=[i - 1 for i in columns])
    except RiskError:
        raise
    except (BadZipFile, ET.ParseError, ValueError, KeyError, IndexError, OSError, RuntimeError,
            TypeError, AttributeError, InvalidFileException, ZlibError, NotImplementedError):
        raise RiskError('Excel 文件损坏或格式不受支持，请另存为 xlsx 后重试。') from None
    finally:
        if workbook is not None:
            workbook.close()


def _bitable_url(url):
    if not isinstance(url, str) or len(url) > 8192 or re.search(r'[\x00-\x20\\]', url):
        raise RiskError('多维表链接格式无效。')
    try:
        parts = urlsplit(url)
        host = parts.hostname or ''
        path = re.fullmatch(r'/(base|wiki)/([A-Za-z0-9]+)', parts.path)
        valid_host = re.fullmatch(r'(?:[A-Za-z0-9](?:[A-Za-z0-9-]*[A-Za-z0-9])?\.)*(?:feishu\.cn|larksuite\.com)', host)
        if (parts.scheme != 'https' or not valid_host or parts.username is not None or parts.password is not None
                or parts.port is not None or parts.netloc.lower() != host or not path):
            raise ValueError
        query = parse_qs(parts.query, keep_blank_values=True, max_num_fields=100)
        if len(query.get('table', [])) != 1 or not re.fullmatch(r'tbl[A-Za-z0-9]+', query['table'][0]):
            raise ValueError
        if 'view' in query and (len(query['view']) != 1 or not re.fullmatch(r'vew[A-Za-z0-9]+', query['view'][0])):
            raise ValueError
    except ValueError:
        raise RiskError('仅支持 HTTPS 飞书/Lark 官方 base 或 wiki 链接，必须显式携带 table，可选 view。') from None
    safe = {'table': query['table'][0]}
    if 'view' in query:
        safe['view'] = query['view'][0]
    return path[1], path[2], safe, urlunsplit(('https', host, parts.path, urlencode(safe), ''))


def _cloud_value(field, value):
    if _empty(value) or field['type'] in {5, 17}:
        return value
    if field['type'] in {1, 13}:
        if isinstance(value, list) and all(isinstance(v, dict) and isinstance(v.get('text'), str) for v in value):
            value = ''.join(v['text'] for v in value)
        if not isinstance(value, str):
            raise RiskError('文本列没有可靠的显示文本，无法导入：' + field['field_name'])
    if field['type'] == 15 and isinstance(value, dict):
        value = value.get('link')
        if not isinstance(value, str):
            raise RiskError('链接列格式无效：' + field['field_name'])
    return value


def preview_bitable(remote, url, *, title_field=None, actor=None):
    """Use only the supplied RiskCloud agent identity and read API operations."""
    if not isinstance(actor, dict) or not manager(actor):
        raise RiskError('导入源表必须提供有效的管理员身份。', 403)
    kind, token, query, safe_url = _bitable_url(url)
    name = '多维表基础模板'
    if kind == 'wiki':
        data = remote.request('GET', 'wiki/v2/spaces/get_node', params={'token': token})
        node = data.get('node') if isinstance(data, dict) else None
        if (not isinstance(node, dict) or node.get('obj_type') != 'bitable'
                or not isinstance(node.get('obj_token'), str) or not re.fullmatch(r'[A-Za-z0-9]+', node['obj_token'])):
            raise RiskError('知识库节点未解析为有效多维表，无法导入。')
        token = node['obj_token']
        if isinstance(node.get('title'), str) and node['title'].strip():
            name = text(node['title'], 120, True)
    path = remote.table_path(token, query['table'])
    originals = remote.all(path + '/fields')
    if not isinstance(originals, list) or not originals:
        raise RiskError('多维表字段未完整返回。', 503)
    from .risk_inspection_cloud import TEMPLATE_APP, TEMPLATE_TABLE
    restricted = actor['level'] == 'L3' and (token, query['table']) != (TEMPLATE_APP, TEMPLATE_TABLE)
    scope_field = None
    if restricted:
        candidates = [f for f in originals if isinstance(f, dict) and f.get('field_name') in SCOPE_NAMES]
        if len(candidates) != 1 or type(candidates[0].get('type')) is not int or candidates[0]['type'] not in {1, 3}:
            raise RiskError('该源表没有唯一可靠的楼栋文本或单选列，楼栋主管不能导入。' + SOURCE_DENIED, 403)
        if not any(can_scope(actor, scope) for scope in SCOPES):
            raise RiskError('当前账号没有可导入的楼栋范围。' + SOURCE_DENIED, 403)
        scope_field = candidates[0]
    fields, warnings, names, keys = [], [], set(), set()
    for original in originals:
        if not isinstance(original, dict):
            raise RiskError('多维表字段格式无效。', 503)
        fname = text(original.get('field_name'), 80, True)
        if _excluded(fname, warnings):
            continue
        fid, ftype = original.get('field_id'), original.get('type')
        if ftype == 13:
            warnings.append('已排除敏感的电话号码类型列：' + fname)
            continue
        if not isinstance(fid, str) or not re.fullmatch(r'[A-Za-z0-9_-]+', fid) or fid in keys or fname in names:
            raise RiskError('多维表字段名称或标识重复、无效。', 503)
        if type(ftype) is not int or ftype not in FIELD_TYPES:
            raise RiskError('不支持关联、公式、人员或其它不可独立复制的字段：' + fname + '。请在源表提供独立文本列后重试。')
        fields.append(_field(fid, fname, ftype, property=original.get('property') or {},
                             is_primary=original.get('is_primary') is True))
        names.add(fname)
        keys.add(fid)
    if len(fields) - 1 + len(SYSTEM_NAMES) > MAX_FIELDS:
        raise RiskError('包含9个系统字段后超过80个字段，请精简源表。')
    if not fields or not any(f['type'] == 1 for f in fields):
        raise RiskError('源表没有可用于事项标题的非敏感文本字段，请先提供排查事项文本列。')
    if any(not isinstance(f['property'], dict) for f in fields):
        raise RiskError('多维表字段属性格式无效。', 503)
    if any(f['type'] == 2 and (f['property'] or {}).get('formatter') for f in fields):
        warnings.append('数字显示格式由基础模板统一规范化，保留原始数值。')
    source = {'kind': 'bitable', 'url': safe_url, 'app_token': token, 'table_id': query['table']}
    body = {'field_names': [f['field_name'] for f in fields if f['type'] not in {5, 17}]}
    if scope_field:
        body['field_names'].append(scope_field['field_name'])
    if 'view' in query:
        source['view_id'] = query['view']
        body['view_id'] = query['view']
        body['automatic_fields'] = False
    rows, cursors, record_ids, cursor, total = [], set(), set(), '', None
    record_count, imported_scopes = 0, set()
    for _ in range(200):
        params = {'page_size': MAX_ITEMS}
        if cursor:
            params['page_token'] = cursor
        data = remote.request('POST', path + '/records/search', params=params, json=body)
        if not isinstance(data, dict) or not isinstance(data.get('items'), list) or type(data.get('has_more')) is not bool:
            raise RiskError('多维表记录未完整返回，不能使用不完整预览。', 503)
        if 'total' in data:
            if type(data['total']) is not int or data['total'] < 0 or total is not None and total != data['total']:
                raise RiskError('多维表记录总数无效或发生变化，请重新预览。', 503)
            total = data['total']
        record_count += len(data['items'])
        if record_count > MAX_ITEMS or total is not None and total > MAX_ITEMS:
            raise RiskError('源表超过100项，请缩小视图或拆分后导入；不会截断数据。')
        for record in data['items']:
            if not isinstance(record, dict) or not isinstance(record.get('fields'), dict):
                raise RiskError('多维表记录格式不完整。', 503)
            rid = record.get('record_id')
            if not isinstance(rid, str) or not rid or rid in record_ids:
                raise RiskError('多维表出现重复或无效记录，已停止预览。', 503)
            record_ids.add(rid)
            values = record['fields']
            if scope_field:
                try:
                    scope = _cloud_value(_field('', scope_field['field_name'], 1), values.get(scope_field['field_name']))
                except RiskError:
                    raise RiskError('源记录的楼栋格式无法明确识别，未导入任何数据。' + SOURCE_DENIED, 403) from None
                scope = SCOPE_LABELS.get(scope.strip()) if isinstance(scope, str) else None
                if scope is None:
                    raise RiskError('源记录的楼栋为空或无法明确识别，未导入任何数据。' + SOURCE_DENIED, 403)
                if not can_scope(actor, scope):
                    continue
                imported_scopes.add(scope)
            rows.append({f['key']: _cloud_value(f, values[f['field_name']]) for f in fields if f['field_name'] in values})
        if not data['has_more']:
            if total is not None and total != record_count:
                raise RiskError('多维表记录总数不一致，不能使用不完整预览。', 503)
            break
        cursor = data.get('page_token')
        if not isinstance(cursor, str) or not cursor or cursor in cursors:
            raise RiskError('多维表分页游标异常，不能使用不完整预览。', 503)
        cursors.add(cursor)
    else:
        raise RiskError('多维表分页未完成，不能使用不完整预览。', 503)
    if not rows:
        raise RiskError('源表或所选视图没有可导入权限范围内的事项。' + (SOURCE_DENIED if restricted else ''), 403 if restricted else 400)
    if restricted:
        source['imported_scopes'] = sorted(imported_scopes)
        warnings.append('仅导入当前全局楼栋权限范围内的记录：' + '、'.join('110站' if s == '110' else s + '楼' for s in sorted(imported_scopes)) +
                        '。其它楼栋记录已排除；此来源信息不授予业务权限。')
    result = _draft(name, fields, rows, source, warnings, title_field)
    if not result['needs_selection'] and not result['items']:
        raise RiskError('源表没有有效的排查事项，未生成草稿。')
    return result
