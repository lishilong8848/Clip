"""Native portal service for plan-convergence review; the original matching rules are retained."""
import io
import hashlib
import json
import threading
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor, wait, FIRST_COMPLETED
from contextlib import closing

import httpx

from . import plan_convergence_auth as auth
from . import plan_convergence_rules as rules
from . import plan_convergence_maintenance as maintenance
from .plan_convergence_compare import compare_scenarios_to_details
from . import plan_convergence_points as points
from .plan_convergence_browser_login import BrowserLogin
from upload_event_module.services.process_lifetime import lower_current_thread_priority

BASE = auth.ZH_BASE + '/api/alarm/alarmBlock'
QUERY_SECONDS = 60


def remaining(deadline):
    seconds = deadline - time.monotonic()
    if seconds <= 0:
        raise TimeoutError('核对读取超时，本次未产生核对结论，请检查 VPN 后重试。')
    return seconds


class PlanConvergenceService:
    def __init__(self, store, ongoing_provider=None):
        self.store = store
        self._ongoing_provider = ongoing_provider
        self._blocks_lock = threading.Lock()
        self._persist_lock = threading.Lock()
        self._details_pool = ThreadPoolExecutor(max_workers=4, thread_name_prefix='PlanDetail', initializer=lower_current_thread_priority)
        self.browser_login = BrowserLogin(store)
        auth.bind_store(store)

    def close(self):
        self._details_pool.shutdown(wait=False, cancel_futures=True)
        self.browser_login.close()

    def bootstrap(self, admin):
        credentials = auth.settings_view()
        counts = {}
        if points.CATALOG_PATH.is_file():
            with closing(points._catalog_conn()) as conn:
                counts = {key: conn.execute('SELECT COUNT(*) FROM ' + table).fetchone()[0]
                          for key, table in [('devices', 'zh_device'), ('rules', 'zh_rules')]}
        return {'is_admin': admin, 'catalog_ready': points.CATALOG_PATH.is_file(),
                'points_ready': points.POINTS_PATH.is_file(), 'counts': counts,
                'zh_ready': bool(credentials['has_token'] and not credentials['expired']),
                'blocks': self.store.get_document('plan_convergence', 'blocks') or {'items': [], 'loaded_at': 0}}

    @staticmethod
    def remote(method, path, payload=None, credentials=None, *, deadline=None):
        headers, cookies = credentials or auth.build_auth()
        seconds = min(25, remaining(deadline)) if deadline is not None else 25
        response = auth.request(method, BASE + path, headers=headers, cookies=cookies,
                                 json=payload,
                                 timeout=httpx.Timeout(connect=min(5, seconds), read=seconds, write=seconds, pool=min(5, seconds)),
                                 follow_redirects=False)
        if deadline is not None:
            remaining(deadline)
        response.raise_for_status()
        result = response.json()
        if not isinstance(result, dict) or result.get('code') != 200 or not result.get('success'):
            raise ValueError('智航返回异常，请检查平台凭证或网络连接')
        if not isinstance(result.get('data'), dict):
            raise ValueError('智航返回数据不完整，已停止核对')
        return result['data']

    def blocks(self, refresh=False, *, deadline=None):
        if not refresh:
            return self.store.get_document('plan_convergence', 'blocks') or {'items': [], 'loaded_at': 0}
        deadline = deadline or time.monotonic() + QUERY_SECONDS
        if not self._blocks_lock.acquire(timeout=remaining(deadline)):
            raise TimeoutError('屏蔽列表读取超时，请稍后重试。')
        try:
            items, seen, cursor, seen_pages = [], set(), None, set()
            credentials = auth.build_auth()
            for page in range(1, 101):
                payload = {} if page == 1 else {'page': page, 'size': 100, 'searchAfter': cursor}
                data = self.remote('POST', '/getAlarmBlock', payload, credentials, deadline=deadline)
                chunk = data.get('content')
                if not isinstance(chunk, list):
                    raise ValueError('智航分页内容不完整')
                signature = tuple(str(row.get('blockId') or row.get('id') or '') for row in chunk if isinstance(row, dict))
                if signature in seen_pages and data.get('hasNext'):
                    raise ValueError('智航重复返回同一页，已停止核对')
                seen_pages.add(signature)
                for row in chunk:
                    if not isinstance(row, dict):
                        raise ValueError('智航屏蔽记录结构无效')
                    key = str(row.get('blockId') or row.get('id') or '')
                    if key and key not in seen:
                        seen.add(key)
                        items.append(row)
                if not data.get('hasNext'):
                    result = {'items': items, 'loaded_at': time.time()}
                    self.store.put_document('plan_convergence', 'blocks', result)
                    return result
                if not chunk:
                    break
                cursor = data.get('searchAfter')
            raise ValueError('智航分页未完整返回，已停止核对')
        finally:
            self._blocks_lock.release()

    def block(self, block_id, *, deadline=None):
        if not str(block_id).isdigit():
            raise ValueError('屏蔽记录 ID 无效')
        data = self.remote('GET', '/getAlarmBlockDetail/' + str(block_id), deadline=deadline)
        if not isinstance(data.get('alarmBlockDetailResultList'), list):
            raise ValueError('屏蔽记录明细不完整，不能核对')
        data.setdefault('blockId', str(block_id))
        return data

    def remote_rows(self, kind, payload):
        paths = {'snapshots': '/getBlockInstanceSnapshot', 'rule-view': '/getRuleView'}
        rows, credentials, signatures = [], auth.build_auth(), set()
        deadline = time.monotonic() + QUERY_SECONDS
        for page in range(1, 501):
            data = self.remote('POST', paths[kind], {**payload, 'page': page, 'size': 100}, credentials, deadline=deadline)
            chunk = data.get('content')
            if not isinstance(chunk, list):
                raise ValueError('智航明细分页结构无效')
            signature = hashlib.sha256(json.dumps(chunk, sort_keys=True, ensure_ascii=False).encode()).digest()
            if signature in signatures and data.get('hasNext'):
                raise ValueError('智航重复返回同一页，已停止读取')
            signatures.add(signature)
            rows.extend(chunk)
            if not data.get('hasNext'):
                return rows
            if not chunk:
                break
        raise ValueError('智航明细分页未完整返回')

    @staticmethod
    def catalog(query):
        kind, kw = query.get('kind', ''), str(query.get('kw') or '')[:200]
        zone, building, floor, room = (str(query.get(key) or '')[:200] for key in ('zone', 'building', 'floor', 'room'))
        objs = [value for value in str(query.get('objs') or '').split(',') if value and value != 'all'] or None
        rooms = [tuple((part.split('|') + [''] * 4)[:4]) for part in str(query.get('rooms') or '').split(',') if part and part != 'all'] or None
        if len(objs or []) > 500 or len(rooms or []) > 200:
            raise ValueError('筛选条件过多，请缩小范围')
        if kind == 'types':
            rows = points.zh_objtypes_new()
        elif kind == 'zones':
            rows = points.zh_zone_list()
        elif kind == 'buildings':
            rows = points.zh_building_list(zone)
        elif kind == 'floors':
            rows = points.zh_floor_list(zone, building)
        elif kind == 'rooms':
            rows = points.zh_room_list(zone, building, floor)
        elif kind == 'devices':
            rows = points.zh_devices_by_rooms(objs, rooms, kw) if rooms else points.zh_devices_new(objs, zone, building, floor, room, kw)
        elif kind == 'rules':
            if query.get('obj'):
                rows = points.zh_rules_for_obj(str(query['obj'])[:200], kw)
            elif query.get('inst'):
                rows = points.zh_rules_for_device(str(query['inst'])[:200], kw)
            elif objs:
                placeholders = ','.join(['?'] * len(objs))
                rows = points._rules_dedup_sql('WHERE classify_model IN (' + placeholders + ')', tuple(objs), kw, 3000)
            elif rooms:
                rows = points.zh_rules_all_by_rooms(rooms, kw, limit=3000)
            else:
                rows = points.zh_rules_all(kw)
        else:
            raise ValueError('目录类型无效')
        if kind not in {'devices', 'rules'} and kw:
            rows = [row for row in rows if kw.casefold() in str(row).casefold()]
        return {'items': rows, 'has_more': kind in {'devices', 'rules'} and len(rows) >= 3000, 'limit': 3000}

    @staticmethod
    def expand(set_id):
        if not rules.get_set(set_id):
            raise ValueError('规则集不存在')
        ids = list(rules.expand_set_items(set_id))
        devices = []
        for start in range(0, len(ids), 900):
            batch = ids[start:start + 900]
            placeholders = ','.join(['%s'] * len(batch))
            devices.extend(rules._exec('SELECT ins_id,inst_name,obj_name,position FROM zh_device WHERE ins_id IN (' + placeholders + ')', tuple(batch)))
        devices.sort(key=lambda row: str(row.get('inst_name') or ''))
        return {'count': len(ids), 'devices': devices}

    def match(self, set_id, payload):
        if not rules.get_set(set_id):
            raise ValueError('规则集不存在')
        details = payload.get('details')
        if details is None:
            details = self.block(payload.get('block_id', ''))['alarmBlockDetailResultList']
        if not isinstance(details, list) or len(details) > 10000 or any(not isinstance(row, dict) for row in details):
            raise ValueError('屏蔽明细结构无效')
        return rules.match_record_to_set(set_id, details)

    @staticmethod
    def parse_excel(content, filename):
        from openpyxl import load_workbook
        if not filename.lower().endswith(('.xlsx', '.xlsm')) or not content.startswith(b'PK'):
            raise ValueError('请选择 xlsx 或 xlsm 文件')
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 200 * 1024 * 1024:
                raise ValueError('Excel 解压内容过大，请拆分场景文件')
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheets, count = [], 0
        try:
            for sheet in workbook:
                iterator = sheet.iter_rows(values_only=True)
                headers = [str(value).strip() if value is not None else '' for value in next(iterator, ())]
                rows = []
                for number, cells in enumerate(iterator, 2):
                    count += 1
                    if count > 10000:
                        raise ValueError('场景 Excel 最多读取 10000 行，请拆分文件')
                    if not any(value is not None and str(value).strip() for value in cells):
                        continue
                    row = {key: str(cells[index]).strip() if index < len(cells) and cells[index] is not None else ''
                           for index, key in enumerate(headers) if key}
                    rows.append({**row, '_excel_row': number})
                sheets.append({'name': sheet.title, 'headers': headers, 'rows': rows})
        finally:
            workbook.close()
        return {'sheets': sheets}

    def compare(self, payload):
        scenarios = payload.get('scenarios')
        if not isinstance(scenarios, list) or len(scenarios) != 1 or not isinstance(scenarios[0], dict):
            raise ValueError('请选择一个场景进行核对')
        scenario = scenarios[0]
        if not scenario.get('scenario_name') or not isinstance(scenario.get('rows'), list) or not scenario['rows']:
            raise ValueError('场景名称或数据为空')
        data = self.block(payload.get('block_id', ''))
        result = compare_scenarios_to_details(scenarios, data['alarmBlockDetailResultList'], data['blockId'])
        return {**result, 'source_detail': data}

    def _manual_check(self, kind, record_id=None, *, records=None, list_method):
        """Explicit user-triggered check; persists a local manual result only for
        user-invoked calls (``records`` is ``None``). The automatic notice flow
        passes ``records=[record]`` and therefore never writes a manual snapshot."""
        internal = records is None
        if internal:
            records = list_method(record_id)
        checked_at = time.time()
        result = self._check_records(records, stat_key='maintenance' if kind == 'repair' else kind)
        if internal:
            self._persist_manual_checks(kind, result['records'], checked_at)
            for row in result['records']:
                row['check_status'] = 'skipped' if maintenance.excluded_buildings(row) else 'ready'
                row['check_source'] = 'manual'
                row['checked_at'] = checked_at
        return result

    def _check_records(self, records, stat_key, *, seconds=QUERY_SECONDS, refresh=True):
        if records and all(maintenance.excluded_buildings(row) for row in records):
            return {'records': [dict(r, hits=[], check_status='skipped', check_error='D/E楼不参与计划收敛匹配') for r in records],
                    'orphan_block_ids': [], 'stats': {stat_key: len(records), 'blocks': 0, 'matched_records': 0, 'orphan_blocks': 0}}
        deadline = time.monotonic() + seconds
        cached = self.blocks()
        blocks = [row for row in self.blocks(refresh=refresh or time.time() - cached.get('loaded_at', 0) > 120, deadline=deadline)['items']
                  if str(row.get('status')) == '1' and not maintenance.excluded_buildings({'name': row.get('blockName', '')})]
        details = {}
        pending, rows = {}, iter(blocks)
        def submit_next():
            row = next(rows, None)
            if row is not None:
                remaining(deadline)
                ident = str(row.get('blockId') or row.get('id'))
                def read_detail():
                    fingerprint = hashlib.sha256(json.dumps(row, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
                    old = self.store.get_document('plan_convergence_details', ident)
                    if not refresh and old and old.get('fingerprint') == fingerprint and time.time() - old['at'] < 120:
                        return old['detail']
                    detail = self.block(ident, deadline=deadline)
                    self.store.put_document('plan_convergence_details', ident, {'fingerprint': fingerprint, 'at': time.time(), 'detail': detail})
                    return detail
                pending[self._details_pool.submit(read_detail)] = ident
        try:
            for _ in range(4):
                submit_next()
            while pending:
                completed, _ = wait(pending, timeout=remaining(deadline), return_when=FIRST_COMPLETED)
                if not completed:
                    raise TimeoutError('核对读取超时，未把不完整数据作为核对结论。')
                for future in completed:
                    details[pending.pop(future)] = future.result()
                for _ in completed:
                    submit_next()
        finally:
            for future in pending:
                future.cancel()
        result, orphans = maintenance.match_records(records, blocks, details)
        return {'records': result, 'orphan_block_ids': orphans,
                'stats': {stat_key: len(records), 'blocks': len(blocks),
                          'matched_records': sum(bool(row['hits']) for row in result), 'orphan_blocks': len(orphans)}}

    def _persist_manual_checks(self, kind, result_rows, checked_at):
        if not result_rows:
            return
        payloads = {}
        with self._persist_lock:
            for row in result_rows:
                key = '{0}:{1}'.format(kind, row['record_id'])
                existing = self.store.get_document('plan_convergence_checks', key)
                # A concurrent newer manual check (larger checked_at) must never be
                # overwritten by this slower, older result.
                if existing and (existing.get('checked_at') or 0) >= checked_at:
                    continue
                payloads[key] = {
                    'work_type': kind,
                    'source': 'manual',
                    'check_source': 'manual',
                    'check_status': 'ready',
                    'checked_at': checked_at,
                    'hits': row.get('hits', []),
                    'fingerprint': maintenance.business_fingerprint(row),
                }
            if payloads:
                # One atomic commit for the whole result batch: no partial write
                # even when the batch contains multiple records.
                self.store.put_documents('plan_convergence_checks', payloads)

    def _overlay_check_results(self, record, kind):
        """Overlay persisted check results without any Zhihang/Feishu request.

        Generic ``json_documents`` lookups per visible ongoing row keep this
        free of history scans or blanket migrations.
        """
        if maintenance.excluded_buildings(record):
            record.update(hits=[], check_status='skipped', check_error='D/E楼不参与计划收敛匹配')
            return
        candidates = []
        cached = self.store.get_document('plan_convergence_checks', '{0}:{1}'.format(kind, record['record_id']))
        if cached:
            fingerprint_ok = cached.get('fingerprint') is not None and \
                cached['fingerprint'] == maintenance.business_fingerprint(record)
            hits = cached.get('hits')
            if fingerprint_ok:
                # A valid empty hits list is still a complete ready result; a
                # missing/malformed hits list marks the payload failed/incomplete.
                status = 'ready' if isinstance(hits, list) else 'failed'
            else:
                status = 'stale'
            candidates.append({
                'source': 'manual',
                'check_source': cached.get('check_source') or 'manual',
                'check_status': status,
                'checked_at': cached.get('checked_at'),
                'hits': hits if status == 'ready' else None,
            })
        if kind == 'repair':
            ref = self.store.get_document('notice_plan_checks', 'latest:' + record['record_id']) or {}
            job = self.store.get_document('notice_plan_checks', ref['job_id']) if ref else None
            if job:
                record['auto_check_status'] = job['status']
                record['auto_checked_at'] = job.get('checked_at')
                if job['status'] == 'ready':
                    fingerprint_ok = maintenance.business_fingerprint(job.get('record') or {}) == \
                        maintenance.business_fingerprint(record)
                    job_hits = None
                    job_result = job.get('result')
                    job_records = job_result.get('records') if isinstance(job_result, dict) else None
                    if isinstance(job_records, list) and len(job_records) == 1 and isinstance(job_records[0], dict):
                        maybe_hits = job_records[0].get('hits')
                        if isinstance(maybe_hits, list):
                            job_hits = maybe_hits
                    if fingerprint_ok:
                        status = 'ready' if job_hits is not None else 'failed'
                    else:
                        status = 'stale'
                    candidates.append({
                        'source': 'auto',
                        'check_source': 'auto',
                        'check_status': status,
                        'checked_at': job.get('checked_at'),
                        'hits': job_hits if status == 'ready' else None,
                    })
        if not candidates:
            return
        ready = [c for c in candidates if c['check_status'] == 'ready']
        selected = max(ready, key=lambda c: (c['checked_at'] or 0)) if ready else None
        if selected is None and candidates:
            # No ready result: surface the most recent stale/failed candidate so
            # the UI can tell the source changed or the payload is incomplete.
            # Incomplete/pending/failed results never attach a fabricated hits list.
            selected = max(candidates, key=lambda c: (c['checked_at'] or 0))
        if selected is None:
            return
        record['check_status'] = selected['check_status']
        record['check_source'] = selected['check_source']
        record['checked_at'] = selected['checked_at']
        if selected['check_status'] == 'ready':
            record['hits'] = selected['hits']

    def maintenance_records(self, record_id=None):
        if self._ongoing_provider is None:
            raise ValueError('本地未结束通告读取入口未接入')
        records = maintenance.ongoing_records(self._ongoing_provider('ALL'), work_type='repair')
        for record in records:
            self._overlay_check_results(record, 'repair')
        if record_id is not None:
            records = [row for row in records if row['record_id'] == str(record_id)]
            if not records:
                raise FileNotFoundError('该检修通告已结束、删除或不在当前未结束列表中，请刷新列表')
        return records

    def change_records(self, record_id=None):
        if self._ongoing_provider is None:
            raise ValueError('本地未结束通告读取入口未接入')
        records = maintenance.ongoing_records(self._ongoing_provider('ALL'), work_type='change')
        for record in records:
            self._overlay_check_results(record, 'change')
        if record_id is not None:
            records = [row for row in records if row['record_id'] == str(record_id)]
            if not records:
                raise FileNotFoundError('该变更通告已结束、删除或不在当前未结束列表中，请刷新列表')
        return records

    def maintenance_check(self, record_id=None, *, records=None):
        return self._manual_check('repair', record_id, records=records, list_method=self.maintenance_records)

    def change_check(self, record_id=None, *, records=None):
        return self._manual_check('change', record_id, records=records, list_method=self.change_records)

    def notice_points(self, kind, record_id, block_id):
        rows = self.change_records(record_id) if kind == 'change' else self.maintenance_records(record_id)
        row = rows[0]
        hit = next((h for h in row.get('hits', []) if str(h['blockId']) == str(block_id)), None)
        if row.get('check_status') != 'ready' or not hit:
            raise ValueError('核对结果已变化，请重新核对后查看点位。')
        points_rows = hit.get('point_rows')
        if points_rows is None:
            points_rows = maintenance.configured_points(self.block(str(block_id)))
            hit['point_rows'] = points_rows
            self._persist_manual_checks(kind, rows, time.time())
        return {'items': points_rows, 'block_name': hit.get('blockName'), 'notice_name': row['name'],
                'checked_at': row.get('checked_at'), 'note': '显示该候选计划实际配置的设备及点位/告警规则；范围为全部时不虚构逐点明细。'}
