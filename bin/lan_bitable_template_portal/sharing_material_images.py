"""Extract file illustrations into the existing permission-protected image store."""
import hashlib
import io
from contextlib import closing
from pathlib import Path

from .learning import LearningError
from .sharing_images import MAX_IMAGES, jpeg


def extract_images(path, digest, store, title=None):
    try:
        return _extract_images(path, digest, store, title)
    except LearningError:
        raise
    except Exception:
        raise LearningError('材料中的图片无法完整读取，请检查文件是否损坏、加密或含不支持的图片格式；原文件和已完成材料保留。', 422) from None


def _extract_images(path, digest, store, title):
    path = Path(path)
    result, seen = [], set()

    def add(content, label, page=None):
        data = jpeg(content)
        identity = hashlib.sha256(data).hexdigest()
        if identity in seen:
            return
        if len(result) >= MAX_IMAGES:
            raise LearningError('材料中的图片页超过24张，请拆分材料；未截断出题。', 413)
        store.save(identity, data)
        result.append({'local_sha256': identity, 'source_sha256': digest, 'revision': digest,
                       'source_url': '', 'label': f'{title or path.name} · {label}',
                       **({'page': page} if page is not None else {})})
        seen.add(identity)

    if path.suffix.lower() == '.pdf':
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(str(path)) as document:
            if len(document) > 100:
                raise LearningError('PDF超过100页，请拆分材料。', 413)
            for index in range(len(document)):
                with closing(document[index]) as page:
                    if not any(page.get_objects(filter=[2, 3])):
                        continue
                    with closing(page.render(scale=min(2, 1800 / max(page.get_size())))) as bitmap:
                        with bitmap.to_pil() as image:
                            buffer = io.BytesIO()
                            image.save(buffer, format='PNG')
                    add(buffer.getvalue(), f'第{index + 1}页', index + 1)
    elif path.suffix.lower() == '.docx':
        from docx import Document
        document = Document(path)
        for relation in document.part.rels.values():
            if relation.reltype.endswith('/image') and not relation.is_external:
                add(relation.target_part.blob, f'图片{len(result) + 1}')
    elif path.suffix.lower() == '.pptx':
        from pptx import Presentation
        def pictures(shapes, slide):
            for shape in shapes:
                if hasattr(shape, 'shapes'):
                    pictures(shape.shapes, slide)
                elif hasattr(shape, 'image'):
                    add(shape.image.blob, f'第{slide}张幻灯片 · {shape.name}', slide)
        for index, slide in enumerate(Presentation(path).slides):
            pictures(slide.shapes, index + 1)
    elif path.suffix.lower() in {'.xlsx', '.xlsm'}:
        from openpyxl import load_workbook
        workbook = load_workbook(path, data_only=True, keep_links=False)
        try:
            for sheet in workbook:
                for index, image in enumerate(sheet._images):
                    add(image._data(), f'{sheet.title} · 图片{index + 1}')
        finally:
            workbook.close()
    return result
