"""Private assistant attachments; originals never become arbitrary filesystem tools."""
import asyncio
import base64
import hashlib
import io
import json
import mimetypes
import multiprocessing
import os
import re
import time
import uuid
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from .lighthouse_ai import AssistantError, safe_data, safe_text

FILE_NAMESPACE = "lighthouse_files"
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_BATCH_BYTES = 100 * 1024 * 1024
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".log", ".json", ".xml"}
ALLOWED_EXTENSIONS = IMAGE_EXTENSIONS | TEXT_EXTENSIONS | {".pdf", ".docx", ".xlsx", ".xlsm"}
GENERATED_EXTENSIONS = {".zip", ".doc", ".xls", ".bin"}


def _ocr_worker(content, sender):
    try:
        if os.name == "nt":
            import ctypes
            kernel = ctypes.WinDLL("kernel32", use_last_error=True)
            kernel.SetPriorityClass(kernel.GetCurrentProcess(), 0x00004000)
        from PIL import Image
        from winocr import recognize_pil
        with Image.open(io.BytesIO(content)) as image:
            image = image.convert("RGB")
            async def recognize():
                return await recognize_pil(image, "zh-Hans-CN")
            result = asyncio.run(recognize())
        sender.send((True, "\n".join(str(line.text or "") for line in result.lines)))
    except Exception:
        sender.send((False, "当前系统图片文字识别不可用，可继续使用图片模型或手动补充文字。"))
    finally:
        sender.close()


def recognize_text(content, timeout=30):
    context = multiprocessing.get_context("spawn")
    receiver, sender = context.Pipe(duplex=False)
    process = context.Process(target=_ocr_worker, args=(content, sender), daemon=True)
    try:
        process.start()
        sender.close()
        if not receiver.poll(timeout):
            raise AssistantError("图片文字识别超时，原图已保留，可继续使用图片模型。", 408)
        ok, result = receiver.recv()
        if not ok:
            raise AssistantError(result, 422)
        return result
    finally:
        if process.is_alive():
            process.terminate()
        if process.pid:
            process.join(timeout=2)
        sender.close()
        receiver.close()


def _check_archive(content):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        info = archive.infolist()
        if len(info) > 5000 or sum(item.file_size for item in info) > MAX_BATCH_BYTES:
            raise AssistantError("文档解压后过大，请拆分文件。", 413)


def safe_document(text):
    text = str(text)
    if len(text) > 100000 or any(len(line) > 6000 for line in text.splitlines()):
        raise AssistantError("文档文字超过完整读取上限，请拆分后上传；原文件仍可下载。", 413)
    return "\n".join(safe_text(line) for line in text.splitlines())


def document_data(value, depth=0):
    if depth > 12:
        raise AssistantError("文档数据层级过深，请简化后上传。")
    if isinstance(value, dict):
        return {key: document_data(item, depth + 1) for key, item in value.items() if key in safe_data({key: None})}
    if isinstance(value, list):
        return [document_data(item, depth + 1) for item in value]
    return safe_data(value)


def extract_text(content, suffix):
    if suffix in TEXT_EXTENSIONS:
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = content.decode(encoding)
            except UnicodeError:
                continue
            if suffix == ".json":
                text = json.dumps(document_data(json.loads(text)), ensure_ascii=False, indent=2)
            return safe_document(text)
        raise AssistantError("无法识别文本编码，请另存为UTF-8。")
    if suffix in IMAGE_EXTENSIONS:
        return safe_document(recognize_text(content))
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise AssistantError("加密PDF无法读取，请上传未加密文件。")
        if len(reader.pages) > 100:
            raise AssistantError("PDF最多100页，请拆分文件。", 413)
        return safe_document("\n".join(page.extract_text() or "" for page in reader.pages))
    _check_archive(content)
    if suffix == ".docx":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            xml = archive.read("word/document.xml")
        root = ElementTree.fromstring(xml)
        return safe_document("\n".join("".join(p.itertext()) for p in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")))
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    lines = []
    try:
        for sheet in workbook.worksheets:
            lines.append("工作表：" + sheet.title)
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= 5000 or sum(map(len, lines)) >= 100000 or any(value is not None for value in row[80:]):
                    raise AssistantError("表格文字超过读取上限，请拆分后上传；原文件仍可下载。", 413)
                lines.append("\t".join("" if value is None else str(value) for value in row[:80]))
        return safe_document("\n".join(lines))
    finally:
        workbook.close()


class LighthouseFiles:
    def __init__(self, store, root=None):
        self.store = store
        self.root = (Path(root) if root else Path(store.db_path).parent / "lighthouse_assistant" / "files").resolve()

    def upload(self, actor, name, content, *, extract=True, source_scopes=None):
        name = re.sub(r"[\x00-\x1f]", "", str(name or "文件").replace("\\", "/").split("/")[-1])[:180]
        suffix = Path(name).suffix.lower()
        if suffix not in ALLOWED_EXTENSIONS and not (not extract and suffix in GENERATED_EXTENSIONS):
            raise AssistantError("支持图片、文本、PDF、Word和Excel文件。")
        if not content or len(content) > MAX_FILE_BYTES:
            raise AssistantError("单文件须为1字节至20MiB。", 413)
        mime = mimetypes.guess_type(name)[0] or "application/octet-stream"
        if suffix in IMAGE_EXTENSIONS:
            from PIL import Image
            with Image.open(io.BytesIO(content)) as image:
                mime = Image.MIME.get(image.format, mime)
                if image.width * image.height > 40000000:
                    raise AssistantError("图片像素过大，请缩小后上传。", 413)
                image.verify()
        identity = uuid.uuid4().hex
        directory = self.root / hashlib.sha256(actor["id"].encode()).hexdigest()[:24]
        directory.mkdir(parents=True, exist_ok=True)
        target = directory / (identity + suffix)
        target.write_bytes(content)
        item = {"id": identity, "owner": actor["id"], "name": name, "mime": mime, "size": len(content),
                "path": str(target), "sha256": hashlib.sha256(content).hexdigest(), "created_at": time.time(), "text": "", "error": "", "extracted": extract}
        if source_scopes:
            item["source_scopes"] = sorted(set(source_scopes))
        if extract:
            try:
                item["text"] = extract_text(content, suffix)
                if not item["text"]:
                    item["error"] = "未提取到可读文字；图片可继续识别，扫描PDF请补充截图。"
            except Exception as exc:
                item["error"] = str(exc) if isinstance(exc, AssistantError) else "文件文字提取未完成，原文件已保留，可补充文字。"
        try:
            self.store.put_document(FILE_NAMESPACE, identity, item)
        except Exception:
            target.unlink(missing_ok=True)
            raise AssistantError("附件保存未完成，请重新上传。", 503) from None
        return self.public(item)

    @staticmethod
    def public(item):
        return {k: item[k] for k in ("id", "name", "mime", "size", "error")} | {"url": "/api/assistant/files/" + item["id"], "is_image": Path(item["name"]).suffix.lower() in IMAGE_EXTENSIONS}

    def get(self, actor, identity):
        if not isinstance(identity, str) or not re.fullmatch(r"[a-f0-9]{32}", identity):
            raise AssistantError("附件编号无效。", 404)
        item = self.store.get_document(FILE_NAMESPACE, identity)
        if not item or item.get("owner") != actor["id"] or item.get("deleted_at"):
            raise AssistantError("附件不存在或没有访问权限。", 404)
        if set(item.get("source_scopes", [])) - set(actor.get("allowed_scopes", actor["scopes"])):
            raise AssistantError("当前账号已无权读取该业务文件。", 403)
        target = Path(item["path"]).resolve()
        if not target.is_relative_to(self.root) or not target.is_file():
            raise AssistantError("附件原文件不可用，请重新上传。", 404)
        return item

    def text(self, actor, identity, offset=0, length=4000):
        item = self.get(actor, identity)
        if not item.get("extracted", True):
            if Path(item["name"]).suffix.lower() in GENERATED_EXTENSIONS:
                raise AssistantError("该文件可下载，但不支持直接提取文字，请另存为PDF或Excel新格式。")
            try:
                item["text"] = extract_text(Path(item["path"]).read_bytes(), Path(item["name"]).suffix.lower())
            except Exception:
                item["error"] = "文件文字提取未完成，可下载原文件核对。"
            item["extracted"] = True
            self.store.put_document(FILE_NAMESPACE, identity, item)
        offset, length = max(0, int(offset)), max(1, min(8000, int(length)))
        text = item.get("text", "")
        return {"file_id": identity, "name": safe_text(item["name"]), "text": text[offset:offset + length], "next_offset": offset + length if offset + length < len(text) else None, "error": item.get("error", "")}

    def context(self, actor, identities):
        if not isinstance(identities, list) or len(identities) > 10 or len(set(identities)) != len(identities):
            raise AssistantError("每次最多附带10个不同文件。")
        items = [self.get(actor, identity) for identity in identities]
        if sum(item["size"] for item in items) > MAX_BATCH_BYTES:
            raise AssistantError("附件合计不得超过100MiB。", 413)
        return [self.text(actor, item["id"], length=max(300, 5000 // max(1, len(items)))) for item in items]

    def image_parts(self, actor, identities):
        from PIL import Image
        parts = []
        for identity in identities:
            if len(parts) >= 5:
                break
            item = self.get(actor, identity)
            if Path(item["name"]).suffix.lower() not in IMAGE_EXTENSIONS:
                continue
            with Image.open(item["path"]) as image:
                image = image.convert("RGB")
                image.thumbnail((1600, 1600))
                output = io.BytesIO()
                image.save(output, format="JPEG", quality=85)
            parts.append({"type": "image_url", "image_url": {"url": "data:image/jpeg;base64," + base64.b64encode(output.getvalue()).decode()}})
        return parts
