"""Manual papers, paginated progress and existing material adapters, entirely offline."""
import asyncio
import copy
import datetime as dt
import io
import json
from pathlib import Path
import random
import socket
import tempfile
import unittest
from unittest.mock import Mock, patch
from types import SimpleNamespace

from bin import test_learning_sharing as fixtures
from bin.test_learning_sharing import SOURCE, ADMIN, questions, TODAY
from lan_bitable_template_portal.learning import LearningError, TZ
from lan_bitable_template_portal.learning_sharing import validate_questions, sample_questions, exercise_card
from lan_bitable_template_portal.learning_sharing_routes import SharingRuntime
from lan_bitable_template_portal.sharing_sources import SharingMaterials, combine_sources, read_material_link, source_url


def manual(count):
    return [{'stem': f'Manual knowledge {i}', 'type': 'interview' if i % 3 == 2 else 'multiple' if i % 3 == 1 else 'single',
             'origin': 'manual', 'options': [] if i % 3 == 2 else ['First', 'Second'],
             'correct': [] if i % 3 == 2 else [0, 1] if i % 3 == 1 else [0],
             'reference_answer': 'Approved reference' if i % 3 == 2 else '', 'explanation': 'Professional experience'} for i in range(count)]


_CONNECT = socket.socket.connect
def local_only(sock, address):
    if isinstance(address, tuple) and address[0] in {'127.0.0.1', '::1'}:
        return _CONNECT(sock, address)
    raise AssertionError('external network forbidden')


class SharingMaterialsTests(unittest.TestCase):
    def setUp(self):
        self.enterContext(patch('socket.socket.connect', local_only))
        self.fixture = fixtures.SharingTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.core, self.learning, self.bid = self.fixture.core, self.fixture.learning, self.fixture.bid

    def draft(self, number, ai=20):
        qs = validate_questions(questions(ai) + manual(number), SOURCE, ['single', 'multiple'], draft=True)
        with self.learning.transaction() as conn:
            batch = self.core.batch(self.bid, conn)
            batch.update(questions=qs, status='draft', version=1)
            self.learning._put('sharing_batch', self.bid, batch, conn)
        return qs

    def publish(self, number, ai=20):
        qs = self.draft(number, ai)
        self.core.publish(self.bid, 1, ADMIN, 'cli_fixtureCurrent')
        return qs, self.core.attempts(self.bid)[0]

    def fields(self, paper, page):
        qs = {q['id']: q for q in self.core.batch(self.bid)['questions']}
        return {qid: 'Answer from experience' if qs[qid]['type'] == 'interview' else [str(i) for i in qs[qid]['correct']]
                for qid in paper['question_ids'][page * 5:page * 5 + 5]}

    def test_manual_zero_less_equal_greater_than_five_assignment_contract(self):
        for count in [0, 1, 4, 5, 6, 11]:
            with self.subTest(count=count):
                qs = validate_questions(questions() + manual(count), SOURCE, ['single', 'multiple'])
                ids = sample_questions(qs, random.Random(1))
                self.assertEqual(len(ids), max(5, count))
                self.assertTrue({q['id'] for q in qs if q.get('origin') == 'manual'} <= set(ids))
                self.assertEqual(len(ids), len(set(ids)))

    def test_limits_experience_and_legacy_evidence_are_not_weakened(self):
        self.assertEqual(len(validate_questions(manual(100), '', ['single'])), 100)
        for qs in [questions(41), manual(101), questions(40) + manual(61)]:
            with self.assertRaises(LearningError):
                validate_questions(qs, SOURCE)
        invalid = questions(5); invalid[0]['evidence'] = ''
        with self.assertRaises(LearningError):
            validate_questions(invalid, SOURCE)
        self.assertTrue(all('origin' not in q for q in validate_questions(questions(5), SOURCE)))
        self.assertEqual(validate_questions([], '', draft=True), [])

    def test_not_enough_ai_to_supplement_manual_rejected_without_assignment(self):
        self.draft(4, 0)
        with self.assertRaises(LearningError):
            self.core.publish(self.bid, 1, ADMIN, 'cli_fixtureCurrent')
        self.assertEqual(self.core.attempts(self.bid), [])

    def test_paged_web_progress_restart_idempotency_and_final_only_completion(self):
        qs, paper = self.publish(11)
        actor = {'id': paper['open_id']}
        self.learning.completion_changed = Mock()
        for page in range(3):
            payload = {'page': page, 'answers': self.fields(paper, page), 'version': page,
                       'operation_id': 'fixture-page-operation-' + str(page), 'submit': page == 2}
            result = self.core.answer(self.bid, payload, actor, lambda _: True)
            self.core.answer(self.bid, payload, actor, lambda _: True)
            stored = self.learning._get('sharing_attempt', paper['id'])
            self.assertEqual(stored['progress_version'], page + 1)
            self.assertEqual(bool(stored['submitted_at']), page == 2)
            if page < 2:
                self.assertIsNone(stored['correct'])
                self.assertEqual(stored['answers'], {})
                self.assertFalse(any(job['kind'] == 'result' for job in self.learning._all('sharing_delivery')))
            self.assertEqual(result['submitted'], page == 2)
        stored = self.learning._get('sharing_attempt', paper['id'])
        self.assertEqual(len(stored['answers']), 11)
        self.assertEqual(len([job for job in self.learning._all('sharing_delivery') if job['kind'] == 'result']), 1)
        with self.assertRaises(LearningError):
            self.core.edit(self.bid, {'version': 2, 'questions': qs}, ADMIN)
        self.learning.completion_changed.assert_called()
        self.core.schedule(dt.datetime(2026, 10, 19, tzinfo=TZ))
        weekly = next(job for job in self.learning._all('sharing_delivery') if job['kind'] == 'weekly')
        self.assertEqual(sum(item['answered'] for item in weekly['items']), 11)

    def test_page_isolation_auth_and_conflicting_save_do_not_overwrite(self):
        _, paper = self.publish(6)
        actor = {'id': paper['open_id']}
        payload = {'page': 0, 'answers': self.fields(paper, 0), 'version': 0, 'operation_id': 'fixture-save-once-0001', 'submit': False}
        for who, allowed in [({'id': 'ou_other'}, lambda _: True), ({**actor, 'shared_account': True}, lambda _: True), (actor, lambda _: False)]:
            with self.assertRaises(LearningError):
                self.core.answer(self.bid, payload, who, allowed)
        self.core.answer(self.bid, payload, actor, lambda _: True)
        for changed in [{**payload, 'operation_id': 'fixture-stale-version-0002'},
                        {**payload, 'answers': {}}, {**payload, 'page': 1}]:
            with self.assertRaises(LearningError):
                self.core.answer(self.bid, changed, actor, lambda _: True)
        final = {'page': 1, 'answers': self.fields(paper, 1), 'version': 1,
                 'operation_id': 'fixture-final-page-0003', 'submit': True}
        self.core.answer(self.bid, final, actor, lambda _: True)
        self.assertTrue(self.learning._get('sharing_attempt', paper['id'])['submitted_at'])

    def test_partial_final_does_not_finish_and_saved_answers_not_leaked(self):
        _, paper = self.publish(11)
        payload = {'page': 2, 'answers': self.fields(paper, 2), 'version': 0, 'operation_id': 'fixture-premature-final', 'submit': True}
        with self.assertRaises(LearningError):
            self.core.answer(self.bid, payload, {'id': paper['open_id']}, lambda _: True)
        detail = self.core.detail(self.bid, {'id': paper['open_id']})
        self.assertEqual(detail['paper']['question_count'], 11)
        self.assertTrue(all('correct' not in q for q in detail['questions']))

    def test_feishu_page_callbacks_return_same_card_progress_then_results(self):
        _, paper = self.publish(6)
        for page in range(2):
            result = self.core.submit({'attempt_id': paper['id'], 'ticket': paper['ticket'], 'page': page,
                'target_page': min(1, page + 1), 'version': page, 'submit': page == 1,
                'operation_id': 'fixture-feishu-page-' + str(page)}, self.fields(paper, page),
                {'union_id': paper['union_id']}, paper['app_id'], lambda _: True)
            self.assertIn('card', result)
            stored = self.learning._get('sharing_attempt', paper['id'])
            self.assertEqual(bool(stored['submitted_at']), page == 1)
        result = self.core.submit({'attempt_id': paper['id'], 'ticket': paper['ticket'], 'target_page': 0}, {},
                                 {'union_id': paper['union_id']}, paper['app_id'], lambda _: True)
        self.assertIn('答题结果', json.dumps(result['card'], ensure_ascii=False))
        self.assertIn('6道题', json.dumps(exercise_card(self.core.batch(self.bid), paper), ensure_ascii=False))

    def test_material_only_requires_explicit_server_people_and_idempotency(self):
        with self.learning.transaction() as conn:
            self.learning._put('person', 'p_current', {'id': 'p_current', 'name': 'Current', 'active': True,
                'login_ids': ['ou_current'], 'scopes': ['A']}, conn, False)
        with self.assertRaises(LearningError):
            self.core.create('', 'fixture-material-create-01', ADMIN, material_urls=['https://vnet.feishu.cn/wiki/Node'])
        kwargs = {'material_urls': ['https://vnet.feishu.cn/wiki/Node'], 'participant_ids': ['p_current'], 'auto_generate': False}
        created = self.core.create('', 'fixture-material-create-02', ADMIN, **kwargs)
        self.assertEqual(self.core.create('', 'fixture-material-create-02', ADMIN, **kwargs), created)
        with self.assertRaises(LearningError):
            self.core.create('', 'fixture-material-create-02', ADMIN, **{**kwargs, 'auto_generate': True})
        with self.assertRaises(LearningError):
            self.core.create('', 'fixture-material-create-03', ADMIN, **{**kwargs, 'participant_ids': ['forged']})

    def test_material_upload_is_local_private_and_full_text_uses_existing_worker(self):
        service = SharingMaterials(self.learning)
        item = service.upload(ADMIN, 'slides.pptx', b'fixture-pptx')
        with patch('openclaw_service.assistant.lighthouse_document_tools.run_local', return_value={
                'text': SOURCE * 8, 'references': [], 'extraction_complete': True}) as parser:
            parsed = service.read(item['id'], ADMIN['id'])
        self.assertEqual(parsed['text'], SOURCE * 8)
        self.assertEqual(parser.call_args.args[0], 'markdown')
        with self.assertRaises(LearningError):
            service.read(item['id'], 'ou_other')
        Path(self.learning._get('sharing_material', item['id'])['path']).write_bytes(b'changed')
        from openclaw_service.assistant.lighthouse_document_tools import checked_file
        with self.assertRaises(Exception):
            checked_file(self.learning._get('sharing_material', item['id']))

    def test_pdf_ocr_fills_all_missing_pages_not_just_first_excerpt(self):
        service = SharingMaterials(self.learning)
        item = service.upload(ADMIN, 'scan.pdf', b'fixture-pdf')
        def parse(kind, payload):
            if kind == 'markdown':
                return {'text': '', 'references': [], 'extraction_complete': False, 'total_pages': 3, 'empty': True}
            return {'text': f'OCR page {payload["page"]}', 'low_confidence_count': 1}
        with patch('openclaw_service.assistant.lighthouse_document_tools.run_local', side_effect=parse) as parser:
            text = service.read(item['id'], ADMIN['id'])
        self.assertEqual(parser.call_count, 4)
        self.assertIn('OCR page 3', text['text'])
        self.assertEqual(text['low_confidence_count'], 3)
        self.assertEqual([ref['page'] for ref in text['references']], [1, 2, 3])
        for ref in text['references']:
            self.assertIn(f'OCR page {ref["page"]}', text['text'][ref['start']:ref['end']])

    def test_real_office_formats_use_existing_full_document_extractor(self):
        import importlib.util
        if not all(importlib.util.find_spec(name) for name in ('docx', 'pptx', 'reportlab')):
            self.skipTest('Office/PDF smoke requires the bundled project runtime')
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from reportlab.pdfgen.canvas import Canvas
        from openclaw_service.assistant.lighthouse_document_tools import document_markdown
        documents = []
        word = Document(); word.add_paragraph('first paragraph'); word.add_paragraph('END_WORD')
        buffer = io.BytesIO(); word.save(buffer); documents.append(('file.docx', buffer.getvalue(), 'END_WORD'))
        book = Workbook(); book.active['A1'] = 'first sheet'; book.create_sheet('last')['A2'] = 'END_SHEET'
        buffer = io.BytesIO(); book.save(buffer)
        for suffix in ('xlsx', 'xlsm'):
            documents.append(('file.' + suffix, buffer.getvalue(), 'END_SHEET'))
        slides = Presentation()
        for text in ('first slide', 'END_SLIDE'):
            slides.slides.add_slide(slides.slide_layouts[1]).shapes.title.text = text
        buffer = io.BytesIO(); slides.save(buffer); documents.append(('file.pptx', buffer.getvalue(), 'END_SLIDE'))
        buffer = io.BytesIO(); canvas = Canvas(buffer)
        for text in ('first page', 'END_PDF'):
            canvas.drawString(50, 700, text); canvas.showPage()
        canvas.save(); documents.append(('file.pdf', buffer.getvalue(), 'END_PDF'))
        def existing_worker(kind, payload):
            self.assertEqual(kind, 'markdown')
            path = Path(payload['path'])
            return document_markdown(path.read_bytes(), path.suffix)
        with patch('openclaw_service.assistant.lighthouse_document_tools.run_local', side_effect=existing_worker):
            for name, content, last in documents:
                with self.subTest(name=name):
                    item = SharingMaterials(self.learning).upload(ADMIN, name, content)
                    parsed = SharingMaterials(self.learning).read(item['id'], ADMIN['id'])
                    self.assertIn(last, parsed['text'])
                    self.assertTrue(parsed['extraction_complete'])

    def test_all_manual_publish_hundred_questions_each_card_bounded_and_restart_progress(self):
        _, paper = self.publish(100, 0)
        self.assertEqual(len(paper['question_ids']), 100)
        batch = self.core.batch(self.bid)
        for page in range(20):
            card = exercise_card(batch, paper, page=page)
            self.assertLess(len(json.dumps(card, ensure_ascii=False).encode()), 28000)
        payload = {'page': 0, 'answers': self.fields(paper, 0), 'version': 0, 'operation_id': 'fixture-restart-answers', 'submit': False}
        self.core.answer(self.bid, payload, {'id': paper['open_id']}, lambda _: True)
        from lan_bitable_template_portal.learning import LearningService
        from lan_bitable_template_portal.learning_sharing import SharingExercises
        restored = SharingExercises(LearningService(self.fixture.folder.name))
        detail = restored.detail(self.bid, {'id': paper['open_id']})
        self.assertEqual(len(detail['paper']['draft_answers']), 5)
        self.assertFalse(detail['paper']['submitted_at'])
        self.assertEqual(detail['paper']['open_id'], paper['open_id'])

    def test_simultaneous_page_saves_use_one_version(self):
        from concurrent.futures import ThreadPoolExecutor
        _, paper = self.publish(6)
        def save(index):
            try:
                self.core.answer(self.bid, {'page': 0, 'answers': self.fields(paper, 0), 'version': 0,
                    'operation_id': f'fixture-racing-save-{index}', 'submit': False}, {'id': paper['open_id']}, lambda _: True)
                return True
            except LearningError as exc:
                self.assertEqual(exc.status, 409)
                return False
        with ThreadPoolExecutor(max_workers=2) as pool:
            self.assertEqual(sorted(pool.map(save, range(2))), [False, True])
        self.assertEqual(self.learning._get('sharing_attempt', paper['id'])['progress_version'], 1)

    def test_completion_hook_includes_actual_dispatch_day_for_legacy_cross_day_batch(self):
        self.learning.completion_changed = Mock()
        self.core.changed_batch({'date': '2026-10-01', 'published_at': '2026-10-13T23:30:00Z'}, '2026-10-20')
        self.assertEqual(self.learning.completion_changed.call_args.args, ('2026-10-14', '2026-10-01', '2026-10-20'))

    def test_sources_do_not_allow_external_urls_or_silent_truncation(self):
        for url in ['http://vnet.feishu.cn/docx/X', 'https://evil.test/wiki/X', 'https://vnet.feishu.cn:443/wiki/X',
                    'https://[invalid/wiki/X']:
            with self.assertRaises(LearningError): source_url(url)
        with self.assertRaises(LearningError): combine_sources([{'text': 'a' * 24001}])
        self.assertEqual(combine_sources([{'text': 'first'}, {'text': 'second'}])['text'], 'first\n\nsecond')

    def test_invalid_manual_origin_is_a_validation_error(self):
        for origin in ([], {}, None, 1):
            with self.subTest(origin=origin), self.assertRaises(LearningError):
                validate_questions([{**manual(1)[0], 'origin': origin}], '', draft=True)

    def test_wiki_material_reads_revision_consistent_full_text_with_qt_reader(self):
        body = SOURCE * 4
        responses = [{'node': {'obj_type': 'docx', 'obj_token': 'Resolved'}},
                     {'document': {'title': 'Material', 'revision_id': 2}}, {'content': body},
                     {'items': [], 'has_more': False},
                     {'document': {'revision_id': 2}}]
        request = Mock(side_effect=[{'code': 0, 'data': data} for data in responses])
        with patch('openclaw_service.assistant.lighthouse_documents.document_reader', return_value=request) as reader:
            material = read_material_link(object(), 'https://vnet.feishu.cn/wiki/Node')
        self.assertEqual(reader.call_count, 1)
        self.assertEqual(material['text'], body)
        self.assertEqual(material['url'], 'https://vnet.feishu.cn/wiki/Node')
        self.assertIn('/Resolved/raw_content', request.call_args_list[2].args[0])
        for data in ({'code': 99991672, 'data': {}}, {'code': 0, 'data': {}}):
            with patch('openclaw_service.assistant.lighthouse_documents.document_reader', return_value=Mock(return_value=data)):
                with self.assertRaises(LearningError): read_material_link(object(), 'https://vnet.feishu.cn/docx/Node')

    def test_generation_material_only_and_combined_keep_verified_sources_and_roster(self):
        from unittest.mock import AsyncMock
        with self.learning.transaction() as conn:
            self.learning._put('person', 'p_current', {'id': 'p_current', 'name': 'Current', 'active': True,
                'login_ids': ['ou_current'], 'scopes': ['A']}, conn, False)
        for minutes in ('', 'https://vnet.feishu.cn/docx/Minutes'):
            bid = self.core.create(minutes, 'fixture-generate-material-' + str(bool(minutes)), ADMIN, ['single', 'multiple'],
                material_urls=['https://vnet.feishu.cn/wiki/Material'], participant_ids=['p_current'])['id']
            generator = AsyncMock(return_value=questions(5))
            runtime = SharingRuntime(SimpleNamespace(service=object(), generate_sharing_questions=generator), lambda: self.learning)
            runtime.core = self.core
            with patch('lan_bitable_template_portal.learning_sharing_routes.read_material_link', return_value={
                    'text': SOURCE, 'title': 'Material', 'sha256': 'fixture-sha', 'extraction_complete': True}), \
                 patch('lan_bitable_template_portal.learning_sharing_routes.read_source', return_value={
                    'text': SOURCE, 'title': 'Minutes', 'participants': [{'open_id': 'ou_other', 'name': 'Other'}]}):
                # A short final model page deliberately ends the generation loop.
                generator.return_value = questions(4)
                asyncio.run(runtime.generate(self.core.batch(bid)))
            batch = self.core.batch(bid)
            self.assertEqual(batch['status'], 'failed', 'four AI questions cannot be published as a complete generation')
            self.assertEqual(self.core.attempts(bid), [])
            # Manual-only drafts still parse the sources, without invoking any model.
            batch.update(status='queued', auto_generate=False)
            with self.learning.transaction() as conn: self.learning._put('sharing_batch', bid, batch, conn)
            generator.reset_mock()
            with patch('lan_bitable_template_portal.learning_sharing_routes.read_material_link', return_value={
                    'text': SOURCE, 'title': 'Material', 'sha256': 'fixture-sha', 'extraction_complete': True}), \
                 patch('lan_bitable_template_portal.learning_sharing_routes.read_source', return_value={
                    'text': SOURCE, 'title': 'Minutes', 'participants': [{'open_id': 'ou_other', 'name': 'Other'}]}):
                asyncio.run(runtime.generate(self.core.batch(bid)))
            batch = self.core.batch(bid)
            self.assertEqual(batch['status'], 'draft')
            self.assertEqual(len(batch['source']['sources']), 2 if minutes else 1)
            self.assertEqual(batch['participants'][0]['open_id'], 'ou_current')
            generator.assert_not_called()

    def test_feishu_manual_matrix_completion_and_readonly_result_pages(self):
        template = self.core.batch(self.bid)
        for count in (0, 4, 5, 11):
            self.bid = 'manual-feishu-' + str(count)
            with self.learning.transaction() as conn:
                self.learning._put('sharing_batch', self.bid, {**copy.deepcopy(template), 'id': self.bid}, conn)
            _, paper = self.publish(count)
            for page in range((len(paper['question_ids']) + 4) // 5):
                final = page == (len(paper['question_ids']) - 1) // 5
                value = {'attempt_id': paper['id'], 'ticket': paper['ticket'], 'page': page,
                    'target_page': page, 'version': page, 'submit': final, 'operation_id': f'fixture-matrix-{count}-page-{page}'}
                response = self.core.submit(value, self.fields(paper, page), {'union_id': paper['union_id']}, paper['app_id'], lambda _: True)
                self.assertTrue(response['ok'])
                current = self.learning._get('sharing_attempt', paper['id'])
                self.assertEqual(bool(current['submitted_at']), final)
                self.assertEqual(self.core.submit(value, self.fields(paper, page), {'union_id': paper['union_id']}, paper['app_id'], lambda _: True)['ok'], True)
            self.assertEqual(len(current['answers']), max(5, count))
            self.assertEqual(len([j for j in self.learning._all('sharing_delivery') if j['kind'] == 'result' and j['batch_id'] == self.bid]), 1)


class SharingMaterialRoutesTests(unittest.TestCase):
    def setUp(self):
        self.fixture = fixtures.SharingRouteTests(); self.fixture.setUp(); self.addCleanup(self.fixture.doCleanups)
        self.client = self.fixture.client
        self.enterContext(patch('socket.socket.connect', local_only))

    def test_linked_material_creation_does_not_widen_minute_urls(self):
        self.fixture.actor = ADMIN
        learning = self.fixture.fixture.learning
        with learning.transaction() as conn:
            learning._put('person', 'p_linked', {'id': 'p_linked', 'name': 'Linked learner',
                'active': True, 'login_ids': ['ou_linked'], 'scopes': ['A']}, conn, False)
        payload = {'url': 'https://vnet.feishu.cn/file/FileToken', 'operation_id': 'fixture-link-create-01',
                   'participant_ids': ['p_linked'], 'auto_generate': False}
        self.assertEqual(self.client.post('/api/learning/sharing', json=payload).status_code, 400)
        payload.update(url='', material_urls=['https://vnet.feishu.cn/file/FileToken'])
        response = self.client.post('/api/learning/sharing', json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        batch = self.fixture.fixture.core.batch(response.json()['data']['id'])
        self.assertEqual(batch['url'], '')
        self.assertEqual(batch['material_urls'], payload['material_urls'])

    def test_material_upload_and_people_are_admin_only_and_csrf_protected(self):
        root = '/api/learning/sharing'
        self.fixture.actor = {'id': 'ou_person0'}
        self.assertEqual(self.client.get(root + '/people').status_code, 403)
        self.assertEqual(self.client.post(root + '/materials', files={'files': ('slide.pptx', b'fixture')}).status_code, 403)
        self.fixture.actor = ADMIN
        self.assertEqual(self.client.post(root + '/materials', files={'files': ('slide.pptx', b'fixture')},
            headers={'origin': 'https://foreign.test'}).status_code, 403)
        self.assertEqual(self.client.post(root + '/materials', files={'files': ('unsupported.exe', b'fixture')}).status_code, 400)
        response = self.client.post(root + '/materials', files={'files': ('slide.pptx', b'fixture')})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['data']['items'][0]['name'], 'slide.pptx')

    def test_web_answers_endpoint_uses_actor_and_complete_paper(self):
        core = self.fixture.fixture.core
        core.publish(self.fixture.fixture.bid, 1, ADMIN, 'cli_fixtureCurrent')
        paper = core.attempts(self.fixture.fixture.bid)[0]
        # Only the allow-check is stubbed; the endpoint still derives the actor from its session.
        from lan_bitable_template_portal.learning_sharing_routes import SharingRuntime
        with patch.object(SharingRuntime, 'allowed', return_value=True):
            self.fixture.actor = {'id': paper['open_id']}
            qmap = {q['id']: q for q in core.batch(paper['batch_id'])['questions']}
            answers = {qid: [str(i) for i in qmap[qid]['correct']] for qid in paper['question_ids']}
            payload = {'page': 0, 'answers': answers, 'version': 0, 'operation_id': 'fixture-route-submit', 'submit': True}
            url = '/api/learning/sharing/' + paper['batch_id'] + '/answers'
            self.assertEqual(self.client.post(url, json=payload, headers={'origin': 'https://foreign.test'}).status_code, 403)
            response = self.client.post(url, json=payload)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertTrue(response.json()['data']['submitted'])
            self.assertEqual(len([j for j in core.learning._all('sharing_delivery') if j['kind'] == 'result']), 1)

    def test_close_receipt_does_not_depend_on_a_second_post_commit_batch_read(self):
        core = self.fixture.fixture.core
        core.publish(self.fixture.fixture.bid, 1, ADMIN, 'cli_fixtureCurrent')
        original = type(core).batch
        def within_transaction(service, identity, conn=None):
            self.assertIsNotNone(conn, 'projection notifications reuse the committed snapshot')
            return original(service, identity, conn)
        with patch.object(type(core), 'batch', within_transaction):
            response = self.client.post('/api/learning/sharing/' + self.fixture.fixture.bid + '/close', json={'version': 2})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(core.batch(self.fixture.fixture.bid)['status'], 'closed')


if __name__ == '__main__':
    unittest.main()
