"""Unit tests (unittest only) for lighthouse_knowledge_extract bounded helpers."""
import io
import sys
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openclaw_service.assistant.lighthouse_ai import AssistantError
from openclaw_service.assistant.lighthouse_knowledge_extract import (
    MAX_BYTES,
    MAX_CHUNKS,
    extract_sections,
    split_sections,
)

from openpyxl import Workbook
from pypdf import PdfWriter


# ---------------------------------------------------------------------------
# in-memory fixtures (no temp business files are ever created)
# ---------------------------------------------------------------------------

DOCX_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def make_docx_body(body_segments):
    markup = [f'<w:document xmlns:w="{DOCX_NS}">', "<w:body>"]
    markup.extend(body_segments)
    markup.append("</w:body></w:document>")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "".join(markup))
    return buffer.getvalue()


def make_docx(paragraphs, with_table=True):
    body = []
    for text in paragraphs:
        body.append(f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>")
    if with_table:
        body.append(
            "<w:tbl><w:tr><w:tc>"
            "<w:p><w:r><w:t>表格中的说明</w:t></w:r></w:p>"
            "</w:tc></w:tr></w:tbl>")
    return make_docx_body(body)


def make_pdf_text(text):
    line = "BT\n/F1 12 Tf\n72 720 Td\n(" + text + ") Tj\nET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        "<< /Length %d >>\nstream\n%s\nendstream" % (len(line.encode("latin-1")), line),
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    lines = ["%PDF-1.4"]
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len("\n".join(lines).encode("latin-1")))
        lines.append("%d 0 obj\n%s\nendobj" % (number, body))
    xref = len("\n".join(lines).encode("latin-1"))
    lines.append("xref")
    lines.append("0 6")
    lines.append("0000000000 65535 f ")
    for offset in offsets:
        lines.append("%010d 00000 n " % offset)
    lines += ["trailer", "<< /Size 6 /Root 1 0 R >>", "startxref", str(xref), "%%EOF"]
    return "\n".join(lines).encode("latin-1")


def make_blank_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_encrypted_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.encrypt("secret")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_xlsx(rows, title="SheetA", formulas=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    if formulas:
        for coordinate, value in formulas.items():
            sheet[coordinate] = value
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_xlsx_multi(sheets):
    """sheets: iterable of (title, rows)."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# extraction tests
# ---------------------------------------------------------------------------

class ExtractTextTests(unittest.TestCase):
    def test_txt_utf8_paragraphs_and_locations(self):
        content = "第一段内容。\n\n第二段内容 abc。\n\n第三段内容。".encode("utf-8")
        sections = extract_sections(content, "说明.txt")
        self.assertEqual([s["text"] for s in sections],
                         ["第一段内容。", "第二段内容 abc。", "第三段内容。"])
        self.assertEqual([s["location"] for s in sections],
                         ["说明.txt:段落1", "说明.txt:段落2", "说明.txt:段落3"])

    def test_txt_gb18030(self):
        content = "你好世界。\n\n第二个段落。".encode("gb18030")
        sections = extract_sections(content, "文档.md")
        self.assertEqual(len(sections), 2)
        self.assertEqual(sections[0]["text"], "你好世界。")
        self.assertTrue(sections[0]["location"].startswith("文档.md:段落"))


class ExtractCsvTests(unittest.TestCase):
    def test_csv_sections_keep_header_context_and_row_range(self):
        raw = "设备,数量,状态\nA机柜,10,正常\nB机柜,20,维护中\n"
        sections = extract_sections(raw.encode("utf-8"), "台账.csv")
        self.assertEqual(len(sections), 1)
        text = sections[0]["text"]
        self.assertIn("设备\t数量\t状态", text)
        self.assertIn("A机柜\t10\t正常", text)
        self.assertIn("B机柜\t20\t维护中", text)
        self.assertTrue(sections[0]["location"].startswith("台账.csv:行2-3"))
        self.assertIn("（列：设备、数量、状态）", sections[0]["location"])

    def test_csv_preserves_original_row_numbers_including_blanks(self):
        # header=row1, blank=row2, data=row3, blank=row4, data=row5
        raw = "设备,数量,状态\n\nA机柜,10,正常\n\nB机柜,20,维护中\n"
        sections = extract_sections(raw.encode("utf-8"), "台账.csv")
        self.assertEqual(len(sections), 1)
        self.assertEqual(sections[0]["location"], "台账.csv:行3-5（列：设备、数量、状态）")
        self.assertIn("A机柜\t10\t正常", sections[0]["text"])
        self.assertIn("B机柜\t20\t维护中", sections[0]["text"])

    def test_csv_header_only_indexes_not_empty(self):
        sections = extract_sections("设备,数量,状态\n".encode("utf-8"), "空表.csv")
        self.assertEqual(len(sections), 1)
        self.assertEqual(sections[0]["text"], "设备\t数量\t状态")
        self.assertEqual(sections[0]["location"], "空表.csv:行1（列：设备、数量、状态）")

    def test_csv_row_limit_raises(self):
        raw = "设备,数量\n" + "\n".join(f"x{i},1" for i in range(20001)) + "\n"
        with self.assertRaises(AssistantError) as ctx:
            extract_sections(raw.encode("utf-8"), "长表.csv")
        self.assertIn("超过20000行", str(ctx.exception))

    def test_csv_column_limit_raises(self):
        header = ",".join("c%d" % i for i in range(101))
        raw = header + "\n"
        with self.assertRaises(AssistantError) as ctx:
            extract_sections(raw.encode("utf-8"), "宽表.csv")
        self.assertIn("超过100列", str(ctx.exception))

    def test_csv_undecodable_raises(self):
        with self.assertRaises(AssistantError) as ctx:
            extract_sections(b"\xc3\x28", "乱码.csv")
        self.assertIn("无法识别文本编码", str(ctx.exception))


class ExtractDocxTests(unittest.TestCase):
    def test_docx_keeps_headings_and_table_paragraphs(self):
        content = make_docx(["标题一", "正文段落二"])
        sections = extract_sections(content, "流程.docx")
        texts = [s["text"] for s in sections]
        self.assertIn("标题一", texts)
        self.assertIn("正文段落二", texts)
        self.assertIn("表格中的说明", texts)
        heading = next(s for s in sections if s["text"] == "标题一")
        body = next(s for s in sections if s["text"] == "正文段落二")
        cell = next(s for s in sections if s["text"] == "表格中的说明")
        self.assertTrue(heading["location"].startswith("流程.docx:段落"))
        self.assertTrue(body["location"].startswith("流程.docx:段落"))
        # Table-cell paragraph keeps its table + row association.
        self.assertTrue(cell["location"].startswith("流程.docx:表格1行1段落"))

    def test_docx_table_rows_keep_row_association(self):
        body = [
            "<w:p><w:r><w:t>开头</w:t></w:r></w:p>",
            "<w:tbl>"
            "<w:tr><w:tc><w:p><w:r><w:t>行1格1</w:t></w:r></w:p></w:tc>"
            "<w:tc><w:p><w:r><w:t>行1格2</w:t></w:r></w:p></w:tc></w:tr>"
            "<w:tr><w:tc><w:p><w:r><w:t>行2格1</w:t></w:r></w:p></w:tc></w:tr>"
            "</w:tbl>",
            "<w:p><w:r><w:t>结尾</w:t></w:r></w:p>",
        ]
        sections = extract_sections(make_docx_body(body), "表格.docx")
        locs = {s["text"]: s["location"] for s in sections}
        self.assertEqual(locs["开头"], "表格.docx:段落1")
        self.assertEqual(locs["行1格1"], "表格.docx:表格1行1段落2")
        self.assertEqual(locs["行1格2"], "表格.docx:表格1行1段落3")
        self.assertEqual(locs["行2格1"], "表格.docx:表格1行2段落4")
        self.assertEqual(locs["结尾"], "表格.docx:段落5")

    def test_docx_paragraph_uses_only_wt_tab_br_not_metadata(self):
        body = [
            "<w:p><w:r><w:t>前置</w:t></w:r>"
            "<w:r><w:tab/></w:r>"
            "<w:r><w:br/><w:t>后置</w:t></w:r>"
            "<w:r><w:instrText>不应出现</w:instrText></w:r>"
            "</w:p>",
        ]
        sections = extract_sections(make_docx_body(body), "格式.docx")
        self.assertEqual(sections[0]["text"], "前置\t\n后置")
        self.assertNotIn("不应出现", sections[0]["text"])

    def test_docx_not_a_zip_raises(self):
        with self.assertRaises(zipfile.BadZipFile):
            extract_sections(b"this is not a docx zip", "坏.docx")


class ExtractXlsxTests(unittest.TestCase):
    def test_xlsx_formula_is_formula_text_and_keeps_context(self):
        content = make_xlsx(
            [["项目", "金额"], ["A", 100], ["B", 200]],
            formulas={"C2": "=SUM(B2:B3)"},
        )
        sections = extract_sections(content, "预算.xlsx")
        joined = "\n".join(s["text"] for s in sections)
        self.assertEqual(len(sections), 1)
        self.assertIn("=SUM(B2:B3)", joined)
        self.assertNotIn("=300", joined)  # never fabricate computed value
        self.assertTrue(sections[0]["location"].startswith("预算.xlsx:工作表「SheetA」行2-3"))
        self.assertIn("（列：项目、金额）", sections[0]["location"])

    def test_xlsx_column_limit_raises(self):
        rows = [["c%d" % i for i in range(101)]]
        with self.assertRaises(AssistantError) as ctx:
            extract_sections(make_xlsx(rows), "wide.xlsx")
        self.assertIn("超过100列", str(ctx.exception))

    def test_xlsx_row_limit_raises(self):
        rows = [[i, "x"] for i in range(20001)]
        with self.assertRaises(AssistantError) as ctx:
            extract_sections(make_xlsx(rows), "long.xlsx")
        self.assertIn("超过20000行", str(ctx.exception))

    def test_xlsx_cumulative_rows_across_sheets_raise(self):
        # Each sheet stays under the row limit; only the whole workbook is over.
        rows = [[f"r{i}"] for i in range(10001)]  # 10001 rows incl. header
        content = make_xlsx_multi([("S1", rows), ("S2", rows)])
        with self.assertRaises(AssistantError) as ctx:
            extract_sections(content, "multi.xlsx")
        self.assertIn("超过20000行", str(ctx.exception))

    def test_xlsx_header_only_indexes_not_empty(self):
        sections = extract_sections(make_xlsx([["项目", "金额"]]), "空表.xlsx")
        self.assertEqual(len(sections), 1)
        self.assertEqual(sections[0]["text"], "项目\t金额")
        self.assertIn("行1", sections[0]["location"])

    def test_xlsx_not_a_zip_raises(self):
        with self.assertRaises(zipfile.BadZipFile):
            extract_sections(b"this is not an xlsx zip", "坏.xlsx")


class ExtractPdfTests(unittest.TestCase):
    def test_pdf_text_page_location(self):
        sections = extract_sections(make_pdf_text("Hello Lighthouse"), "手册.pdf")
        self.assertEqual(len(sections), 1)
        self.assertEqual(sections[0]["text"], "Hello Lighthouse")
        self.assertEqual(sections[0]["location"], "手册.pdf:第1页")

    def test_pdf_scanned_blank_raises_need_ocr(self):
        with self.assertRaises(AssistantError) as ctx:
            extract_sections(make_blank_pdf(), "扫描.pdf")
        self.assertIn("OCR", str(ctx.exception))
        self.assertIn("不支持PDF内置OCR", str(ctx.exception))

    def test_pdf_encrypted_raises(self):
        with self.assertRaises(AssistantError) as ctx:
            extract_sections(make_encrypted_pdf(), "加密.pdf")
        self.assertIn("加密PDF", str(ctx.exception))

    def test_pdf_more_than_300_pages_raises(self):
        writer = PdfWriter()
        for _ in range(301):
            writer.add_blank_page(width=100, height=100)
        buffer = io.BytesIO()
        writer.write(buffer)
        with self.assertRaises(AssistantError) as ctx:
            extract_sections(buffer.getvalue(), "多页.pdf")
        self.assertIn("超过300页", str(ctx.exception))


class ExtractImageTests(unittest.TestCase):
    def test_image_uses_provided_ocr(self):
        seen = []

        def ocr(content):
            seen.append(content)
            return "图片里识别的设备文字"

        sections = extract_sections(b"fake-image-bytes", "照片.png", ocr=ocr)
        self.assertEqual(seen, [b"fake-image-bytes"])
        self.assertEqual(sections[0]["text"], "图片里识别的设备文字")
        self.assertEqual(sections[0]["location"], "照片.png:图片OCR")


# ---------------------------------------------------------------------------
# validation & bound tests
# ---------------------------------------------------------------------------

class ValidationTests(unittest.TestCase):
    def test_empty_content_raises(self):
        with self.assertRaises(AssistantError) as ctx:
            extract_sections(b"", "a.txt")
        self.assertIn("为空", str(ctx.exception))

    def test_bad_extension_raises(self):
        with self.assertRaises(AssistantError) as ctx:
            extract_sections(b"x" * 10, "archive.zip")
        self.assertIn("不支持的文件类型", str(ctx.exception))

    def test_bad_extension_without_suffix_raises(self):
        with self.assertRaises(AssistantError):
            extract_sections(b"x", "noext")

    def test_oversized_original_raises(self):
        blob = b"a" * (MAX_BYTES + 1)
        with self.assertRaises(AssistantError) as ctx:
            extract_sections(blob, "big.txt")
        self.assertIn("100MiB", str(ctx.exception))

    def test_total_extracted_char_limit_raises_not_truncates(self):
        # "段落甲乙丙丁" = 6 chars x 400_000 = 2.4M chars (>2M, still <20MiB)
        huge = ("段落甲乙丙丁" * 400_000).encode("utf-8")
        with self.assertRaises(AssistantError) as ctx:
            extract_sections(huge, "超大.txt")
        self.assertIn("200万字符", str(ctx.exception))


# ---------------------------------------------------------------------------
# split tests
# ---------------------------------------------------------------------------

class SplitTests(unittest.TestCase):
    def _long_text(self):
        paragraphs = ["句子%02d xxx yyy zzz" % i for i in range(40)]
        return "\n\n".join(paragraphs)

    def test_multichunk_exact_overlap(self):
        text = self._long_text()
        parts = split_sections([{"text": text, "location": "d.txt:段落1"}],
                               size=120, overlap=20)
        self.assertGreater(len(parts), 2)
        for first, second in zip(parts, parts[1:]):
            # A trailing remainder shorter than the overlap window cannot
            # carry a full overlap; assert exact overlap for full chunks only.
            if min(len(first["text"]), len(second["text"])) >= 20:
                self.assertEqual(second["text"][:20], first["text"][-20:])

    def test_every_part_keeps_source_location_anchor(self):
        text = self._long_text()
        parts = split_sections([{"text": text, "location": "手册.txt:第1段"}],
                               size=140, overlap=15)
        self.assertGreater(len(parts), 1)
        for index, part in enumerate(parts, 1):
            self.assertEqual(part["location"], f"手册.txt:第1段（片段{index}）")

    def test_single_paragraph_also_overlaps(self):
        text = "词".join("片段%02d" % i for i in range(40))
        parts = split_sections([{"text": text, "location": "f:段落1"}],
                               size=60, overlap=12)
        self.assertGreater(len(parts), 2)
        for first, second in zip(parts, parts[1:]):
            self.assertEqual(second["text"][:12], first["text"][-12:])

    def test_chunk_count_bound_raises(self):
        huge = [{"text": "字" * (MAX_CHUNKS + 1), "location": "f"}]
        with self.assertRaises(AssistantError) as ctx:
            split_sections(huge, size=1, overlap=0)
        self.assertIn(str(MAX_CHUNKS), str(ctx.exception))

    def test_default_chunks_fit_local_bge_context(self):
        text = "公司费用报销流程。" * 180
        parts = split_sections([{"text": text, "location": "f"}])
        self.assertTrue(all(len(part["text"]) <= 420 for part in parts))
        self.assertGreater(len(parts), 1)

    def test_empty_paragraph_skipped(self):
        parts = split_sections([{"text": "   \n\n ", "location": "f"}])
        self.assertEqual(parts, [])


if __name__ == "__main__":
    unittest.main(verbosity=2)
